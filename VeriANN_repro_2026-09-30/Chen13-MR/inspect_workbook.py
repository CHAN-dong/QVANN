import openpyxl
from pathlib import Path
w=openpyxl.load_workbook(Path(__file__).resolve().parents[2]/'work9.xlsx')
print([(s.title,s.max_row,s.max_column,len(s._charts)) for s in w])
for s in list(w)[:4]:
 print(s.title,'FORMULAS',[(c.coordinate,c.value) for row in s for c in row if c.data_type=='f'])
s=w['SIFT1M']
print('MERGES',[str(r) for r in s.merged_cells.ranges])
for r in [8,9,12,16,20,26,27,60,64,71,75,81,85,93,97,104,108,114,118,128,129,135]:
 print(r,[(s.cell(r,c).coordinate,s.cell(r,c).value) for c in range(1,12) if s.cell(r,c).value is not None])
for s in list(w)[4:]:
 if 'index' in s.title.lower() or '索引' in s.title:
  print('INDEX SHEET',s.title,list(s.values)[:5])
