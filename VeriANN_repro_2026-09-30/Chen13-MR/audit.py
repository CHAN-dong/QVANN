"""Independent Python audit of wire proofs, snapshots, raw samples and workbook.
Run after benchmark completion to avoid contending with timed native queries.
"""
import csv, hashlib, json, math, mmap, statistics, struct, subprocess, tempfile
from pathlib import Path
import numpy as np
import openpyxl
from run import ROOT, OUT, STUDY, DATA, KS, SIZES, sha, write
from fill import INSERTS, mapped

def h(b):return hashlib.sha256(b).digest()
def pack(*v):return struct.pack('<'+'I'*len(v),*v)
def nh(lo,hi,content):return h(pack(0x4e4f4445)+lo.astype('<f4').tobytes()+hi.astype('<f4').tobytes()+content)
def anchor(n,d,fo,lc,root):return h(pack(0x4d523031,n,d,fo,lc)+root)
def sigcheck(root,sig,pub):
 with tempfile.TemporaryDirectory() as folder:
  p=Path(folder);(p/'pub.der').write_bytes(bytes.fromhex('302a300506032b6570032100')+pub)
  (p/'message').write_bytes(root);(p/'signature').write_bytes(sig)
  subprocess.run(['openssl','pkeyutl','-verify','-pubin','-inkey',str(p/'pub.der'),'-keyform','DER','-rawin','-in',str(p/'message'),'-sigfile',str(p/'signature')],check=True,capture_output=True)

def inspect_snapshot(path,metadata,basepath=None):
 f=path.open('rb');m=mmap.mmap(f.fileno(),0,access=mmap.ACCESS_READ)
 magic,n,d,fo,lc,count,root=struct.unpack_from('<7I',m)
 assert magic==0x4d523031 and root==0
 dtype=np.dtype([('id','<u4'),('v','<f4',(d,))]);stride=4+4*d
 records=np.ndarray((n,),dtype=dtype,buffer=m,offset=28)
 assert np.array_equal(np.sort(records['id']),np.arange(n,dtype='u4'))
 pos=28+n*stride;nodes=[]
 for ix in range(count):
  begin,end,nc=struct.unpack_from('<3I',m,pos);pos+=12
  children=list(struct.unpack_from('<'+'I'*nc,m,pos));pos+=4*nc
  lo=np.frombuffer(m,dtype='<f4',count=d,offset=pos).copy();pos+=4*d
  hi=np.frombuffer(m,dtype='<f4',count=d,offset=pos).copy();pos+=4*d
  content=m[pos:pos+32];digest=m[pos+32:pos+64];pos+=64
  assert 0<=begin<end<=n and all(ix<c<count for c in children)
  nodes.append((begin,end,children,lo,hi,content,digest))
 assert pos+96==len(m)
 signature=m[pos:pos+64];pub=m[pos+64:pos+96]
 for ix in range(count-1,-1,-1):
  begin,end,children,lo,hi,content,digest=nodes[ix]
  if children:
   assert 2<=len(children)<=fo
   assert nodes[children[0]][0]==begin and nodes[children[-1]][1]==end
   assert all(nodes[a][1]==nodes[b][0] for a,b in zip(children,children[1:]))
   lower=np.minimum.reduce([nodes[c][3] for c in children]);upper=np.maximum.reduce([nodes[c][4] for c in children])
   actual=h(pack(2,len(children))+b''.join(nodes[c][6] for c in children))
  else:
   assert end-begin<=lc
   vectors=records['v'][begin:end];lower=vectors.min(axis=0);upper=vectors.max(axis=0)
   actual=h(pack(1,end-begin)+m[28+begin*stride:28+end*stride])
  assert np.array_equal(lo,lower) and np.array_equal(hi,upper)
  assert actual==content and nh(lo,hi,content)==digest
 a=anchor(n,d,fo,lc,nodes[root][6]);assert a.hex()==metadata['root'];sigcheck(a,signature,pub)
 assert len(m)==metadata['index_bytes']
 if basepath:
  original=fvecs(basepath,n)
  assert original.shape==(n,d)
  for begin in range(0,n,16384):
   end=min(n,begin+16384);ids=records['id'][begin:end]
   assert np.array_equal(records['v'][begin:end].view('<u4'),original[ids].view('<u4'))
  del ids
 del records;del vectors
 m.close();f.close()
 return dict(N=n,D=d,fanout=fo,leaf_capacity=lc,root=a,pub=pub,node_count=count)

def fvecs(path,count=None):
 d=int(np.fromfile(path,dtype='<i4',count=1)[0]);mm=np.memmap(path,dtype='<f4',mode='r').reshape(-1,d+1)
 if count is not None:assert len(mm)>=count;mm=mm[:count]
 assert np.all(mm[:,0].view('<i4')==d)
 return mm[:,1:]

def inspect_proof(path,a,querypath,basepath):
 # Reads from a stream, never the server tree. Only independent oracle below
 # accesses raw vectors, after the proof-only completeness check has succeeded.
 f=path.open('rb');pos=0
 def get(n):
  nonlocal pos
  b=f.read(n);assert len(b)==n;pos+=n;return b
 def u():return struct.unpack('<I',get(4))[0]
 magic,n,d,k=u(),u(),u(),u();assert magic==0x564f3031 and n==a['N'] and d==a['D']
 root=get(32);sig=get(64);assert root==a['root'];sigcheck(root,sig,a['pub'])
 q=fvecs(querypath,101)[100].copy();assert get(32)==h(q.astype('<f4').tobytes())
 claimed=np.frombuffer(get(k*4),dtype='<u4').copy()
 best_d=np.empty(0);best_id=np.empty(0,dtype='u4');minbound=math.inf;opened=0;closed=0
 dtype=np.dtype([('id','<u4'),('v','<f4',(d,))]);seen=np.zeros(n,dtype=bool)
 def node(depth=0):
  nonlocal best_d,best_id,minbound,opened,closed
  assert depth<64
  tag=get(1)[0]
  if tag==0:
   lo=np.frombuffer(get(d*4),dtype='<f4').copy();hi=np.frombuffer(get(d*4),dtype='<f4').copy();content=get(32)
   assert np.all(np.isfinite(lo)) and np.all(np.isfinite(hi)) and np.all(lo<=hi)
   delta=np.maximum(np.maximum(lo.astype('f8')-q,q.astype('f8')-hi),0);minbound=min(minbound,float(np.sum(delta*delta)));closed+=1
  elif tag==1:
   count=u();assert 0<count<=a['leaf_capacity'];b=get(count*(d*4+4));v=np.frombuffer(b,dtype=dtype)
   ids=v['id'];assert len(np.unique(ids))==count and np.all(ids<n) and not np.any(seen[ids]);seen[ids]=True
   assert np.all(np.isfinite(v['v']))
   delta=v['v'].astype('f8')-q.astype('f8');dd=np.sum(delta*delta,axis=1)
   combined_d=np.concatenate([best_d,dd]);combined_id=np.concatenate([best_id,ids]);order=np.lexsort((combined_id,combined_d))[:k];best_d=combined_d[order];best_id=combined_id[order]
   lo=v['v'].min(axis=0);hi=v['v'].max(axis=0);content=h(pack(1,count)+b);opened+=count
  else:
   assert tag==2;count=u();assert 2<=count<=a['fanout'];children=[node(depth+1) for _ in range(count)]
   lo=np.minimum.reduce([x[0] for x in children]);hi=np.maximum.reduce([x[1] for x in children]);content=h(pack(2,count)+b''.join(x[2] for x in children))
  return lo,hi,nh(lo,hi,content)
 tree=node();assert not f.read(1);f.close()
 assert anchor(n,d,a['fanout'],a['leaf_capacity'],tree[2])==root
 assert len(best_d)==k and np.array_equal(best_id,claimed) and minbound>best_d[-1]
 # Numpy float64 blocked exhaustive search, independently coded and no MR-tree.
 b=fvecs(basepath,n);oracle_d=np.empty(0);oracle_id=np.empty(0,dtype='u4')
 for begin in range(0,n,16384):
  end=min(n,begin+16384);delta=b[begin:end].astype('f8')-q.astype('f8');dd=np.sum(delta*delta,axis=1)
  ds=np.concatenate([oracle_d,dd]);ids=np.concatenate([oracle_id,np.arange(begin,end,dtype='u4')]);order=np.lexsort((ids,ds))[:k];oracle_d=ds[order];oracle_id=ids[order]
 assert np.array_equal(claimed,oracle_id)
 return dict(k=k,query_id=100,proof_bytes=pos,opened_vectors=opened,closed_subtrees=closed,independent_exact_match=True)

def audit():
 assert (STUDY/'RUN_FINISHED.json').exists()
 compatibility=json.loads((STUDY/'ground-truth-compatibility.json').read_text())
 assert compatibility['all_shared_recalls_one'] and len(compatibility['records'])==13
 reports=[];total_samples=0;records={}
 for path in sorted(OUT.glob('*/measurement.json')):
  r=json.loads(path.read_text());folder=path.parent
  if r['dataset'].startswith('smoke') or r['kind']=='pilot':continue
  records[r['dataset'],r['N'],r['kind']]=r
  assert r['native']['tamper_tests']==9*len(r['points'])
  assert sha(ROOT/'Chen13-MR/mr.cpp')==r['source_sha256'] and sha(ROOT/'Chen13-MR/mr')==r['binary_sha256']
  assert sha(folder/'samples.csv')==r['samples_sha256'] and sha(folder/'serving.snapshot')==r['snapshot_sha256']
  with (folder/'samples.csv').open() as f:rows=list(csv.DictReader(f))
  for p in r['points']:
   ss=[x for x in rows if int(x['k'])==p['k']]
   assert len(ss)==60 and {(int(x['query_id']),int(x['repeat'])) for x in ss}=={(q,t) for q in range(100,120) for t in range(3)}
   assert all(x['valid']=='1' and float(x['recall'])==1 for x in ss)
   for metric in ['query_ms','search_ms','encode_ms','verify_ms','vo_bytes','examined','expanded']:
    assert math.isclose(statistics.mean(float(x[metric]) for x in ss),p[metric],rel_tol=1e-13)
   assert math.isclose(p['query_ms'],p['search_ms']+p['encode_ms'],rel_tol=1e-13)
  a=inspect_snapshot(folder/'serving.snapshot',r['native'],Path(r['command'][1]))
  proof=inspect_proof(folder/'example.vo',a,Path(r['command'][2]),Path(r['command'][1]))
  reports.append(dict(dataset=r['dataset'],N=r['N'],kind=r['kind'],snapshot_valid=True,proof=proof,samples=len(rows)))
  total_samples+=len(rows);print('AUDITED',folder.name,flush=True)
 old=openpyxl.load_workbook(STUDY/'work9_before_chen13.xlsx');new=openpyxl.load_workbook(ROOT.parent/'work9.xlsx');preserved=0
 for name in old.sheetnames:
  a=old[name];b=new[name]
  for row in a:
   for x in row:
    if isinstance(x,openpyxl.cell.cell.MergedCell):continue
    y=b.cell(mapped(x.row),x.column) if name in [v[0] for v in DATA.values()] else b[x.coordinate]
    assert x.value==y.value and x._style==y._style,(name,x.coordinate,'value/style')
    assert (x.comment.text if x.comment else None)==(y.comment.text if y.comment else None),(name,x.coordinate,'comment')
    preserved+=1
 # Every expected actual group is present; no NYTimes/GloVe K/size runs.
 expected={(ds,n,'overall') for ds,(_,n) in DATA.items()}
 expected|={('bigann10m',n,'size') for n in SIZES}
 expected|={('deep1m' if n<=1_000_000 else 'deep10m',n,'size') for n in SIZES if n!=1_000_000}
 assert {(x['dataset'],x['N'],x['kind']) for x in reports}==expected
 assert total_samples==1380,(total_samples,'4 overall + 10 extra K points + 9 extra size points = 23*60')
 verified_cells=0
 def check(ws,row,col,expected):
  nonlocal verified_cells
  got=ws.cell(row,col).value
  assert (got is None if expected is None else isinstance(got,(int,float)) and math.isclose(got,expected,rel_tol=1e-13,abs_tol=1e-10)),(ws.title,row,col,got,expected)
  verified_cells+=1
 for ds,(name,n) in DATA.items():
  ws=new[name];r=records[ds,n,'overall'];points={p['k']:p for p in r['points']};p=points[10]
  for col,key,factor in [(5,'recall',100),(6,'query_ms',1),(7,'vo_kb',1),(8,'verify_ms',1)]:check(ws,13,col,p[key]*factor)
  check(ws,13,9,r['native']['index_bytes']/1e6);check(ws,13,10,r['native']['build_seconds'])
  for row,key in [(29,'vo_kb'),(43,'verify_ms'),(57,'query_ms')]:
   check(ws,row,26,100);check(ws,row+1,26,p[key])
   for rr in [row,row+1]:
    for col in range(6,26):check(ws,rr,col,None)
  for row,key in [(108,'vo_kb'),(120,'verify_ms'),(131,'query_ms')]:
   for col,k in enumerate(KS,6):check(ws,row,col,points[k][key] if ds in ['sift1m','deep1m'] else None)
  for row,key in [(72,'vo_kb'),(84,'verify_ms'),(95,'query_ms')]:
   for col,size in enumerate(SIZES,6):
    expected=None
    if ds in ['sift1m','deep1m']:
     source='bigann10m' if ds=='sift1m' else ('deep10m' if size>1_000_000 else 'deep1m')
     rr=r if ds=='deep1m' and size==1_000_000 else records[source,size,'size'];expected=next(p for p in rr['points'] if p['k']==10)[key]
    check(ws,row,col,expected)
 report=dict(passed=True,groups=reports,total_measured_samples=total_samples,negative_cases=207,shared_float32_ground_truth_agreement=True,old_cells_preserved=preserved,new_baseline_cells_verified=verified_cells,workbook_sha256=sha(ROOT.parent/'work9.xlsx'))
 write(STUDY/'independent-audit.json',report);print(json.dumps({k:v for k,v in report.items() if k!='groups'},indent=2))

if __name__=='__main__':
 import sys
 if '--smoke' in sys.argv:
  checks=[]
  for f in sorted(OUT.glob('smoke*/measurement.json')):
   r=json.loads(f.read_text());a=inspect_snapshot(f.parent/'serving.snapshot',r['native'],Path(r['command'][1]))
   checks.append(inspect_proof(f.parent/'example.vo',a,Path(r['command'][2]),Path(r['command'][1])))
  write(STUDY/'smoke-independent-audit.json',dict(passed=True,proofs=checks));print('Independent smoke audits passed:',len(checks))
 else:audit()
