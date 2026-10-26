"""Add Chen13 rows without changing any previous baseline/forecast cell."""
from copy import copy
import json, shutil, math
from pathlib import Path
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import ScatterChart, Reference, Series
from run import ROOT, STUDY, OUT, DATA, KS, SIZES, sha, write

BASE=STUDY/'work9_before_chen13.xlsx'
SOURCE=ROOT.parent/'work9.xlsx'
OUTPUT=STUDY/'work9_with_chen13.xlsx'
INSERTS=[(12,1),(27,2),(39,2),(51,2),(64,1),(75,1),(85,1),(97,1),(108,1),(118,1)]
NAME='Chen13-MR'
def mapped(r):return r+sum(n for after,n in INSERTS if r>after)

def fill():
 assert BASE.exists(),'Run creates the backup before filling'
 if (STUDY/'RUN_FINISHED.json').exists():
  compatibility=json.loads((STUDY/'ground-truth-compatibility.json').read_text())
  assert compatibility['all_shared_recalls_one'],'Resolve reference-precision differences before reporting 100%'
 wb=openpyxl.load_workbook(BASE)
 records={}
 for f in OUT.glob('*/measurement.json'):
  r=json.loads(f.read_text());r['record_path']=str(f)
  if r['dataset'].startswith('smoke') or r['kind']=='pilot':continue
  records[r['dataset'],r['N'],r['kind']]=r
 mappings=[];lineage=[];newcoords={};oldnames=list(wb.sheetnames)
 for ds,(sheet,n) in DATA.items():
  old=wb[sheet];new=wb.create_sheet(sheet+'__chen',wb.index(old))
  for key,value in old.column_dimensions.items():new.column_dimensions[key]=copy(value)
  for key,value in old.row_dimensions.items():
   value=copy(value);value.index=mapped(key);new.row_dimensions[mapped(key)]=value
  for row in old:
   for c in row:
    if isinstance(c,openpyxl.cell.cell.MergedCell):continue
    v=new.cell(mapped(c.row),c.column,c.value);v._style=copy(c._style)
    if c.comment:v.comment=copy(c.comment)
    if c.hyperlink:v.hyperlink=copy(c.hyperlink)
    if c.value is not None or c.has_style or c.comment:mappings.append([sheet,c.coordinate,v.coordinate,c.value])
  for a in old.merged_cells.ranges:new.merge_cells(start_row=mapped(a.min_row),end_row=mapped(a.max_row),start_column=a.min_col,end_column=a.max_col)
  new.sheet_format=copy(old.sheet_format);new.sheet_properties=copy(old.sheet_properties)
  new.sheet_view.showGridLines=old.sheet_view.showGridLines;new.freeze_panes=old.freeze_panes
  new.sheet_state=old.sheet_state
  new.print_options=copy(old.print_options);new.page_margins=copy(old.page_margins);new.page_setup=copy(old.page_setup)
  wb.remove(old);new.title=sheet
  new.column_dimensions['Z'].width=19
  def row(after,count=1):
   first=mapped(after)+1
   for rr in range(first,first+count):
    for col in range(4,27):
     c=new.cell(rr,col);c._style=copy(new.cell(mapped(after),min(col,25))._style)
     c.font=Font(name='Arial',size=10,color='17365D');c.fill=PatternFill('solid',fgColor='DDEBF7')
   if count==2:new.merge_cells(start_row=first,end_row=first+1,start_column=4,end_column=4)
   elif after!=12:new.merge_cells(start_row=first,end_row=first,start_column=4,end_column=5)
   new.cell(first,4,NAME)
   return first
  def put(rr,col,value,record=None,point=None,note=''):
   c=new.cell(rr,col,value);c.number_format='0.000'
   if isinstance(value,(int,float)):
    c.font=Font(name='Arial',size=10,color='006100');c.fill=PatternFill('solid',fgColor='E2F0D9')
   details=note+'\n高维明文 MR-tree 改编；精确 kNN，不含隐私保护；空间批量装载并非作者 R*-tree 源码。'
   if record:details+='\n来源：'+record['record_path']+'\n'+json.dumps(dict(N=record['N'],dataset=record['dataset'],kind=record['kind'],samples=(point or {}).get('samples'),point=point,native=record['native']),ensure_ascii=False)
   c.comment=Comment(details,'Chen13 实验来源')
   lineage.append([sheet,c.coordinate,value,record['record_path'] if record else None,note])
  main=records.get((ds,n,'overall'));points={p['k']:p for p in main['points']} if main else {};p=points.get(10)
  overall=row(12)
  for col,key in [(5,'recall'),(6,'query_ms'),(7,'vo_kb'),(8,'verify_ms')]:put(overall,col,p[key]*(100 if key=='recall' else 1) if p else None,main,p,'四数据集默认 k=10；20条×3次，单线程实测。精确召回100%，不是95%匹配点。' if p else '正式实测未完成，留空。')
  for col,key,factor in [(9,'index_bytes',1e-6),(10,'build_seconds',1)]:put(overall,col,main['native'][key]*factor if main else None,main,p,'一次全新单线程构建；无压缩完整服务快照；构建排除最终快照写盘、GT及输入加载。')
  curves=[]
  for after,header,key,label in [(27,16,'vo_kb','VO size'),(39,28,'verify_ms','verify time'),(51,40,'query_ms','query time')]:
   rr=row(after,2);curves.append(rr);new.cell(rr,5,'Recall rate');new.cell(rr+1,5,label)
   c=new.cell(mapped(header),26,100);c._style=copy(new.cell(mapped(header),25)._style)
   c.comment=Comment('新增精确方法100%参考点；其他已有方法及原15–99%列保持原值。','Chen13')
   for col in range(6,26):
    for x in [rr,rr+1]:put(x,col,None,note='不适用：精确算法无可调低召回档位。仅Z列为100%实测参考点，不能作为等召回95%比较。')
   put(rr,26,p['recall']*100 if p else None,main,p,'精确方法唯一召回点。')
   put(rr+1,26,p[key] if p else None,main,p,'沿用默认 k=10 的同一组60次真实测量；非重复实验。')
  srows=[];krows=[]
  for after,key in [(64,'vo_kb'),(75,'verify_ms'),(85,'query_ms')]:
   rr=row(after);srows.append(rr)
   for col,size in enumerate(SIZES,6):
    if ds not in ['sift1m','deep1m']:put(rr,col,None,note='按用户要求，此数据集不做 Dataset size 实验。');continue
    source='bigann10m' if ds=='sift1m' else ('deep10m' if size>1_000_000 else 'deep1m')
    r=main if ds=='deep1m' and size==1_000_000 else records.get((source,size,'size'))
    sp=next((a for a in r['points'] if a['k']==10),None) if r else None
    put(rr,col,sp[key] if sp else None,r,sp,'20条×3次、单线程、每个N重新精确GT；SIFT使用BIGANN前缀。' if sp else '正式实测未完成，留空。')
  for after,key in [(97,'vo_kb'),(108,'verify_ms'),(118,'query_ms')]:
   rr=row(after);krows.append(rr)
   for col,k in enumerate(KS,6):
    kp=points.get(k) if ds in ['sift1m','deep1m'] else None
    put(rr,col,kp[key] if kp else None,main if kp else None,kp,'每个k独立20条×3次实测；精确Recall@k=100%。' if kp else ('按用户要求，此数据集不做K实验。' if ds not in ['sift1m','deep1m'] else '正式实测未完成，留空。'))
  abl=mapped(old.max_row)+1
  for col in range(4,10):new.cell(abl,col)._style=copy(new.cell(abl-1,col)._style)
  new.merge_cells(start_row=abl,end_row=abl,start_column=4,end_column=5);new.cell(abl,4,NAME+'（外部参考）')
  for col,key in [(6,'recall'),(7,'query_ms'),(8,'vo_kb'),(9,'verify_ms')]:put(abl,col,p[key] if p else None,main,p,'外部baseline参考，不是Ours的消融变体。Recall以0–1计。')
  newcoords[sheet]=dict(overall=overall,curves=curves,size=srows,k=krows,ablation=abl)

 tab=wb.create_sheet('Chen13实验说明')
 notes=[['项目','说明'],['论文','Chen, Hu, Xu. Authenticating Top-k Queries in Location-based Services with Confidentiality. PVLDB 7(1), 49–60, 2013.'],
 ['实现范围','第5.1.1节与图6(a)明文MR-tree：精确best-first kNN、剩余frontier MBB、Merkle重建与签名根；多维原坐标、非空间权重置零。'],
 ['改编与限制','不是作者完整隐私MR/EMR/PD复现：空间批量装载代替R*-tree；fanout16、leaf64；SHA256与Ed25519；无PPB/PLB、无Power Diagram、无投影/PQ。对高维效果只解释为这一明文改编版本。'],
 ['默认值','仅新增Chen13默认精确100%；原VLAH/ANNProof/Ours约95%及V3DB低召回默认、所有已实测/预测值均保持。'],
 ['召回曲线','精确方法无低召回参数；新增Z列100%唯一参考点。新行原15–99%栏留空表示不适用，不能把该参考点说成等召回对比。'],
 ['实验协议','20条查询IDs100..119×3次，单线程；四数据集默认k10；K和规模只测SIFT/DEEP；每个规模独立全扫描GT、每次真实VO验证通过。'],
 ['实测机器','当前WSL约62GiB可用总内存，进程地址空间上限50GiB；不是512GiB机器实测。历史V3DB的512GiB预测条件与标记保持原样。环境记录见build-environment.json。'],
 ['K','1/5/10/20/50/100；SIFT1M与DEEP1M。'],['规模','100K/500K/1M/5M/10M；SIFT使用原表BIGANN前缀，与单独SIFT1M默认不同；DEEP同源前缀。'],
 ['单位','查询/验证ms；VO十进制KB，含根/签名/查询绑定/结果ID；索引十进制MB；构建秒。查询=搜索+VO序列化；验证含解码、签名、Merkle、排名与完整性。'],
 ['索引','完整无压缩服务快照：原向量、ID、节点拓扑/MBB/哈希、元数据/签名/公钥；不重复算原文件、查询VO。构建一次，排除输入/GT和最终快照I/O。'],
 ['正确性','独立全扫描逐查询逐k对齐ID与距离次序；同距按ID。2/100/128/256维测试含重复点/零向量；每k执行9项真实序列化篡改拒绝测试。验证器没有索引或原数据访问。'],
 ['共享GT口径','全部23个参数点、每点20条查询另行核对原项目float32逐维L2+ID排序，与本实现float64精确答案完全一致。故表中100%同时满足原项目GT口径。详见ground-truth-compatibility.json。'],
 ['消融表','新增行仅外部baseline参考，不是Ours消融。'],['实测/未测','只写measurement.json中的完成组；未完成留空，不以预测冒充实测。NYTimes/GloVe的K/规模按用户要求不测。'],
 ['源码',str(ROOT/'Chen13-MR/mr.cpp')],['备份',str(BASE)],['原有来源坐标','历史来源附表记录的是当时坐标，保持原始档案；本次新增Chen13位置映射给出插行前→当前坐标。'],
 ['状态','全部完成' if (STUDY/'RUN_FINISHED.json').exists() else '正在实测；仅完成组填入'],['新增行',json.dumps(newcoords,ensure_ascii=False)]]
 for r in notes:tab.append(r)
 tab.column_dimensions['A'].width=23;tab.column_dimensions['B'].width=145
 tab=wb.create_sheet('Chen13实测记录');tab.append(['Dataset','N','Kind','k','Recall','Samples','Query ms','Verify ms','VO KB','Index MB','Build s','Snapshot I/O s','Process peak RSS GiB','Mean examined vectors','Source'])
 for r in records.values():
  for p in r['points']:tab.append([r['dataset'],r['N'],r['kind'],p['k'],p['recall'],p['samples'],p['query_ms'],p['verify_ms'],p['vo_kb'],r['native']['index_bytes']/1e6,r['native']['build_seconds'],r['native']['snapshot_write_seconds'],r['native']['peak_rss_KiB']/2**20,p['examined'],r['record_path']])
 tab=wb.create_sheet('Chen13位置映射');tab.append(['Sheet','Before cell','Current cell','Original value'])
 for r in mappings:tab.append(r)
 tab=wb.create_sheet('Chen13单元格来源');tab.append(['Sheet','Cell','Value','Record','Note'])
 for r in lineage:tab.append(r)
 tab=wb.create_sheet('Chen13召回参考');tab.append(['Dataset','Recall %','VO KB','Verify ms','Query ms','Note'])
 for ds,(sheet,n) in DATA.items():
  r=records.get((ds,n,'overall'))
  if r:
   p=next(p for p in r['points'] if p['k']==10);tab.append([sheet,100,p['vo_kb'],p['verify_ms'],p['query_ms'],'精确唯一点，不能冒充较低召回参数曲线'])
 for i,(col,title) in enumerate([(3,'VO size (KB)'),(4,'Verification (ms)'),(5,'Query (ms)')]):
  chart=ScatterChart();chart.title=title+' — exact reference';chart.x_axis.title='Actual Recall@10 (%)';chart.y_axis.title=title
  for rr in range(2,tab.max_row+1):
   series=Series(Reference(tab,min_col=col,min_row=rr,max_row=rr),Reference(tab,min_col=2,min_row=rr,max_row=rr),title=tab.cell(rr,1).value);series.marker.symbol='circle';series.graphicalProperties.line.noFill=True;chart.series.append(series)
  tab.add_chart(chart,f'H{1+16*i}')
 if '索引构建实测记录' in wb:
  tab=wb['索引构建实测记录']
  for ds,(sheet,n) in DATA.items():
   r=records.get((ds,n,'overall'))
   if not r:continue
   a=r['native'];auth=a['nodes']*64+96;search=a['index_bytes']-auth
   tab.append([sheet,NAME,n,a['D'],1,1,search,auth,a['index_bytes']/1e6,None,None,a['build_seconds'],a['peak_rss_KiB']/2**20,'实测；RSS含查询/VO；搜索与认证构建未分拆',16,64,r['record_path']])
   for c in tab[tab.max_row]:c.comment=Comment('Chen13：总构建实测，子阶段未单独计时；RSS含查询与篡改测试，非仅构建。search bytes含向量/ID/拓扑/MBB；auth bytes含双哈希/签名公钥。','Chen13')

 # Audit before writing: existing main cells/styles/comments must survive mapping.
 original=openpyxl.load_workbook(BASE);checked=0
 for ds,(sheet,_) in DATA.items():
  a=original[sheet];b=wb[sheet]
  for row in a:
   for x in row:
    if isinstance(x,openpyxl.cell.cell.MergedCell):continue
    y=b.cell(mapped(x.row),x.column)
    assert x.value==y.value,(sheet,x.coordinate,'value')
    assert x._style==y._style,(sheet,x.coordinate,'style')
    assert (x.comment.text if x.comment else None)==(y.comment.text if y.comment else None),(sheet,x.coordinate,'comment')
    checked+=1
 for name in oldnames[4:]:
  a=original[name];b=wb[name]
  for row in a:
   for x in row:
    if isinstance(x,openpyxl.cell.cell.MergedCell):continue
    y=b[x.coordinate];assert x.value==y.value and x._style==y._style,(name,x.coordinate)
    assert (x.comment.text if x.comment else None)==(y.comment.text if y.comment else None),(name,x.coordinate,'comment')
 wb.save(OUTPUT)
 check=openpyxl.load_workbook(OUTPUT)
 for row in lineage:
  got=check[row[0]][row[1]].value;expected=row[2]
  assert (math.isclose(got,expected,rel_tol=1e-14,abs_tol=1e-12) if isinstance(expected,(int,float)) else got==expected),(row[0],row[1],got,expected)
 marker=STUDY/'last-written.txt';allowed={sha(BASE)}
 if marker.exists():allowed.add(marker.read_text().strip())
 assert sha(SOURCE) in allowed,'work9 changed externally; output retained, do not overwrite user edits'
 shutil.copy2(OUTPUT,SOURCE);marker.write_text(sha(SOURCE))
 write(STUDY/'fill-audit.json',dict(original_main_cells_preserved=checked,new_metric_and_status_cells=len(lineage),measurement_groups=len(records),new_rows=newcoords,output_sha256=sha(SOURCE),finished=(STUDY/'RUN_FINISHED.json').exists()))
 print(json.dumps(json.loads((STUDY/'fill-audit.json').read_text()),ensure_ascii=False,indent=2))

if __name__=='__main__':fill()
