"""Actual million-vector ArXiv overall/recall curves/Table4; isolated checkpoints."""
import argparse,csv,hashlib,json,math,os,shutil,statistics as st,subprocess,sys,time
from collections import defaultdict
from copy import copy
from pathlib import Path
import numpy as np
import work9_resource_matched as u
R=u.ROOT;S=R/'experiments/openai-arxiv-20261024';O=R/'results/openai-arxiv-20261024';F=R/'dataset/openai_arxiv1m';KEY='openai_arxiv1m';N=1000000;D=1536;M=16;EFC=200;BUILD_THREADS=8
TARGETS=[.15,.20,.25,.30,.35,.40,.45,.50,.55,.60,.65,.70,.75,.80,.85,.90,.95,.97,.98,.99,1.]
GRID=[1,2,3,4,6,8,10,12,16,24,32,48,64,96,128,192,256,384,512,768,1024,1536,2048,4096,8192]
NAMES={1:'IVF-ADS',2:'NSW-ADS',3:'IVF-ADS+(i)',4:'NSW-ADS+(i)',5:'(i)+(ii)',6:'QVANN'}
S.mkdir(parents=True,exist_ok=True);O.mkdir(parents=True,exist_ok=True);u.STUDY,u.RESULTS=S,O
def event(**x):u.event(**x)
def rows(path):return list(csv.DictReader(Path(path).open()))
def writecsv(path,rr):
    if not rr:return
    with Path(path).open('w',newline='') as f:w=csv.DictWriter(f,fieldnames=list(rr[0]));w.writeheader();w.writerows(rr)
def execute(cmd,out,marker='PASS.txt',timeout=86400):
    out=Path(out)
    if (out/marker).exists():return
    if out.exists() and any(out.iterdir()):
        archived=out.with_name(out.name+'-incomplete-'+str(time.time_ns()));out.rename(archived)
    out.mkdir(parents=True,exist_ok=True)
    meta=dict(command=list(map(str,cmd)),binary_sha256=u.sha(cmd[0]),threads=1,queries=20,repeats=3,query_offset=100)
    u.write(out/'environment.json',meta)
    u.command(['/usr/bin/time','-v','-o',out/'resources.txt',*cmd],S/'logs'/(out.name+'.log'),bounded=True,timeout=timeout)
    (out/marker).write_text('Native invocation completed; additional row/proof checks are performed before aggregation.\n')
def common(out,n=N):return ['--dataset',F,'--base',n,'--M',M,'--ef-construction',EFC,'--build-threads',BUILD_THREADS,'--queries',20,'--query-offset',100,'--repeats',3,'--k',10,'--output',out]
def checkpoint():
    filled=[]
    for file in sorted((S/'records').glob('*.json')) if (S/'records').exists() else []:filled.append(u.read(file))
    u.write(S/'progress.json',dict(time=time.strftime('%FT%T%z'),records=filled))
def save(name,rec):u.write(S/'records'/(name+'.json'),rec);checkpoint()
def aggregate(out):
    rr=rows(out/'per_query.csv');assert all(r['verified']=='1' for r in rr)
    attacks=rows(out/'tamper_tests.csv');assert attacks and all(r['rejected']=='1' for r in attacks)
    assert (out/'query0.vo').stat().st_size==int(rr[0]['vo_bytes'])
    groups=defaultdict(list)
    for r in rr:groups[int(r['ef_search']),int(r['probes'])].append(r)
    result=[]
    for (e,p),rs in groups.items():
        assert len(rs)==60 and {(int(r['query']),int(r['repeat'])) for r in rs}=={(q,t) for q in range(100,120) for t in range(3)}
        mean=lambda c:st.fmean(float(r[c]) for r in rs)
        result.append(dict(status='measured',ef=e,probes=p,recall=mean('recall_at_k'),query_ms=mean('search_ms')+mean('vo_generation_ms'),vo_kb=mean('vo_bytes')/1000,verify_ms=mean('verification_ms'),samples=60,source=str(out/'per_query.csv'),source_sha256=u.sha(out/'per_query.csv')))
    return result
def width():
    file=S/'lsh-width.json'
    if file.exists():return u.read(file)['width']
    b=np.memmap(F/'sift_base.fvecs',dtype='<f4',shape=(N,D+1))[:,1:];ids=np.random.default_rng(20261024).integers(0,N,(4096,2));diff=b[ids[:,0]].astype(np.float64)-b[ids[:,1]].astype(np.float64);w=float(np.median(np.sqrt((diff*diff).sum(axis=1)))/4)
    u.write(file,dict(width=w,seed=20261024,pairs=4096,rule='base-only median pairwise L2/4'));return w
def pilot(method,grid,probes,tag):
    out=O/'pilots'/(method+'-'+tag)
    binary=R/'build'/('multidata_pilot_'+method)
    cache=R/('VLAH-25TrustCom' if method=='VLAH' else 'ANNProof-24FGCS')/'cache'
    cmd=[binary,*common(out),'--cache',cache,'--ef-search',','.join(map(str,grid)),'--probes',','.join(map(str,probes)),'--lsh-width',width()]
    execute(cmd,out,'PILOT_PASS.txt');points=[];groups=defaultdict(list)
    for r in rows(out/'pilot.csv'):groups[int(r['ef_search']),int(r['probes'])].append(r)
    for (ef,p),rr in groups.items():
        assert len(rr)==20
        points.append(dict(ef=ef,probes=p,recall=st.fmean(float(r['recall_at_k']) for r in rr),valid=all(r['valid']=='1' for r in rr),source=str(out/'pilot.csv')))
    return points
def select(points,target,method):
    good=[p for p in points if p['valid'] and p['recall']>=target-1e-12]
    return min(good,key=lambda p:(p['probes'],p['ef']) if method=='VLAH' else (p['ef'],p['probes'])) if good else None
def baselines():
    for method in ['Ours','ANNProof','VLAH']:
        record=S/'records'/(method+'-curve.json')
        if record.exists():continue
        points=[]
        if method=='VLAH':
            for ef in [50,100,200,400,800,1600]:
                points+=pilot(method,[ef],[1,2,4,6,8,12,16,24,32,40,48,50],f'e{ef}')
                if max(p['recall'] for p in points if p['valid'])>=.99-1e-12:break
        else:
            for phase,grid in enumerate([GRID if method=='Ours' else [x for x in GRID if x>=10],[16384,32768]]):
                points+=pilot(method,grid,[1],f'phase{phase}')
                if max(p['recall'] for p in points if p['valid'])>=.99-1e-12:break
        u.write(S/(method+'-calibration.json'),points)
        selected={t:select(points,t,method) for t in TARGETS};assert selected[.95],(method,'95% recall unreachable')
        settings=sorted({(p['ef'],p['probes']) for p in selected.values() if p});measured={}
        batches=defaultdict(list)
        if method=='VLAH':
            for ef,p in settings:batches[ef].append(p)
        else:batches['all']=[e for e,p in settings]
        for key,values in batches.items():
            out=O/'measurements'/(method+'-curve-'+str(key))
            if method=='Ours':binary=R/'build/veriann_lut_benchmark';cache=R/'ANNProof-24FGCS/cache'
            else:
                folder=R/('ANNProof-24FGCS' if method=='ANNProof' else 'VLAH-25TrustCom');binary=folder/'build'/('annproof' if method=='ANNProof' else 'vlah');cache=folder/'cache'
            efs=values if key=='all' else [key];probes=[1] if key=='all' else values
            execute([binary,*common(out),'--cache',cache,'--ef-search',','.join(map(str,efs)),'--probes',','.join(map(str,probes)),'--lsh-width',width()],out)
            for r in aggregate(out):measured[r['ef'],r['probes']]=r
        curve=[]
        for t,p in selected.items():
            item=dict(measured[p['ef'],p['probes']],target=t) if p else dict(status='unreachable_in_grid',target=t,reason='No calibrated configuration reached target; not a prediction.')
            if p:assert abs(item['recall']-p['recall'])<1e-10
            curve.append(item)
        save(method+'-curve',dict(method=method,kind='curve',points=curve))
        save(method+'-overall',dict(next(x for x in curve if x['target']==.95),method=method,kind='overall'))
def ablations():
    chosen={}
    for v in [1,2]:
        record=S/f'ablation-{v}-calibration.json'
        if record.exists():chosen[v]=u.read(record);continue
        grid=[1,2,4,8,16,24,32,48,64,96,128,192,256] if v==1 else [10,16,24,32,48,64,96,128,192,256,384,512,768,1024,1536,2048,4096,8192]
        out=O/'pilots'/f'ablation-{v}'
        execute([R/'build/work9_ablation',*common(out),'--cache',R/'ANNProof-24FGCS/cache','--variants',v,'--nlist',256,'--pilot',1,'--ef-search',','.join(map(str,grid))],out,'PILOT_PASS.txt')
        groups=defaultdict(list)
        for r in rows(out/'pilot.csv'):groups[int(r['parameter'])].append(r)
        points=[dict(parameter=p,recall=st.fmean(float(r['recall_at_k']) for r in rr)) for p,rr in groups.items()]
        good=[p for p in points if p['recall']>=.95-1e-12];assert good,(v,'ablation recall unreachable')
        chosen[v]=min(good,key=lambda p:p['parameter']);u.write(record,dict(chosen[v],points=points,source=str(out/'pilot.csv')))
    ours=u.read(S/'records/Ours-overall.json');chosen[5]=dict(parameter=ours['ef'],recall=ours['recall'])
    for pair,key in [([1,3],1),([2,4],2),([5,6],5)]:
        out=O/'ablations'/('-'.join(map(str,pair)));p=chosen[key]['parameter']
        if all((S/'records'/f'ablation-{v}.json').exists() for v in pair):continue
        execute([R/'build/work9_ablation',*common(out),'--cache',R/'ANNProof-24FGCS/cache','--variants',','.join(map(str,pair)),'--nlist',256,'--pilot',0,'--ef-search',p],out)
        answers=[]
        for v in pair:
            record=aggregate(out/f'variant-{v}');assert len(record)==1;rec=record[0];assert abs(rec['recall']-chosen[key]['recall'])<1e-10
            answers.append(rows(out/f'variant-{v}'/'answers.csv'))
            save(f'ablation-{v}',dict(rec,variant=v,method=NAMES[v],kind='ablation',parameter=p))
        assert answers[0]==answers[1],('paired ablation answers differ',pair)
def mr():
    if (S/'records/MR-tree-overall.json').exists():return
    out=O/'measurements/MR-tree'
    execute([R/'Chen13-MR/mr',F/'sift_base.fvecs',F/'sift_query.fvecs',N,out,'10',20,3,64],out)
    rr=rows(out/'samples.csv');assert len(rr)==60 and all(r['valid']=='1' and float(r['recall'])==1 for r in rr)
    native=u.read(out/'native.json');assert native['tamper_tests']>=9
    mean=lambda c:st.fmean(float(r[c]) for r in rr)
    save('MR-tree-overall',dict(method='MR-tree',kind='overall',status='measured',recall=1.,query_ms=mean('query_ms'),vo_kb=mean('vo_bytes')/1000,verify_ms=mean('verify_ms'),samples=60,index_mb=native['index_bytes']/1e6,index_seconds=native['build_seconds'],source=str(out/'samples.csv'),source_sha256=u.sha(out/'samples.csv')))
def v3db():
    sys.path.insert(0,str(R/'V3DB-26VLDB/scripts'));import work9_calibrate as c
    c.DATA[KEY]=('OpenAI-ArXiv-1M',N,D,1024,8);c.STUDY=S/'V3DB';c.STUDY.mkdir(exist_ok=True)
    if (S/'records/V3DB-curve.json').exists():return
    calibration=c.calibrate(KEY,pq_m=128);model=f'{KEY}-{N}-M128';cfg=calibration['model']['config'];points=[p for p in calibration['points'] if p['k']==10]
    # Retain author's IVF-PQ; low recall is reported honestly, never exact-rerank.
    candidates=[p for p in points if p['n_probe']<=32];measured=[];stops=[]
    gate_points=[]
    for p in candidates:
        gate=u.static_gate(cfg,p['n_probe'],10);gate_points.append(dict(point=p,gate=gate))
        if gate['status']!='passed':stops.append(dict(probe=p['n_probe'],status='resource_limit',reason=gate));continue
        out=O/'V3DB'/f'{model}-p{p["n_probe"]}-k10-q1-r1-t1'
        env=u.environment();env['V3DB_RESULTS_DIR']=str(O/'V3DB')
        if not (out/'per_query.csv').exists():
            cmd=[sys.executable,R/'V3DB-26VLDB/scripts/work9_measure.py','--model',model,'--probe',p['n_probe'],'--k',10,'--queries',1,'--repeats',1,'--attempt-large']
            log=S/'logs'/f'V3DB-p{p["n_probe"]}.log'
            with log.open('w') as f:result=subprocess.run(list(map(str,cmd)),cwd=R,env=env,stdout=f,stderr=subprocess.STDOUT,preexec_fn=u.limit)
            if result.returncode:stops.append(dict(probe=p['n_probe'],status='failed',log=str(log)));continue
        rr=rows(out/'per_query.csv');assert len(rr)==1 and rr[0]['verified']=='1';r=rr[0]
        measured.append(dict(method='V3DB',status='measured',ef=p['n_probe'],probes=p['n_probe'],recall=p['recall'],query_ms=float(r['query_ms']),vo_kb=float(r['vo_bytes'])/1000,verify_ms=float(r['verify_ms']),samples=1,source=str(out/'per_query.csv'),source_sha256=u.sha(out/'per_query.csv'),recall_scope='20 search-only queries; query100 one real proof/timing, not average'))
        # At most three fixed single-query proof points: cost grows rapidly.
        if len(measured)>=3:break
    u.write(S/'V3DB-gates.json',gate_points)
    curve=[]
    for t in TARGETS:
        good=[p for p in measured if p['recall']>=t-1e-12]
        curve.append(dict(min(good,key=lambda p:p['probes']),target=t) if good else dict(status='unmeasured_or_resource_limit',target=t,reason='No verified actual proof at this recall; see V3DB-gates.json and calibration.'))
    save('V3DB-curve',dict(method='V3DB',kind='curve',points=curve,stops=stops))
    default=next((p for p in measured if p['recall']>=.4),measured[-1] if measured else dict(method='V3DB',status='resource_limit',reason=stops))
    save('V3DB-overall',dict(default,kind='overall',method='V3DB'))
def indexes():
    # Separate actual fresh builds, up to8 offline threads; never treat cache-load time
    # from the query benchmark as index construction time.
    import work9_index_study as ix
    for method in ['ANNProof','Ours','VLAH']:
        file=S/'records'/(method+'-index.json')
        if file.exists():continue
        directory=S/'index-build';directory.mkdir(exist_ok=True);binary=directory/method;out=O/'indexes'/method
        if not binary.exists():
            original=R/({'ANNProof':'ANNProof-24FGCS/annproof.cpp','VLAH':'VLAH-25TrustCom/vlah.cpp','Ours':'build/veriann_lut_benchmark.cpp'}[method])
            prefix=original.read_text().split('int main(int argc,char**argv)try{')[0] if method=='Ours' else '#define main original_main_unused\n#include "'+str(original)+'"\n#undef main\n'
            main={'ANNProof':ix.ANN,'VLAH':ix.VLAH,'Ours':ix.OURS}[method].replace('build_threads=1','build_threads=8').replace('omp_set_num_threads(1)','omp_set_num_threads(8)')
            source=directory/(method+'.cpp');source.write_text(prefix+ix.COMMON.replace('\\"threads\\":1','\\"threads\\":8')+main)
            cmd=u.read(R/'experiments/work9-quantization-20261023/build.json')['command']
            cmd=[str(source) if str(x).endswith('/scripts/quantization_refinement_probe.cpp') else x for x in cmd];cmd[cmd.index('-o')+1]=str(binary)
            u.command(cmd,S/'logs'/(method+'-index-compile.log'),timeout=300)
            u.write(directory/(method+'-build.json'),dict(command=cmd,original_sha256=u.sha(original),harness_sha256=u.sha(source),binary_sha256=u.sha(binary),threads=8))
        cmd=[binary,*common(out),'--cache',out/'cache','--lsh-width',width()]
        execute(cmd,out)
        rec=u.read(out/'offline.json');assert rec['total_bytes']==(out/'serving.snapshot').stat().st_size
        save(method+'-index',dict(method=method,kind='index',status='measured',index_mb=rec['total_bytes']/1e6,index_seconds=rec['construct_seconds'],threads=8,source=str(out/'offline.json'),source_sha256=u.sha(out/'offline.json')))
def initialize():
    assert (F/'provenance.json').exists(),'Run prepare_openai_arxiv_20261024.py first'
    for binary in ['build/veriann_lut_benchmark','build/multidata_pilot_Ours','build/multidata_pilot_ANNProof','build/multidata_pilot_VLAH','build/work9_ablation','ANNProof-24FGCS/build/annproof','VLAH-25TrustCom/build/vlah','Chen13-MR/mr']:
        assert (R/binary).exists(),binary
    u.write(S/'protocol.json',dict(dataset=KEY,N=N,D=D,M=M,efConstruction=EFC,k=10,query_ids=list(range(100,120)),repeats=3,query_threads=1,index_build_threads=BUILD_THREADS,index_note='New ArXiv offline graph construction uses up to8 threads; MR-tree construction serial. These new construction times must not be directly pooled with historical single-thread build timings. Query and verification remain single-threaded. V3DB original training helper uses16 coarse-training threads and8 one-thread PQ workers; report separately.',targets=TARGETS,distance='squared L2 on unchanged float32 source embeddings',GT='Exact exhaustive oracle recomputed independently per implementation for this1M prefix. Calibration and measurement use same20 query IDs; not held-out accuracy claims.',defaults='Non-V3DB minimal calibrated operating point reaching>=95%; MR-tree exact100%. V3DB original integer IVF-PQ, M128,K256,nlist1024, low-recall feasible probes, one fixed query100 proof per feasible point, separate20-query search recall.',V3DB='At most3 actual fixed single-query proof points; 50GiB process address-space gate. No high-dimensional extrapolation or 512GiB forecasts.',table4=NAMES,raw='All native CSVs preserved. No monotonic smoothing or fabricated values. ANNProof prior coefficients, if displayed, are explicitly labeled derived and accompanied by raw sheets/CSV.',paper=str(next((R.parent/'papers').glob('QVANN*6*.pdf'))),provenance=u.read(F/'provenance.json')))
def report():
    complete_v3db_index()
    records={p.stem:u.read(p) for p in (S/'records').glob('*.json')};overall=[];curves=[];ab=[]
    def parameters(point):
        if point.get('method')=='V3DB' and point.get('status')=='measured':
            point=dict(point);point.setdefault('model',f'{KEY}-{N}-M128')
            metadata=R/'V3DB-26VLDB/work9-cache'/point['model']/'model.json'
            if metadata.exists():point.setdefault('config',u.read(metadata)['config'])
        return point
    for method in ['ANNProof','VLAH','V3DB','MR-tree','Ours']:
        r=records.get(method+'-overall',dict(method=method,status='pending'));r=parameters(dict(r))
        if method+'-index' in records:
            r.update({k:records[method+'-index'][k] for k in ['index_mb','index_seconds']});r['index_provenance']=records[method+'-index']
        overall.append(r)
        if method+'-curve' in records:curves.extend(parameters(dict(p,method=method)) for p in records[method+'-curve']['points'])
        elif method=='MR-tree' and r.get('status')=='measured':curves.append(dict(r,target=1.,ef=0,probes=0,parameter_policy='Exact search; only the actual 100% endpoint exists.'))
    ab=[records.get('ablation-'+str(v),dict(method=NAMES[v],variant=v,status='pending')) for v in [1,3,2,4,5,6]]
    # JSON preserves heterogeneous status/provenance; CSV uniform metrics.
    for name,data in [('overall',overall),('recall_curves',curves),('table4',ab)]:
        u.write(S/(name+'.json'),data);writecsv(S/(name+'.csv'),[dict({k:r.get(k,'') for k in ['method','status','recall','query_ms','vo_kb','verify_ms','ef','probes','target','index_mb','index_seconds','samples','source','source_sha256']},dataset=KEY,N=N,D=D,k=10,graph_M=M if r['method'] not in ['V3DB','MR-tree','IVF-ADS','IVF-ADS+(i)'] else '',graph_efConstruction=EFC if r['method'] not in ['V3DB','MR-tree','IVF-ADS','IVF-ADS+(i)'] else '',query_threads=1,model=r.get('model',''),PQ_M=r.get('config',{}).get('M',''),PQ_K=r.get('config',{}).get('K',''),PQ_nlist=r.get('config',{}).get('n_list','')) for r in data])
    def latex(name,data,caption):
        note=r'ANNProof$^{*}$ uses derived display factors $Q\times1.2$, $\mathrm{VO}\times2$, $V\times1.5$; other costs are raw measurements.' if name.endswith('adjusted') else 'Costs are raw measurements.'
        if name.startswith('overall'):note+=' V3DB uses one fixed-query proof and20-query search recall; dashes denote unavailable measurements.'
        lines=[r'\begin{table}[t]',r'\centering',r'\caption{'+caption+r' OpenAI-ArXiv-1M ($d=1536$, $k=10$). '+note+'}',r'\label{tab:arxiv_'+name+'}',r'\footnotesize',r'\setlength{\tabcolsep}{3pt}',r'\resizebox{\columnwidth}{!}{%',r'\begin{tabular}{@{}lrrrr@{}}',r'\toprule',r'Method & $R$ (\%) & $Q$ (ms) & VO (KB) & $V$ (ms) \\',r'\midrule']
        for r in data:
            values=[f'{r[k]*(100 if k=="recall" else 1):.2f}' if r.get('status') in ['measured','derived_display'] and k in r else '--' for k in ['recall','query_ms','vo_kb','verify_ms']]
            label='QVANN' if r['method']=='Ours' else r['method'];lines.append(' & '.join([label]+values)+r' \\')
        lines.extend([r'\bottomrule',r'\end{tabular}%',r'}',r'\end{table}','']);(S/(name+'.tex')).write_text('\n'.join(lines))
    latex('overall',overall,'Overall performance.');latex('table4',ab,'Ablation study corresponding to Table4.')
    adjusted=[]
    for r in overall:
        display=dict(r)
        if r['method']=='ANNProof' and r.get('status')=='measured':
            display.update(method=r'ANNProof$^{*}$',status='derived_display')
            for key,factor in [('query_ms',1.2),('vo_kb',2),('verify_ms',1.5)]:display[key]=r[key]*factor
        adjusted.append(display)
    latex('overall_adjusted',adjusted,'Comparison using the historical display convention.');u.write(S/'overall_adjusted.json',adjusted)
    fill_workbook(overall,curves,ab)
def complete_v3db_index():
    default=S/'records/V3DB-overall.json';selected=u.read(default) if default.exists() else {};model_name=selected.get('model',f'{KEY}-{N}-M128')
    record=S/'records/V3DB-index.json';folder=R/'V3DB-26VLDB/work9-cache'/model_name
    if record.exists() or not (folder/'model.json').exists() or not (folder/'snapshot.json').exists():return
    model=u.read(folder/'model.json');snapshot=u.read(folder/'snapshot.json');assert 'coarse_reused' not in model,'Reused training is not a full index build measurement'
    out=O/'indexes/V3DB';out.mkdir(parents=True,exist_ok=True)
    cfg=model['config'];groupsfile=np.load(folder/'groups.npz');groups=[groupsfile[str(i)] for i in range(cfg['n_list'])]
    arrays={name:np.load(folder/(name+'.npy'),mmap_mode='r') for name in ['centers','books','codes']}
    # Identical serving dtypes/rounding to work9_measure.load_model, not the
    # floating training centroids stored before integer witness conversion.
    arrays['centers']=np.asarray(arrays['centers'],dtype=np.int64);arrays['books']=np.rint(arrays['books']).astype(np.int64)
    with (out/'search.snapshot').open('wb') as f:
        np.savez(f,**arrays,group_offsets=np.cumsum([0]+[len(g) for g in groups],dtype=np.int64),group_ids=np.concatenate(groups).astype(np.int64),parameters=np.array([N,D,cfg['n_list'],cfg['M'],cfg['K'],cfg['capacity']],dtype=np.int64),scale=np.array([model['vmin'],model['vmax']],dtype=np.float64))
    with (out/'auth.snapshot').open('wb') as f:np.savez(f,list_roots=np.asarray(snapshot['list_roots'],dtype=np.uint64),roots=np.asarray([snapshot['root'],snapshot['codebook_root']],dtype=np.uint64))
    total=sum((out/name).stat().st_size for name in ['search.snapshot','auth.snapshot'])
    save('V3DB-index',dict(method='V3DB',model=model_name,kind='index',status='measured',index_mb=total/1e6,index_seconds=model['train_seconds']+snapshot['snapshot_build_seconds'],training_seconds=model['train_seconds'],authentication_build_seconds=snapshot['snapshot_build_seconds'],threads='coarse16; PQ8 workers x1; snapshot/proofs1',source=str(folder/'model.json'),source_sha256=u.sha(folder/'model.json'),snapshot_source=str(folder/'snapshot.json'),snapshot_source_sha256=u.sha(folder/'snapshot.json'),serving_files={name:dict(bytes=(out/name).stat().st_size,sha256=u.sha(out/name)) for name in ['search.snapshot','auth.snapshot']},scope='Actual fresh author-compatible model training and Merkle snapshot construction, excludes query circuits/proofs. Offline thread counts differ from ADS8-thread builds; report explicitly.'))
def fill_workbook(overall,curves,ab):
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Font,PatternFill
    from openpyxl.chart import ScatterChart,Series,Reference
    source=R.parent/'work9.xlsx';backup=S/'work9_before_openai_arxiv.xlsx'
    if not backup.exists():shutil.copy2(source,backup)
    wb=openpyxl.load_workbook(source);sheet='OpenAI-ArXiv-1M'
    if sheet in wb:del wb[sheet]
    ws=wb.create_sheet(sheet);template=wb['SIFT1M']
    for row in template.iter_rows(min_row=1,max_row=58,max_col=26):
        for c in row:
            cell=ws.cell(c.row,c.column);cell._style=copy(c._style)
            if isinstance(c.value,str) or c.row in [17,31,45]:cell.value=c.value
    for k,dim in template.column_dimensions.items():ws.column_dimensions[k].width=dim.width
    for row in [17,31,45]:
        ws.cell(row,5,'Target Recall (%) - actual recall is reported below')
        for col,target in enumerate(TARGETS,6):ws.cell(row,col,target*100).number_format='0.00'
    for row in [7,20,34,48]:ws.cell(row,4,sheet)
    ws['A1']='OpenAI-ArXiv-1M:1536D,L2. ms/decimal KB. ANNProof* inherits historical DERIVED display factors Qx1.2,VOx2,Vx1.5; raw native results are in ArXiv原始实测 and CSV. Other rows raw; no predictions.'
    ws['A4']='Non-V3DB defaults>=95% Recall@10; MR exact100%. V3DB lower recall/one fixed proof.20 queries x3; query/verify1thread. Offline graph builds up to8threads (MR serial); compare build times only with matching resources.'
    for row in [9,10,11,12,13,21,23,25,27,29,35,37,39,41,43,49,51,53,55,57]:
        for col in range(5,27):
            if row in [9,10,11,12,13] or col>=6:ws.cell(row,col).value=None
    names={'VLAH':9,'ANNProof':10,'Ours':11,'V3DB':12,'MR-tree':13}
    def put(row,col,value,record):
        cell=ws.cell(row,col);cell.value=value;cell.number_format='0.00';cell.font=Font(name='Calibri',size=11,color='000000');cell.fill=PatternFill(fill_type=None);cell.comment=Comment(json.dumps(record,ensure_ascii=False)[:30000],'ArXiv actual-data provenance')
    for r in overall:
        row=names[r['method']];ws.cell(row,4,r['method']+('*' if r['method']=='ANNProof' else ''))
        if r.get('status')=='measured':
            for col,key in [(5,'recall'),(6,'query_ms'),(7,'vo_kb'),(8,'verify_ms'),(9,'index_mb'),(10,'index_seconds')]:
                factor={'query_ms':1.2,'vo_kb':2.,'verify_ms':1.5}.get(key,1.) if r['method']=='ANNProof' else 1.
                if key in r:put(row,col,r[key]*(100 if key=='recall' else factor),dict(r,display_factor=factor,derived=bool(factor!=1)))
        else:
            ws.cell(row,11,r.get('status','pending'))
            for col,key in [(9,'index_mb'),(10,'index_seconds')]:
                if key in r:put(row,col,r[key],r.get('index_provenance',r))
    loc={'VLAH':21,'ANNProof':23,'Ours':25,'V3DB':27,'MR-tree':29}
    for method,row in loc.items():
        for offset in [0,14,28]:ws.cell(row+offset,4,method+('*' if method=='ANNProof' else ''))
    for r in curves:
        col=6+TARGETS.index(r['target'])
        if r.get('status')!='measured':
            for offset in [0,14,28]:put(loc[r['method']]+offset,col,None,r);put(loc[r['method']]+offset+1,col,None,r)
            continue
        for offset,key in [(0,'vo_kb'),(14,'verify_ms'),(28,'query_ms')]:
            factor={'query_ms':1.2,'vo_kb':2.,'verify_ms':1.5}[key] if r['method']=='ANNProof' else 1.
            put(loc[r['method']]+offset,col,r['recall']*100,r);put(loc[r['method']]+offset+1,col,r[key]*factor,dict(r,display_factor=factor,derived=bool(factor!=1)))
    mr=next(r for r in overall if r['method']=='MR-tree')
    if mr.get('status')=='measured':
        for row,key in [(29,'vo_kb'),(43,'verify_ms'),(57,'query_ms')]:put(row,26,100,mr);put(row+1,26,mr[key],mr)
    ws.cell(63,4,'Paper Table4: ablation study (raw actual results)')
    for col,name in [(4,'Variant'),(5,'Parameter'),(6,'R (%)'),(7,'Q (ms)'),(8,'VO (KB)'),(9,'V (ms)')]:ws.cell(64,col,name)
    for row,r in enumerate(ab,65):
        ws.cell(row,4,r['method']);ws.cell(row,10,r.get('status'))
        if r.get('status')=='measured':
            for col,key in [(5,'parameter'),(6,'recall'),(7,'query_ms'),(8,'vo_kb'),(9,'verify_ms')]:put(row,col,r[key]*(100 if key=='recall' else 1),r)
    # Historical ANNProof factors are an explicitly separate derived comparison,
    # never relabeled as raw real VO/timing. Main new dataset remains measured.
    ws.cell(74,4,'ANNProof* historical-coefficient derived display; not raw measurements')
    ap=next(r for r in overall if r['method']=='ANNProof')
    if ap.get('status')=='measured':
        for col,key,factor in [(5,'recall',100),(6,'query_ms',1.2),(7,'vo_kb',2),(8,'verify_ms',1.5)]:put(75,col,ap[key]*factor,dict(ap,derived=True,factors='Q*1.2,VO*2,V*1.5'))
    auditname='ArXiv实测与来源'
    if auditname in wb:del wb[auditname]
    audit=wb.create_sheet(auditname);audit.append(['kind','method','status','actualRecall','Qms','VO_KB','Vms','parameter','probes','target','source','record'])
    for kind,data in [('overall',overall),('curve',curves),('Table4',ab)]:
        for r in data:audit.append([kind,r['method'],r.get('status'),r.get('recall'),r.get('query_ms'),r.get('vo_kb'),r.get('verify_ms'),r.get('ef',r.get('parameter')),r.get('probes'),r.get('target'),r.get('source'),json.dumps(r,ensure_ascii=False)[:30000]])
    rawname='ArXiv原始实测'
    if rawname in wb:del wb[rawname]
    rawtab=wb.create_sheet(rawname);rawtab.append(['Kind','Method','Actual Recall@10','Query ms RAW','VO KB RAW','Verify ms RAW','Parameter','Probes','Target','Status','Source'])
    for kind,data in [('overall',overall),('curve',curves),('Table4',ab)]:
        for r in data:rawtab.append([kind,r['method'],r.get('recall'),r.get('query_ms'),r.get('vo_kb'),r.get('verify_ms'),r.get('ef',r.get('parameter')),r.get('probes'),r.get('target'),r.get('status'),r.get('source')])
    chartname='ArXivRecall曲线'
    if chartname in wb:del wb[chartname]
    charttab=wb.create_sheet(chartname);charttab.append(['Method','Actual Recall@10','Query ms','VO KB','Verify ms','ef/E/probes','proof samples','source'])
    unique={}
    for r in curves:
        if r.get('status')=='measured':unique[r['method'],r.get('model',''),r['ef'],r['probes']]=r
    for r in sorted(unique.values(),key=lambda r:(r['method'],r['recall'],r['ef'],r['probes'])):
        ap=r['method']=='ANNProof';cfg=r.get('config',{});parameter=f"{r['ef']}/{r['probes']}"+(f" M={cfg['M']},K={cfg['K']}" if cfg else '');charttab.append([r['method']+('*' if ap else ''),r['recall'],r['query_ms']*(1.2 if ap else 1),r['vo_kb']*(2 if ap else 1),r['verify_ms']*(1.5 if ap else 1),parameter,r['samples'],r['source']])
    for metric,col in [('Query time (ms)',3),('VO size (KB)',4),('Verification time (ms)',5)]:
        chart=ScatterChart();chart.title='OpenAI-ArXiv-1M: '+metric;chart.x_axis.title='Actual Recall@10';chart.x_axis.scaling.min=0;chart.x_axis.scaling.max=1;chart.x_axis.numFmt='0%';chart.y_axis.title=metric;chart.y_axis.scaling.logBase=10;chart.width=22;chart.height=13
        for method in ['ANNProof*','VLAH','V3DB','MR-tree','Ours']:
            indices=[i for i in range(2,charttab.max_row+1) if charttab.cell(i,1).value==method]
            if indices:
                series=Series(Reference(charttab,min_col=col,min_row=min(indices),max_row=max(indices)),Reference(charttab,min_col=2,min_row=min(indices),max_row=max(indices)),title=method)
                series.marker.symbol='circle';series.marker.size=4;chart.series.append(series)
        charttab.add_chart(chart,'J'+str(1+(col-3)*27))
    output=S/'work9_with_openai_arxiv.xlsx';tmp=output.with_suffix('.tmp.xlsx');wb.save(tmp);check=openpyxl.load_workbook(tmp,read_only=True);assert sheet in check.sheetnames;check.close();tmp.replace(output)
    try:shutil.copy2(output,source)
    except PermissionError:event(action='workbook_locked',output=str(output))
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('stage',choices=['initialize','baselines','ablations','mr','v3db','indexes','report','all'],nargs='?',default='all');a=p.parse_args()
    if a.stage=='all':
        initialize()
        for phase in [baselines,ablations,mr,v3db,indexes]:
            try:phase()
            except Exception as e:
                import traceback
                event(action='phase_failed',phase=phase.__name__,error=str(e));(S/(phase.__name__+'-failure.txt')).write_text(traceback.format_exc());raise
            finally:report()
        u.write(S/'FINISHED.json',dict(passed=True,scope='all requested native overall/curves/Table4 phases completed'))
    else:
        globals()[a.stage]()
