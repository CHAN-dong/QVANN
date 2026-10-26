from pathlib import Path
import json
import openpyxl
from work9_nytimes import ROOT,row_map,sha,write,STUDY
base=openpyxl.load_workbook(ROOT/'experiments/work9-v3db-20261016/work9_before_v3db.xlsx')
current=openpyxl.load_workbook(ROOT.parent/'work9.xlsx')
changes=[];style_changes=[]
for s in base:
    if s.title not in current:changes.append([s.title,'missing sheet']);continue
    for row in s:
        for c in row:
            if c.value is None:continue
            r=row_map(c.row) if s.title in ['SIFT1M','GIST1M','GloVe-100','DEEP1M'] else c.row
            got=current[s.title].cell(r,c.column)
            if got.value!=c.value:changes.append([s.title,c.coordinate,got.coordinate,c.value,got.value])
            if got._style!=c._style or got.number_format!=c.number_format:style_changes.append([s.title,c.coordinate,got.coordinate])
print(json.dumps(dict(sheets=current.sheetnames,changed_original_values=len(changes),changes=changes[:30]),ensure_ascii=False))
assert not changes and not style_changes,(changes[:10],style_changes[:10])
write(STUDY/'source_reconciliation.json',dict(status='PASS',source_sha256=sha(ROOT.parent/'work9.xlsx'),
    archived_source_sha256=sha(STUDY/'work9_before_nytimes.xlsx'),changed_original_values=0,changed_original_styles=0,
    reason='Current workbook hash differs from last writer marker, but every nonempty original cell value/style/number format matches; archived before replacement.'))
