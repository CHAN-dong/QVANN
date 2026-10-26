"""Conditional 512-GiB forecasts for previously unmeasured V3DB cells only."""
import argparse,csv,json,math,shutil,statistics,subprocess,sys,time
from collections import defaultdict
from pathlib import Path
from copy import copy
import numpy as np
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill,Font
from work9_resource_matched import ROOT,DATA,sha,write,config as assumed_config,model_name

STUDY=ROOT/'experiments/work9-v3db-512GiB-20261021'
SOURCE=ROOT.parent/'work9.xlsx';BACKUP=STUDY/'work9_before_512GiB.xlsx'
CACHE=ROOT/'V3DB-26VLDB/work9-cache'
OLD=ROOT/'experiments/work9-v3db-reaudit-20261018'
NATIVE='/root/.cache/v3db-target/release/work9-inspect-verifier'
sys.path.insert(0,str(ROOT/'V3DB-26VLDB/scripts'))
from work9_forecast import features,pi_count,signature

def read(p):return json.loads(Path(p).read_text())
def collect():
 refs={}
 dirs=['work9-v3db-20261016','work9-v3db-reaudit-20261018','work9-v3db-single-20261020','work9-v3db-lowrecall-20261020']
 for name in dirs:
  for f in sorted((ROOT/'results'/name).glob('*/per_query.csv')):
   rows=list(csv.DictReader(f.open()))
   r=next((r for r in rows if r['verified']=='1' and int(r['query'])==100 and int(r['repeat'])==0),None)
   if not r:continue
   vk=f.parent/'trusted_verifier.bin'
   if not vk.exists():continue
   meta=f.parent/'result.json'
   if meta.exists() and 'config' in read(meta):c=read(meta)['config'];p=read(meta)['n_probe'];k=read(meta)['k']
   else:
    import re
    match=re.match(r'(.+)-p(\d+)-k(\d+)-q\d+-r\d+-t1$',f.parent.name)
    if not match:continue
    model,p,k=match.groups();p=int(p);k=int(k)
    mf=CACHE/model/'model.json'
    if not mf.exists():continue
    c=read(mf)['config']
   if c['dataset'] not in DATA or k!=10:continue
   key=signature(c,p,k)+(c['N'],c['source'])
   refs[key]=dict(config=c,n_probe=p,k=k,source=str(f),source_sha256=sha(f),verifier=str(vk),gates=int(r['gates']),
    query_ms=float(r['query_ms']),verify_ms=float(r['verify_ms']),vo_kb=float(r['vo_bytes'])/1000,
    build_seconds=float(r['build_seconds']),prove_seconds=float(r['prove_seconds']),proof_phase_rss_GiB=int(r['rss_bytes'])/2**30,
    sample_policy='fixed query100 repeat0, one vote per configuration; latest study overrides identical shape/source')
 refs=list(refs.values());info=[json.loads(x) for x in subprocess.check_output([NATIVE,*[r['verifier'] for r in refs]],text=True).splitlines()]
 bypath={x['path']:x for x in info}
 for r in refs:
  v=bypath[r['verifier']];r['degree_bits']=v['degree_bits'];r['domain']=2**v['degree_bits']
  r['public_fields']=v['public_inputs'];assert r['public_fields']==pi_count(r['config'],r['n_probe'],10)
  r['body_bytes']=r['vo_kb']*1000-8*r['public_fields']
 write(STUDY/'proof-calibration-samples.json',refs)
 return refs

class Model:
 def __init__(self,refs):
  self.refs=refs;old=read(OLD/'forecast-model.json');self.old=old
  self.scales=np.array(old['feature_scales']);self.coef=np.array(old['gate_coefficients'])
  self.exact={}
  for f in [OLD/'gate-model-training.json',OLD/'gate-observations.json']:
   if f.exists():
    for r in read(f):
     c=r['config'];p=c.get('n_probe',r.get('n_probe'));k=c.get('k',10)
     if p is not None and k==10:self.exact[signature(c,p,10)]=r['gates']
  for folder in ['work9-v3db-single-20261020','work9-v3db-lowrecall-20261020','work9-resource-matched-20261020']:
   for f in (ROOT/'experiments'/folder/'gates').glob('*.json'):
    r=read(f)
    if 'gates' in r and r.get('k')==10:self.exact[signature(r['config'],r['n_probe'],10)]=r['gates']
  for r in refs:self.exact[signature(r['config'],r['n_probe'],10)]=r['gates']
  self.ap=statistics.median(r['prove_seconds']/(r['domain']*r['degree_bits']) for r in refs)
  self.ab=statistics.median(r['build_seconds']/(r['domain']*r['degree_bits']) for r in refs)
  self.ah=statistics.median(max(0,r['query_ms']/1000-r['build_seconds']-r['prove_seconds'])/self.work(r['config'],r['n_probe']) for r in refs)
  large=[r for r in refs if r['degree_bits']>=19]
  self.mem=statistics.median(r['proof_phase_rss_GiB']/r['domain'] for r in large)
  self.mem_upper=float(np.quantile([r['proof_phase_rss_GiB']/r['domain'] for r in large],.95))
  self.body15=old['proof_body_bytes_at_b15']
  self.body_slope=statistics.median((r['body_bytes']-self.body15)/(r['degree_bits']-15) for r in refs if r['degree_bits']>15)
  self.hashcoef=statistics.median(r['median_ms']/r['fields'] for r in old['public_input_hash_microbenchmark'])*1.5
  self.verify={b:statistics.median(max(1,r['verify_ms']-self.hashcoef*r['public_fields']) for r in refs if r['degree_bits']==b) for b in sorted({r['degree_bits'] for r in refs})}
  err=[abs(float(features(r['config'],r['n_probe'])/self.scales@self.coef)/r['gates']-1) for r in refs]
  self.report=dict(memory_capacity_GiB=512,threads=1,single_core_speed='same as local reference; more RAM alone provides no speedup',
   proof_shapes=len(refs),observed_domain_bits=sorted(self.verify),gate_model='unchanged audited structural model; exact observations override it',
   gate_model_median_relative_error_on_new_shapes=statistics.median(err),gate_model_max_relative_error_on_new_shapes=max(err),gate_relative_planning_margin=.10,
   proof_seconds_per_B_logB=self.ap,compile_seconds_per_B_logB=self.ab,host_seconds_per_T_logT=self.ah,
   memory_proof_GiB_per_B=self.mem,memory_upper_coefficient_large_anchor_p95=self.mem_upper,
   memory_planning='1.5 * p95-large-anchor RSS-per-row * B_guard + 32 GiB; B_guard=nextpow2(1.10*G). This is a declared engineering scenario, not a measured/proven upper bound.',
   reserve_reason='Proof-phase RSS does not cover all builder/transient peaks. 50% temporary/allocator allowance plus 32GiB OS/data reserve is applied.',
   body15_bytes=self.body15,body_bytes_per_domain_bit=self.body_slope,verify_core_ms_by_observed_bin=self.verify,hash_ms_per_public_field=self.hashcoef,
   limitations=['No proof has been run on a 512GiB machine. Only conditional forecasts, not measured results or guaranteed hardware requirements.',
    'Exact gate counts are before final builder padding. 10% row margin accounts for model/padding uncertainty; it is not a statistical confidence interval.',
    'Domains above the measured maximum are extrapolated. RAM planning is not a mathematical maximum-memory proof.',
    'Recall is taken only from existing 20-query search calibration and is never predicted upward.',
    'Absent calibrated index/recall for a large prefix: no matched-recall main-table forecast is created. Fixed-budget hypothetical values are listed separately.'])
  timing_errors=[]
  for r in refs:
   B=r['domain'];b=r['degree_bits']
   predicted=((self.ap+self.ab)*B*b+self.ah*self.work(r['config'],r['n_probe']))*1000
   timing_errors.append(abs(predicted/r['query_ms']-1))
  self.report['query_model_median_relative_error_on_calibration_shapes']=statistics.median(timing_errors)
  self.report['query_model_p95_relative_error_on_calibration_shapes']=float(np.quantile(timing_errors,.95))
  self.report['calibration_fit_warning']='in-sample diagnostic only; does not validate large-domain extrapolation'
  write(STUDY/'model.json',self.report)
 @staticmethod
 def work(c,p):
  t=p*max(c['capacity']*c['M'],c['M']*c['K']);return t*math.log2(max(2,t))
 def predict(self,c,p,k=10):
  g=self.exact.get(signature(c,p,10));exact=g is not None
  if g is None:g=max(1,int(features(c,p)/self.scales@self.coef))
  b=math.ceil(math.log2(g));guard=math.ceil(math.log2(g*1.10));B=2**b;P=pi_count(c,p,k)
  mem=self.mem*B;planned=1.5*self.mem_upper*2**guard+32
  query=((self.ap+self.ab)*B*b+self.ah*self.work(c,p))*1000
  core=self.verify[b] if b in self.verify else self.verify[max(self.verify)]*b/max(self.verify)
  verify=core+self.hashcoef*P
  vo=(max(50000,self.body15+self.body_slope*(b-15))+8*P)/1000
  return dict(status='conditional_forecast' if planned<=512 else 'outside_512GiB_scenario',config=c,n_probe=p,k=k,gates=g,gate_source='exact observed static count' if exact else 'structural estimate',
   domain_bits=b,guard_domain_bits=guard,proof_phase_memory_estimate_GiB=mem,planned_machine_memory_GiB=planned,fits_512GiB_scenario=planned<=512,
   query_ms=query,verify_ms=verify,vo_kb=vo,
   declared_sensitivity=dict(query_ms=[query/2,query*2],verify_ms=[verify/2,verify*2],vo_kb=[vo*.85,vo*1.15]),
   label='512GiB条件预测，未在512GiB机器实测',recall_scope='existing search calibration only')

def candidates(model):
 rows=[]
 for f in sorted(CACHE.glob('*/calibration.json')):
  cal=read(f);c=cal['model']['config']
  if c['dataset'] not in DATA or c['N']!=DATA[c['dataset']]['N'] or c['source']!=c['dataset']:continue
  for point in cal['points']:
   if point['k']!=10:continue
   r=model.predict(c,point['n_probe']);r.update(model=f.parent.name,recall=point['recall'],recall_source=str(f),recall_queries=20)
   rows.append(r)
 write(STUDY/'curve-candidates.json',rows);return rows

def plan(model):
 rows=candidates(model);wb=openpyxl.load_workbook(BACKUP);decisions=[];summary=[]
 for ds,c in DATA.items():
  ws=wb[c['sheet']];targets=[]
  for col in range(6,ws.max_column+1):
   target=ws.cell(16,col).value
   if not isinstance(target,(int,float)):continue
   if not any(isinstance(ws.cell(r,col).value,str) for r in [26,27,38,39,50,51]):continue
   good=[r for r in rows if r['config']['dataset']==ds and r['recall']>=target/100-1e-12]
   fit=[r for r in good if r['fits_512GiB_scenario']]
   chosen=min(fit,key=lambda r:(r['query_ms'],r['recall'])) if fit else None
   nearest=min(good,key=lambda r:r['planned_machine_memory_GiB']) if good else None
   reason='existing calibrated configuration fits declared 512GiB planning envelope' if chosen else ('all calibrated configurations reaching this target exceed the declared 512GiB planning envelope' if good else 'no existing calibrated configuration reaches target; recall cannot be invented')
   decisions.append(dict(kind='curve',dataset=ds,sheet=ws.title,column=col,target=target/100,choice=chosen,reason=reason,smallest_memory_reaching_target=nearest))
   if chosen:targets.append(target)
  summary.append(dict(dataset=ds,new_curve_targets=targets,max_calibrated_recall_fitting_scenario=max([r['recall'] for r in rows if r['config']['dataset']==ds and r['fits_512GiB_scenario']],default=None)))
  for col,n in enumerate([100000,500000,1000000,5000000,10000000],6):
   if not any(isinstance(ws.cell(r,col).value,str) for r in [64,75,85]):continue
   if ds not in ['sift1m','deep1m'] and n>c['N']:
    decisions.append(dict(kind='size',dataset=ds,sheet=ws.title,column=col,N=n,choice=None,reason='requested size exceeds actual available corpus; not a RAM question'));continue
   source='bigann10m' if ds=='sift1m' else ('deep10m' if ds=='deep1m' and n>1000000 else ds)
   name=model_name(ds,n,source);f=CACHE/name/'calibration.json'
   if f.exists():
    cal=read(f);cfg=cal['model']['config'];points=[p for p in cal['points'] if p['k']==10 and p['recall']>=c['target']-1e-12]
    p=min(points,key=lambda p:p['n_probe']) if points else None
    pred=model.predict(cfg,p['n_probe']) if p else None
    if pred:pred.update(model=name,recall=p['recall'],recall_source=str(f),recall_queries=20)
    chosen=pred if pred and pred['fits_512GiB_scenario'] else None
    reason='existing prefix-index calibration reaches the existing default target' if chosen else 'prefix target not reached or exceeds 512GiB planning envelope'
    decisions.append(dict(kind='size',dataset=ds,sheet=ws.title,column=col,N=n,choice=chosen,reason=reason,resource_assessment=pred))
   else:
    cfg=assumed_config(ds,n,source);p=8 if ds=='sift1m' else c['probe']
    hypothesis=model.predict(cfg,p);hypothesis.update(parameter_assumption=True,recall=None,reason='fixed query budget hypothesis only; no trained/calibrated index for this prefix. Not an equal/default-recall result.')
    decisions.append(dict(kind='size',dataset=ds,sheet=ws.title,column=col,N=n,choice=None,reason='no index/recall calibration; matched-recall performance is unsupported',fixed_budget_hypothesis=hypothesis))
 write(STUDY/'decisions.json',decisions);write(STUDY/'summary.json',summary)
 return decisions,summary

def fill(decisions,model):
 wb=openpyxl.load_workbook(BACKUP);changes=[];orange=PatternFill('solid',fgColor='FFF2CC')
 def put(ws,row,col,value,evidence,forecast):
  cell=ws.cell(row,col)
  assert not isinstance(cell.value,(int,float)),('refuse to overwrite real measurement',ws.title,cell.coordinate)
  old=cell.value;cell.value=value
  if value is not None:
   cell.fill=orange;cell.font=Font(name='Calibri',size=11,color='9C5700')
   cell.number_format='0.000"*"' if forecast else '0.0"†"'
  cell.comment=Comment(('512GiB条件预测（*），不是实测。仍假设单线程、相近单核速度、无换页。' if forecast and value is not None else '†为同配置20条搜索校准召回，不是预测召回。' if value is not None else '按用户要求留空；具体原因如下。')+'\n'+json.dumps(evidence,ensure_ascii=False,indent=2)[:29000],'512GiB预测来源')
  changes.append(dict(sheet=ws.title,cell=cell.coordinate,before=old,value=value,estimated=forecast and value is not None,source=evidence))
 for d in decisions:
  ws=wb[d['sheet']];col=d['column'];r=d['choice']
  if d['kind']=='curve':
   for row,key in [(26,'vo_kb'),(38,'verify_ms'),(50,'query_ms')]:
    for rr,value,is_prediction in [(row,r['recall']*100 if r else None,False),(row+1,r[key] if r else None,True)]:
     if isinstance(ws.cell(rr,col).value,str) or ws.cell(rr,col).value is None:put(ws,rr,col,value,d,is_prediction)
  else:
   for row,key in [(64,'vo_kb'),(75,'verify_ms'),(85,'query_ms')]:
    if isinstance(ws.cell(row,col).value,str) or ws.cell(row,col).value is None:put(ws,row,col,r[key] if r else None,d,True)
 for c in DATA.values():
  ws=wb[c['sheet']];ws['A4']=str(ws['A4'].value or '')+'\n橙色*为512GiB单线程条件预测，†为既有20查询搜索校准召回；不是512GiB机器实测。超规划内存/无校准依据的缺项留空，见V3DB512GiB说明。'
  alignment=copy(ws['A4'].alignment);alignment.wrap_text=True;ws['A4'].alignment=alignment
  ws.row_dimensions[4].height=max(60,ws.row_dimensions[4].height or 15)
 for name in ['V3DB512GiB说明','V3DB512GiB预测记录','V3DB512GiB单元格来源','V3DB512GiB固定预算规模']:
  if name in wb:del wb[name]
 notes=wb.create_sheet('V3DB512GiB说明')
 for row in [
  ['项目','说明'],['性质','用户假设512GiB（约550GB十进制）内存机器：本文件填入条件预测，不是该机器实测。512GB十进制约476.8GiB，硬件口径应区分。'],
  ['原未测原因','此前本机约62GiB可见RAM，证明子进程50GiB地址空间上限，另有200万关系元组/2^20证明域预检预算；跳过不等于实际出证OOM，更不等于512GiB不可运行。'],
  ['显示','橙色数值后*是512GiB条件预测；†为同配置已有20条搜索校准的实际召回率，未经20次证明。现有实测单元格不覆盖。'],
  ['CPU与时间','保持单线程及与本机相近的单核性能，假设足够RAM且不换页；不将扩容当作CPU加速。Query仍含搜索、见证、约束/编译、真实证明和压缩的完整流程。'],
  ['可行性规则',model.report['memory_planning']],
  ['规则性质','该内存余量是公开的工程规划假设，不是实测上界或可运行保证。RSS校准仅覆盖证明阶段；构建峰值、分配器、版本/线程差异仍可能改变结论。'],
  ['模型范围',f'按{len(model.refs)}个固定query100样本校准。实际观测证明域位数{sorted(model.verify)}；更大档位为外推。时间/验证声明0.5–2倍、VO±15%敏感性范围，不是置信区间。'],
  ['曲线参数','仅填此前未测位置。候选来自现有PQ/probe参数及实际20查询搜索校准，选满足目标且进入512GiB规划包络的最低预测查询时间配置。不同PQ配置的索引成本不能当成主表默认索引成本。'],
  ['留空','达到目标的所有已校准配置都超规划内存；或现有校准未达到目标；或数据集容量不足；或规模前缀未训练/校准。这些原因在单元格批注分别记录，不能都解释成物理内存不够。'],
  ['规模假设','无真实前缀索引/召回校准的5M/10M不在主表冒充默认召回预测。另页提供固定预算资源/性能假设供参考；其召回率未知，不是可比实验数据。'],
  ['默认与其他方案','所有默认、索引、K、消融和其他方法全部数据保持不变。只补V3DB原缺失点并将不支持的缺项留空。'],
  ['原表备份',str(BACKUP)],['模型与来源',str(STUDY)]]:notes.append(row)
 notes.column_dimensions['A'].width=23;notes.column_dimensions['B'].width=145
 tab=wb.create_sheet('V3DB512GiB预测记录');tab.append(['dataset','panel','target/N','action','calibrated recall','M','K','probe','gates','bits','guard bits','proof RSS GiB estimate','machine RAM GiB plan','Query ms forecast','Verify ms forecast','VO KB forecast','reason'])
 for d in decisions:
  r=d['choice'] or d.get('resource_assessment') or d.get('smallest_memory_reaching_target') or {}
  c=r.get('config',{})
  tab.append([d['dataset'],d['kind'],d.get('target',d.get('N')),'预测填入' if d['choice'] else '留空',r.get('recall'),c.get('M'),c.get('K'),r.get('n_probe'),r.get('gates'),r.get('domain_bits'),r.get('guard_domain_bits'),r.get('proof_phase_memory_estimate_GiB'),r.get('planned_machine_memory_GiB'),r.get('query_ms') if d['choice'] else None,r.get('verify_ms') if d['choice'] else None,r.get('vo_kb') if d['choice'] else None,d['reason']])
 tab=wb.create_sheet('V3DB512GiB单元格来源');tab.append(['sheet','cell','before','after','estimated','evidence'])
 for c in changes:tab.append([c['sheet'],c['cell'],c['before'],c['value'],c['estimated'],json.dumps(c['source'],ensure_ascii=False)[:32000]])
 tab=wb.create_sheet('V3DB512GiB固定预算规模');tab.append(['dataset','N','M','K','probe assumed','RSS proof GiB estimate','RAM GiB planning','fits scenario','Query ms hypothetical','Verify ms hypothetical','VO KB hypothetical','recall','warning'])
 for d in decisions:
  r=d.get('fixed_budget_hypothesis')
  if not r:continue
  c=r['config'];fit=r['fits_512GiB_scenario']
  tab.append([d['dataset'],d['N'],c['M'],c['K'],r['n_probe'],r['proof_phase_memory_estimate_GiB'],r['planned_machine_memory_GiB'],fit,r['query_ms'] if fit else None,r['verify_ms'] if fit else None,r['vo_kb'] if fit else None,'未知/未校准',r['reason']])
 output=STUDY/'work9_512GiB_forecast.xlsx';wb.save(output)
 before=openpyxl.load_workbook(BACKUP);after=openpyxl.load_workbook(output);allowed={(c['sheet'],c['cell']) for c in changes}|{(c['sheet'],'A4') for c in DATA.values()};preserved=0
 for cfg in DATA.values():
  a=before[cfg['sheet']];b=after[cfg['sheet']]
  for row in a:
   for x in row:
    if (a.title,x.coordinate) in allowed:continue
    y=b[x.coordinate]
    assert x.value==y.value and x._style==y._style,(a.title,x.coordinate)
    assert (x.comment.text if x.comment else None)==(y.comment.text if y.comment else None)
    preserved+=1
 for c in changes:
  got=after[c['sheet']][c['cell']].value
  assert got==c['value'] or isinstance(c['value'],float) and math.isclose(got,c['value'],rel_tol=1e-12)
  if c['estimated']:assert c['source']['choice']['planned_machine_memory_GiB']<=512
 assert sha(SOURCE)==sha(BACKUP),'user workbook changed during analysis; merge rather than overwrite'
 shutil.copy2(output,SOURCE);assert sha(SOURCE)==sha(output)
 write(STUDY/'cell-updates.json',changes)
 write(STUDY/'validation.json',dict(status='PASS',forecast_metric_cells=sum(c['estimated'] for c in changes),calibrated_recall_cells=sum(not c['estimated'] and c['value'] is not None for c in changes),blanked_cells=sum(c['value'] is None for c in changes),unrelated_main_cells_preserved=preserved,workbook_sha256=sha(SOURCE),updated_original=True))

if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('action',choices=['plan','fill']);a=p.parse_args();STUDY.mkdir(parents=True,exist_ok=True)
 if not BACKUP.exists():shutil.copy2(SOURCE,BACKUP)
 refs=collect();model=Model(refs);decisions,summary=plan(model)
 print(json.dumps(dict(calibration_shapes=len(refs),memory_coefficient=model.mem,upper_memory_coefficient=model.mem_upper,summary=summary,size_decisions=[dict(dataset=d['dataset'],N=d['N'],filled=bool(d['choice']),reason=d['reason']) for d in decisions if d['kind']=='size']),ensure_ascii=False,indent=2))
 if a.action=='fill':fill(decisions,model);print(json.dumps(read(STUDY/'validation.json'),ensure_ascii=False))
