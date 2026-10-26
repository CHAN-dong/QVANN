"""Audit replacement scope, dataset provenance and completed real measurements."""
import csv
import json
from pathlib import Path
import numpy as np
import openpyxl
from work9_nytimes import ROOT,STUDY,FOLDER,N,D,NAME,sha,write,row_map


def main():
    provenance=json.loads((FOLDER/'provenance.json').read_text())
    assert provenance['status']=='PASS'
    for item in provenance['files'].values():assert sha(FOLDER/item['file'])==item['sha256']
    base=np.memmap(FOLDER/'base.fvecs',dtype='<f4',mode='r',shape=(N,D+1))
    assert np.all(base.view('<i4')[:,0]==D)
    for start in range(0,N,8192):
        block=base[start:start+8192,1:].astype(np.float64)
        assert np.isfinite(block).all() and np.allclose(np.sum(block*block,axis=1),1,atol=3e-7,rtol=0)
    ids=np.load(FOLDER/'base_original_ids.npy')
    assert len(ids)==N and len(set(ids.tolist()))==N
    assert np.all(ids[1:]>ids[:-1]) and ids.max()<290000
    wb=openpyxl.load_workbook(STUDY/'work9_nytimes.xlsx')
    assert NAME in wb and 'GIST1M' not in wb
    old=openpyxl.load_workbook(ROOT/'experiments/work9-v3db-20261016/work9_before_v3db.xlsx')
    preserved=0
    for sheet in ['SIFT1M','GloVe-100','DEEP1M']:
        for row in old[sheet]:
            for c in row:
                if c.value is None:continue
                got=wb[sheet].cell(row_map(c.row),c.column)
                assert got.value==c.value,(sheet,c.coordinate,got.coordinate,c.value,got.value)
                assert got._style==c._style and got.number_format==c.number_format
                preserved+=1
    for tab in wb:
        if tab.title=='NYTimes实验说明':continue
        for row in tab:
            assert not any(isinstance(c.value,str) and (c.value.lower()=='gist1m' or c.value.lower().startswith('gist1m-')) for c in row[:3]),(tab.title,row[0].row)
    verified=[]
    for f in sorted(STUDY.glob('evaluation-*.json')):
        job=json.loads(f.read_text());out=Path(job['output'])
        assert (out/'PASS.txt').exists()
        groups={}
        for r in csv.DictReader((out/'per_query.csv').open()):
            assert r['verified']=='1'
            groups.setdefault((r['k'],r['ef_search'],r['probes']),[]).append(r)
        for key,rows in groups.items():
            assert len(rows)==60 and {(int(r['query']),int(r['repeat'])) for r in rows}=={(q,t) for q in range(100,120) for t in range(3)}
        raw=list(csv.DictReader((out/'per_query.csv').open()))
        assert (out/'query0.vo').stat().st_size==int(raw[0]['vo_bytes'])
        attacks=list(csv.DictReader((out/'tamper_tests.csv').open()))
        assert attacks and all(r['rejected']=='1' for r in attacks)
        verified.append(dict(method=job['method'],purpose=job['purpose'],N=job['base'],workpoints=len(groups),samples=len(raw)))
    ablations=[]
    for f in sorted(STUDY.glob('ablation-evaluation-*.json')):
        records=json.loads(f.read_text());reference=None
        for record in records:
            folder=Path(record['output'])
            assert sha(folder/'per_query.csv')==record['source_sha256']
            rows=list(csv.DictReader((folder/'per_query.csv').open()))
            answers=list(csv.DictReader((folder/'answers.csv').open()))
            assert len(rows)==len(answers)==60 and all(r['verified']=='1' for r in rows)
            assert {(int(r['query']),int(r['repeat'])) for r in rows}=={(q,t) for q in range(100,120) for t in range(3)}
            for csvkey,key,scale in [('recall_at_k','recall',1),('server_total_ms','query_ms',1),('verification_ms','verify_ms',1),('vo_bytes','vo_kb',1000)]:
                mean=sum(float(r[csvkey])/scale for r in rows)/len(rows)
                assert abs(mean-record[key])<max(1e-9,abs(mean)*1e-12)
            attacks=list(csv.DictReader((folder/'tamper_tests.csv').open()))
            assert len(attacks)>=12 and all(r['rejected']=='1' for r in attacks)
            if reference is not None:
                for a,b in zip(reference,answers):assert all(a[k]==b[k] for k in ['query','repeat','parameter','answer'])
            else:reference=answers
            ablations.append(dict(variant=record['variant'],samples=len(rows),tamper_tests=len(attacks)))
    pending=sum(c.value in ['待测','待校准'] for row in wb[NAME] for c in row)
    proved=[]
    for f in (ROOT/'results/work9-v3db-20261016').glob('*/result.json'):
        result=json.loads(f.read_text())
        if result['status']!='passed':continue
        rows=list(csv.DictReader((f.parent/'per_query.csv').open()))
        assert len(rows)==60 and result['query_count']==20 and result['repeats']==3
        assert {(int(r['query']),int(r['repeat'])) for r in rows}=={(q,t) for q in range(100,120) for t in range(3)}
        assert all(r['verified']=='1' for r in rows)
        for key in ['recall','query_ms','verify_ms']:
            mean=sum(float(r[key]) for r in rows)/len(rows)
            assert abs(mean-result[key])<max(1e-9,abs(mean)*1e-12)
        assert abs(sum(int(r['vo_bytes'])/1000 for r in rows)/len(rows)-result['vo_kb'])<1e-8
        assert (f.parent/'query0.proof').stat().st_size==int(rows[0]['vo_bytes'])
        assert all(json.loads((f.parent/'tamper.json').read_text()).values())
        proved.append(result['job'])
    limits=sum(isinstance(c.value,str) and (c.value in ['未达目标','超预算未测'] or c.value.startswith('内存不足')) for row in wb[NAME] for c in row)
    state=dict(status='PASS',dataset='NYTimes-256 nonzero-base variant',N=N,d=D,
        archived_workbook_sha256=sha(STUDY/'work9_before_nytimes.xlsx'),
        other_main_original_cells_preserved=preserved,verified_jobs=verified,pending_cells=pending,
        verified_controlled_ablation=ablations,
        limitations_cells=limits,all_requested_numeric_measurements_complete=pending==0 and limits==0,
        V3DB_verified_20x3_workpoints=proved,
        replacement_copy=str(STUDY/'work9_nytimes.xlsx'))
    write(STUDY/'validation.json',state);print(json.dumps(state,indent=2))


if __name__=='__main__':main()
