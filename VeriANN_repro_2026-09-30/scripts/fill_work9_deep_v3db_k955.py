"""Conditional forecasts only for DEEP V3DB K curves at matched Recall@k."""
import json,math,shutil,sys
from pathlib import Path
import numpy as np
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment,Font,PatternFill
R=Path(__file__).resolve().parents[1];S=R/'experiments/work9-deep-v3db-k955-20261022';SOURCE=R.parent/'work9.xlsx';BACKUP=S/'work9_before_deep_v3db_k955.xlsx'
sys.path.insert(0,str(R/'V3DB-26VLDB/scripts'))
from work9_forecast import features,pi_count
from work9_resource_matched import sha,write
assert not (S/'APPLIED.json').exists(),'Already applied'
assert sha(SOURCE)==sha(BACKUP),'Workbook externally changed since backup'
cal=json.loads((S/'calibration.json').read_text());cfg=cal['config']
MROOT=R/'experiments/work9-v3db-512GiB-20261021';model=json.loads((MROOT/'model.json').read_text());gate=json.loads((R/'experiments/work9-v3db-reaudit-20261018/forecast-model.json').read_text())
coefs=np.array(gate['gate_coefficients']);scales=np.array(gate['feature_scales'])
w=openpyxl.load_workbook(BACKUP);old=openpyxl.load_workbook(BACKUP);s=w['DEEP1M'];predictions=[];changes=[]
for point in cal['selected']:
 k,p=point['k'],point['n_probe'];assert point['recall']>=.955-1e-12
 G=max(1,int(features(cfg,p)/scales@coefs));b=math.ceil(math.log2(G));guard=math.ceil(math.log2(G*1.10));B=2**b
 P=pi_count(cfg,p,k)
 planned=1.5*model['memory_upper_coefficient_large_anchor_p95']*2**guard+32
 rss=model['memory_proof_GiB_per_B']*B
 t=p*max(cfg['capacity']*cfg['M'],cfg['M']*cfg['K']);work=t*math.log2(max(2,t))
 query=((model['proof_seconds_per_B_logB']+model['compile_seconds_per_B_logB'])*B*b+model['host_seconds_per_T_logT']*work)*1000
 bins={int(x):v for x,v in model['verify_core_ms_by_observed_bin'].items()}
 core=bins[b] if b in bins else bins[max(bins)]*b/max(bins)
 verify=core+model['hash_ms_per_public_field']*P
 vo=(max(50000,model['body15_bytes']+model['body_bytes_per_domain_bit']*(b-15))+8*P)/1000
 pred=dict(**point,target_recall=.955,config=cfg,gates=G,gate_source='existing structural gate predictor; static gadget top_k only registers outputs, so pre-build gate estimate is k-independent at fixed probe; actual k used for public input count',domain_bits=b,guard_domain_bits=guard,public_fields=P,proof_phase_rss_estimate_GiB=rss,planned_machine_memory_GiB=planned,fits_512GiB_scenario=planned<=512,query_ms=query,vo_kb=vo,verify_ms=verify,kind='conditional prediction; not proof/query/verification measurement',threads=1,sensitivity={'query_ms':[query/2,query*2],'verify_ms':[verify/2,verify*2],'vo_kb':[vo*.85,vo*1.15]},source_model_sha256=sha(MROOT/'model.json'),calibration_sha256=sha(S/'calibration.json'))
 predictions.append(pred);col=6+[1,5,10,20,50,100].index(k)
 for row,key in [(107,'vo_kb'),(119,'verify_ms'),(130,'query_ms')]:
  c=s.cell(row,col);value=pred[key] if planned<=512 else None
  reason='DEEP1M V3DB条件预测，非实测性能：每个k独立选择最小整数n_probe，使同20条查询校准Recall@k≥95.5%。不把固定Recall@10配置误作所有k的相同召回。'
  reason+=' CPU单线程、512GiB机器规划、相近单核速度、无换页；出证工作含约束/编译、见证及证明，不含客户端验证。沿用33个历史证明形状的模型，超出已实测最大证明域2^20部分为外推。'
  reason+=' 公共输入计数使用实际k，不硬编码k=10；未使用ANNProof系数，未作单调平滑。每个k独立调参可以导致性能非单调。'
  reason+=' 原始模型给出query/verify×0.5–2、VO±15%的情景敏感范围，不是置信区间。超512GiB规划位置留空。'
  changes.append(dict(cell=c.coordinate,before=c.value,after=value,metric=key,source=pred,old_comment=c.comment.text if c.comment else ''))
  c.value=value;c.fill=PatternFill('solid',fgColor='FCE4D6');c.font=Font(color='C00000');c.number_format='0.000"†预测"'
  c.comment=Comment(reason+'\n'+json.dumps(pred,ensure_ascii=False,indent=2)+'\n校准明细='+str(S/'per-query-recall.csv')+'\n\n历史批注（非当前来源）：\n'+changes[-1]['old_comment'],'V3DB预测来源')
for row in [107,119,130]:
 s.cell(row,4).value='V3DB†预测\nRecall@k≥95.5%';s.cell(row,4).alignment=Alignment(wrap_text=True,vertical='center')
s['A1']=str(s['A1'].value or '')+' | DEEP V3DB K曲线最新口径：逐k校准Recall@k≥95.5%，橙色†预测为512GiB单线程条件预测。实际召回/参数见DEEP V3DB K955预测。'
s['A1'].alignment=Alignment(wrap_text=True);s.row_dimensions[1].height=max(70,s.row_dimensions[1].height or 15)
def tab(name,headers,rows):
 t=w.create_sheet(name);t.append(headers)
 for row in rows:t.append(row)
 t.freeze_panes='A2';t.auto_filter.ref=t.dimensions
 for c in t[1]:c.font=Font(bold=True)
 for col in t.columns:t.column_dimensions[col[0].column_letter].width=23 if col[0].column<17 else 70
 for row in t:
  for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
 return t
data=tab('DEEP V3DB K955预测',['k','目标Recall@k','实测搜索校准Recall@k','n_probe','PQ M','PQ K码字数','静态门数估计','证明域bits','内存规划GiB','预测query ms','预测VO KB','预测verify ms','query下界ms','query上界ms','verify下界ms','verify上界ms','VO下界KB','VO上界KB','性质'],[[p['k'],.955,p['recall'],p['n_probe'],cfg['M'],cfg['K'],p['gates'],p['domain_bits'],p['planned_machine_memory_GiB'],p['query_ms'] if p['fits_512GiB_scenario'] else None,p['vo_kb'] if p['fits_512GiB_scenario'] else None,p['verify_ms'] if p['fits_512GiB_scenario'] else None,*p['sensitivity']['query_ms'],*p['sensitivity']['verify_ms'],*p['sensitivity']['vo_kb'],'性能为条件预测；召回为20查询搜索校准'] for p in predictions])
for row in data.iter_rows(min_row=2):
 row[1].number_format=row[2].number_format='0.0%'
notes=tab('DEEP V3DB K955说明',['项目','说明'],[
 ['用户选择','每个K分别调参，使实际搜索校准Recall@k尽量达到至少95.5%，不是固定n_probe=16。'],
 ['数据与量化','DEEP1M，N=1000000，D=96，n_list=1024，capacity=2048，PQ M=96、K=1024。只调整n_probe，不重新训练量化器。'],
 ['校准','同20条query100..119，以原始向量精确L2 top-k为真值；整数PQ距离排序，无原向量精排。密扫n_probe1..64。调参与报告为相同查询集，不声明独立留出性能。'],
 ['离散召回','k=1的20条查询只能每次增加5个百分点；k=5为1个百分点。因此≥95.5%的实际校准分别为100%、96%，其余四点95.5%。不能把所有实际召回强改95.5%。'],
 ['非单调','K=5需probe34，K=10只需16，因为两个Recall@k定义不同。达到阈值所需探测预算并不必然随k增长，性能不应强制单调。'],
 ['模型','原512GiB模型：预构建门数回归、nextpow2域、B logB证明/编译、关系规模log规模见证成本；公开输入与VO按实际k计数。'],
 ['k的影响','作者set_based.rs中top_k用于注册公开输出，候选排序与评分仍覆盖同一批全部候选；预构建静态门数沿用k10结构模型。最终构建/公共输入哈希的不确定性由原10%规划余量覆盖，非逐k完整出证验证。'],
 ['资源','按512GiB单线程条件预测，内存=1.5×大档RSS/行p95×带10%门数余量的二次幂域+32GiB。规划余量不是数学保证，超过512GiB则主表留空。'],
 ['标记','只修改DEEP主表V3DB的18个K曲线性能格、3个方法标签与顶部说明；其他方法、默认表、Recall与规模曲线保持原值。主表†预测不能表述为实测。'],
 ['系数','ANNProof的1.2/2/1.5系数不适用于V3DB；本次不乘该系数。'],
 ['备份',str(BACKUP)],['校准与拟合来源',str(S)],['历史来源','旧页保留为历史；本次18格以本页、预测明细与单元格批注为准。']])
tab('DEEP V3DB K955修订位置',['单元格','修订前','修订后','指标','k','n_probe','实际校准Recall@k'],[[x['cell'],x['before'],x['after'],x['metric'],x['source']['k'],x['source']['n_probe'],x['source']['recall']] for x in changes])
for pos,title in enumerate(['DEEP V3DB K955说明','DEEP V3DB K955预测','DEEP V3DB K955修订位置'],4):w.move_sheet(title,offset=pos-w.sheetnames.index(title))
allowed={x['cell'] for x in changes}|{'A1','D107','D119','D130'};preserved=0
for os in old:
 for row in os:
  for c in row:
   if os.title=='DEEP1M' and c.coordinate in allowed:continue
   assert c.value==w[os.title][c.coordinate].value,(os.title,c.coordinate);preserved+=1
assert sha(SOURCE)==sha(BACKUP)
tmp=S/'work9_deep_v3db_k955.xlsx';w.save(tmp);saved=openpyxl.load_workbook(tmp)
for x in changes:
 v=saved['DEEP1M'][x['cell']].value
 assert v==x['after'] or isinstance(v,(int,float)) and math.isclose(v,x['after'],rel_tol=1e-13,abs_tol=1e-9)
shutil.copy2(tmp,SOURCE)
write(S/'predictions.json',predictions);write(S/'cell-changes.json',changes);write(S/'APPLIED.json',dict(passed=True,numeric_cells_written=len(changes),other_cell_values_preserved=preserved,all_points_fit_512GiB=all(p['fits_512GiB_scenario'] for p in predictions),workbook_sha256=sha(SOURCE),backup_sha256=sha(BACKUP),model_sha256=sha(MROOT/'model.json')))
print('PREDICTIONS',json.dumps([{k:p[k] for k in ['k','n_probe','recall','query_ms','vo_kb','verify_ms','planned_machine_memory_GiB']} for p in predictions]),flush=True)
print('APPLIED',len(changes),'forecast cells; preserved',preserved,'other cell values',flush=True)
