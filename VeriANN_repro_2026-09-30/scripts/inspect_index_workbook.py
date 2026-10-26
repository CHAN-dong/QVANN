from pathlib import Path
import openpyxl
root=Path(__file__).resolve().parents[1]
w=openpyxl.load_workbook(root.parent/'work9.xlsx')
for s in w:
    if s.title not in ['SIFT1M','NYTimes-256','GloVe-100','DEEP1M'] and s.max_row>300:continue
    print('SHEET',s.title,s.max_row,s.max_column)
    for row in s:
        if any(isinstance(c.value,str) and any(t in c.value.lower() for t in ['index','construct','索引']) for c in row):
            r=row[0].row
            for rr in s.iter_rows(min_row=max(1,r-1),max_row=min(s.max_row,r+7),max_col=16):print([(c.coordinate,c.value) for c in rr if c.value is not None])
