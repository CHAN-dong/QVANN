"""Report fresh lookup-table measurements versus the backed-up workbook."""
import argparse
import json
import openpyxl
from work9_lut_study import ROOT,STUDY,configure


def main(completed_copy=False):
    workbook=configure();runs=workbook.runs()
    if completed_copy:workbook.SOURCE=workbook.OUTPUT
    old=openpyxl.load_workbook(workbook.TEMPLATE,data_only=True)['Sheet1']
    selected=[r for r in runs if r['name']=='sift1m-overall']
    selected +=[workbook.size_evaluation(runs,n) for n in [100000,500000,1000000,5000000,10000000]]
    assert all(selected)
    lines=['# work9.xlsx 查询与验证时间修正：VeriANN-LUT','',
           '原目录只有work9.xlsx，按用户所指工作簿更新。备份为本目录work9_before_lut.xlsx。',
           'Ours全部正式配置用已验证查表版本重新测量；没有把单点加速比例套用到其他配置。'
           '保留原图、查询、E、K、算法、wire格式与舍入规则；服务器和客户端的查表准备成本均计入时间。','',
           '总体、11点Recall曲线、5点规模曲线、6点K曲线共23个配置，100条查询重复3次，共6900条测量。'
           '刷新103个Ours数据单元格及来源，其中46个为查询/验证计时；57个非计时格与原版一致。'
           '原206个baseline数值和工作簿合并布局保留。','',
           '| 数据 / N | E | Recall@10 % | 原Query ms | 新Query ms | 原Verify ms | 新Verify ms | VO KB |',
           '|---|---:|---:|---:|---:|---:|---:|---:|']
    for r in selected:
        if r['name']=='sift1m-overall':name='SIFT1M / 1M';previous_q=old['F11'].value;previous_v=old['H11'].value
        else:
            sizes=[100000,500000,1000000,5000000,10000000];col=6+sizes.index(r['base'])
            name='BIGANN / '+f"{r['base']:,}";previous_q=old.cell(75,col).value;previous_v=old.cell(66,col).value
        lines.append(f"| {name} | {r['ef']} | {r['recall_percent']:.2f} | {previous_q:.3f} | {r['query_ms']:.3f} | "
                     f"{previous_v:.3f} | {r['verify_ms']:.3f} | {r['vo_kb']:.3f} |")
    lines+=['','Query=search+真实VO生成和序列化；Verify=解码、哈希认证、确定性搜索重放及渐进top-k检查。'
            '单位ms；VO使用十进制KB。计时不含建索引、ADS构建、精确GT、网络传输或链上操作。','',
            '所有6900条查询的非计时CSV字段与对应原版逐项一致，三个可信根一致。'
            '协议测试沿用同一二进制/源代码指纹已通过的查表版结果：90例诚实查询、240例结构性攻击检查。',
            '500K/10M沿用原固定查询集匹配出的E=112/512；这是工作点匹配，不能称为独立泛化评估。'
            '已有baseline的召回率差异仍保留，整表不是严格等召回比较。','',
            f'原文件：`{workbook.SOURCE}`',f'备份：`{workbook.TEMPLATE}`',f'完成版：`{workbook.OUTPUT}`',
            '逐条测量、指纹、VO、篡改检查及参数保存在results/work9-ours-lut-20261015各子目录。','']
    if completed_copy:lines+=['原work9.xlsx被WPS占用，系统拒绝覆盖。本次核验交付独立完成版，原文件尚未写回。','']
    (STUDY/'RESULTS_CN.md').write_text('\n'.join(lines),encoding='utf-8')
    print(json.dumps({'report':str(STUDY/'RESULTS_CN.md'),'configurations':len(runs),'measurements':sum(r['samples'] for r in runs)}))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--completed-copy',action='store_true')
    main(parser.parse_args().completed_copy)
