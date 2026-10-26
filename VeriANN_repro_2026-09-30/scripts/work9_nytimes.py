"""Replace the active GIST panel, retaining archived evidence; sequential real runs."""
import argparse
import csv
import hashlib
import json
import math
import os
from pathlib import Path
import shutil
import signal
import statistics
import subprocess
import sys
import time

ROOT=Path(__file__).resolve().parents[1]
STUDY=ROOT/'experiments/work9-nytimes-20261017'
RESULTS=ROOT/'results/work9-nytimes-20261017'
FOLDER=ROOT/'dataset/nytimes256'
SOURCE=ROOT.parent/'work9.xlsx'
URL='https://ann-benchmarks.com/nytimes-256-angular.hdf5'
NAME='NYTimes-256'
KEY='nytimes256'
SOURCE_N=290000
N=289761  # Exclude exactly the 239 source base rows with undefined angular norm.
D=256
TARGETS=[.60,.65,.70,.75,.80,.85,.90,.95,.97,.98,.99]


def sha(p):
    with Path(p).open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()


def write(p,obj):
    p.parent.mkdir(parents=True,exist_ok=True)
    temp=p.with_suffix(p.suffix+'.tmp')
    temp.write_text(json.dumps(obj,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
    temp.replace(p)


def event(**obj):
    obj['time']=time.strftime('%Y-%m-%dT%H:%M:%S%z')
    STUDY.mkdir(parents=True,exist_ok=True)
    with (STUDY/'events.jsonl').open('a') as f:f.write(json.dumps(obj)+'\n')
    print(json.dumps(obj),flush=True)


def archive():
    STUDY.mkdir(parents=True,exist_ok=True)
    backup=STUDY/'work9_before_nytimes.xlsx'
    if not backup.exists():shutil.copy2(SOURCE,backup)
    if not (STUDY/'gist_archive_manifest.json').exists():
        records=[]
        for base in [ROOT/'experiments',ROOT/'results']:
            for p in base.rglob('*'):
                if p.is_file() and 'gist' in str(p.relative_to(base)).lower():
                    r=dict(path=str(p.relative_to(ROOT)),bytes=p.stat().st_size)
                    if p.suffix in ['.csv','.json','.jsonl','.txt']:r['sha256']=sha(p)
                    records.append(r)
        write(STUDY/'gist_archive_manifest.json',dict(status='archived in place; excluded from active NYTimes evaluation',
            workbook=str(backup),workbook_sha256=sha(backup),files=records))


def prepare():
    import h5py
    import numpy as np
    from prepare_extra_datasets import write_vectors,finalize
    FOLDER.mkdir(parents=True,exist_ok=True)
    raw=FOLDER/'raw/nytimes-256-angular.hdf5'
    if not raw.exists():
        part=raw.with_suffix(raw.suffix+'.part')
        if not part.exists():
            subprocess.run(['curl','--fail','--location','--retry','3','--output',str(part),URL],check=True)
        # A partial download is never accepted based on the filename alone.
        with h5py.File(part,'r') as h:
            assert h['train'].shape==(SOURCE_N,D) and h['test'].shape==(10000,D)
            for key in ['train','test','neighbors']:
                for start in range(0,len(h[key]),8192):assert np.isfinite(h[key][start:start+8192]).all()
        part.rename(raw)
    if (FOLDER/'provenance.json').exists():
        record=json.loads((FOLDER/'provenance.json').read_text())
        assert record['source']['sha256']==sha(raw)
        for item in record['files'].values():assert sha(FOLDER/item['file'])==item['sha256']
        return
    record=dict(name=KEY,source=dict(url=URL,sha256=sha(raw),bytes=raw.stat().st_size),
        metric='cosine via unit-normalized squared L2',original_metric='angular',normalized=True,
        benchmark_metric_argument='l2',files={},
        ground_truth='Recompute exact top-k on prepared normalized float32 base for each N and queries100..119; supplied source neighbors retained as reference only.')
    with h5py.File(raw,'r') as h:
        assert h['train'].shape==(SOURCE_N,D) and h['test'].shape==(10000,D)
        assert h['neighbors'].shape==(10000,100)
        base=h['train'][:]
        ids=np.flatnonzero(np.sum(base.astype(np.float64)**2,axis=1)>0)
        assert len(ids)==N
        np.save(FOLDER/'base_original_ids.npy',ids)
        record['files']['train']=write_vectors(base[ids],FOLDER/'base.fvecs','<f4',True)
        queries=h['test'][:].astype('<f4')
        norms=np.sqrt(np.sum(queries.astype(np.float64)**2,axis=1))
        zeros=np.flatnonzero(norms==0)
        assert not set(zeros)&set(range(100,120))
        norms[norms==0]=1
        queries*=np.asarray(1/norms,dtype='<f4')[:,None]
        record['files']['test']=write_vectors(queries,FOLDER/'query.fvecs','<f4',False)
        record['zero_policy']=dict(base_zero_count=SOURCE_N-N,
            base_original_ids_file='base_original_ids.npy',base_id_map_sha256=sha(FOLDER/'base_original_ids.npy'),
            query_zero_rows=zeros.tolist(),query_policy='preserve original query row IDs; zero queries excluded from evaluation; IDs100..119 all nonzero',
            source_neighbors='Original HDF5 neighbors refer to unfiltered original IDs; NEVER use as prepared ground truth')
        record['hdf5_attributes']={k:str(v) for k,v in h.attrs.items()}
    finalize(FOLDER,record)
    event(action='dataset_prepared',N=N,d=D,source_sha256=sha(raw))


def row_map(r):
    return r+sum(n for after,n in [(11,1),(24,2),(34,2),(44,2),(56,1),(66,1),(75,1),(86,1),(96,1),(105,1)] if r>after)


def overlay(wb):
    """Called on every V3DB workbook refresh; never copy old GIST numbers to NYTimes."""
    import openpyxl
    from openpyxl.comments import Comment
    import work9_multidata as common
    if 'GIST1M' in wb:
        ws=wb['GIST1M'];ws.title=NAME
    elif NAME in wb:ws=wb[NAME]
    else:raise RuntimeError('Missing replacement template')
    mapped=any(ws.cell(12,c).value=='V3DB' for c in [4,5])
    rm=row_map if mapped else lambda r:r
    # Remove ALL old baseline, Ours and controlled-ablation measurements.
    addresses=[]
    for a in common.coordinates():
        col,row=openpyxl.utils.cell.coordinate_from_string(a)
        addresses.append(f'{col}{rm(row)}')
    for addr in addresses:
        ws[addr]='待测';ws[addr].comment=None
    import run_work9_ablation as abl
    header=abl.table(ws)
    for r in range(header+1,header+7):
        for c in range(6,10):ws.cell(r,c,'待测').comment=None
    # Added V3DB rows are separately identified through the provenance table.
    if 'V3DB单元格来源' in wb:
        for cells in wb['V3DB单元格来源'].iter_rows(min_row=2):
            if cells[0].value=='GIST1M' and isinstance(cells[1].value,str):
                ws[cells[1].value]='待测';ws[cells[1].value].comment=None
    for row in ws:
        for cell in row:
            if isinstance(cell,openpyxl.cell.cell.MergedCell):continue
            cell.comment=None
            if isinstance(cell.value,str):
                cell.value=cell.value.replace('GIST1M',NAME).replace('gist1m',KEY)
                if 'N=1,000,000' in cell.value:cell.value=cell.value.replace('N=1,000,000',f'N={N:,}')
                cell.value=cell.value.replace('d=960','d=256')
    for old in [54,55,56,64,65,66,73,74,75]:
        for c in range(7,11):ws.cell(rm(old),c,'不适用')
    if mapped:
        for old in [56,66,75]:
            for c in range(7,11):ws.cell(row_map(old)+1,c,'不适用')
    ws['A4']=f'NYTimes-256: N={N:,} nonzero base rows; normalized angular; query IDs100..119 x3; single thread; GIST results archived externally.'
    # Blank archived GIST rows without shifting chart references for other datasets.
    for tab in wb:
        if tab is ws:continue
        for row in tab:
            if any(isinstance(c.value,str) and (c.value.lower() in ['gist1m','gist-960-euclidean'] or c.value.lower().startswith('gist1m-')) for c in row[:3]):
                for c in row:
                    if not isinstance(c,openpyxl.cell.cell.MergedCell):c.value=None;c.comment=None
        def chart_text(chart):
            try:return ''.join(t.t for p in chart.title.tx.rich.p for t in p.r)
            except (AttributeError,TypeError):return ''
        tab._charts=[ch for ch in tab._charts if 'GIST' not in chart_text(ch).upper()]
    provenance=[];measurements=[];curvepoints=[]
    def put(row,col,record,metric,method):
        value=record['recall']*100 if metric=='recall_percent' else record[metric]
        cell=ws.cell(row,col,value);cell.number_format='0.000'
        cell.comment=Comment(f"NYTimes-256 N={record.get('N',N)}, d=256; {method}; normalized angular; queries100..119 x3; single thread. Actual recall={record['recall']}; source={record['output']}/per_query.csv",'NYTimes实测来源')
        provenance.append([NAME,cell.coordinate,value,metric,method,record['output']+'/per_query.csv'])
    for f in sorted(STUDY.glob('evaluation-*.json')):
        job=json.loads(f.read_text());records=common.measurements(Path(job['output']))
        if job.get('filter'):records=[r for r in records if all(r[k]==v for k,v in job['filter'].items())]
        for r in records:
            assert r['samples']==60
            r['N']=job['base'];measurements.append([job['method'],job['purpose'],job['base'],r['k'],r['ef'],r['probes'],r['recall'],r['query_ms'],r['verify_ms'],r['vo_kb'],60,r['output']])
        m=job['method']
        if job['purpose']=='curve':
            for col,p in enumerate(job['selection'],6):
                r=next(r for r in records if (r['ef'],r['probes'])==(p['ef'],p['probes']))
                for (rr,mr),metric in zip(common.CURVE_ROWS[m],['vo_kb','verify_ms','query_ms']):
                    put(rm(rr),col,r,'recall_percent',m);put(rm(mr),col,r,metric,m)
                curvepoints.append([m,p['target'],r['recall'],r['query_ms'],r['verify_ms'],r['vo_kb']])
            p=job['selection'][7];r=next(r for r in records if (r['ef'],r['probes'])==(p['ef'],p['probes']))
            for c,metric in zip(range(5,9),['recall_percent','query_ms','vo_kb','verify_ms']):put(rm({'VLAH':9,'ANNProof':10,'Ours':11}[m]),c,r,metric,m)
        elif job['purpose']=='k':
            for col,k in enumerate([1,5,10,20,50,100],6):
                r=next(r for r in records if r['k']==k)
                for row,metric in zip(common.K_ROWS[m],['vo_kb','verify_ms','query_ms']):put(rm(row),col,r,metric,m)
        elif job['purpose']=='size':
            assert job['base']==100000 and len(records)==1
            for row,metric in zip(common.SIZE_ROWS[m],['vo_kb','verify_ms','query_ms']):put(rm(row),6,records[0],metric,m)
    for f in sorted(STUDY.glob('ablation-evaluation-*.json')):
        for r in json.loads(f.read_text()):
            for c,metric in [(6,'recall'),(7,'query_ms'),(8,'vo_kb'),(9,'verify_ms')]:put(header+r['variant'],c,dict(r,N=N),metric,r['name'])
    # Actual V3DB data will be filled independently once its new model is tuned.
    fill_v3db(wb,ws,put,mapped)
    for title,rows in [('NYTimes单元格来源',[['sheet','cell','value','metric','method','raw CSV'],*provenance]),
                       ('NYTimes实测记录',[['method','purpose','N','k','ef/E','probes','recall','query ms','verify ms','VO KB','samples','output'],*measurements]),
                       ('NYTimes召回曲线',[['method','target recall','actual recall','query ms','verify ms','VO KB'],*curvepoints])]:
        if title in wb:del wb[title]
        tab=wb.create_sheet(title)
        for row in rows:tab.append(row)
        tab.freeze_panes='A2'
        if title=='NYTimes召回曲线' and curvepoints:
            from openpyxl.chart import ScatterChart,Series,Reference
            for i,(label,col) in enumerate([('Query ms',4),('Verify ms',5),('VO KB',6)]):
                chart=ScatterChart();chart.title='NYTimes-256: '+label
                chart.x_axis.title='Actual Recall@10';chart.y_axis.title=label
                chart.x_axis.numFmt='0%';chart.x_axis.scaling.min=.6;chart.x_axis.scaling.max=1
                chart.width=22;chart.height=12
                for method in ['ANNProof','VLAH','Ours']:
                    idx=[j+2 for j,r in enumerate(curvepoints) if r[0]==method]
                    if idx:
                        chart.series.append(Series(Reference(tab,min_col=col,min_row=min(idx),max_row=max(idx)),Reference(tab,min_col=3,min_row=min(idx),max_row=max(idx)),title=method))
                tab.add_chart(chart,f'H{1+i*24}')
    title='NYTimes实验说明'
    if title in wb:del wb[title]
    tab=wb.create_sheet(title)
    for row in [['项目','说明'],['替换','用户选定NYTimes-256；GIST1M未删除，归档见work9-nytimes-20261017/work9_before_nytimes.xlsx及gist_archive_manifest.json。'],
                ['数据','公开ANN-Benchmarks NYTimes-256 angular，源290000 base；统一剔除239个零向量后289761条。查询行号不变，所用20条均非零；单位归一化为float32，以平方L2排序。原ID映射及零向量审计保留。'],
                ['规模','保留原100K/500K/1M/5M/10M规模列；仅100K适用，其余不适用。总体/召回/k/受控消融用全部289761条清洗基库。'],
                ['协议','20条查询IDs100..119×3次；查询、生成VO、验证单线程；同工作负载调参，不是独立留出准确率估计。'],
                ['公平性','所有方法重测；ANNProof保持原完整优化配置，不能以消融版本替换主baseline。无法保证Ours在新数据上占优。'],
                ['历史说明','辅助表中原GIST数据行清空，原始历史完整保留于外部归档。新数据不得由旧GIST结果重命名取得。'],
                ['进度','待测仅由完整60次且验证通过的实测记录替代；失败或目标不可达明确标注。']]:tab.append(row)


def fill_v3db(wb,ws,put,mapped):
    from openpyxl.comments import Comment
    if not mapped:return
    selected_file=ROOT/'experiments/work9-v3db-20261016/selected-models.json'
    selected=json.loads(selected_file.read_text()).get(KEY) if selected_file.exists() else None
    if not selected:return
    model=ROOT/'V3DB-26VLDB/work9-cache'/selected['model']
    cal=json.loads((model/'calibration.json').read_text())
    points=[p for p in cal['points'] if p['k']==10]
    choices=[p for p in points if p['recall']>=.95-1e-12]
    overall=min(choices,key=lambda p:p['n_probe']) if choices else None
    def result(modelname,p,k):
        f=ROOT/'results/work9-v3db-20261016'/f'{modelname}-p{p}-k{k}-q20-r3-t1/result.json'
        return json.loads(f.read_text()) if f.exists() else None
    def setmetric(row,col,p,k,metric,modelname=None):
        r=result(modelname or selected['model'],p,k)
        if r and r['status']=='passed':
            put(row,col,r,metric,'V3DB')
        elif r:
            ws.cell(row,col,'内存不足(50GiB上限)' if r.get('actual_attempt') else '超预算未测')
    if overall:
        p=overall['n_probe'];ws.cell(12,5,overall['recall']*100)
        ws.cell(12,5).comment=Comment(f'仅20条查询的搜索校准召回率；不是60次证明均完成的声明。来源：{model}/calibration.json','V3DB校准')
        for c,m in [(6,'query_ms'),(7,'vo_kb'),(8,'verify_ms')]:setmetric(12,c,p,10,m)
        for old,m in [(86,'vo_kb'),(96,'verify_ms'),(105,'query_ms')]:
            for c,k in enumerate([1,5,10,20,50,100],6):setmetric(row_map(old)+1,c,p,k,m)
        base=next((row[0].row for row in ws if row[3].value=='V3DB (baseline)'),None)
        if base:
            ws.cell(base,6,overall['recall'])
            for c,m in [(7,'query_ms'),(8,'vo_kb'),(9,'verify_ms')]:setmetric(base,c,p,10,m)
    for c,target in enumerate(TARGETS,6):
        eligible=[p for p in points if p['recall']>=target-1e-12]
        for old,m in [(24,'vo_kb'),(34,'verify_ms'),(44,'query_ms')]:
            row=row_map(old)+1
            if not eligible:ws.cell(row,c,'未达目标');ws.cell(row+1,c,'未达目标');continue
            p=min(eligible,key=lambda p:p['n_probe']);ws.cell(row,c,p['recall']*100);setmetric(row+1,c,p['n_probe'],10,m)
            ws.cell(row,c).comment=Comment(f'仅20条查询的搜索校准召回率；来源：{model}/calibration.json。时间和VO只取完整60次验证通过的结果。','V3DB校准')
    for f in (ROOT/'V3DB-26VLDB/work9-cache').glob(f'{KEY}-100000*/calibration.json'):
        small=json.loads(f.read_text());cfg=small['model']['config']
        if cfg['M']!=selected['M'] or cfg['K']!=selected['K']:continue
        eligible=[p for p in small['points'] if p['k']==10 and p['recall']>=.95-1e-12]
        if eligible:
            p=min(eligible,key=lambda p:p['n_probe'])
            for old,m in [(56,'vo_kb'),(66,'verify_ms'),(75,'query_ms')]:setmetric(row_map(old)+1,6,p['n_probe'],10,m,f.parent.name)


def refresh():
    sys.path.insert(0,str(ROOT/'V3DB-26VLDB/scripts'))
    import work9_fill
    result=work9_fill.refresh()
    try:shutil.copy2(work9_fill.OUTPUT,STUDY/'work9_nytimes.xlsx')
    except PermissionError:pass
    return result


def run_methods():
    import run_work9_multidata as runner
    import work9_multidata as common
    # Reuse measured protocol logic, isolate all filenames from old studies.
    for module in [runner,common]:
        module.STUDY=STUDY;module.RESULTS=RESULTS
        module.DATASETS={KEY:(NAME,N,D)}
    runner.checkpoint=refresh
    w=runner.width(KEY)
    for m in ['ANNProof','VLAH','Ours']:
        p=runner.calibrate(KEY,m,N,.99,w)
        runner.evaluate(KEY,m,N,'curve',p,w)
        runner.evaluate(KEY,m,N,'k',[p[7]],w)
        p=runner.calibrate(KEY,m,100000,.95,w)
        runner.evaluate(KEY,m,100000,'size',p,w)


def run_ablations():
    import run_work9_ablation as ab
    ab.STUDY=STUDY;ab.RESULTS=RESULTS/'controlled-ablation'
    original=ab.launch
    def launch(*args,**kwargs):
        kwargs.setdefault('n',N)
        return original(*args,**kwargs)
    ab.launch=launch
    for pair in ab.PAIRS:
        f=STUDY/f'ablation-evaluation-{pair[0]}-{pair[1]}.json'
        if not f.exists():
            point=ab.calibrate(KEY,pair)
            out=ab.launch(KEY,pair,[point['parameter']],f'evaluation-{pair[0]}-{pair[1]}')
            rows=ab.audit_pair(out,pair,point)
            for r in rows:r['dataset']=KEY
            write(f,rows)
        refresh()


def run_v3db():
    sys.path.insert(0,str(ROOT/'V3DB-26VLDB/scripts'))
    import work9_calibrate as ca
    selected_file=ca.STUDY/'selected-models.json'
    selection=json.loads(selected_file.read_text())
    if KEY not in selection:
        trials=[]
        for m,k in [(8,256),(64,256),(128,256),(256,256),(256,512),(256,1024)]:
            name=f'{KEY}-{N}'+(f'-M{m}' if m!=8 else '')+(f'-K{k}' if k!=256 else '')
            f=ca.CACHE/name/'calibration.json'
            if not f.exists():
                with (STUDY/f'v3db-tune-M{m}-K{k}.log').open('a') as log:
                    subprocess.run([sys.executable,str(ca.REPRO/'scripts/work9_calibrate.py'),'--dataset',KEY,'--M',str(m),'--K',str(k)],stdout=log,stderr=subprocess.STDOUT,check=True)
            cal=json.loads(f.read_text());best=max((p for p in cal['points'] if p['k']==10),key=lambda p:p['recall'])
            trials.append(dict(M=m,K=k,model=name,best_grid_recall10=best['recall'],best_probe=best['n_probe']))
            if best['recall']>=.99-1e-12:break
        eligible=[p for p in trials if p['best_grid_recall10']>=.99-1e-12]
        choice=min(eligible,key=lambda p:(p['M'],p['K'])) if eligible else max(trials,key=lambda p:p['best_grid_recall10'])
        selection[KEY]=dict(choice,trials=trials,query_count=20,repeats=3,policy='same tuning policy as other datasets; no held-out accuracy claim')
        write(selected_file,selection)
    refresh()
    subprocess.run([sys.executable,str(ca.REPRO/'scripts/work9_run_all.py'),'--sampling','20x3','--proof-threads','1','--datasets',KEY],check=True)


def main():
    p=argparse.ArgumentParser();p.add_argument('action',choices=['archive','prepare','fill','run','start']);p.add_argument('--wait-pid',type=int)
    p.add_argument('--take-over-pid',type=int)
    args=p.parse_args();STUDY.mkdir(parents=True,exist_ok=True);RESULTS.mkdir(parents=True,exist_ok=True)
    if args.action=='start':
        command=[sys.executable,str(Path(__file__).resolve()),'run']
        if args.take_over_pid:command+=['--take-over-pid',str(args.take_over_pid)]
        env=dict(os.environ,PYTHONUNBUFFERED='1',OMP_NUM_THREADS='1',OPENBLAS_NUM_THREADS='1',RAYON_NUM_THREADS='1')
        with (STUDY/'runner.log').open('a') as log:
            child=subprocess.Popen(command,env=env,stdin=subprocess.DEVNULL,stdout=log,stderr=subprocess.STDOUT,start_new_session=True)
        write(STUDY/'runner_pid.json',dict(pid=child.pid,command=command))
        print(json.dumps(dict(pid=child.pid,command=command)));return
    if args.action=='archive':archive();return
    if args.action=='fill':refresh();return
    if args.take_over_pid:
        proc=Path(f'/proc/{args.take_over_pid}/cmdline')
        if proc.exists():
            cmd=proc.read_bytes().split(b'\0')
            assert any(b'work9_measure.py' in x for x in cmd),cmd
            model=cmd[cmd.index(b'--model')+1].decode();probe=int(cmd[cmd.index(b'--probe')+1]);k=int(cmd[cmd.index(b'--k')+1])
            folder=ROOT/'results/work9-v3db-20261016'/f'{model}-p{probe}-k{k}-q20-r3-t1'
            csvfile=folder/'per_query.csv'
            def count():
                return len(list(csv.DictReader(csvfile.open()))) if csvfile.exists() else 0
            before=count();last=0
            while proc.exists() and proc.read_bytes():
                if (folder/'result.json').exists() or count()>before:
                    # Stop only after the in-flight sample flushed successfully;
                    # the next job invocation resumes all completed query/repeat pairs.
                    os.kill(args.take_over_pid,signal.SIGTERM)
                    event(action='proof_checkpoint_preserved',model=model,completed=count(),resume='after NYTimes CPU benchmarks')
                    break
                if time.monotonic()-last>60:
                    event(action='waiting_query_boundary',pid=args.take_over_pid,completed=before);last=time.monotonic()
                time.sleep(.2)
            for _ in range(100):
                if not proc.exists() or not proc.read_bytes():break
                time.sleep(.1)
            else:raise RuntimeError('Existing worker did not exit; refusing concurrent measurements')
    if args.wait_pid:
        while Path(f'/proc/{args.wait_pid}').exists():
            event(action='waiting_existing_proof',pid=args.wait_pid);time.sleep(60)
    archive();prepare()
    write(STUDY/'replacement.json',dict(old_dataset='gist1m',new_dataset=KEY,N=N,d=D,
        selection='explicit user choice NYTimes-256',queries=20,repeats=3,threads=1,
        source=URL,archive=str(STUDY/'work9_before_nytimes.xlsx'),old_results='archived, not destroyed'))
    refresh()
    if args.action=='prepare':return
    run_methods();run_ablations();run_v3db();refresh()
    write(STUDY/'RUN_FINISHED.json',dict(status='NYTimes requested jobs executed; inspect cell limitations separately'))
    # Resume the already-authorized remaining datasets with checkpoint reuse.
    subprocess.run([sys.executable,str(ROOT/'V3DB-26VLDB/scripts/work9_run_all.py'),
        '--sampling','20x3','--proof-threads','1','--datasets','sift1m','glove100','deep1m'],check=True)


if __name__=='__main__':main()
