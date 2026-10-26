"""Actual low-recall ANNProof rerun, never replaced by a dominating graph."""
import csv,hashlib,json,math,shutil,statistics
from pathlib import Path
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill,Font,Alignment
import work9_resource_matched as u
R=u.ROOT;S=R/'experiments/work9-nytimes-low-retest-20261021';O=R/'results/work9-nytimes-low-retest-20261021'
SOURCE=R.parent/'work9.xlsx';BACKUP=S/'work9_before_nytimes_low_retest.xlsx'
u.STUDY=S;u.RESULTS=O;S.mkdir(parents=True,exist_ok=True)
assert not (S/'APPLIED.json').exists(),'Already applied'
if not BACKUP.exists():shutil.copy2(SOURCE,BACKUP)
assert u.sha(SOURCE)==u.sha(BACKUP),'External workbook changes detected'
binary=R/'ANNProof-24FGCS/build/annproof';efs=[26,51,80,114]
cmd=[binary,'--dataset',R/'dataset/nytimes256','--base',289761,'--output',O,'--cache',R/'experiments/work9-annproof-low-adjusted-20261021/cache','--queries',20,'--query-offset',100,'--repeats',3,'--build-threads',1,'--M',4,'--ef-construction',40,'--k',10,'--ef-search',','.join(map(str,efs)),'--probes',1,'--metric','l2','--proof','multiproof']
if not (S/'MEASURED.json').exists():
 u.write(S/'protocol.json',dict(command=list(map(str,cmd)),binary_sha256=u.sha(binary),threads=1,scope='fresh full-query/proof/verification run; graph fixed M4/efc40; no trend prediction',query_ids=list(range(100,120)),repeats=3,source_sha256={str(p.relative_to(R)):u.sha(p) for p in [R/'ANNProof-24FGCS/annproof.cpp',R/'src/paper_repro_common.hpp',R/'src/verifiable_ann_common.hpp']}))
 u.command(cmd,S/'measurement.log',bounded=True,timeout=3600)
 rows=list(csv.DictReader((O/'per_query.csv').open()));attacks=list(csv.DictReader((O/'tamper_tests.csv').open()))
 assert len(rows)==240 and all(r['verified']=='1' for r in rows)
 assert len(attacks)==8 and all(r['rejected']=='1' for r in attacks)
 assert (O/'query0.vo').stat().st_size==int(rows[0]['vo_bytes'])
 points=[]
 for ef,target in zip(efs,[.4,.45,.5,.55]):
  rr=[r for r in rows if int(r['ef_search'])==ef]
  assert len(rr)==60 and {(int(r['query']),int(r['repeat'])) for r in rr}=={(q,t) for q in range(100,120) for t in range(3)}
  points.append(dict(ef=ef,target=target,recall=statistics.fmean(float(r['recall_at_k']) for r in rr),query_ms=statistics.fmean(float(r['search_ms'])+float(r['vo_generation_ms']) for r in rr),vo_kb=statistics.fmean(int(r['vo_bytes'])/1000 for r in rr),verify_ms=statistics.fmean(float(r['verification_ms']) for r in rr),samples=60))
 assert all(abs(p['recall']-t)<=.025+1e-9 for p,t in zip(points,[.4,.45,.5,.55]))
 for key in ['recall','vo_kb']:
  assert all(b[key]>=a[key] for a,b in zip(points,points[1:]))
 u.write(S/'MEASURED.json',dict(points=points,csv=str(O/'per_query.csv'),csv_sha256=u.sha(O/'per_query.csv'),tamper_rejections=8,wire_size_checked=True))
m=u.read(S/'MEASURED.json');print(json.dumps(m,ensure_ascii=False),flush=True)
w=openpyxl.load_workbook(BACKUP);old=openpyxl.load_workbook(BACKUP);s=w['NYTimes-256'];events=[];factors={'query_ms':1.2,'vo_kb':2.,'verify_ms':1.5}
def put(row,col,value,p,reused=False):
 c=s.cell(row,col);reason='本次重新实测，固定M=4、efConstruction=40；N=289761，D=256，L2，k=10，query100..119，20条×3次，单线程，完整multiproof和验证。'
 if reused:reason+=' 该低目标列复用已测约40%点，并非实测达到列头15–35%。'
 reason+=' 召回率为实际值；性能主表按用户既定query×1.2、VO×2、verify×1.5调整，非原始测量数值，未作趋势拟合。60%起使用不同图配置，不与本段声称固定参数单调。'
 events.append(dict(cell=c.coordinate,before=c.value,after=value,reason=reason,point=p))
 c.value=value;c.comment=Comment(reason+'\n'+json.dumps(p,ensure_ascii=False)+'\n原始CSV='+m['csv']+'\nSHA256='+m['csv_sha256']+'\n\n历史批注（非本次数据来源）：\n'+(c.comment.text if c.comment else ''),'Codex')
 c.fill=PatternFill('solid',fgColor='E2F0D9');c.font=Font(color='000000');c.number_format='0.000'
for col,p in zip(range(11,15),m['points']):
 for r in [23,37,51]:put(r,col,p['recall']*100,p)
 for r,key in [(24,'vo_kb'),(38,'verify_ms'),(52,'query_ms')]:put(r,col,p[key]*factors[key],p)
for col in range(6,11):
 p=m['points'][0]
 for r in [23,37,51]:put(r,col,p['recall']*100,p,True)
 for r,key in [(24,'vo_kb'),(38,'verify_ms'),(52,'query_ms')]:put(r,col,p[key]*factors[key],p,True)
for r in [23,37,51]:s.cell(r,4).value='ANNProof*（分图配置）'
s['A1']=str(s['A1'].value)+' | NYTimes低召回本次覆盖：40–55%目标为固定M4/efc40真实重测，非预测；60%起不同图配置。'
s['A1'].alignment=Alignment(wrap_text=True);s.row_dimensions[1].height=65
a=w.create_sheet('NYTimes低召回重新实测');a.append(['目标Recall@10','实际Recall@10','M','efConstruction','efSearch','原始查询ms','原始VO KB','原始验证ms','调整后查询ms','调整后VO KB','调整后验证ms','样本数','来源'])
for p in m['points']:
 a.append([p['target'],p['recall'],4,40,p['ef'],p['query_ms'],p['vo_kb'],p['verify_ms'],p['query_ms']*1.2,p['vo_kb']*2,p['verify_ms']*1.5,60,m['csv']])
for row in range(2,6):
 for col in [1,2]:a.cell(row,col).number_format='0.0%'
a.append(['说明','84%仅是默认M16图的较低搜索预算结果，不是算法整体召回下限；本次保留低召回点，不再被更高召回配置支配替换。'])
a.append(['口径','原始实测与系数调整后数值分列；不是预测。60%起不同图配置的成本可低于本段，不应用虚构单调值代替。'])
a.append(['验证','240条完整查询均验证成功；8项篡改均被拒绝；实际VO文件大小核对通过。'])
a.append(['备份',str(BACKUP)]);a.freeze_panes='A2';a.auto_filter.ref='A1:M5'
for col in a.columns:a.column_dimensions[col[0].column_letter].width=23
a.column_dimensions['M'].width=80
for row in a:
 for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
b=w.create_sheet('NYTimes低召回修订位置');b.append(['单元格','修订前','修订后','性质'])
for e in events:b.append([e['cell'],e['before'],e['after'],'新实测实际召回/用户指定系数调整性能；非预测'])
allowed={e['cell'] for e in events}|{'A1','D23','D37','D51'};preserved=0
for os in old:
 for row in os:
  for c in row:
   if os.title=='NYTimes-256' and c.coordinate in allowed:continue
   assert c.value==w[os.title][c.coordinate].value,(os.title,c.coordinate);preserved+=1
assert u.sha(SOURCE)==u.sha(BACKUP)
tmp=S/'work9_nytimes_low_retested.xlsx';w.save(tmp);saved=openpyxl.load_workbook(tmp)
for e in events:assert math.isclose(saved['NYTimes-256'][e['cell']].value,e['after'],rel_tol=1e-13,abs_tol=1e-10)
shutil.copy2(tmp,SOURCE)
u.write(S/'changes.json',events);u.write(S/'APPLIED.json',dict(passed=True,cells_written=len(events),unrelated_values_preserved=preserved,points=m['points'],workbook_sha256=u.sha(SOURCE),backup_sha256=u.sha(BACKUP),raw_csv_sha256=m['csv_sha256']))
print('APPLIED',len(events),'cells;',preserved,'unrelated values preserved',flush=True)
