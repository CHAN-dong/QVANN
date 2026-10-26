"""Workbook and evidence utilities for the three additional dataset experiments."""
import csv
from datetime import datetime
import hashlib
import json
import math
from pathlib import Path
import shutil
import statistics
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.chart import ScatterChart, Series, Reference

ROOT=Path(__file__).resolve().parents[1]
STUDY=ROOT/'experiments/work9-multidata-20261015'
RESULTS=ROOT/'results/work9-multidata-20261015'
SOURCE=ROOT.parent/'work9.xlsx'
BACKUP=STUDY/'work9_before_multidata.xlsx'
OUTPUT=STUDY/'work9_completed_multidata.xlsx'
DATASETS={'gist1m':('GIST1M',1000000,960),'glove100':('GloVe-100',1183514,100),'deep1m':('DEEP1M',1000000,96)}
TARGETS=[.60,.65,.70,.75,.80,.85,.90,.95,.97,.98,.99]
METHODS=['ANNProof','VLAH','Ours']
QUERY_COUNT=20
QUERY_OFFSET=100
REPEATS=3
INDEX_M=16
INDEX_EFC=200
CURVE_ROWS={'VLAH':[(19,20),(29,30),(39,40)],'ANNProof':[(21,22),(31,32),(41,42)],'Ours':[(23,24),(33,34),(43,44)]}
SIZE_ROWS={'VLAH':[54,64,73],'ANNProof':[55,65,74],'Ours':[56,66,75]}
K_ROWS={'VLAH':[84,94,103],'ANNProof':[85,95,104],'Ours':[86,96,105]}

def sha(p):
    h=hashlib.sha256()
    with Path(p).open('rb') as f:
        while b:=f.read(8*1024*1024):h.update(b)
    return h.hexdigest()

def event(**kw):
    STUDY.mkdir(parents=True,exist_ok=True);kw['time']=datetime.now().astimezone().isoformat()
    with (STUDY/'events.jsonl').open('a') as f:f.write(json.dumps(kw)+'\n')
    print(json.dumps(kw),flush=True)

def coordinates():
    cells=[f'{c}{r}' for r in [9,10,11] for c in 'EFGH']
    cells +=[f'{openpyxl.utils.get_column_letter(c)}{r}' for r in list(range(19,25))+list(range(29,35))+list(range(39,45)) for c in range(6,17)]
    cells +=[f'{c}{r}' for r in sum(SIZE_ROWS.values(),[]) for c in 'FGHIJ']
    cells +=[f'{c}{r}' for r in sum(K_ROWS.values(),[]) for c in 'FGHIJK']
    assert len(cells)==309;return cells

def deliver():
    marker=STUDY/'last_written_sha256.txt'
    allowed={sha(BACKUP)}
    if marker.exists():allowed.add(marker.read_text().strip())
    if sha(SOURCE) not in allowed:
        event(action='original_changed',completed_copy=str(OUTPUT));return False
    try:shutil.copy2(OUTPUT,SOURCE)
    except PermissionError:
        event(action='workbook_locked',completed_copy=str(OUTPUT));return False
    assert sha(SOURCE)==sha(OUTPUT);marker.write_text(sha(SOURCE)+'\n');return True

def initialize():
    STUDY.mkdir(parents=True,exist_ok=True);RESULTS.mkdir(parents=True,exist_ok=True)
    if OUTPUT.exists():return
    shutil.copy2(SOURCE,BACKUP)
    wb=openpyxl.load_workbook(BACKUP)
    assert all(name not in wb.sheetnames for name,_,_ in DATASETS.values())
    for key,(name,n,d) in DATASETS.items():
        ws=wb.copy_worksheet(wb['Sheet1']);ws.title=name
        for addr in coordinates():
            c=ws[addr];c.value='待测';c.comment=Comment('等待本数据集实测；未沿用 SIFT 数据。','实验记录')
        for row in sum(SIZE_ROWS.values(),[]):
            for col in [9,10]:ws.cell(row,col,'不适用').comment=Comment(f'该规模超过 {name} 已准备库的容量 {n:,}；按用户要求保留规模点。','实验范围')
        for addr in ['D7','D18','D28','D38','D82','D92','D101']:ws[addr]=f'{name} (N={n:,}, d={d})'
        for addr in ['D52','D62','D71']:ws[addr]=f'{name}: real prefix subsets; 5M/10M not applicable'
        ws['A4']='New dataset measurements; PREVIEW: 20 distinct queries, 3 repeats; actual recall and raw source in cell comments.'
    notes=wb.create_sheet('新增数据实验说明')
    for row in [
        ['项目','说明'],['范围','复制原表结构，各测 VLAH、ANNProof、Ours；原 SIFT 主表保留。'],
        ['规模','保留 100K/500K/1M/5M/10M；超过当前数据集容量的 5M、10M 标为不适用。'],
        ['查询','快速预览：同一数据集三个方案均 query IDs 100..119，20 条查询，计时重复3次，单线程。'],
        ['召回率','曲线目标60/65/70/75/80/85/90/95/97/98/99%；主表填实际值。总体与规模以约95%为目标。'],
        ['参数选择','使用同一查询工作负载校准达到召回率的档位；非独立留出召回估计，不插值，不外推。'],
        ['计时','Query=search+真实VO生成/序列化；Verify=解码+认证+重放；Ours 包含 LUT 准备。时间ms，VO KB=bytes/1000。'],
        ['索引','新增三数据集统一 M=16, efConstruction=200。ANNProof/Ours 共用同一缓存 HNSW 图；VLAH 独立 LSH 分桶图。原 SIFT 配置保留。'],
        ['LSH宽度','按各数据集库向量固定抽样的成对L2距离中位数/4确定；不使用查询或真值，所有该数据集实验固定同一宽度。'],
        ['GloVe','使用已归一化的基库与查询，平方L2表示cosine排序；总体/召回/k使用全部1,183,514条。'],
        ['K曲线','k=1/5/10/20/50/100，固定总体95%档参数；baseline ef 至少100以支持k=100；实际各k召回率另记。'],
        ['复现范围','ANNProof/VLAH为已有算法重实现，不宣称作者代码或区块链完整性能复现。'],
        ['进度','待测单元格仅在真实测量、诚实验证和篡改检测通过后填入。']]:notes.append(row)
    notes.column_dimensions['A'].width=20;notes.column_dimensions['B'].width=120
    wb.create_sheet('新增数据测量记录').append(['dataset','method','experiment','N','k','efSearch/E','probes','LSH width','Recall %','Query ms','Verify ms','VO KB','samples','output'])
    wb.create_sheet('新增数据单元格来源').append(['sheet','cell','value','metric','k','efSearch/E','probes','raw CSV'])
    wb.save(OUTPUT);deliver();event(action='workbook_initialized',sheets=[v[0] for v in DATASETS.values()])

def measurements(out):
    assert (out/'PASS.txt').exists()
    groups={}
    with (out/'per_query.csv').open() as f:
        for row in csv.DictReader(f):
            if not QUERY_OFFSET<=int(row['query'])<QUERY_OFFSET+QUERY_COUNT:continue
            assert row['verified']=='1';key=(int(row['k']),int(row['ef_search']),int(row['probes']))
            groups.setdefault(key,[]).append(row)
    records=[]
    for (k,ef,probes),rows in groups.items():
        records.append(dict(k=k,ef=ef,probes=probes,samples=len(rows),output=str(out),
             recall=statistics.fmean(float(r['recall_at_k']) for r in rows),
             query_ms=statistics.fmean(float(r['search_ms'])+float(r['vo_generation_ms']) for r in rows),
             verify_ms=statistics.fmean(float(r['verification_ms']) for r in rows),
             vo_kb=statistics.fmean(int(r['vo_bytes'])/1000 for r in rows)))
    return records

def checkpoint():
    wb=openpyxl.load_workbook(OUTPUT);allcells={};allruns=[];curvepoints=[]
    for name,_,_ in DATASETS.values():wb[name]['A4']='PREVIEW: 20 distinct queries (IDs 100..119), 3 repeats; same subset for all methods.'
    for row in wb['新增数据实验说明'].iter_rows(min_row=2):
        if row[0].value=='查询':row[1].value='快速预览：三个方案统一 query IDs 100..119，20 条查询×3 次重复；已完成的100查询原始记录保留，仅提取相同20条统计。'
    if '预览精度' not in [r[0].value for r in wb['新增数据实验说明'].iter_rows(min_row=2)]:
        wb['新增数据实验说明'].append(['预览精度','20条查询用于快速预览。k=10时平均Recall的离散步长为0.5个百分点；3次重复增加计时样本，不增加独立查询数。正式结论需扩大查询样本。'])
    if '最小档位与起点' not in [r[0].value for r in wb['新增数据实验说明'].iter_rows(min_row=2)]:
        wb['新增数据实验说明'].append(['最小档位与起点','共享60%至99%的目标阈值，图中使用实际Recall。若最小合法参数已超过某些阈值，这些列共享该实测点；例如DEEP的ANNProof在efSearch=10时已达71%，其60/65/70%列均为71%。实际起点会受合法参数下限影响。'])

    for row in wb['新增数据实验说明'].iter_rows(min_row=2):
        if row[0].value=='索引':row[1].value=f'新增三数据集统一 M={INDEX_M}, efConstruction={INDEX_EFC}。ANNProof/Ours 共用同一缓存 HNSW 图；VLAH 独立 LSH 分桶图。原 SIFT 配置保留。'
    if '离散档位' not in [r[0].value for r in wb['新增数据实验说明'].iter_rows(min_row=2)]:
        wb['新增数据实验说明'].append(['离散档位','目标为统一召回阈值，主表填实际值。若相邻目标命中同一已测配置，重复引用该实测点；不插值、不把实际召回率改写为目标值。'])
    for file in sorted(STUDY.glob('evaluation-*.json')):
        job=json.loads(file.read_text());name=DATASETS[job['dataset']][0];ws=wb[name]
        records=measurements(Path(job['output']));method=job['method'];purpose=job['purpose']
        if job.get('filter'):records=[r for r in records if all(r[k]==v for k,v in job['filter'].items())]
        def put(row,col,r,metric):
            value=r['recall']*100 if metric=='recall_percent' else r[metric]
            c=ws.cell(row,col,value);c.number_format='0.00' if metric=='recall_percent' else '0.000'
            c.comment=Comment(f"{name}; {method}; N={job['base']}; M={INDEX_M}; efConstruction={INDEX_EFC}; k={r['k']}; efSearch/E={r['ef']}; probes={r['probes']}; LSH width={job['width'] if method=='VLAH' else 'not used'}\nActual recall={r['recall']:.4%}; samples={r['samples']}; PREVIEW queries100..119 repeated3.\nQuery includes real VO generation; verify includes authentication and replay; KB=1000 bytes.\nParameters matched on same workload; no held-out accuracy claim.\n{r['output']}/per_query.csv",'实测来源')
            allcells[(name,c.coordinate)]=[name,c.coordinate,value,metric,r['k'],r['ef'],r['probes'],r['output']+'/per_query.csv']
        if purpose=='curve':
            selection=job['selection']
            for col,point in enumerate(selection,6):
                r=next(r for r in records if (r['ef'],r['probes'])==(point['ef'],point['probes']))
                assert r['recall']+1e-12>=point['target']
                curvepoints.append([name,method,point['target'],r['recall'],r['query_ms'],r['verify_ms'],r['vo_kb'],r['ef'],r['probes']])
                for (rr,mr),metric in zip(CURVE_ROWS[method],['vo_kb','verify_ms','query_ms']):put(rr,col,r,'recall_percent');put(mr,col,r,metric)
            p=selection[7];r=next(r for r in records if (r['ef'],r['probes'])==(p['ef'],p['probes']))
            row={'VLAH':9,'ANNProof':10,'Ours':11}[method]
            for col,metric in zip(range(5,9),['recall_percent','query_ms','vo_kb','verify_ms']):put(row,col,r,metric)
        elif purpose=='k':
            for col,k in enumerate([1,5,10,20,50,100],6):
                r=next(r for r in records if r['k']==k)
                for row,metric in zip(K_ROWS[method],['vo_kb','verify_ms','query_ms']):put(row,col,r,metric)
        elif purpose=='size':
            assert len(records)==1;col={100000:6,500000:7,1000000:8}[job['base']]
            for row,metric in zip(SIZE_ROWS[method],['vo_kb','verify_ms','query_ms']):put(row,col,records[0],metric)
        for r in records:allruns.append([name,method,purpose,job['base'],r['k'],r['ef'],r['probes'],job['width'] if method=='VLAH' else None,r['recall']*100,r['query_ms'],r['verify_ms'],r['vo_kb'],r['samples'],r['output']])
    for sheet,rows in [('新增数据测量记录',allruns),('新增数据单元格来源',list(allcells.values()))]:
        s=wb[sheet]
        if s.max_row>1:s.delete_rows(2,s.max_row-1)
        for row in rows:s.append(row)
        s.freeze_panes='D2';s.auto_filter.ref=s.dimensions
        for c in s[1]:c.font=Font(bold=True)
    title='新增数据召回曲线'
    if title in wb.sheetnames:del wb[title]
    chartdata=wb.create_sheet(title)
    chartdata.append(['dataset','method','Target Recall','Actual Recall','Query ms','Verify ms','VO KB','efSearch/E','probes'])
    groups={}
    for row in curvepoints:
        chartdata.append(row);groups.setdefault((row[0],row[1]),[]).append(chartdata.max_row)
    chartdata.freeze_panes='E2'
    for row in chartdata.iter_rows(min_row=2):row[2].number_format=row[3].number_format='0.0%'
    for i,(name,_,_) in enumerate(DATASETS.values()):
        for j,(metric,col,unit) in enumerate([('Query time',5,'ms'),('VO size',7,'KB'),('Verify time',6,'ms')]):
            chart=ScatterChart();chart.title=f'{name}: Recall vs. {metric}'
            chart.x_axis.title='Actual Recall@10';chart.x_axis.scaling.min=.60;chart.x_axis.scaling.max=1
            chart.x_axis.numFmt='0%';chart.x_axis.majorUnit=.05
            chart.y_axis.title=metric+' ('+unit+')';chart.y_axis.scaling.min=0;chart.width=22;chart.height=12
            for index,method in enumerate(METHODS):
                indices=groups.get((name,method),[])
                if not indices:continue
                series=Series(Reference(chartdata,min_col=col,min_row=min(indices),max_row=max(indices)),Reference(chartdata,min_col=4,min_row=min(indices),max_row=max(indices)),title=method)
                series.marker.symbol=['triangle','circle','square'][index];series.marker.size=5;series.smooth=False
                chart.series.append(series)
            if chart.series:chartdata.add_chart(chart,'L'+str(1+(i*3+j)*24))
    wb.save(OUTPUT)
    status=dict(measured_cells=len(allcells),expected_numeric_cells=3*(309-18),not_applicable_cells=54,original_updated=deliver(),output_sha256=sha(OUTPUT))
    (STUDY/'workbook_progress.json').write_text(json.dumps(status,indent=2)+'\n')
    event(action='checkpoint',**status)

if __name__=='__main__':initialize();checkpoint()
