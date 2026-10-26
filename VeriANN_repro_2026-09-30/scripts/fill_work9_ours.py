"""Add only measured Ours cells to the user's current baseline workbook."""
import argparse
import csv
from datetime import datetime
import hashlib
import json
import math
from pathlib import Path
import shutil
import statistics
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
import fill_work9 as common

ROOT=Path(__file__).resolve().parents[1]
STUDY=ROOT/'experiments/work9-ours-20261015'
SOURCE=ROOT.parent/'work9.xlsx'
TEMPLATE=STUDY/'work9_before_ours.xlsx'
OUTPUT=STUDY/'work9_completed_with_ours.xlsx'


def initialize():
    STUDY.mkdir(parents=True,exist_ok=True)
    if not TEMPLATE.exists():
        shutil.copy2(SOURCE,TEMPLATE)
        (STUDY/'source_sha256.txt').write_text(common.sha(TEMPLATE)+'\n')


def runs():
    common.STUDY=STUDY
    return common.load_runs()


def size_evaluation(measurements,n):
    matched=[r for r in measurements if r['name']==f'size-{n}-matched-recall']
    if matched:return min(matched,key=lambda r:(abs(r['recall_percent']-95),r['ef']))
    initial=[r for r in measurements if r['name']==f'size-{n}-evaluation']
    return initial[0] if len(initial)==1 else None


def main(update=False):
    initialize();measurements=runs();wb=openpyxl.load_workbook(TEMPLATE);ws=wb['Sheet1']
    is_lut=any(r['manifest']['method']=='VeriANN-LUT' for r in measurements)
    for name in ['Ours实验说明','Ours配置与测量','Ours单元格来源']:
        if name in wb.sheetnames:del wb[name]
    cells=[]
    def find(name,k=10,E=None):
        matches=[r for r in measurements if r['name']==name and r['k']==k and (E is None or r['ef']==E)]
        return matches[0] if len(matches)==1 else None
    def set_value(row,col,run,metric):
        cell=ws.cell(row,col);value=run[metric] if run else '待测'
        cell.value=value;cell.number_format='0.00' if metric=='recall_percent' else '0.000'
        cell.font=Font(name='Calibri',size=11,color='006100' if run else '9C6500')
        if run:
            text=(f"VeriANN 全部可组合优化；实测 {metric}\nN={run['base']}, k={run['k']}, E={run['ef']}\n"
                  f"M=8, efConstruction=100；E 为唯一 expanded-node 预算，不是 HNSW efSearch。\n"
                  f"实际 Recall@K={run['recall_percent']:.4f}%\n"
                  f"queries={run['queries']}, offset={run['query_offset']}, repeats={run['repeats']}\n"
                  f"数据={common.windows(run['dataset'])}\n"
                  f"CSV={common.windows(Path(run['output'])/'per_query.csv')}\n"
                  "Query=search+真实VO序列化/生成；Verify=解码+认证+重放+渐进top-k检查；KB=1000 bytes。")
            if is_lut:text+='\n版本：VeriANN-LUT；量化查表准备成本已包含在服务器搜索和客户端验证计时内。'
            if 'matched-recall' in run['name'] or run.get('selection_basis')=='fixed_matched_workpoint':text+='\n选择口径：沿用固定正式查询集上补测E网格选出的约95%工作点；不是独立泛化评估。'
        else:text='没有通过验证的该配置实测结果；不使用估算或插值。'
        cell.comment=Comment(text,'Ours 实测来源')
        cells.append({'cell':cell.coordinate,'value':value,'metric':metric,
                      'run':run['output'] if run else None,'k':run['k'] if run else None,
                      'E':run['ef'] if run else None})
    overall=find('sift1m-overall')
    for col,metric in [(5,'recall_percent'),(6,'query_ms'),(7,'vo_kb'),(8,'verify_ms')]:set_value(11,col,overall,metric)
    curve=sorted([r for r in measurements if r['name']=='sift1m-recall-curve'],key=lambda r:(r['recall_percent'],r['ef']))
    assert len(curve) in [0,11]
    for i in range(11):
        run=curve[i] if curve else None
        for rr,mr,metric in [(23,24,'vo_kb'),(33,34,'verify_ms'),(43,44,'query_ms')]:
            set_value(rr,6+i,run,'recall_percent');set_value(mr,6+i,run,metric)
    for i,n in enumerate([100000,500000,1000000,5000000,10000000]):
        run=size_evaluation(measurements,n)
        for row,metric in [(56,'vo_kb'),(66,'verify_ms'),(75,'query_ms')]:set_value(row,6+i,run,metric)
    for i,k in enumerate([1,5,10,20,50,100]):
        run=find('sift1m-k-curve',k=k)
        for row,metric in [(86,'vo_kb'),(96,'verify_ms'),(105,'query_ms')]:set_value(row,6+i,run,metric)
    assert len(cells)==103
    ws['A4']='Baselines and optimized VeriANN: actual measurements; source, configuration and Recall are in comments/extra sheets.'
    notes=wb.create_sheet('Ours实验说明');notes.append(['项目','口径 / 范围'])
    for key,value in [
        ('测试日期',datetime.now().isoformat()),
        ('最终组合','SQ4 top-1 traversal；SQ4→Delta4→6→raw；Direct-content Expansion-Centric Merkle；BFS refinement ordering；共享路径multiproof。'),
        ('实现优化','预打包64B SQ4和32B Delta；稀疏查询状态及稀疏multiproof；证明坐标由已披露targets按canonical顺序推导，仅传输32B sibling hashes。'),
        ('排除的历史分支','primary/residual progressive adjacency 在原研究中等召回更差，文档明确停止；MAC与Merkle、Node-ID/SQ4排序与BFS、top2/certified与SQ4 top-1均是互斥对照，不能同时启用。'),
        ('图参数','M=8, efConstruction=100；复用ANNProof相同数据前缀的HNSW缓存。未改用旧SIFT100K headline的M=16。'),
        ('E','原始算法唯一expanded-node预算，包含upper layers；CSV为兼容统计沿用ef_search列名，其语义是E，不是HNSW efSearch。'),
        ('时间与单位','Query Time=search_ms+vo_generation_ms；Verify=真实VO解码+Merkle认证+遍历重放+渐进top-k检查；VO为真实wire bytes/1000。网络和区块链访问不包含在内。'),
        ('正式查询','query IDs 100–199，共100条不同查询，3次重复；300测量/配置。查询、生成和验证均单线程。'),
        ('初始独立校准','query IDs 0–99；选择校准Recall@10达到95%的最小已测E。原校准和原正式测试全部保留。'),
        ('约95%工作点匹配','若原规模正式Recall不在94%–96%范围，固定query IDs 100–199补测E网格，取实测Recall最接近95%的工作点。这些补测点使用了正式查询结果选择参数，不是独立泛化评估；每个单元格明确标注来源。'),
        ('规模来源','100K、500K、1M、5M、10M统一使用BIGANN/SIFT1B前缀；其他主表、Recall和K曲线使用SIFT1M。'),
        ('整体与K配置','整体E由SIFT1M独立校准决定；K曲线使用同一E。Recall曲线固定11个E，三个面板使用相同实测点，未插值。'),
        ('真实VO与旧计费差异','包含header、query binding、答案ID、计数、refinement位置与proof计数；不是旧实验仅按payload+32B hashes计算的估算VO。'),
        ('正确性','先在原SIFT100K/M16/E64缓存上回归；原版与优化版的轨迹、答案、量化上下界与认证叶一致。端到端回归Recall@10=99.2%。'),
        ('客户端','仅使用VO、查询、可信root和量化参数；不读取服务端原始数据集、图、SQ代码数组或服务器trace。'),
        ('已有baseline','保留原206个baseline数据单元格。其校准只有20条查询，实际Recall存在偏离（如ANNProof BIGANN1M 87.6%）；不能把全表当作严格等召回性能比较。'),
        ('备份',common.windows(TEMPLATE)),
    ]:notes.append([key,value])
    if is_lut:
        for row in notes.iter_rows(min_row=2):
            if row[0].value=='实现优化':row[1].value+='；SQ4打分/SQ4与SQ6下界按查询查表复用，保持逐维累加与向外舍入。'
            if row[0].value=='初始独立校准':row[1].value='沿用上次独立校准确定的E、k、图和查询。此次重跑全部正式配置以更新计时，未重新挑选参数。'
            if row[0].value=='约95%工作点匹配':row[1].value='500K与10M沿用此前固定查询集匹配出的E=112与512；此次仅更新相同工作点的实测计时。原选择使用了正式查询集，不是独立泛化评估。'
        notes.append(['本次版本','VeriANN-LUT；服务器和客户端各自准备查询查表，准备成本计入search_ms和verification_ms；每条查询与旧版逐项核对非计时字段。'])
        notes.append(['更新时间口径','总体、11点召回曲线、5点规模曲线、6点K曲线全部重新实测；未将47%等单点加速比例用于缩放旧数据。'])
    notes.column_dimensions['A'].width=23;notes.column_dimensions['B'].width=115
    for row in notes.iter_rows(min_row=2):row[1].alignment=Alignment(wrap_text=True,vertical='top')
    config=wb.create_sheet('Ours配置与测量');config.append(['experiment','N','dataset','k','E','queries','offset','repeats','samples',
        'Recall@K %','Query mean ms','Query P50 ms','Query P95 ms','Verify mean ms','Verify P95 ms',
        'VO KB','Search ms','VO generation ms','Authentication ms','Replay ms','Peak RSS MiB',
        'Index build or load ms','ADS construction ms','Exact GT ms','ADS bytes','CSV folder'])
    for r in measurements:
        s=next(s for s in r['summary'] if int(s['k'])==r['k'] and int(s['ef_search'])==r['ef'])
        config.append([r['name'],r['base'],common.windows(r['dataset']),r['k'],r['ef'],r['queries'],r['query_offset'],r['repeats'],r['samples'],
            r['recall_percent'],r['query_ms'],r['query_p50_ms'],r['query_p95_ms'],r['verify_ms'],r['verify_p95_ms'],r['vo_kb'],
            r['search_ms'],r['vo_generation_ms'],r['authentication_ms'],r['replay_ms'],float(s['peak_rss_mib']),float(s['index_ms']),
            float(s['ads_ms']),float(s['ground_truth_ms']),int(s['ads_bytes']),common.windows(r['output'])])
    config.freeze_panes='F2';config.auto_filter.ref=config.dimensions
    for col in range(1,27):config.column_dimensions[openpyxl.utils.get_column_letter(col)].width=22
    config.column_dimensions['A'].width=34;config.column_dimensions['Z'].width=100
    sources=wb.create_sheet('Ours单元格来源');sources.append(['cell','value','metric','k','E','per_query.csv'])
    for c in cells:sources.append([c['cell'],c['value'],c['metric'],c['k'],c['E'],common.windows(Path(c['run'])/'per_query.csv') if c['run'] else '待测'])
    sources.column_dimensions['F'].width=115;sources.freeze_panes='D2';sources.auto_filter.ref=sources.dimensions
    for sheet in [notes,config,sources]:
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='276749')
    wb.save(OUTPUT);check=openpyxl.load_workbook(OUTPUT,data_only=True);original=openpyxl.load_workbook(TEMPLATE,data_only=True)
    old_audit=json.loads((ROOT/'experiments/work9-20261015/workbook_audit.json').read_text())
    for c in old_audit['cells']:assert check['Sheet1'][c['cell']].value==original['Sheet1'][c['cell']].value
    assert set(check['Sheet1'].merged_cells.ranges)==set(original['Sheet1'].merged_cells.ranges)
    for c in cells:
        value=check['Sheet1'][c['cell']].value
        assert value==c['value'] or (isinstance(value,float) and math.isclose(value,c['value'],rel_tol=1e-12))
    completed=sum(isinstance(c['value'],(int,float)) for c in cells)
    audit={'ours_numeric_cells':completed,'ours_target_cells':103,'baseline_numeric_cells_preserved':206,
           'pending_cells':[c['cell'] for c in cells if c['run'] is None],'source_sha256':common.sha(TEMPLATE),
           'output_sha256':common.sha(OUTPUT),'cells':cells}
    (STUDY/'workbook_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    if update:
        assert completed==103,'refusing final update with missing measurements'
        marker=STUDY/'last_written_sha256.txt'
        allowed={common.sha(TEMPLATE)}
        if marker.exists():allowed.add(marker.read_text().strip())
        if common.sha(SOURCE) not in allowed:raise RuntimeError('original workbook changed during the study; completed copy saved without overwriting edits')
        shutil.copy2(OUTPUT,SOURCE);assert common.sha(SOURCE)==common.sha(OUTPUT)
        marker.write_text(common.sha(SOURCE)+'\n')
    print(json.dumps({'output':str(OUTPUT),'ours_filled':completed,'ours_total':103,
                      'baseline_cells_preserved':206,'original_updated':update},ensure_ascii=False))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--update-original',action='store_true')
    main(parser.parse_args().update_original)
