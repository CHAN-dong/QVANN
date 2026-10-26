"""Update only the three recall curves, retaining raw evidence and other experiments."""
import csv
import json
import math
import shutil
from pathlib import Path
import statistics
import openpyxl
from openpyxl.comments import Comment
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.styles import Font, PatternFill
from work9_recall99 import ROOT, STUDY, TARGETS, sha

def main():
    source=ROOT.parent/'work9.xlsx'
    backup=STUDY/'work9_before_recall99.xlsx'
    wb=openpyxl.load_workbook(backup)
    ws=wb['Sheet1']
    selected=json.loads((STUDY/'selection.json').read_text())
    records=[]
    changes={}
    rowsets={'VLAH':[(19,20),(29,30),(39,40)],
             'ANNProof':[(21,22),(31,32),(41,42)],
             'Ours':[(23,24),(33,34),(43,44)]}
    for method,pairs in rowsets.items():
        points=json.loads((STUDY/(method+'-evaluation.json')).read_text())
        for col,choice in enumerate(selected[method],6):
            p=next(p for p in points if p['parameter']==choice['parameter'])
            p['target']=choice['target'];records.append(p)
            text=(f"SIFT1M N=1000000, k=10; {method}; "
                  f"efSearch={'not applicable (E)' if method=='Ours' else p['ef']}; "
                  f"E={p['parameter'] if method=='Ours' else 'n/a'}; probes={p['probes']}\n"
                  f"Target Recall={p['target']:.2%}; actual Recall@10={p['recall']:.4%}\n"
                  "Queries 100..199; 3 repeats; single thread; M=8, efConstruction=100.\n"
                  "Parameters selected on these same 100 queries: matched workload operating points, not held-out accuracy estimates.\n"
                  "Query=search+VO generation/serialization (including LUT setup for Ours); "
                  "verify=VO decoding/authentication/replay; network/blockchain excluded; KB=1000 bytes.\n"
                  f"Raw measurements: {p['output']}/per_query.csv")
            for (rr,mr),metric in zip(pairs,['vo_kb','verify_ms','query_ms']):
                for row,value,key in [(rr,p['recall']*100,'recall_percent'),(mr,p[metric],metric)]:
                    c=ws.cell(row,col,value);c.number_format='0.00' if row==rr else '0.000'
                    c.comment=Comment(text,'Measured recall curve')
                    changes[c.coordinate]=dict(value=value,metric=key,**p)
    for row in [15,25,35]:
        ws.cell(row,5,'Target Recall (%)')
        for col,target in enumerate(TARGETS,6):ws.cell(row,col,target*100).number_format='0'
    # Keep existing cell-level provenance consistent with the updated main table.
    for name in ['单元格来源','Ours单元格来源']:
        sheet=wb[name]
        for cells in sheet.iter_rows(min_row=2):
            entry=changes.get(cells[0].value)
            if not entry:continue
            values=[cells[0].value,entry['value'],entry['metric'],10]
            if name=='单元格来源':values +=[entry['ef'],entry['probes'],entry['output']+'/per_query.csv']
            else:values +=[entry['parameter'],entry['output']+'/per_query.csv']
            for c,v in zip(cells,values):c.value=v
    # Existing configuration/resource sheets retain the old run history, explicitly marked.
    for name in ['实验配置与实测召回','资源与尾延迟','Ours配置与测量']:
        sheet=wb[name]
        for row in sheet.iter_rows(min_row=2):
            for cell in row[:2]:
                if cell.value=='sift1m-recall-curve':cell.value='historical-recall-curve (superseded by Recall99 measurements)'
    for name in ['实验说明','Ours实验说明']:
        sheet=wb[name]
        for row in sheet.iter_rows(min_row=2):
            if row[0].value and ('Recall曲线' in str(row[0].value) or '召回曲线' in str(row[0].value)):
                row[1].value='已由 Recall99 measurements 页的统一目标召回率实测曲线替代；旧配置页仅保留历史。'
        sheet.append(['Recall99 更新','主表三条 Recall 曲线统一目标 60/65/70/75/80/85/90/95/97/98/99%；填实际召回率，无插值。配置、原始数据与三张图见 Recall99 measurements。'])
    data=wb.create_sheet('Recall99 measurements')
    data.append(['method','Target Recall','Actual Recall@10','Query ms','Verify ms','VO KB','efSearch or E','probes','samples','raw CSV'])
    for p in records:data.append([p['method'],p['target'],p['recall'],p['query_ms'],p['verify_ms'],p['vo_kb'],p['parameter'] if p['method']!='VLAH' else p['ef'],p['probes'],p['samples'],p['output']+'/per_query.csv'])
    for row in data.iter_rows(min_row=2,max_row=34):
        row[1].number_format=row[2].number_format='0.0%'
    data.freeze_panes='D2';data.auto_filter.ref='A1:J34'
    for col in 'ABCDEFGHIJ':data.column_dimensions[col].width=22
    data.column_dimensions['J'].width=100
    for c in data[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='234E70')
    notes=[
        'SIFT1M; k=10; M=8; efConstruction=100; identical query IDs 100..199; 3 repeats; single-thread query and verification.',
        'Common target thresholds; main table and charts use actual measured recall. Each parameter is the smallest tested parameter reaching its threshold, not a global optimum.',
        'Pilot and evaluation use the SAME 100 queries to match the workload recall. This is not an independent held-out recall estimate. No interpolation or extrapolation.',
        'VLAH uses efSearch=100 throughout this new curve (previous curve: 50), varying probes; ANNProof varies efSearch; Ours varies E using validated LUT implementation.',
        'Query time includes actual VO generation/serialization; verification includes client decoding/authentication/replay; KB=1000 bytes; no network/blockchain time.',
        'ANNProof and VLAH are disclosed algorithmic reimplementations, not strict author-code/full blockchain reproductions.',
        'Historical configuration/resource rows were retained and marked superseded; current curve provenance is this sheet and main-cell comments.',
        '60% is the common target start; zero-recall Ours points were removed from these comparison curves. Discrete parameters cause small differences in actual starting recalls.',
    ]
    for i,note in enumerate(notes,37):data.cell(i,1,note)
    for i,(metric,column,unit) in enumerate([('Query time',4,'ms'),('VO size',6,'KB'),('Verify time',5,'ms')]):
        chart=ScatterChart();chart.title='Recall@10 vs. '+metric;chart.x_axis.title='Actual Recall@10';chart.y_axis.title=metric+' ('+unit+')'
        chart.x_axis.scaling.min=.60;chart.x_axis.scaling.max=1.0;chart.x_axis.numFmt='0%';chart.x_axis.majorUnit=.05
        chart.y_axis.scaling.min=0
        chart.width=22;chart.height=12
        for j,method in enumerate(rowsets):
            start=2+j*11;end=start+10
            series=Series(Reference(data,min_col=column,min_row=start,max_row=end),Reference(data,min_col=3,min_row=start,max_row=end),title=method)
            series.marker.symbol=['circle','triangle','square'][j];series.marker.size=5;series.smooth=False
            chart.series.append(series)
        data.add_chart(chart,'L'+str(1+i*24))
    output=STUDY/'work9_completed_recall99.xlsx';wb.save(output)
    check=openpyxl.load_workbook(output)
    old=openpyxl.load_workbook(backup)
    assert set(map(str,check['Sheet1'].merged_cells.ranges))==set(map(str,old['Sheet1'].merged_cells.ranges))
    allowed=set(changes)|{f'{openpyxl.utils.get_column_letter(c)}{r}' for r in [15,25,35] for c in range(5,17)}
    for row in old['Sheet1']:
        for c in row:
            if c.coordinate not in allowed: assert check['Sheet1'][c.coordinate].value==c.value,c.coordinate
    for addr,p in changes.items():assert math.isclose(check['Sheet1'][addr].value,p['value'],rel_tol=1e-12,abs_tol=1e-12)
    for method in rowsets:
        pts=[p for p in records if p['method']==method]
        assert len(pts)==11 and pts[-1]['recall']>=.99-1e-12 and pts[0]['recall']>=.6-1e-12
    audit=dict(status='PASS',curve_cells=len(changes),measurements=sum(p['samples'] for p in records),
               output_sha256=sha(output),backup_sha256=sha(backup),points=records,charts=3,original_updated=False)
    delivery=ROOT.parent/'work9_recall99_updated.xlsx';shutil.copy2(output,delivery)
    allowed={sha(backup),sha(output)}
    if (STUDY/'validation.json').exists():allowed.add(json.loads((STUDY/'validation.json').read_text())['output_sha256'])
    if sha(source) not in allowed:raise RuntimeError('Workbook changed; completed copy saved, preserving concurrent edits.')
    try:
        shutil.copy2(output,source);assert sha(source)==sha(output);audit['original_updated']=True
    except PermissionError as e:audit['write_error']=str(e)
    (STUDY/'validation.json').write_text(json.dumps(audit,indent=2)+'\n')
    lines=['# Recall curves: common targets 60% to 99%', '',*notes,'', '| Method | Actual start | Actual end | Query ms at >=99% | Verify ms | VO KB | Parameter |','|---|---:|---:|---:|---:|---:|---|']
    for method in rowsets:
        pts=[p for p in records if p['method']==method];p=pts[-1]
        lines.append(f"| {method} | {pts[0]['recall']:.1%} | {p['recall']:.1%} | {p['query_ms']:.3f} | {p['verify_ms']:.3f} | {p['vo_kb']:.3f} | ef/E={p['ef']}, probes={p['probes']} |")
    lines+=['',f'198 curve cells verified; {audit["measurements"]} formal measurements; original workbook updated: {audit["original_updated"]}.',f'Backup: {backup}']
    (STUDY/'RESULTS.md').write_text('\n'.join(lines)+'\n')
    print(json.dumps({k:v for k,v in audit.items() if k!='points'}))

if __name__=='__main__':main()
