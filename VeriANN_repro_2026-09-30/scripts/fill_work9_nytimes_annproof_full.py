"""Publish fresh NYTimes ANNProof results with explicit scenario factors."""
import json,math,shutil
from pathlib import Path
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment,Font,PatternFill
import work9_resource_matched as u
R=u.ROOT;S=R/'experiments/work9-nytimes-annproof-full-20261021';O=R/'results/work9-nytimes-annproof-full-20261021'
SOURCE=R.parent/'work9.xlsx';BACKUP=S/'work9_before_full_retest.xlsx'
assert not (S/'APPLIED.json').exists(),'Already applied'
record=u.read(S/'FINISHED.json');assert u.sha(SOURCE)==record['backup_sha256'],'Workbook changed during run'
w=openpyxl.load_workbook(BACKUP);original=openpyxl.load_workbook(BACKUP);s=w['NYTimes-256'];events=[]
factors={'query_ms':1.2,'vo_kb':2.,'verify_ms':1.5};rawrows=[]
ch=record['curve_graph'];curvebyef={p['ef']:p for p in record['curve_measurements']}
def put(r,c,v,p,kind,extra=''):
 cell=s.cell(r,c);events.append(dict(cell=cell.coordinate,before=cell.value,after=v,kind=kind,point=p))
 cell.value=v;cell.number_format='0.000';cell.font=Font(color='000000');cell.fill=PatternFill('solid',fgColor='FFF2CC' if kind=='用户系数调整' else 'E2F0D9')
 cell.comment=Comment('本次完整重测：NYTimes-256，N=289761（规模表另标N=100000），query100..119，20条×3次，单线程，完整证明生成和验证。\n性质='+kind+'；非趋势预测。性能数值为新实测原始均值乘用户指定系数：query×1.2，VO×2，verify×1.5。召回率、索引大小和构建时间不乘系数。\n'+extra+'\n来源='+json.dumps(p,ensure_ascii=False)+'\n\n历史批注（非当前来源）：\n'+(cell.comment.text if cell.comment else ''),'Codex')
def metric(r,c,p,key,extra=''):
 put(r,c,p[key]*factors.get(key,100 if key=='recall' else 1),p,'用户系数调整' if key in factors else '实际召回率',extra)
for col,sel in enumerate(record['curve_selection'],6):
 p=curvebyef[sel['choice']['ef']] if sel['choice'] else None
 extra=f'召回曲线全段固定M={ch["M"]}, efConstruction={ch["efc"]}；只调efSearch。目标={sel["target"]*100:g}%。实际召回按实测；同一点复用不算独立新测。默认总表及K/规模表独立采用M16/efc200达到约95%，不与低图曲线声称相同索引。'
 for r in [23,37,51]:
  if p:metric(r,col,p,'recall',extra)
  else:put(r,col,None,sel,'未达目标',extra)
 for r,key in [(24,'vo_kb'),(38,'verify_ms'),(52,'query_ms')]:
  if p:metric(r,col,p,key,extra)
  else:put(r,col,None,sel,'未达目标',extra)
default=next(p for p in record['default_k'] if p['k']==10);assert default['recall']>=.95-1e-12
for col,key in [(5,'recall'),(6,'query_ms'),(7,'vo_kb'),(8,'verify_ms')]:metric(10,col,default,key,'默认配置M16/efc200；搜索预算重校准至至少95%。')
for p in record['default_k']:
 col=6+[1,5,10,20,50,100].index(p['k'])
 for r,key in [(105,'vo_kb'),(117,'verify_ms'),(128,'query_ms')]:metric(r,col,p,key,'K曲线固定默认图和默认搜索预算；实际Recall@k见来源，不声称每个k均等于95%。')
p=record['size'][0];assert p['recall']>=.95-1e-12
for r,key in [(69,'vo_kb'),(81,'verify_ms'),(92,'query_ms')]:metric(r,6,p,key,'NYTimes只适用N=100000，其余规模超过原数据容量，保持不适用。')
ix=record['index'];put(10,9,ix['total_bytes']/1e6,ix,'索引实测','完整服务快照decimal MB；不乘性能系数。')
put(10,10,ix['construct_seconds'],ix,'索引实测','全新单线程图+ADS构建，非缓存载入；不乘性能系数。')
for r in [10,69,81,92,105,117,128]:s.cell(r,4).value='ANNProof*'
for r in [23,37,51]:
 s.cell(r,4).value=f'ANNProof*\nM{ch["M"]}/efc{ch["efc"]}'
 s.cell(r,4).alignment=Alignment(wrap_text=True,vertical='center')
s['A1']=str(original['NYTimes-256']['A1'].value)+' | 最新NYTimes ANNProof完整重测：召回曲线固定一张图，仅调efSearch；本方案旧预测已覆盖。*性能含用户系数，原始值见NYTimes ANNProof全量重测。'
s['A1'].alignment=Alignment(wrap_text=True);s.row_dimensions[1].height=85
def add(title,header,rows):
 t=w.create_sheet(title);t.append(header)
 for row in rows:t.append(row)
 t.freeze_panes='A2';t.auto_filter.ref=t.dimensions
 for c in t[1]:c.font=Font(bold=True)
 for col in t.columns:t.column_dimensions[col[0].column_letter].width=22 if col[0].column<13 else 80
 for row in t:
  for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
 return t
for kind,pp in [('recall-fixed-graph',record['curve_measurements']),('default-and-k',record['default_k']),('size100k',record['size'])]:
 for p in pp:rawrows.append([kind,p['N'],p['k'],p['M'],p['efc'],p['ef'],p['recall'],p['query_ms'],p['vo_kb'],p['verify_ms'],p['samples'],'新实测',p['output']])
add('NYTimes ANNProof全量重测',['实验','N','k','M','efConstruction','efSearch','实际Recall','原始query ms','原始VO KB','原始verify ms','样本数','性质','来源'],rawrows)
add('NYTimes ANNProof重测系数',['实验','N','k','M','efConstruction','efSearch','实际Recall','调整query ms ×1.2','调整VO KB ×2','调整verify ms ×1.5','样本数','性质','来源'],[r[:7]+[r[7]*1.2,r[8]*2,r[9]*1.5]+r[10:11]+['用户系数调整；不是原始实测值']+r[12:] for r in rawrows])
add('NYTimes ANNProof重测说明',['项目','说明'],[
 ['异常原因','旧曲线56%左右使用M4/efc40，60%改用M8/efc80；建图配置切换造成真实成本下降。此前平滑/支配替换未提供固定图曲线。'],
 ['本次修正',f'全部Recall目标用同一张M{ch["M"]}/efc{ch["efc"]}图，仅调efSearch；禁止不同图混用，禁止预测补齐新实测曲线。'],
 ['固定图选参','在能覆盖约40%低召回的候选图中，比较95%校准点的访问记录体积代理；候选与评分保存在GRAPH_SELECTED.json。不是全局最优性声明。最终性能全部以完整证明重测，不使用pilot耗时。'],
 ['默认及K规模','默认仍用标准M16/efc200并重校准95%；全新索引与该默认查询使用相同图文件。K使用默认预算，规模100k重校准95%。低召回扫描图独立，须按参数解释95%曲线点与默认表差别。'],
 ['资源与重复','20条query100..119×3次；单线程；每组8项篡改测试；实际VO文件长度核对。索引独立新建1次，查询均全新测量。'],
 ['校准范围','参数校准与报告使用相同query100..119，不是独立留出测试，不声明对其他查询的泛化召回率。'],
 ['曲线横轴','按实际Recall行绘图，不把目标列头当作实际召回率。整数搜索预算可跳过某些召回率；复用同一配置的目标列绘图时去重。'],
 ['性能系数','主表query×1.2，VO×2，verify×1.5；原始均值独立保存，系数只应用一次，召回率与索引不乘。'],
 ['时间波动','固定图消除跨配置跳变，但实测时间均值不要求严格单调；不人为修改新测均值。'],
 ['覆盖范围','只覆盖NYTimes主表ANNProof所有适用实验；其余方案、数据集、历史原始附表保留。超容量规模保持不适用。'],
 ['历史页','之前的NYTimes低召回、趋势预测和ANNProof记录页保留为历史，本次主表以本次全量重测及单元格来源为准。'],
 ['完整备份',str(BACKUP)],['实验索引',str(S/'FINISHED.json')]])
add('NYTimes ANNProof重测位置',['单元格','旧值','新值','性质','来源'],[[e['cell'],e['before'],e['after'],e['kind'],json.dumps(e['point'],ensure_ascii=False)] for e in events])
for title in ['NYTimes ANNProof全量重测','NYTimes ANNProof重测系数']:
 for row in w[title].iter_rows(min_row=2):
  row[6].number_format='0.0%'
  for col in [7,8,9]:row[col].number_format='0.000000'
for pos,title in enumerate(['NYTimes ANNProof重测说明','NYTimes ANNProof全量重测','NYTimes ANNProof重测系数','NYTimes ANNProof重测位置'],4):
 w.move_sheet(title,offset=pos-w.sheetnames.index(title))
allowed={e['cell'] for e in events}|{'A1'}|{f'D{r}' for r in [10,23,37,51,69,81,92,105,117,128]};preserved=0
for os in original:
 for row in os:
  for c in row:
   if os.title=='NYTimes-256' and c.coordinate in allowed:continue
   assert c.value==w[os.title][c.coordinate].value,(os.title,c.coordinate);preserved+=1
assert u.sha(SOURCE)==u.sha(BACKUP)
tmp=S/'work9_nytimes_annproof_all_retested.xlsx';w.save(tmp);saved=openpyxl.load_workbook(tmp)
for e in events:
 v=saved['NYTimes-256'][e['cell']].value
 assert v==e['after'] or isinstance(v,(float,int)) and math.isclose(v,e['after'],rel_tol=1e-13,abs_tol=1e-9)
shutil.copy2(tmp,SOURCE)
u.write(S/'cell-changes.json',events);u.write(S/'APPLIED.json',dict(passed=True,written_cells=len(events),unrelated_cell_values_preserved=preserved,workbook_sha256=u.sha(SOURCE),backup_sha256=u.sha(BACKUP),raw_points=len(rawrows),raw_samples=sum(r[10] for r in rawrows)))
print('APPLIED',len(events),'cells;',len(rawrows),'distinct points;',sum(r[10] for r in rawrows),'samples')
