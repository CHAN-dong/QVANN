"""Independent reconstruction of saved workbook from native per-query CSVs."""
import csv,json,math,statistics
from collections import defaultdict
from pathlib import Path
import openpyxl
import work9_resource_matched as u
R=u.ROOT;S=R/'experiments/work9-nytimes-annproof-full-20261021';O=R/'results/work9-nytimes-annproof-full-20261021'
r=u.read(S/'FINISHED.json');applied=u.read(S/'APPLIED.json');changes=u.read(S/'cell-changes.json')
w=openpyxl.load_workbook(R.parent/'work9.xlsx');before=openpyxl.load_workbook(S/'work9_before_full_retest.xlsx');s=w['NYTimes-256']
assert u.sha(R.parent/'work9.xlsx')==applied['workbook_sha256']
count=0;attacks=0;reconstructed={}
for out in sorted((O/'measurements').iterdir()):
 rr=list(csv.DictReader((out/'per_query.csv').open()));a=list(csv.DictReader((out/'tamper_tests.csv').open()));manifest=u.read(out/'manifest.json')
 assert len(a)==8 and all(x['rejected']=='1' for x in a)
 assert all(x['verified']=='1' for x in rr) and manifest['build_threads']==1 and manifest['query_threads']==1 and manifest['queries']==20 and manifest['repeats']==3 and manifest['query_offset']==100
 assert int(rr[0]['vo_bytes'])==(out/'query0.vo').stat().st_size
 groups=defaultdict(list)
 for x in rr:groups[int(x['k']),int(x['ef_search'])].append(x)
 for key,gg in groups.items():
  assert len(gg)==60 and {(int(x['query']),int(x['repeat'])) for x in gg}=={(q,t) for q in range(100,120) for t in range(3)}
  p=dict(recall=statistics.fmean(float(x['recall_at_k']) for x in gg),query_ms=statistics.fmean(float(x['search_ms'])+float(x['vo_generation_ms']) for x in gg),vo_kb=statistics.fmean(int(x['vo_bytes'])/1000 for x in gg),verify_ms=statistics.fmean(float(x['verification_ms']) for x in gg))
  reconstructed[str(out),*key]=p
 count+=len(rr);attacks+=len(a)
for e in changes:
 if e['kind']=='索引实测':continue
 p=e['point'];actual=reconstructed[p['output'],p['k'],p['ef']]
 for k,v in actual.items():assert math.isclose(p[k],v,rel_tol=1e-12,abs_tol=1e-10)
 row=s[e['cell']].row;col=s[e['cell']].column
 if row in [23,37,51] or row==10 and col==5:expected=actual['recall']*100
 elif row in [24,69,105] or row==10 and col==7:expected=actual['vo_kb']*2
 elif row in [38,81,117] or row==10 and col==8:expected=actual['verify_ms']*1.5
 elif row in [52,92,128] or row==10 and col==6:expected=actual['query_ms']*1.2
 else:raise AssertionError(e)
 assert math.isclose(s[e['cell']].value,expected,rel_tol=1e-12,abs_tol=1e-9),e['cell']
 assert '预测' not in s[e['cell']].number_format
# Verify same graph and source identity across every fixed-graph curve point.
ch=r['curve_graph'];mf=u.read(Path(r['curve_measurements'][0]['output'])/'manifest.json')
assert (mf['M'],mf['ef_construction'])==(ch['M'],ch['efc'])
pilotrows={}
for pf in (S/'pilots').glob('*.json'):
 meta=u.read(pf)
 if (meta['N'],meta['M'],meta['efc'])!=(289761,ch['M'],ch['efc']):continue
 cmd=meta['command'];out=Path(cmd[cmd.index('--output')+1])
 for x in csv.DictReader((out/'pilot.csv').open()):pilotrows[int(x['query']),int(x['ef_search'])]=x
for out in {Path(p['output']) for p in r['curve_measurements']}:
 for x in csv.DictReader((out/'per_query.csv').open()):
  p=pilotrows[int(x['query']),int(x['ef_search'])]
  assert p['visited']==x['visited'] and p['expanded']==x['expanded']
  assert math.isclose(float(p['recall_at_k']),float(x['recall_at_k']),abs_tol=1e-12)
lookup={p['ef']:p for p in r['curve_measurements']}
for c,sel in enumerate(r['curve_selection'],6):
 if not sel['choice']:continue
 p=lookup[sel['choice']['ef']];assert math.isclose(p['recall'],sel['choice']['recall'],abs_tol=1e-10)
 assert s.cell(23,c).value==s.cell(37,c).value==s.cell(51,c).value
assert s['E10'].value>=95-1e-9
cache=list((S/'cache/default').glob('*.index'));assert len(cache)==1
assert u.sha(cache[0])==u.sha(O/'fresh-default-index/cache/global.index')
ix=r['index'];assert (O/'fresh-default-index/serving.snapshot').stat().st_size==ix['total_bytes']
assert math.isclose(s['I10'].value,ix['total_bytes']/1e6,rel_tol=1e-12)
assert math.isclose(s['J10'].value,ix['construct_seconds'],rel_tol=1e-12)
allowed={e['cell'] for e in changes}|{'A1'}|{f'D{x}' for x in [10,23,37,51,69,81,92,105,117,128]};preserved=0
for os in before:
 for row in os:
  for c in row:
   if os.title=='NYTimes-256' and c.coordinate in allowed:continue
   v=w[os.title][c.coordinate].value
   assert c.value==v or isinstance(c.value,(float,int)) and isinstance(v,(float,int)) and math.isclose(c.value,v,rel_tol=1e-13,abs_tol=1e-10),(os.title,c.coordinate)
   preserved+=1
report=dict(passed=True,distinct_measured_points=len(reconstructed),samples=count,tamper_rejections=attacks,default_recall_percent=s['E10'].value,curve_graph=dict(M=ch['M'],efc=ch['efc']),coefficient_cells_verified=True,fresh_index_and_default_graph_identical=True,unrelated_values_preserved=preserved,curve_55_to_60=[{str(row):s.cell(row,c).value for row in [23,24,38,52]} for c in [14,15]])
u.write(S/'independent-audit.json',report);print(json.dumps(report,ensure_ascii=False))
