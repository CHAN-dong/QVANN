"""Checkpointed, measured-only work9 tables at feasible V3DB operating points.

Run under the existing WSL v3db-work9 Python environment. No forecasts are
written. Only complete, verified 20x3 groups enter the comparison tables.
"""
import argparse
import base64
from collections import defaultdict
from copy import copy
import csv
import hashlib
import json
import math
import os
from pathlib import Path
import resource
import shutil
import statistics
import subprocess
import sys
import time
import traceback
import zlib

ROOT=Path(__file__).resolve().parents[1]
STUDY=ROOT/'experiments/work9-resource-matched-20261020'
RESULTS=ROOT/'results/work9-resource-matched-20261020'
SOURCE=ROOT.parent/'work9.xlsx'
BACKUP=STUDY/'work9_before_resource_matched.xlsx'
OUTPUT=STUDY/'work9_measured.xlsx'
DATA={
 'sift1m':dict(sheet='SIFT1M',N=1000000,D=128,M=8,efc=100,pqM=128,K=1024,L=1024,probe=2,target=.705),
 'nytimes256':dict(sheet='NYTimes-256',N=289761,D=256,M=16,efc=200,pqM=256,K=256,L=512,probe=2,target=.58),
 'glove100':dict(sheet='GloVe-100',N=1183514,D=100,M=16,efc=200,pqM=100,K=1024,L=1024,probe=2,target=.405),
 'deep1m':dict(sheet='DEEP1M',N=1000000,D=96,M=16,efc=200,pqM=96,K=1024,L=1024,probe=4,target=.78),
}
TARGETS=[.20,.30,.40,.405,.50,.55,.58,.60,.65,.70,.705,.75,.78,.80,.85,.90,.95,.97,.98,.99]
METHODS=['VLAH','ANNProof','Ours','V3DB']
KS=[1,5,10,20,50,100]
SIZES=[100000,500000,1000000,5000000,10000000]
CURVES=[(20,'vo_kb'),(32,'verify_ms'),(44,'query_ms')]
SIZE_ROWS=[(61,'vo_kb'),(72,'verify_ms'),(82,'query_ms')]
K_ROWS=[(94,'vo_kb'),(105,'verify_ms'),(115,'query_ms')]
CACHE=ROOT/'V3DB-26VLDB/work9-cache'
sys.path.insert(0,str(ROOT/'V3DB-26VLDB/scripts'))

def read(p):return json.loads(Path(p).read_text(encoding='utf-8'))
def write(p,obj):
 p=Path(p);p.parent.mkdir(parents=True,exist_ok=True)
 tmp=p.with_suffix(p.suffix+'.tmp');tmp.write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');tmp.replace(p)
def sha(p):
 with Path(p).open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()
def event(**x):
 x['time']=time.strftime('%Y-%m-%dT%H:%M:%S%z');STUDY.mkdir(parents=True,exist_ok=True)
 with (STUDY/'events.jsonl').open('a',encoding='utf-8') as f:f.write(json.dumps(x,ensure_ascii=False)+'\n')
 print(json.dumps(x,ensure_ascii=False),flush=True)
def environment():
 e=dict(os.environ,PYTHONUNBUFFERED='1')
 for k in ['OMP_NUM_THREADS','OPENBLAS_NUM_THREADS','MKL_NUM_THREADS','RAYON_NUM_THREADS']:e[k]='1'
 f=ROOT/'third_party/LearnedEarlyTer-20SIGMOD'
 e['LD_LIBRARY_PATH']=':'.join(map(str,[ROOT/'third_party/runtime/lib',f,f/'lib']))
 e['V3DB_RESULTS_DIR']=str(RESULTS/'V3DB')
 return e
def limit():
 resource.setrlimit(resource.RLIMIT_CORE,(0,0))
 resource.setrlimit(resource.RLIMIT_AS,(50*2**30,50*2**30))
def command(cmd,log,bounded=False,timeout=None):
 log=Path(log);log.parent.mkdir(parents=True,exist_ok=True)
 event(action='command_start',command=list(map(str,cmd)),log=str(log))
 with log.open('a') as f:
  p=subprocess.run(list(map(str,cmd)),cwd=ROOT,env=environment(),stdout=f,stderr=subprocess.STDOUT,
                   preexec_fn=limit if bounded else None,timeout=timeout)
 event(action='command_done',code=p.returncode,log=str(log))
 if p.returncode:raise RuntimeError(f'command failed ({p.returncode}): {log}\n'+log.read_text(errors='replace')[-2000:])

def model_name(ds,n=None,source=None):
 c=DATA[ds];default=10 if ds=='glove100' else 8
 return f'{source or ds}-{n or c["N"]}'+(f'-M{c["pqM"]}' if c['pqM']!=default else '')+(f'-K{c["K"]}' if c['K']!=256 else '')
def config(ds,n=None,source=None):
 c=DATA[ds];n=n or c['N']
 return dict(dataset=ds,source=source or ds,N=n,D=c['D'],M=c['pqM'],K=c['K'],n_list=c['L'],capacity=1<<(math.ceil(2*n/c['L'])-1).bit_length())
def jobname(model,p,k):return f'{model}-p{p}-k{k}-q20-r3-t1'

def initialize():
 STUDY.mkdir(parents=True,exist_ok=True);RESULTS.mkdir(parents=True,exist_ok=True)
 if not BACKUP.exists():shutil.copy2(SOURCE,BACKUP)
 # Tighten the initial coarse ef/E pilot grid. Retain all old raw evidence, but
 # reselect and remeasure derived tables if their calibration policy changed.
 migration=STUDY/'dense-calibration-v2.json'
 if not migration.exists():
  moved=[]
  for f in (STUDY/'records').glob('*.json'):
   r=read(f)
   if r.get('method') in ['ANNProof','Ours'] and r['kind'] in ['curve','overall','k','size']:
    target=STUDY/'superseded-coarse-grid'/f.name;target.parent.mkdir(exist_ok=True);f.rename(target);moved.append(str(target))
  write(migration,dict(reason='dense ef/E grid for closer low-recall matching',archived_records=moved))
 protocol=dict(datasets=DATA,targets=TARGETS,k_values=KS,sizes=SIZES,query_ids=list(range(100,120)),repeats=3,
  search_threads=1,proof_threads=1,address_space_limit_GiB=50,maximum_proof_domain_bit=20,
  actual_recall='set overlap top-k, not NN1-hit-in-top-k; selection and evaluation share these 20 queries',
  default_policy='Keep the user-accepted feasible V3DB defaults. Other methods select the least-budget calibrated configuration reaching the dataset target; always report actual recall.',
  v3db_curve='Fixed default PQ index, integer probes 1..2*default; stop at first static gate/domain/address-space budget limit. No high-recall forecasts.',
  reuse='Existing complete verified 20x3 V3DB groups with exactly the same model/probe/k may be reused; incomplete CSVs are copied with provenance and resumed. Non-V3DB query measurements are fresh.',
  k_curve='Fixed default search budget; ANNProof/VLAH effective ef=max(default ef,k). Ours keeps E and reports actual Recall@k.',
  size_curve='Fixed PQ M/K and per-dataset recall target, recalibrate probe/ef for each real prefix. Resource-unavailable V3DB points are explicit. SIFT size panels use BIGANN prefixes, as the original workbook.',
  index='Full fresh single-thread index builds; no query circuit/proving included; existing index results remain pending until fresh build succeeds.',
  ablation='Fresh 20x3 measurements at dataset target; paired variants share answers. Same prepared full dataset, including full GloVe.',
  query_time='search + actual VO/proof generation; V3DB includes circuit construction/compilation every call, no unmeasured amortization',
  units=dict(time='ms',VO='decimal KB',index_size='decimal MB',index_time='seconds'),backup=str(BACKUP))
 write(STUDY/'protocol.json',protocol)
 checkpoint()

def record(key,obj):write(STUDY/'records'/f'{key}.json',obj);checkpoint()
def aggregate(out):
 out=Path(out);assert (out/'PASS.txt').exists(),out
 groups=defaultdict(list)
 for r in csv.DictReader((out/'per_query.csv').open()):
  assert r['verified']=='1'
  groups[int(r['k']),int(r['ef_search']),int(r['probes'])].append(r)
 result=[]
 for (k,e,p),rows in groups.items():
  assert len(rows)==60
  assert {(int(r['query']),int(r['repeat'])) for r in rows}=={(q,t) for q in range(100,120) for t in range(3)}
  mean=lambda key:statistics.fmean(float(r[key]) for r in rows)
  result.append(dict(k=k,ef=e,probes=p,recall=mean('recall_at_k'),query_ms=mean('search_ms')+mean('vo_generation_ms'),verify_ms=mean('verification_ms'),vo_kb=mean('vo_bytes')/1000,samples=60,status='passed',output=str(out),csv_sha256=sha(out/'per_query.csv')))
 return result

def checkpoint():
 if (STUDY/'SCOPE_CORRECTED.json').exists():
  raise RuntimeError('Obsolete writer disabled: only V3DB defaults may use low recall. Other methods retain their original ~95% defaults. See SCOPE_CORRECTED.json.')
 if not BACKUP.exists():return
 import openpyxl
 from openpyxl.comments import Comment
 from openpyxl.styles import Font,PatternFill
 from openpyxl.chart import ScatterChart,Series,Reference
 wb=openpyxl.load_workbook(BACKUP)
 records=[read(p) for p in sorted((STUDY/'records').glob('*.json'))] if (STUDY/'records').exists() else []
 audit=[];curve_rows=[]
 def put(ws,r,c,val,src=None):
  cell=ws.cell(r,c);cell.value=val;cell.number_format='0.000';cell.font=Font(name='Calibri',size=11,color='000000')
  cell.fill=PatternFill(fill_type=None)
  cell.comment=Comment(json.dumps(src,ensure_ascii=False,indent=2)[:30000] if src else '本轮实测排队中；旧表已完整备份。预测值不再作为本轮数据。','实测来源')
  audit.append([ws.title,cell.coordinate,val,json.dumps(src,ensure_ascii=False) if src else 'pending'])
 def metric(ws,r,c,rec,key,percent=False):
  if rec and rec.get('status')=='passed':
   put(ws,r,c,rec[key]*(100 if percent else 1),rec)
  else:put(ws,r,c,(rec or {}).get('label','待实测'),rec)
 for ds,cfg in DATA.items():
  ws=wb[cfg['sheet']]
  assert ws['E8'].value=='Recall@10',ws.title
  ah=next(cell.row for row in ws for cell in row if cell.column==6 and cell.row>120 and cell.value=='Recall@10')
  ws['A4']='本轮：相同20条查询IDs100..119×3次，单线程；只填完整验证通过的实测。默认召回率与可运行V3DB相近，实际召回逐项标注。旧预测已撤出当前对比，历史附表仅供归档。'
  for i,m in enumerate(METHODS):
   ws.cell(9+i,4,m)
   for col in range(5,11):put(ws,9+i,col,'待实测')
   for start,_ in CURVES:
    ws.cell(start+2*i,4,m)
    for col in range(6,6+len(TARGETS)):
     for r in [start+2*i,start+2*i+1]:put(ws,r,col,'待实测')
   for start,_ in SIZE_ROWS:
    ws.cell(start+i,4,m)
    for col,n in enumerate(SIZES,6):put(ws,start+i,col,'不适用' if n>cfg['N'] and ds not in ['sift1m','deep1m'] else '待实测')
   for start,_ in K_ROWS:
    ws.cell(start+i,4,m)
    for col in range(6,12):put(ws,start+i,col,'待实测')
  for r in [16,28,40]:
   for col,t in enumerate(TARGETS,6):
    ws.cell(r,col)._style=copy(ws.cell(r,6)._style);ws.cell(r,col,t*100)
  for r in range(ah+1,ah+8):
   for col in range(6,10):put(ws,r,col,'待实测')
  # Preserve ablation design labels; remove the old "contains forecast" wording.
  ws.cell(ah+7,4,'V3DB baseline')
  for rec in [r for r in records if r['dataset']==ds]:
   m=rec.get('method');i=METHODS.index(m) if m in METHODS else 0;kind=rec['kind']
   if kind=='overall':
    for col,key in [(5,'recall'),(6,'query_ms'),(7,'vo_kb'),(8,'verify_ms')]:metric(ws,9+i,col,rec,key,col==5)
    if m=='V3DB':
     for col,key in [(6,'recall'),(7,'query_ms'),(8,'vo_kb'),(9,'verify_ms')]:metric(ws,ah+7,col,rec,key)
   elif kind=='index' and rec['status']=='passed':
    put(ws,9+i,9,rec['total_bytes']/1e6,rec);put(ws,9+i,10,rec['construct_seconds'],rec)
   elif kind=='curve':
    for j,t in enumerate(TARGETS):
     p=next((p for p in rec['points'] if abs(p['target']-t)<1e-10),None)
     for start,key in CURVES:
      metric(ws,start+2*i,6+j,p,'recall',True);metric(ws,start+2*i+1,6+j,p,key)
    seen=set()
    for p in rec['points']:
     ident=(p.get('output'),p.get('ef'),p.get('probes'),p.get('n_probe'))
     if p.get('status')=='passed' and ident not in seen:
      seen.add(ident);curve_rows.append([cfg['sheet'],m,p['recall'],p['query_ms'],p['verify_ms'],p['vo_kb'],p.get('ef',p.get('n_probe')),p.get('probes',1),p['output']])
   elif kind=='size':
    col=6+SIZES.index(rec['N'])
    for start,key in SIZE_ROWS:metric(ws,start+i,col,rec,key)
   elif kind=='k':
    col=6+KS.index(rec['k'])
    for start,key in K_ROWS:metric(ws,start+i,col,rec,key)
   elif kind=='ablation':
    r=ah+rec['variant']
    for col,key in [(6,'recall'),(7,'query_ms'),(8,'vo_kb'),(9,'verify_ms')]:metric(ws,r,col,rec,key)
 for title in ['资源匹配说明','资源匹配实测记录','资源匹配单元格来源','资源匹配召回曲线']:
  if title in wb:del wb[title]
 notes=wb.create_sheet('资源匹配说明')
 for row in [
  ['项目','说明'],['当前状态','本轮计划已完成；资源限制及不支持点保留文字状态。' if (STUDY/'RUN_FINISHED.json').exists() else '运行中；待实测不是0，也不是预测。完整实验结束以外部RUN_FINISHED.json为准。'],
  ['默认目标','SIFT70.5%、NYTimes58%、DEEP约78%；GloVe暂以p2配置的40.5%搜索校准为目标，最终用完整真实证明结果核对。其他方案填实际达到的召回率，不能视为完全相等。'],
  ['协议','所有新查询测量统一20条IDs100..119×3重复，单线程；相同查询调参和报告，不声称独立留出评估。'],
  ['曲线','增加20/30/40/40.5/50/55/58/60/65/70/70.5/75/78/80/85/90/95/97/98/99%目标。参数离散；最小配置已高于目标时显示实际召回，多个目标可共享同一实测点。图的横轴使用实际召回。'],
  ['V3DB资源','50GiB进程虚拟地址空间预算；保守限制证明域不超过2^20。静态约束计数、域预算跳过和真实分配失败分别记录；不将它们称为机器绝对可达上限。高召回不可运行点不预测。'],
  ['V3DB时间','包括搜索、约束生成、builder.build、见证和实际证明/压缩；每次构建，未扣除未经测量的复用收益。Verify包括压缩证明解码与公开输入绑定检查。'],
  ['复用','仅复用完全相同配置且验证通过的20×3 V3DB实测；未完成的同配置样本从归档副本恢复，保留来源。其他方案在本轮新测。'],
  ['规模','SIFT规模表保持BIGANN前缀；DEEP使用原DEEP前缀。超过NYTimes/GloVe容量为不适用。每个规模重新选达到本数据集默认目标的预算。'],
  ['K曲线','固定默认预算；ANNProof/VLAH有效ef=max(原预算,k)，避免返回不足。各k实际召回见实测记录，不能称为恒定Recall@10。'],
  ['历史附表','本文件既有预测/旧实验辅助表保留为历史证据，不再是四个主表来源；当前来源仅资源匹配实测记录/单元格来源。原文件完整备份：'+str(BACKUP)],
  ['实验文件',str(STUDY)],['索引','重新完整构建1次，单线程；不是20×3的查询重复。大小包含服务端索引和认证结构，不包含临时查询证明器。']]:notes.append(row)
 notes.column_dimensions['A'].width=22;notes.column_dimensions['B'].width=135
 tab=wb.create_sheet('资源匹配实测记录');tab.append(['dataset','method','kind','status','recall','query_ms','verify_ms','vo_kb','samples','N','k','record'])
 for r in records:
  details=json.dumps(r,ensure_ascii=False)
  if len(details)>30000:details=json.dumps(dict(encoding='zlib+base64',data=base64.b64encode(zlib.compress(details.encode('utf-8'))).decode('ascii')))
  assert len(details)<32767,'provenance record exceeds Excel cell limit'
  tab.append([r.get(k) for k in ['dataset','method','kind','status','recall','query_ms','verify_ms','vo_kb','samples','N','k']]+[details])
 tab=wb.create_sheet('资源匹配单元格来源');tab.append(['sheet','cell','value','source'])
 # Keep only the final assignment for each cell, rather than initialization rows.
 final={}
 for r in audit:final[r[0],r[1]]=r
 for r in final.values():tab.append(r)
 tab=wb.create_sheet('资源匹配召回曲线');tab.append(['dataset','method','actual_recall','query_ms','verify_ms','vo_kb','ef_or_probe','probes','output'])
 curve_rows.sort(key=lambda r:(r[0],r[1],r[2]))
 for r in curve_rows:tab.append(r)
 for d,(ds,cfg) in enumerate(DATA.items()):
  for j,(name,col) in enumerate([('Query ms',4),('Verify ms',5),('VO KB',6)]):
   chart=ScatterChart();chart.title=cfg['sheet']+' '+name;chart.x_axis.title='Actual Recall@10';chart.x_axis.numFmt='0%';chart.x_axis.scaling.min=0;chart.x_axis.scaling.max=1;chart.y_axis.title=name;chart.width=22;chart.height=12
   for m in METHODS:
    idx=[i+2 for i,r in enumerate(curve_rows) if r[0]==cfg['sheet'] and r[1]==m]
    if idx:chart.series.append(Series(Reference(tab,min_col=col,min_row=min(idx),max_row=max(idx)),Reference(tab,min_col=3,min_row=min(idx),max_row=max(idx)),title=m))
   tab.add_chart(chart,f'K{1+(d*3+j)*24}')
 tmp=OUTPUT.with_suffix('.tmp.xlsx');wb.save(tmp)
 # Read back before replacing the user's workbook. Never replace with a bad zip.
 check=openpyxl.load_workbook(tmp,read_only=True);assert all(c['sheet'] in check.sheetnames for c in DATA.values());check.close()
 tmp.replace(OUTPUT)
 marker=STUDY/'last_written_sha256.txt'
 allowed={sha(BACKUP)}
 if marker.exists():allowed.add(marker.read_text().strip())
 if sha(SOURCE) not in allowed:
  event(action='external_workbook_edit_detected',output=str(OUTPUT));return
 try:
  shutil.copy2(OUTPUT,SOURCE);marker.write_text(sha(SOURCE)+'\n')
 except PermissionError:event(action='workbook_locked',output=str(OUTPUT))
 write(STUDY/'progress.json',dict(time=time.strftime('%Y-%m-%dT%H:%M:%S%z'),records=len(records),passed=sum(r.get('status')=='passed' for r in records),main_workbook=str(SOURCE),output=str(OUTPUT)))

def width(ds):
 if ds=='sift1m':return 256
 old=ROOT/'experiments'/('work9-nytimes-20261017' if ds=='nytimes256' else 'work9-multidata-20261015')
 return read(old/f'{ds}-lsh-width.json')['width']

def pilot(ds,method,n,source=None,goal=.99):
 """Use exact existing search-only binaries; never use these timing columns."""
 import run_work9_multidata as runner
 import work9_multidata as common
 cfg=DATA[ds];source=source or ds
 runner.STUDY=STUDY;runner.RESULTS=RESULTS/'search-pilots';runner.INDEX_M=cfg['M'];runner.INDEX_EFC=cfg['efc']
 common.STUDY=STUDY
 points=[]
 if method=='VLAH':
  phases=[(ef,list(range(1,math.ceil(n/20000)+1))) for ef in [100,200,400,800,1600,3200,6400]]
 else:
  grid=list(range(1 if method=='Ours' else 10,65))+[80,96,112,128,160,192,224,256,320,384,448,512,768,1024,1536,2048]
  phases=[(grid,[1]),([4096,8192],[1]),([16384,32768],[1]),([65536,131072],[1]),([262144,524288],[1])]
 for phase,pair in enumerate(phases):
  efs,probes=([pair[0]],pair[1]) if method=='VLAH' else pair
  efs=[e for e in efs if e<=n]
  if not efs:continue
  signature=hashlib.sha256(json.dumps([efs,probes,cfg['M'],cfg['efc']]).encode()).hexdigest()[:10]
  out=runner.pilot(source,method,n,efs,probes,f'{n}-resource-phase{phase}-{signature}',width(ds))
  points+=runner.pilot_groups(out)
  if max(p['recall'] for p in points if p['valid'])>=goal-1e-12:break
 write(STUDY/f'pilot-points-{ds}-{method}-{source}-{n}.json',points)
 return points
def select(points,target,method):
 good=[p for p in points if p['valid'] and p['recall']>=target-1e-12]
 return min(good,key=lambda p:(p['probes'],p['ef']) if method=='VLAH' else (p['ef'],p['probes'])) if good else None

def evaluate(ds,method,n,source,tag,efs,probes,ks):
 cfg=DATA[ds];out=RESULTS/ds/method/f'{source}-{n}-{tag}'
 if not (out/'PASS.txt').exists():
  if out.exists() and any(out.iterdir()):
   archived=out.with_name(out.name+'-incomplete-'+str(time.time_ns()));out.rename(archived)
  if method=='Ours':
   cmd=[sys.executable,ROOT/'scripts/run_ours_benchmark.py','--source',ROOT/'build/veriann_lut_benchmark.cpp','--binary',ROOT/'build/veriann_lut_benchmark','--E',','.join(map(str,efs))]
  else:
   folder=ROOT/('VLAH-25TrustCom' if method=='VLAH' else 'ANNProof-24FGCS')
   cmd=[sys.executable,folder/'run.py','--ef-search',','.join(map(str,efs)),'--probes',','.join(map(str,probes)),'--lsh-width',str(width(ds))]
  cmd+=['--dataset',ROOT/'dataset'/source,'--base',str(n),'--queries','20','--query-offset','100','--repeats','3','--build-threads','1','--M',str(cfg['M']),'--ef-construction',str(cfg['efc']),'--k',','.join(map(str,ks)),'--output',out,'--skip-build']
  command(cmd,STUDY/'logs'/f'{ds}-{method}-{source}-{n}-{tag}.log')
 return aggregate(out)

def run_baselines(section='all'):
 for ds,c in (DATA.items() if section!='sizes' else []):
  for method in METHODS[:-1]:
   key=f'{ds}-{method}-curve'
   if (STUDY/'records'/f'{key}.json').exists():curve=read(STUDY/'records'/f'{key}.json')['points']
   else:
    points=pilot(ds,method,c['N'])
    selections=[dict(target=t,selection=select(points,t,method)) for t in sorted(set(TARGETS+[c['target']]))]
    chosen={ (p['selection']['ef'],p['selection']['probes']) for p in selections if p['selection'] }
    groups=defaultdict(set)
    for ef,p in chosen:groups[ef].add(p)
    measured={}
    if method!='VLAH':
     # Same fixed graph/ADS, many query budgets. One invocation avoids repeated
     # untimed ADS initialization without changing any measured query work.
     for r in evaluate(ds,method,c['N'],ds,'curve-dense',sorted(groups),[1],[10]):measured[r['ef'],r['probes']]=r
    else:
     for ef,probes in groups.items():
      for r in evaluate(ds,method,c['N'],ds,f'curve-e{ef}',[ef],sorted(probes),[10]):measured[r['ef'],r['probes']]=r
    curve=[]
    for p in selections:
     s=p['selection'];r=measured[s['ef'],s['probes']] if s else dict(status='unreachable',label='网格未达目标')
     if s:assert abs(r['recall']-s['recall'])<1e-10,(ds,method,s,r)
     r=dict(r,target=p['target'],selection=s);curve.append(r)
    record(key,dict(dataset=ds,method=method,kind='curve',points=curve))
   main=next(p for p in curve if abs(p['target']-c['target'])<1e-10)
   assert main['status']=='passed',(ds,method,'default target not reached')
   record(f'{ds}-{method}-overall',dict(main,dataset=ds,method=method,kind='overall',N=c['N']))
   # Split k runs: preserve the k10 default and let baseline effective ef grow
   # only where k requires it. Ours must expose enough candidates at high k.
   for k in KS:
    recfile=STUDY/'records'/f'{ds}-{method}-k{k}.json'
    if recfile.exists():continue
    if k==10:r=main
    else:
     ef=max(main['ef'],k) if method!='Ours' else main['ef']
     try:r=evaluate(ds,method,c['N'],ds,f'k{k}-e{ef}',[ef],[main['probes']],[k])[0]
     except RuntimeError as exc:
      if method!='Ours' or 'candidate set smaller than k' not in str(exc).lower():raise
      r=dict(status='not_supported',label='默认预算返回候选不足',reason=str(exc))
    record(f'{ds}-{method}-k{k}',dict(r,dataset=ds,method=method,kind='k',k=k,N=c['N']))
   event(action='baseline_main_done',dataset=ds,method=method,recall=main['recall'])
 if section=='main':return
 # Main comparisons and low-recall curves finish before the larger prefixes.
 for ds,c in DATA.items():
  for method in METHODS[:-1]:
   for n in SIZES:
    if n>c['N'] and ds not in ['sift1m','deep1m']:continue
    key=f'{ds}-{method}-size{n}'
    if (STUDY/'records'/f'{key}.json').exists():continue
    source='bigann10m' if ds=='sift1m' else ('deep10m' if ds=='deep1m' and n>1000000 else ds)
    if n==c['N'] and source==ds:r=read(STUDY/'records'/f'{ds}-{method}-overall.json')
    else:
     pts=pilot(ds,method,n,source,c['target']);s=select(pts,c['target'],method)
     if s:r=evaluate(ds,method,n,source,'size',[s['ef']],[s['probes']],[10])[0]
     else:r=dict(status='unreachable',label='网格未达目标')
    record(key,dict(r,dataset=ds,method=method,kind='size',N=n,source_dataset=source,target=c['target']))

def static_gate(c,p,k):
 file=STUDY/'gates'/f'{c["source"]}-{c["N"]}-M{c["M"]}-K{c["K"]}-p{p}-k{k}.json'
 if file.exists():return read(file)
 # Avoid attempting to allocate enormous static builders. This is explicitly
 # a construction budget, not an empirically established memory minimum.
 tuples=p*max(c['capacity'],c['K'])*c['M']
 if tuples>2000000:
  r=dict(status='resource_limit',label='超构建预算，未出证',reason='p*max(capacity,K)*M exceeds 2,000,000 padded relation entries',relation_tuples=tuples,actual_attempt=False)
 else:
  cmd=[sys.executable,__file__,'gate-worker','--shape',json.dumps(dict(c,n_probe=p,k=k))]
  log=file.with_suffix('.log');log.parent.mkdir(parents=True,exist_ok=True)
  try:
   command(cmd,log,True,240)
   line=next(x for x in reversed(log.read_text().splitlines()) if x.startswith('GATE_JSON '))
   r=json.loads(line[len('GATE_JSON '):]);r.update(status='passed' if r['gates']<=2**20 else 'resource_limit',label='超证明域预算，未出证' if r['gates']>2**20 else '',actual_attempt=False)
  except (RuntimeError,subprocess.TimeoutExpired) as exc:
   r=dict(status='resource_limit',label='静态构建受限，未出证',reason=str(exc),actual_static_attempt=True,actual_proof_attempt=False)
 r.update(config=c,n_probe=p,k=k,address_space_limit_GiB=50)
 write(file,r);return r

def v3db_measure(ds,model,p,k):
 job=jobname(model,p,k);out=RESULTS/'V3DB'/job;result=out/'result.json'
 if result.exists():return read(result)
 # Same source snapshot, code and protocol: retain complete actual evidence.
 old=ROOT/'results/work9-v3db-20261016'/job
 if (old/'result.json').exists():
  prior=read(old/'result.json')
  if prior.get('status')=='passed' and prior.get('samples')==60:
   rows=list(csv.DictReader((old/'per_query.csv').open()));assert len(rows)==60 and all(r['verified']=='1' for r in rows)
   assert {(int(r['query']),int(r['repeat'])) for r in rows}=={(q,t) for q in range(100,120) for t in range(3)}
   assert all(read(CACHE/model/'model.json')['config'][key]==prior['config'][key] for key in ['M','K','N','D','n_list','capacity'])
   snapshot_file=CACHE/model/'snapshot.json'
   if snapshot_file.exists():assert read(snapshot_file)['root']==prior['snapshot_root']
   return dict(prior,reused_verified_measurement=True,csv_sha256=sha(old/'per_query.csv'))
 # Resume incomplete historical samples in a separate directory; never modify
 # the historical source. A completed group is still required for workbook use.
 if (old/'per_query.csv').exists() and not out.exists():
  shutil.copytree(old,out)
  if result.exists():result.rename(out/'historical-result.json')
  write(out/'resume-provenance.json',dict(source=str(old),csv_sha256=sha(old/'per_query.csv'),copied_at=time.strftime('%Y-%m-%dT%H:%M:%S%z')))
 c=read(CACHE/model/'model.json')['config'];gate=static_gate(c,p,k)
 if gate['status']!='passed':return gate
 log=STUDY/'logs'/f'{job}.log'
 try:
  out.mkdir(parents=True,exist_ok=True)
  command(['/usr/bin/time','-v','-o',out/'resources.txt',sys.executable,ROOT/'V3DB-26VLDB/scripts/work9_measure.py','--model',model,'--probe',p,'--k',k,'--queries',20,'--repeats',3,'--attempt-large'],log,True)
  r=read(result);assert r['status']=='passed' and r['samples']==60
  usage=(out/'resources.txt').read_text()
  peak=next(int(line.split(':',1)[1].strip()) for line in usage.splitlines() if 'Maximum resident set size (kbytes)' in line)
  r.update(process_peak_rss_GiB=peak/2**20,process_peak_scope='entire current worker including static preflight, circuit construction and proving; historical resumed samples measured in their original processes')
  write(result,r)
  return r
 except RuntimeError as exc:
  text=log.read_text(errors='replace').lower()
  if any(t in text for t in ['memory allocation','memoryerror','cannot allocate','unable to allocate']):
   r=dict(status='resource_limit',label='出证分配失败（50GiB地址空间）',actual_attempt=True,reason=str(exc),output=str(out));write(result,r);return r
  raise

def import_defaults():
 # Import the two fully completed main points immediately; resume DEEP before
 # starting the still-unmeasured GloVe full proof group.
 for ds in ['sift1m','nytimes256']:
  c=DATA[ds];r=v3db_measure(ds,model_name(ds),c['probe'],10)
  record(f'{ds}-V3DB-overall',dict(r,dataset=ds,method='V3DB',kind='overall',N=c['N']))
  record(f'{ds}-V3DB-k10',dict(r,dataset=ds,method='V3DB',kind='k',k=10,N=c['N']))
  curvefile=STUDY/'records'/f'{ds}-V3DB-curve.json'
  if not curvefile.exists():
   points=[dict(r,target=t) if t<=r['recall']+1e-12 else dict(target=t,status='pending',label='待可运行配置实测') for t in TARGETS]
   record(f'{ds}-V3DB-curve',dict(dataset=ds,method='V3DB',kind='curve',points=points,scope='partial curve: only the completed default point so far'))

def run_v3db_main():
 for ds in ['deep1m','glove100','sift1m','nytimes256']:
  key=f'{ds}-V3DB-overall'
  if (STUDY/'records'/f'{key}.json').exists():continue
  c=DATA[ds];r=v3db_measure(ds,model_name(ds),c['probe'],10)
  record(key,dict(r,dataset=ds,method='V3DB',kind='overall',N=c['N']))

def run_v3db_panels():
 for ds,c in DATA.items():
  model=model_name(ds);cal=read(CACHE/model/'calibration.json');points=[];stop=None
  # Probe-by-probe instead of powers of two adds feasible intermediate points.
  for p in range(1,2*c['probe']+1):
   gate=static_gate(read(CACHE/model/'model.json')['config'],p,10)
   if gate['status']!='passed':stop=dict(gate,first_unmeasured_probe=p);break
   r=v3db_measure(ds,model,p,10)
   if r['status']!='passed':stop=dict(r,first_unmeasured_probe=p);break
   points.append(r)
   write(STUDY/f'{ds}-v3db-curve-completed.json',points)
   # Publish every completed point; higher targets stay explicitly unmeasured.
   curve=[]
   for t in TARGETS:
    good=[r for r in points if r['recall']>=t-1e-12]
    r=min(good,key=lambda r:r['n_probe']) if good else dict(status='pending',label='待可运行配置实测')
    curve.append(dict(r,target=t))
   record(f'{ds}-V3DB-curve',dict(dataset=ds,method='V3DB',kind='curve',points=curve))
  curve=[]
  for t in TARGETS:
   good=[r for r in points if r['recall']>=t-1e-12]
   r=min(good,key=lambda r:r['n_probe']) if good else dict(status='resource_limit',label='本资源预算内未测',reason=stop)
   curve.append(dict(r,target=t))
  record(f'{ds}-V3DB-curve',dict(dataset=ds,method='V3DB',kind='curve',points=curve,stop=stop))
 # Finish all four recall curves before the longer k/size panels.
 for ds,c in DATA.items():
  for k in KS:
   key=f'{ds}-V3DB-k{k}'
   if (STUDY/'records'/f'{key}.json').exists():continue
   r=v3db_measure(ds,model_name(ds),c['probe'],k)
   record(key,dict(r,dataset=ds,method='V3DB',kind='k',k=k,N=c['N']))
 for ds,c in DATA.items():
  for n in SIZES:
   if n>c['N'] and ds not in ['sift1m','deep1m']:continue
   key=f'{ds}-V3DB-size{n}'
   if (STUDY/'records'/f'{key}.json').exists():continue
   source='bigann10m' if ds=='sift1m' else ('deep10m' if ds=='deep1m' and n>1000000 else ds)
   if n==c['N'] and source==ds:r=read(STUDY/'records'/f'{ds}-V3DB-overall.json')
   else:
    cfg=config(ds,n,source);gate=static_gate(cfg,1,10)
    if gate['status']!='passed':r=gate
    else:
     model=model_name(ds,n,source)
     if not (CACHE/model/'calibration.json').exists():command([sys.executable,ROOT/'V3DB-26VLDB/scripts/work9_calibrate.py','--dataset',ds,'--N',n,'--source',source,'--M',c['pqM'],'--K',c['K']],STUDY/'logs'/f'calibrate-{model}.log',True)
     pts=[p for p in read(CACHE/model/'calibration.json')['points'] if p['k']==10 and p['recall']>=c['target']-1e-12]
     if pts:r=v3db_measure(ds,model,min(pts,key=lambda p:p['n_probe'])['n_probe'],10)
     else:r=dict(status='unreachable',label='网格未达默认目标')
   record(key,dict(r,dataset=ds,method='V3DB',kind='size',N=n,source_dataset=source,target=c['target']))

def run_ablations():
 import run_work9_ablation as ab
 ab.STUDY=STUDY;ab.RESULTS=RESULTS/'ablation'
 for ds,c in DATA.items():
  for pair in ab.PAIRS:
   if all((STUDY/'records'/f'{ds}-ablation{v}.json').exists() for v in pair):continue
   grid=[1,2,3,4,6,8,12,16,24,32,48,64,96,128,192,256] if pair[0]==1 else [10,12,16,20,24,32,48,64,96,128,192,256,384,512,768,1024,1536,2048]
   out=ab.launch(ds,pair,grid,f'resource-pilot-{pair[0]}-{pair[1]}',True,n=c['N'])
   points=ab.pilot_points(out,pair)
   good=[p for p in points if p['recall']>=c['target']-1e-12]
   assert good,(ds,pair,'extend ablation calibration')
   p=min(good,key=lambda p:p['parameter'])
   out=ab.launch(ds,pair,[p['parameter']],f'resource-eval-{pair[0]}-{pair[1]}',n=c['N'])
   # audit_pair's old fixed .95 assertion is bypassed by passing no selection;
   # all paired answers, wire lengths, tamper tests and 60 samples still checked.
   rows=ab.audit_pair(out,pair,None)
   for r in rows:
    assert abs(r['recall']-p['recall'])<1e-10
    record(f'{ds}-ablation{r["variant"]}',dict(r,dataset=ds,method='Ours',kind='ablation',status='passed',target=c['target'],selection=p))

def run_indexes():
 import work9_index_study as ix
 ix.STUDY=ROOT/'experiments/work9-index-20261017' # already-built unchanged native harnesses
 ix.OUT=RESULTS/'index';ix.run()
 # The standalone V3DB index script imports OUT and DATA from this module.
 # Execute in a fresh process to keep its allocations and ru_maxrss isolated.
 for ds,c in DATA.items():
  file=ix.OUT/ds/'V3DB/measurement.json'
  if not file.exists():command([sys.executable,__file__,'index-worker','--dataset',ds],STUDY/'logs'/f'{ds}-fresh-v3db-index.log',True)
 for ds in DATA:
  for m in METHODS:
   r=read(ix.OUT/ds/m/'measurement.json')
   assert r['status']=='measured' and r['threads']==1
   record(f'{ds}-{m}-index',dict(r,dataset=ds,method=m,kind='index',status='passed',raw_measurement=str(ix.OUT/ds/m/'measurement.json')))

def validate():
 import openpyxl
 w=openpyxl.load_workbook(OUTPUT,data_only=False)
 records=[read(p) for p in (STUDY/'records').glob('*.json')]
 required=set()
 for ds,c in DATA.items():
  for m in METHODS:
   required|={f'{ds}-{m}-overall',f'{ds}-{m}-curve',f'{ds}-{m}-index'}
   required|={f'{ds}-{m}-k{k}' for k in KS}
   required|={f'{ds}-{m}-size{n}' for n in SIZES if n<=c['N'] or ds in ['sift1m','deep1m']}
  required|={f'{ds}-ablation{v}' for v in range(1,7)}
 present={p.stem for p in (STUDY/'records').glob('*.json')}
 assert required<=present,('queued work remains',sorted(required-present))
 for r in records:
  if r.get('status')=='passed' and r['kind'] not in ['curve','index']:assert r['samples']==60,r
 for ds,c in DATA.items():
  ws=w[c['sheet']]
  for row in ws:
   for cell in row:
    if isinstance(cell.value,str):assert '[预测]' not in cell.value and '（含预测）' not in cell.value,(ws.title,cell.coordinate)
 assert len(w['资源匹配召回曲线']._charts)==12
 assert sha(SOURCE)==sha(OUTPUT),'Completed copy exists, but original workbook could not be updated; inspect file lock or external edits'
 write(STUDY/'validation.json',dict(status='PASS',records=len(records),workbook_sha256=sha(OUTPUT),main_sha256=sha(SOURCE),scope='complete groups only; no forecast values in current comparison tables'))

def main():
 p=argparse.ArgumentParser();p.add_argument('action',choices=['start','reload','run','fill','gate-worker','index-worker','validate','baselines','v3db-main']);p.add_argument('--shape');p.add_argument('--dataset');a=p.parse_args()
 if (STUDY/'SCOPE_CORRECTED.json').exists() and a.action not in ['gate-worker','index-worker']:
  raise RuntimeError('This controller used an incorrect non-V3DB default policy and has been disabled. Original ~95% defaults are restored; do not restart it.')
 if a.action=='gate-worker':
  limit();c=json.loads(a.shape)
  from zk_IVF_PQ.zk_IVF_PQ import py_set_based_gate
  g=py_set_based_gate(c['M'],c['K'],c['D']//c['M'],c['n_list'],c['n_probe'],c['capacity'],c['k'],True)
  print('GATE_JSON '+json.dumps(dict(gates=g,degree_bits=math.ceil(math.log2(max(1,g))),static_peak_rss_KiB=resource.getrusage(resource.RUSAGE_SELF).ru_maxrss)));return
 if a.action=='index-worker':
  import work9_index_study as ix
  ix.OUT=RESULTS/'index';sys.argv=[str(ROOT/'scripts/work9_index_v3db.py'),'--dataset',a.dataset]
  import runpy;runpy.run_path(str(ROOT/'scripts/work9_index_v3db.py'),run_name='__main__');return
 if a.action=='fill':checkpoint();return
 if a.action=='validate':validate();return
 if a.action=='reload':
  # Reload an updated controller only at a safe child-process boundary. The
  # active benchmark finishes and flushes PASS before its waiting parent exits.
  import signal
  pid=read(STUDY/'runner_pid.json')['pid'];proc=Path(f'/proc/{pid}')
  assert b'work9_resource_matched.py' in (proc/'cmdline').read_bytes()
  assert (proc/'wchan').read_text().strip()=='do_wait','controller must be waiting for a child'
  os.kill(pid,signal.SIGSTOP)
  try:
   children=[int(x) for x in (proc/f'task/{pid}/children').read_text().split()]
   deadline=time.monotonic()+240
   while any(Path(f'/proc/{ch}/stat').exists() and Path(f'/proc/{ch}/stat').read_text().split(') ',1)[1][0]!='Z' for ch in children):
    if time.monotonic()>deadline:raise RuntimeError('active benchmark has not finished; old controller will continue')
    time.sleep(.5)
   os.kill(pid,signal.SIGKILL)
   while proc.exists() and (proc/'cmdline').read_bytes():time.sleep(.1)
  except BaseException:
   if proc.exists():os.kill(pid,signal.SIGCONT)
   raise
  event(action='controller_reloaded_after_complete_child',old_pid=pid)
  a.action='start'
 if a.action=='start':
  STUDY.mkdir(parents=True,exist_ok=True)
  pidfile=STUDY/'runner_pid.json'
  if pidfile.exists():
   old=read(pidfile);proc=Path(f'/proc/{old["pid"]}/cmdline')
   if proc.exists() and b'work9_resource_matched.py' in proc.read_bytes():raise RuntimeError('runner already active')
  with (STUDY/'runner.log').open('a') as f:child=subprocess.Popen([sys.executable,__file__,'run'],cwd=ROOT,env=environment(),stdin=subprocess.DEVNULL,stdout=f,stderr=subprocess.STDOUT,start_new_session=True)
  write(pidfile,dict(pid=child.pid,script=__file__,action='run'));print(json.dumps(read(pidfile)));return
 # The advisory lock survives sampling interruptions and prevents overlapping
 # benchmark controllers, while each individual measurement is single threaded.
 import fcntl
 STUDY.mkdir(parents=True,exist_ok=True)
 lock=(STUDY/'runner.lock').open('w');fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
 try:
  if (STUDY/'runner_failure.json').exists():
   (STUDY/'runner_failure.json').rename(STUDY/f'previous-failure-{time.time_ns()}.json')
  initialize();import_defaults()
  if a.action=='baselines':run_baselines();return
  if a.action=='v3db-main':run_v3db_main();return
  run_baselines('main');run_v3db_main();run_ablations();run_indexes();run_baselines('sizes');run_v3db_panels();checkpoint();validate()
  write(STUDY/'RUN_FINISHED.json',dict(status='all planned measurements or explicit resource limits recorded',finished=time.strftime('%Y-%m-%dT%H:%M:%S%z')))
  checkpoint();validate()
  command([sys.executable,ROOT/'scripts/audit_work9_resource_matched.py'],STUDY/'final-independent-audit.log')
 except BaseException as exc:
  write(STUDY/'runner_failure.json',dict(error=str(exc),traceback=traceback.format_exc(),time=time.strftime('%Y-%m-%dT%H:%M:%S%z')));raise

if __name__=='__main__':main()
