"""Independent final workbook/raw-data cross-check; run after completion in WSL."""
import json
import math
from pathlib import Path
import statistics
import openpyxl
import run_work9_ablation as r

def main():
    final=json.loads((r.STUDY/'final_validation.json').read_text())
    assert final['status']=='PASS' and final['workbook_sha256']==r.sha(r.SOURCE)
    wb=openpyxl.load_workbook(r.SOURCE);original=openpyxl.load_workbook(r.BACKUP)
    assert wb.sheetnames[:len(original.sheetnames)]==original.sheetnames
    records=[];targets=set();attacks=0;paired=0;artifacts={}
    for dataset,(name,dim) in r.DATASETS.items():
        header=r.table(wb[name])
        for col in [7,8,9]:targets.add((name,header,col))
        for variant in [2,4]:
            row=header+variant
            col=4 if original[name].cell(row,4).value is not None else 5
            assert wb[name].cell(row,col).value==r.NAMES[variant]
            targets.add((name,row,col))
        for pair in r.PAIRS:
            point=json.loads((r.STUDY/f'selection-{dataset}-{pair[0]}-{pair[1]}.json').read_text())
            folder=r.RESULTS/dataset/f'formal-{pair[0]}-{pair[1]}'
            audit=json.loads((folder/'pair_audit.json').read_text());paired+=audit['paired_answer_comparisons']
            assert audit['status']=='PASS'
            for v in pair:
                path=folder/f'variant-{v}';raw=r.rows(path/'per_query.csv');manifest=json.loads((path/'manifest.json').read_text())
                assert len(raw)==60 and all(row['verified']=='1' for row in raw)
                assert manifest['base']==1000000 and manifest['M']==16 and manifest['ef_construction']==200
                assert manifest['queries']==20 and manifest['repeats']==3 and manifest['hash']=='SHA-256'
                assert manifest['query_threads']==1 and manifest['query_offset']==100
                assert len({(x['query'],x['repeat']) for x in raw})==60
                metrics={6:statistics.fmean(float(x['recall_at_k']) for x in raw),
                         7:statistics.fmean(float(x['search_ms'])+float(x['vo_generation_ms']) for x in raw),
                         8:statistics.fmean(int(x['vo_bytes']) for x in raw)/1000,
                         9:statistics.fmean(float(x['verification_ms']) for x in raw)}
                assert metrics[6]>=.95-1e-12
                for col,value in metrics.items():
                    cell=wb[name].cell(header+v,col)
                    assert math.isclose(cell.value,value,rel_tol=1e-8,abs_tol=1e-8),(name,cell.coordinate,cell.value,value)
                    assert cell.comment and r.sha(path/'per_query.csv') in cell.comment.text
                    targets.add((name,header+v,col))
                tests=r.rows(path/'tamper_tests.csv');assert all(x['rejected']=='1' for x in tests);attacks+=len(tests)
                records.append(dict(dataset=dataset,variant=v,metrics=metrics))
                graph=r.ROOT/'ANNProof-24FGCS/cache'/f"{manifest['dataset_sha256'][:16]}_M16_efc200.index"
                artifacts[str(graph.relative_to(r.ROOT))]=graph
                if v in [1,3]:
                    ivf=r.ROOT/'ANNProof-24FGCS/cache'/f"{manifest['dataset_sha256'][:16]}_ivf256_train65536_iter12_seed42_v1.bin"
                    artifacts[str(ivf.relative_to(r.ROOT))]=ivf
    assert len(records)==24 and paired==720 and attacks==final['tamper_tests']
    preserved=0
    for ws in original:
        for row in ws:
            for cell in row:
                if (ws.title,cell.row,cell.column) not in targets:
                    assert r.same_value(cell.value,wb[ws.title][cell.coordinate].value),(ws.title,cell.coordinate,'changed outside ablation')
                    preserved+=1
    info=dict(status='PASS',metric_cells=96,formal_measurements=1440,paired_answer_comparisons=paired,
              tamper_tests=attacks,preserved_original_cells=preserved,workbook_sha256=r.sha(r.SOURCE),
              cache_artifact_sha256={key:r.sha(p) for key,p in artifacts.items()})
    r.write_json(r.STUDY/'independent_audit.json',info)
    lines=['# 实测消融结果','',
           '四数据集均为1M，k=10，20条不同查询×3次；表中时间为ms，VO为十进制KB。',
           'Query包含搜索与完整VO生成；Verify包含解码、认证和重放。',
           '参数按同一组查询约95%召回率校准，属于快速预览；没有使用独立留出查询。','',
           '| Dataset | Variant | Recall@10 | Query (ms) | VO (KB) | Verify (ms) |',
           '| --- | --- | ---: | ---: | ---: | ---: |']
    for row in records:
        m=row['metrics'];lines.append(f"| {r.DATASETS[row['dataset']][0]} | {r.NAMES[row['variant']]} | {m[6]:.3f} | {m[7]:.4f} | {m[8]:.4f} | {m[9]:.4f} |")
    lines+=['',f'完成96个指标单元格，1440条诚实查询验证，{attacks}项篡改检测，720组成对答案一致性核对。',
            f'保留原工作簿消融区域以外的{preserved}个单元格值。',
            '',f'Workbook SHA256: `{info["workbook_sha256"]}`',
            '', '原始记录及模块定义见本目录 README.md 与工作簿“消融实验说明”。']
    (r.STUDY/'RESULTS.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    print(json.dumps({k:v for k,v in info.items() if k!='cache_artifact_sha256'},ensure_ascii=False))

if __name__=='__main__':main()
