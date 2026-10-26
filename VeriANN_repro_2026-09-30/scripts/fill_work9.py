"""Fill only evidence-backed work9.xlsx cells, preserving the template and provenance."""
import argparse
from collections import defaultdict
from datetime import datetime
import csv
import hashlib
import json
import math
from pathlib import Path
import shutil
import statistics
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

ROOT=Path(__file__).resolve().parents[1]
STUDY=ROOT/'experiments/work9-20261015'
ORIGINAL=ROOT.parent/'work9.xlsx'
TEMPLATE=STUDY/'work9_original.xlsx'
OUTPUT=STUDY/'work9_completed.xlsx'

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def windows(path):
    text=str(path)
    if text.startswith('/mnt/') and len(text)>7:
        return text[5].upper()+':\\'+text[7:].replace('/','\\')
    return text

def load_runs():
    runs=[]
    for file in sorted(STUDY.glob('*.json')):
        record=json.loads(file.read_text())
        if 'output' not in record:continue
        folder=Path(record['output'])
        if not (folder/'PASS.txt').exists():continue
        manifest=json.loads((folder/'manifest.json').read_text())
        with (folder/'tamper_tests.csv').open() as f:
            assert all(r['rejected']=='1' for r in csv.DictReader(f))
        groups=defaultdict(list)
        with (folder/'per_query.csv').open() as f:
            for row in csv.DictReader(f):
                assert row['verified']=='1'
                groups[(int(row['k']),int(row['ef_search']),int(row['probes']))].append(row)
        for (k,ef,probes),rows in groups.items():
            assert len(rows)==record['queries']*record['repeats']
            def percentile(values, fraction):
                ordered=sorted(values)
                return ordered[max(0,math.ceil(len(ordered)*fraction)-1)]
            query_times=[float(r['search_ms'])+float(r['vo_generation_ms']) for r in rows]
            measurements={
                'recall_percent':statistics.fmean(float(r['recall_at_k'])*100 for r in rows),
                'query_ms':statistics.fmean(float(r['search_ms'])+float(r['vo_generation_ms']) for r in rows),
                'verify_ms':statistics.fmean(float(r['verification_ms']) for r in rows),
                'vo_kb':statistics.fmean(int(r['vo_bytes'])/1000 for r in rows),
                'search_ms':statistics.fmean(float(r['search_ms']) for r in rows),
                'vo_generation_ms':statistics.fmean(float(r['vo_generation_ms']) for r in rows),
                'authentication_ms':statistics.fmean(float(r['authentication_ms']) for r in rows),
                'replay_ms':statistics.fmean(float(r['replay_ms']) for r in rows),
                'query_p50_ms':percentile(query_times,.5),
                'query_p95_ms':percentile(query_times,.95),
                'verify_p95_ms':percentile([float(r['verification_ms']) for r in rows],.95),
            }
            assert all(math.isfinite(v) and v>=0 for v in measurements.values())
            run=dict(record,manifest=manifest,k=k,ef=ef,probes=probes,samples=len(rows),**measurements)
            runs.append(run)
    return runs

def main(update_original=False):
    STUDY.mkdir(parents=True,exist_ok=True)
    if not TEMPLATE.exists():
        shutil.copy2(ORIGINAL,TEMPLATE)
        (STUDY/'original_sha256.txt').write_text(sha(TEMPLATE)+'\n')
    wb=openpyxl.load_workbook(TEMPLATE)
    ws=wb['Sheet1']
    runs=load_runs()
    cells=[]
    def find(method,name,k=None,ef=None,probes=None):
        candidates=[r for r in runs if r['method']==method and r['name']==name]
        if k is not None:candidates=[r for r in candidates if r['k']==k]
        if ef is not None:candidates=[r for r in candidates if r['ef']==ef]
        if probes is not None:candidates=[r for r in candidates if r['probes']==probes]
        return candidates[0] if len(candidates)==1 else None
    def set_cell(row,col,run,metric):
        cell=ws.cell(row,col)
        if run is None:
            cell.value='待测'
            cell.comment=Comment('尚无该条件的通过验证实测结果；未填估算、插值或旧结果。','实验记录')
            cell.font=Font(name=cell.font.name or 'Calibri',size=cell.font.sz or 11,color='9C6500')
        else:
            cell.value=run[metric]
            cell.number_format='0.00' if metric=='recall_percent' else '0.000'
            m=run['manifest']
            text=(f"实测，{metric}\n方法={run['method']}\n数据={run['dataset']}\n"
                  f"N={run['base']}, k={run['k']}, efSearch={run['ef']}, probes={run['probes']}\n"
                  f"M={m['M']}, efc={m['ef_construction']}, bucket={m['bucket_size']}, shard={m['shard_size']}\n"
                  f"query_offset={run['query_offset']}, queries={run['queries']}, repeats={run['repeats']}\n"
                  f"实际 Recall@K={run['recall_percent']:.4f}%\n"
                  f"原始 CSV={windows(Path(run['output'])/'per_query.csv')}\n"
                  "Query Time=search+VO generation; KB=1000 bytes; integrity verification includes replay.")
            cell.comment=Comment(text,'实测来源')
        cells.append({'cell':cell.coordinate,'value':cell.value,'metric':metric,
                      'run':run['output'] if run else None,
                      'k':run['k'] if run else None,'ef':run['ef'] if run else None,
                      'probes':run['probes'] if run else None})
    for method,row in [('VLAH-25TrustCom',9),('ANNProof-24FGCS',10)]:
        run=find(method,'sift1m-recall-curve',k=10,ef=50 if row==9 else 100,probes=41 if row==9 else 1)
        for col,metric in [(5,'recall_percent'),(6,'query_ms'),(7,'vo_kb'),(8,'verify_ms')]:set_cell(row,col,run,metric)
    for method,pairs in [('VLAH-25TrustCom',[(19,20),(29,30),(39,40)]),
                         ('ANNProof-24FGCS',[(21,22),(31,32),(41,42)])]:
        curve=sorted([r for r in runs if r['method']==method and r['name']=='sift1m-recall-curve'],
                     key=lambda r:(r['recall_percent'],r['ef'],r['probes']))
        assert len(curve) in (0,11)
        for col in range(6,17):
            run=curve[col-6] if curve else None
            for (rr,mr),metric in zip(pairs,['vo_kb','verify_ms','query_ms']):
                set_cell(rr,col,run,'recall_percent');set_cell(mr,col,run,metric)
    sizes=[100000,500000,1000000,5000000,10000000]
    for method,rows in [('VLAH-25TrustCom',[54,64,73]),('ANNProof-24FGCS',[55,65,74])]:
        for col,n in enumerate(sizes,6):
            run=find(method,f'size-{n}-evaluation',k=10)
            for row,metric in zip(rows,['vo_kb','verify_ms','query_ms']):set_cell(row,col,run,metric)
    for row in [52,62,71]:
        ws.cell(row,4,'BIGANN / SIFT1B prefix subsets (same source)')
        ws.cell(row,4).alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws.row_dimensions[row].height=30
    ks=[1,5,10,20,50,100]
    for method,rows in [('VLAH-25TrustCom',[84,94,103]),('ANNProof-24FGCS',[85,95,104])]:
        for col,k in enumerate(ks,6):
            run=find(method,'sift1m-k-curve',k=k)
            for row,metric in zip(rows,['vo_kb','verify_ms','query_ms']):set_cell(row,col,run,metric)
    ws['A3']='Query/verify time: ms; VO: KB (1000 bytes); recall: %; Query Time includes VO generation.'
    ws['A4']='Baselines: disclosed algorithmic reimplementations; source/configuration/actual Recall are in comments and extra sheets.'
    ws['E80']='K vs. VO size (actual Recall@K logged)'
    ws['E90']='K vs. Verify time'
    ws['E99']='K vs. Query time (includes VO generation)'
    for col in range(6,17):ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width=13

    notes=wb.create_sheet('实验说明')
    notes.append(['项目','定义 / 范围'])
    for key,value in [
        ('测试日期',datetime.now().isoformat()),
        ('Query Time','逐条 search_ms + vo_generation_ms 的平均值；不含客户端验证、网络或链上访问。'),
        ('VO Size','真实序列化 VO 字节数 / 1000。KB 与原始 CSV 的 KiB(1024字节)不同，已正确换算。'),
        ('Verify Time','真实客户端解码、哈希认证和搜索重放；不是仅验证 Merkle hash。'),
        ('Recall','Recall@K 为与精确 top-k 的交集大小 / K；主表写百分数，不填0到1小数。'),
        ('查询与重复','正式测量使用 query IDs 100–199，共100条不同查询、3次重复；300样本/配置。'),
        ('独立校准','规模点使用 query IDs 0–19 校准到至少95% recall，正式测量与校准查询不重叠。正式实测召回可能不同，已记录。'),
        ('总体配置','SIFT1M, k=10, M=8, efc=100；ANNProof ef=100；VLAH ef=50, probes=41, bucket=20000, w=256。'),
        ('Recall曲线','每方法11个实测配置；按实测召回排序，同一个配置同时填VO/验证/查询三张关系表，未插值。'),
        ('K曲线','k=1,5,10,20,50,100；两方法 ef=100；VLAH probes=41；实际 Recall@K 见配置页，非固定召回实验。'),
        ('规模曲线来源','BIGANN/SIFT1B 前缀子集；不与 SIFT1M 拼接。原表5M/10M超出SIFT1M总量，所以三处标题已注明来源。'),
        ('规模条件','M=8, efc=100；在独立校准集上选达95% recall的最小测试ef/probes。这是统一校准门槛，并非测试召回完全相同；实际测试召回另列，应结合其差异解读规模曲线。'),
        ('规模比较限制','校准仅20条查询，测试召回可能明显低于95%；例如BIGANN 1M ANNProof校准95.5%，正式87.6%。规模表不是严格等召回性能比较。'),
        ('资源与尾延迟','额外记录查询P50/P95、验证P95和峰值RSS；经验分位数采用nearest-rank。索引阶段可为首次构建或缓存加载，二者不得混用。'),
        ('线程','查询及验证单线程。BIGANN首次建索引可用8线程；索引一经生成即缓存并记录承诺。'),
        ('复现范围','ANNProof/Faiss和VLAH的链下算法重实现；VLAH部分缺失参数按README补充。未宣称论文原程序或性能优势的严格数值复现。'),
        ('区块链','可信根在本地模拟；未测区块链交易、共识、网络、根获取和动态更新。'),
        ('原表备份',windows(TEMPLATE)),
        ('原始实验日志',windows(STUDY)),
        ('待测','没有通过验证的实测数据的单元格保留“待测”，不使用估计、插值或复制其他规模数据。'),
    ]:notes.append([key,value])
    notes.column_dimensions['A'].width=22;notes.column_dimensions['B'].width=105
    notes.freeze_panes='B2'
    for row in notes.iter_rows(min_row=2):row[1].alignment=Alignment(wrap_text=True,vertical='top')

    config=wb.create_sheet('实验配置与实测召回')
    headers=['method','experiment','dataset','N','k','efSearch','probes','M','efc','bucket','shard',
             'queries','offset','repeats','samples','Recall@K %','Query ms','Verify ms','VO KB','Search ms',
             'VO generation ms','Auth ms','Replay ms','Build threads','Output directory']
    config.append(headers)
    for r in runs:
        m=r['manifest']
        config.append([r['method'],r['name'],r['dataset'],r['base'],r['k'],r['ef'],r['probes'],m['M'],
                       m['ef_construction'],m['bucket_size'],m['shard_size'],r['queries'],r['query_offset'],
                       r['repeats'],r['samples'],r['recall_percent'],r['query_ms'],r['verify_ms'],r['vo_kb'],
                       r['search_ms'],r['vo_generation_ms'],r['authentication_ms'],r['replay_ms'],
                       m.get('build_threads',1),windows(r['output'])])
    config.freeze_panes='F2';config.auto_filter.ref=config.dimensions
    for col in range(1,len(headers)+1):config.column_dimensions[openpyxl.utils.get_column_letter(col)].width=20
    config.column_dimensions['Y'].width=100

    sources=wb.create_sheet('单元格来源')
    sources.append(['单元格','数值','指标','k','efSearch','probes','原始 CSV'])
    for cell in cells:
        sources.append([cell['cell'],cell['value'],cell['metric'],cell['k'],cell['ef'],cell['probes'],
                        windows(Path(cell['run'])/'per_query.csv') if cell['run'] else '待测'])
    sources.freeze_panes='D2';sources.auto_filter.ref=sources.dimensions
    sources.column_dimensions['G'].width=120
    resources=wb.create_sheet('资源与尾延迟')
    resources.append(['method','experiment','N','k','efSearch','probes','samples',
                      'Query P50 ms','Query P95 ms','Verify P95 ms','Peak RSS MiB',
                      'Index build or load ms','ADS construction ms','Exact GT ms','ADS bytes'])
    for r in runs:
        summary=next(s for s in r['summary'] if (int(s['k']),int(s['ef_search']),int(s['probes']))==
                     (r['k'],r['ef'],r['probes']))
        resources.append([r['method'],r['name'],r['base'],r['k'],r['ef'],r['probes'],r['samples'],
                          r['query_p50_ms'],r['query_p95_ms'],r['verify_p95_ms'],
                          float(summary['peak_rss_mib']),float(summary['index_ms']),
                          float(summary['ads_ms']),float(summary['ground_truth_ms']),int(summary['ads_bytes'])])
    resources.freeze_panes='H2';resources.auto_filter.ref=resources.dimensions
    for col in range(1,16):resources.column_dimensions[openpyxl.utils.get_column_letter(col)].width=23
    resources['H1'].comment=Comment('经验分位数采用排序后 nearest-rank；Query=search+VO generation。','统计口径')
    resources['L1'].comment=Comment('本次进程的索引阶段耗时；缓存命中时为加载，不能当作首次建图耗时。首次建图记录在 calibration 实验。','统计口径')
    for sheet in [notes,config,sources,resources]:
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='234E70')
    wb.save(OUTPUT)
    # Reopen and verify all 206 target cells retain exactly the intended values.
    check=openpyxl.load_workbook(OUTPUT,data_only=True)
    assert len(cells)==206
    for c in cells:assert check['Sheet1'][c['cell']].value==c['value'] or (
        isinstance(c['value'],float) and math.isclose(check['Sheet1'][c['cell']].value,c['value'],rel_tol=1e-12))
    completed=sum(isinstance(c['value'],(int,float)) for c in cells)
    audit={'completed_numeric_cells':completed,'total_target_cells':len(cells),
           'pending_cells':[c['cell'] for c in cells if c['run'] is None],
           'original_backup_sha256':sha(TEMPLATE),'output_sha256':sha(OUTPUT),'cells':cells}
    (STUDY/'workbook_audit.json').write_text(json.dumps(audit,indent=2,ensure_ascii=False)+'\n')
    if update_original:
        if sha(ORIGINAL)!=sha(TEMPLATE):
            raise RuntimeError('Original workbook changed during experiments; completed copy is saved, refusing to overwrite concurrent edits.')
        shutil.copy2(OUTPUT,ORIGINAL)
        assert sha(ORIGINAL)==sha(OUTPUT)
    print(json.dumps({'output':windows(OUTPUT),'filled':completed,'total':len(cells),
                      'original_updated':update_original},ensure_ascii=False))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--update-original',action='store_true')
    main(parser.parse_args().update_original)
