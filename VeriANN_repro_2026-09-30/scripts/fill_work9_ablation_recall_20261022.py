"""Apply and independently audit fresh ablation Recall@10 measurements."""
import csv
import hashlib
import json
import math
import shutil
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
STUDY = ROOT / 'experiments/work9-ablation-recall-20261022'
SOURCE = ROOT.parent / 'work9.xlsx'
BACKUP = STUDY / 'work9_before_ablation_recall.xlsx'

def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def read(path):
    return json.loads(Path(path).read_text())

assert not (STUDY / 'APPLIED.json').exists(), 'Already applied'
assert sha(SOURCE) == sha(BACKUP), 'Workbook changed since recall run began'
assert read(STUDY / 'RUN_FINISHED.json')['passed']
records = read(STUDY / 'recall-results.json')
protocol = read(STUDY / 'protocol.json')
assert len(records) == 24
wb = openpyxl.load_workbook(BACKUP)
allowed = set()
for rec in records:
    cell = wb[rec['sheet']][rec['cell']]
    cell.value = rec['hits'] / (rec['samples'] * 10)
    assert math.isclose(cell.value, rec['recall'], abs_tol=1e-12)
    cell.number_format = '0.0%'
    cell.comment = Comment(
        '2026-10-22 Recall@10 独立补测（实测，不是预测）。\n'
        f"N={rec['N']}, D={rec['D']}, k=10；query IDs=100..119；20条查询×3次重复；单线程。\n"
        f"原消融配置不变，variant={rec['variant']} ({rec['label']})，搜索参数={rec['parameter']}。\n"
        f"命中总数={rec['hits']}/600；三个重复使用同一组20条查询，不是60条独立查询。\n"
        '以原数据的精确穷举 top-10 为真值；新返回结果逐条与原性能实验 answers.csv 对齐。\n'
        '本次只补测检索召回率，不重新测量证明或性能，G/H/I 原值保留。\n'
        f"新记录：{rec['output']}/recall.csv\nSHA256={rec['csv_sha256']}\n"
        f"原性能来源：{rec['source_performance_csv']}", 'Codex')
    allowed.add((rec['sheet'], rec['cell']))

for ds, cfg in protocol['mapping'].items():
    ws = wb[cfg['sheet']]
    cell = ws.cell(cfg['header'], 6)
    cell.comment = Comment(
        'Recall@10 = 返回结果与精确 top-10 交集大小 / 10。\n'
        '下方6组 Ours 消融配置于2026-10-22重新实测：20条查询×3次重复，单线程，保持各行原配置。\n'
        '百分比展示；并非对新数据重调参。外部 baseline 行仍为其原始来源，并未在本次重测。\n'
        '详见“消融Recall补测说明”和“消融Recall补测记录”。', 'Codex')
    allowed.add((ws.title, cell.coordinate))
    # Consistent percentage display for the two existing reference rows only.
    for row in range(cfg['header'] + 7, cfg['header'] + 9):
        ref = ws.cell(row, 6)
        assert isinstance(ref.value, (int, float)) and 0 <= ref.value <= 1
        ref.number_format = '0.0%'
        allowed.add((ws.title, ref.coordinate))

notes = wb.create_sheet('消融Recall补测说明')
for row in [
    ['项目', '说明'],
    ['日期', '2026-10-22'],
    ['范围', '四个数据集×6组 Ours 消融配置，共24组；写入原消融表 Recall@10 列，统一百分比展示。'],
    ['实测口径', '每组20条查询（ID100..119）×3次重复，k=10，单线程。共1440次检索。重复不增加独立查询数量。'],
    ['数据规模', 'SIFT1M、DEEP1M、GloVe-100均沿用原消融实验的前1M向量；NYTimes-256为289761向量。'],
    ['真值', '对对应数据规模重新进行精确穷举，计算 top-10 交集；新检索结果与原性能实验答案逐条一致。'],
    ['结果', '新实测均值与旧记录一致；未为了差异修改或预测召回率。'],
    ['性能', '本次为Recall补测；不重新执行证明阶段，不更新原查询/VO/验证性能。'],
    ['外部baseline', '消融表下方外部baseline仅调整Recall显示格式；其值、原来源和原测量/预测状态均保留。'],
    ['配对检查', '1=3、2=4、5=6 的检索结果逐查询核对，共720项通过。'],
    ['协议', str(STUDY / 'protocol.json')],
    ['原表备份', str(BACKUP)],
]: notes.append(row)
details = wb.create_sheet('消融Recall补测记录')
details.append(['数据集', '配置编号', '配置', 'N', 'D', '搜索参数', 'Recall@10', '命中总数', '查询次数', '独立查询数', '原表单元格', '实测CSV', 'SHA256', '原性能CSV'])
for r in records:
    details.append([r['sheet'], r['variant'], r['label'], r['N'], r['D'], r['parameter'], r['hits']/600,
                    r['hits'], r['samples'], r['distinct_queries'], r['cell'], r['output']+'/recall.csv', r['csv_sha256'], r['source_performance_csv']])
    details.cell(details.max_row, 7).number_format = '0.0%'
for ws in [notes, details]:
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='305496')
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 24
    for row in ws:
        for c in row:
            c.alignment = Alignment(vertical='top', wrap_text=True)
notes.column_dimensions['B'].width = 100

review = STUDY / 'work9_ablation_recall_updated.xlsx'
wb.save(review)
new = openpyxl.load_workbook(review)
old = openpyxl.load_workbook(BACKUP)
checked = 0
for ws in old:
    target = new[ws.title]
    assert (ws.max_row, ws.max_column) == (target.max_row, target.max_column)
    for row in ws:
        for c in row:
            nc = target[c.coordinate]
            if (ws.title, c.coordinate) not in allowed:
                assert c.value == nc.value, (ws.title, c.coordinate, 'unrelated value')
                assert (c.comment.text if c.comment else None) == (nc.comment.text if nc.comment else None), (ws.title, c.coordinate, 'unrelated comment')
                assert c.number_format == nc.number_format
                checked += 1
            elif c.column != 6 or c.row <= 130:
                raise AssertionError('Unexpected scope')

executions = pairs = 0
for ds, cfg in protocol['mapping'].items():
    output = ROOT / 'results/work9-ablation-recall-20261022' / ds
    truth = {int(r['query']): set(map(int, r['answer'].split())) for r in csv.DictReader((output/'ground_truth.csv').open())}
    assert set(truth) == set(range(100, 120))
    rows = list(csv.DictReader((output/'recall.csv').open()))
    assert len(rows) == 360
    indexed = {}
    for p in cfg['points']:
        assert sha(p['source']) == p['source_sha256']
        rr = [r for r in rows if int(r['variant']) == p['variant']]
        assert {(int(r['query']), int(r['repeat'])) for r in rr} == {(q,t) for q in range(100,120) for t in range(3)}
        original = {(int(r['query']),int(r['repeat'])): list(map(int,r['answer'].split())) for r in csv.DictReader((Path(p['source']).parent/'answers.csv').open())}
        hits = 0
        for r in rr:
            q,t = int(r['query']),int(r['repeat'])
            ans = list(map(int,r['answer'].split()))
            assert len(ans) == len(set(ans)) == 10
            assert ans == original[q,t]
            h = len(set(ans) & truth[q])
            assert h == int(r['hits']) and math.isclose(h/10,float(r['recall_at_k']),abs_tol=1e-12)
            hits += h
            indexed[p['variant'],q,t] = ans
            executions += 1
        cell = new[cfg['sheet']][p['cell']]
        assert math.isclose(cell.value,hits/600,abs_tol=1e-12) and cell.number_format == '0.0%'
    for a,b in [(1,3),(2,4),(5,6)]:
        for q in range(100,120):
            for t in range(3):
                assert indexed[a,q,t] == indexed[b,q,t]
                pairs += 1
assert executions == 1440 and pairs == 720
audit = dict(passed=True, fresh_query_executions=executions, paired_answer_checks=pairs,
             unrelated_cells_preserved=checked, measured_recall_cells=24, new_workbook_sha256=sha(review),
             backup_sha256=sha(BACKUP), workbook=str(SOURCE))
(STUDY/'AUDIT.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
assert sha(SOURCE) == sha(BACKUP)
shutil.copy2(review, SOURCE)
assert sha(SOURCE) == audit['new_workbook_sha256']
(STUDY/'APPLIED.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
print(json.dumps(audit,ensure_ascii=False,indent=2))
