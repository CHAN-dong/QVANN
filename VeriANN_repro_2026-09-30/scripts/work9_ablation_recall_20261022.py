"""Fresh 20x3 recall runs, preserving the ablation performance measurements."""
import csv,json,math,re,shutil,statistics
from pathlib import Path
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment,Font,PatternFill
import work9_resource_matched as u
R=u.ROOT;S=R/'experiments/work9-ablation-recall-20261022';O=R/'results/work9-ablation-recall-20261022';SOURCE=R.parent/'work9.xlsx';BACKUP=S/'work9_before_ablation_recall.xlsx'
u.STUDY=S;u.RESULTS=O;S.mkdir(parents=True,exist_ok=True)
NAMES={1:'Merkle+IVF',2:'Merkle+HNSW',3:'Merkle+IVF+(i)',4:'Merkle+HNSW+(i)',5:'(i)+(ii)',6:'(i)+(ii)+(iii)'}
DATA={'sift1m':'SIFT1M','nytimes256':'NYTimes-256','glove100':'GloVe-100','deep1m':'DEEP1M'}
assert not (S/'APPLIED.json').exists(),'Already applied'
if not BACKUP.exists():shutil.copy2(SOURCE,BACKUP)
w=openpyxl.load_workbook(BACKUP);mapping={}
for ds,name in DATA.items():
 s=w[name];h=next(c.row for row in s for c in row if c.column==6 and c.row>130 and c.value=='Recall@10');points=[]
 for v in range(1,7):
  cell=s.cell(h+v,6);source=Path(re.findall(r'/mnt/d/[^\s]+?per_query\.csv',cell.comment.text)[-1]);manifest=u.read(source.parent/'manifest.json');raw=list(csv.DictReader(source.open()));params={int(x['ef_search']) for x in raw}
  assert len(params)==1 and manifest['queries']==20 and manifest['repeats']==3
  points.append(dict(variant=v,cell=cell.coordinate,source=str(source),parameter=params.pop(),manifest=manifest,source_sha256=u.sha(source),old_recall=cell.value))
 assert len({(p['manifest']['base'],p['manifest']['M'],p['manifest']['ef_construction'],p['manifest']['dataset_sha256']) for p in points})==1
 mapping[ds]=dict(sheet=name,header=h,points=points)
u.write(S/'protocol.json',dict(scope='Fresh recall-only searches through unchanged ablation search functions; no new performance or proof measurements',queries=20,repeats=3,query_offset=100,k=10,threads=1,variant_pairs=[[1,3],[2,4],[5,6]],mapping=mapping))
binary=S/'ablation_recall_probe';source=R/'scripts/ablation_recall_probe.cpp'
if not binary.exists():
 original=R/'src/experiments/work9_ablation.cpp';code=original.read_text();signature='int main(int argc,char**argv)try{';assert code.count(signature)==1
 (S/'ablation_recall_protocol_include.cpp').write_text(code.replace(signature,'int archived_ablation_protocol_main(int argc,char**argv)try{'))
 cmd=u.read(R/'experiments/work9-ablation-20261015/build_command.json');cmd=[str(source) if str(x).endswith('/src/experiments/work9_ablation.cpp') else str(x) for x in cmd];cmd[cmd.index('-o')+1]=str(binary)
 cmd.insert(1,'-I'+str(S))
 u.command(cmd,S/'build.log',timeout=300)
 u.write(S/'build.json',dict(command=cmd,binary_sha256=u.sha(binary),harness_sha256=u.sha(source),search_source_sha256=u.sha(R/'src/experiments/work9_ablation.cpp')))
records=[];paired=0
for ds,m in mapping.items():
 ps=m['points'];cfg=ps[0]['manifest'];out=O/ds;cache=R/'ANNProof-24FGCS/cache'
 stem=cfg['dataset_sha256'][:16];graph=cache/f'{stem}_M{cfg["M"]}_efc{cfg["ef_construction"]}.index';ivf=cache/f'{stem}_ivf256_train65536_iter12_seed42_v1.bin'
 assert graph.exists() and ivf.exists()
 before={str(p):u.sha(p) for p in [graph,ivf]}
 if not (out/'PASS.txt').exists():
  cmd=[binary,'--dataset',R/'dataset'/ds,'--base',cfg['base'],'--cache',cache,'--output',out,'--queries',20,'--query-offset',100,'--repeats',3,'--build-threads',1,'--M',cfg['M'],'--ef-construction',cfg['ef_construction'],'--k',10,'--ef-search',','.join(str(p['parameter']) for p in ps)]
  u.command(cmd,S/(ds+'.log'),bounded=True,timeout=3600)
  u.write(out/'environment.json',dict(command=list(map(str,cmd)),binary_sha256=u.sha(binary),cache_sha256=before))
 assert before=={str(p):u.sha(p) for p in [graph,ivf]}
 identity=u.read(out/'identity.json');assert identity['dataset_sha256']==cfg['dataset_sha256']
 rows=list(csv.DictReader((out/'recall.csv').open()));assert len(rows)==360
 truth={int(r['query']):set(map(int,r['answer'].split())) for r in csv.DictReader((out/'ground_truth.csv').open())};by={}
 for p in ps:
  v=p['variant'];rr=[r for r in rows if int(r['variant'])==v];assert len(rr)==60
  oldanswers={(int(r['query']),int(r['repeat'])):list(map(int,r['answer'].split())) for r in csv.DictReader((Path(p['source']).parent/'answers.csv').open())}
  assert {(int(r['query']),int(r['repeat'])) for r in rr}==set(oldanswers)
  for r in rr:
   q,t=int(r['query']),int(r['repeat']);answer=list(map(int,r['answer'].split()));hits=len(set(answer)&truth[q]);assert len(answer)==len(set(answer))==10
   assert answer==oldanswers[q,t],(ds,v,q,t,'fresh search differs from original performance experiment')
   assert hits==int(r['hits']) and math.isclose(hits/10,float(r['recall_at_k']),abs_tol=1e-12)
   assert int(r['parameter'])==p['parameter'];by[v,q,t]=answer
  rec=dict(dataset=ds,sheet=m['sheet'],variant=v,label=NAMES[v],N=cfg['base'],D=identity['D'],parameter=p['parameter'],recall=statistics.fmean(float(r['recall_at_k']) for r in rr),hits=sum(int(r['hits']) for r in rr),samples=60,distinct_queries=20,output=str(out),csv_sha256=u.sha(out/'recall.csv'),source_performance_csv=p['source'],cell=p['cell'])
  assert math.isclose(rec['recall'],statistics.fmean(float(r['recall_at_k']) for r in csv.DictReader(Path(p['source']).open())),abs_tol=1e-12)
  records.append(rec)
 for a,b in [(1,3),(2,4),(5,6)]:
  for q in range(100,120):
   for t in range(3):assert by[a,q,t]==by[b,q,t];paired+=1
 print('RECALL_COMPLETE',ds,[(p['variant'],p['recall']) for p in records if p['dataset']==ds],flush=True)
u.write(S/'recall-results.json',records);u.write(S/'RUN_FINISHED.json',dict(passed=True,measured_variant_points=24,fresh_query_executions=1440,paired_answer_checks=paired))
print('ALL_RECALL_EXPERIMENTS_COMPLETE',flush=True)
