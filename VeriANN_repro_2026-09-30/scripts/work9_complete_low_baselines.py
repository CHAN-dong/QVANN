"""Fill only the missing GloVe/DEEP peer low-recall curves; freeze all defaults."""
import argparse,json,shutil,sys,time
from collections import defaultdict
from pathlib import Path
import openpyxl
from openpyxl.comments import Comment
import work9_resource_matched as util

ROOT=util.ROOT;SOURCE=ROOT.parent/'work9.xlsx'
STUDY=ROOT/'experiments/work9-low-baselines-20261021'
OUT=ROOT/'results/work9-low-baselines-20261021'
BASE=STUDY/'work9_before_missing_low_curves.xlsx'
DATASETS=['glove100','deep1m'];METHODS=['VLAH','ANNProof','Ours']
TARGETS=[.15,.20,.25,.30,.35,.40,.45,.50,.55]
util.STUDY=STUDY;util.RESULTS=OUT
read=util.read;write=util.write

def fill():
 wb=openpyxl.load_workbook(BASE);allowed_cells=set();filled=0
 records=[read(f) for f in sorted((STUDY/'records').glob('*.json'))]
 for ds in DATASETS:
  ws=wb[util.DATA[ds]['sheet']]
  for i,method in enumerate(METHODS):
   rec=next((r for r in records if r['dataset']==ds and r['method']==method),None)
   for j,t in enumerate(TARGETS):
    col=next(c.column for c in ws[16] if isinstance(c.value,(int,float)) and abs(c.value-t*100)<1e-8)
    for start,metric in [(20,'vo_kb'),(32,'verify_ms'),(44,'query_ms')]:
     for rr in [start+2*i,start+2*i+1]:allowed_cells.add((ws.title,rr,col))
     if not rec:continue
     r=next(r for r in rec['points'] if abs(r['target']-t)<1e-10)
     assert r['status']=='passed' and r['samples']==60
     for rr,value in [(start+2*i,r['recall']*100),(start+2*i+1,r[metric])]:
      cell=ws.cell(rr,col);cell.value=value
      cell.comment=Comment('仅补充低召回率曲线；默认约95%配置及其数据保持不变。20条查询IDs100..119×3次，单线程，完整VO生成与验证。目标和实际召回不同，填实际值。\n'+json.dumps(r,ensure_ascii=False),'低召回补测来源');filled+=1
 for name in ['低召回缺项补测说明','低召回缺项实测记录']:
  if name in wb:del wb[name]
 tab=wb.create_sheet('低召回缺项补测说明')
 for row in [['项目','说明'],['补测范围','GloVe-100、DEEP1M上的VLAH、ANNProof、Ours三方案，15/20/25/30/35/40/45/50/55%目标曲线，共六组。'],['协议','每个参数点20条查询IDs100..119×3次，单线程；校准与计时使用相同查询，不宣称独立留出评估。仅V3DB维持用户要求的单次出证口径。'],['参数离散','用最低预算的有效配置达到目标，保留实际召回。最小合法预算已超过低目标时多个目标共用同一实测点，不捏造更低召回。'],['保持原样','四数据集默认表、索引、规模、K、消融、全部V3DB数据及原60–99%曲线均与本轮前备份一致。'],['状态','全部补测完成' if (STUDY/'RUN_FINISHED.json').exists() else f'已完成{len(records)}/6组曲线'],['备份',str(BASE)]]:tab.append(row)
 tab.column_dimensions['B'].width=140
 tab=wb.create_sheet('低召回缺项实测记录');tab.append(['dataset','method','target','actual recall','query_ms','verify_ms','VO KB','ef/E','probes','samples','source'])
 for rec in records:
  for r in rec['points']:tab.append([rec['dataset'],rec['method'],r['target'],r['recall'],r['query_ms'],r['verify_ms'],r['vo_kb'],r['ef'],r['probes'],r['samples'],r['output']])
 base=openpyxl.load_workbook(BASE);checks=0
 for cfg in util.DATA.values():
  a=base[cfg['sheet']];b=wb[cfg['sheet']]
  for row in a:
   for x in row:
    if (a.title,x.row,x.column) in allowed_cells:continue
    y=b[x.coordinate]
    assert x.value==y.value and x._style==y._style,(a.title,x.coordinate)
    assert (x.comment.text if x.comment else None)==(y.comment.text if y.comment else None)
    checks+=1
 output=STUDY/'work9_low_curves_completed.xlsx';wb.save(output)
 marker=STUDY/'last-written.txt';allowed={util.sha(BASE)}
 if marker.exists():allowed.add(marker.read_text().strip())
 if util.sha(SOURCE) in allowed:
  try:shutil.copy2(output,SOURCE);marker.write_text(util.sha(SOURCE))
  except PermissionError:pass
 write(STUDY/'progress.json',dict(groups_completed=len(records),groups_total=6,filled_metric_cells=filled,untouched_cells_checked=checks,delivered=util.sha(SOURCE)==util.sha(output),time=time.strftime('%Y-%m-%dT%H:%M:%S%z')))

def work():
 import fcntl
 lock=(STUDY/'runner.lock').open('w');fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
 if not BASE.exists():shutil.copy2(SOURCE,BASE)
 write(STUDY/'protocol.json',dict(datasets=DATASETS,methods=METHODS,targets=TARGETS,queries=20,repeats=3,query_offset=100,threads=1,scope='only missing low-recall peer curves; all defaults frozen'))
 for ds in DATASETS:
  cfg=util.DATA[ds]
  for method in METHODS:
   f=STUDY/'records'/f'{ds}-{method}.json'
   if f.exists():continue
   pilots=util.pilot(ds,method,cfg['N'],goal=max(TARGETS))
   selections=[dict(target=t,point=util.select(pilots,t,method)) for t in TARGETS]
   assert all(r['point'] for r in selections),(ds,method,'extend pilot grid')
   groups=defaultdict(set)
   for r in selections:groups[r['point']['ef']].add(r['point']['probes'])
   measured={}
   batches=[(sorted(groups),[1],'low-all')] if method!='VLAH' else [([e],sorted(p),f'low-e{e}') for e,p in groups.items()]
   for efs,probes,tag in batches:
    for r in util.evaluate(ds,method,cfg['N'],ds,tag,efs,probes,[10]):measured[r['ef'],r['probes']]=r
   points=[]
   for item in selections:
    p=item['point'];r=measured[p['ef'],p['probes']]
    assert abs(r['recall']-p['recall'])<1e-10 and r['recall']>=item['target']-1e-12
    points.append(dict(r,target=item['target'],selection=p))
   write(f,dict(dataset=ds,method=method,points=points));fill()
 write(STUDY/'RUN_FINISHED.json',dict(status='six missing peer low-recall curve groups completed'));fill()

if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('action',choices=['run','fill']);a=p.parse_args();STUDY.mkdir(parents=True,exist_ok=True)
 try:
  if a.action=='fill':fill()
  else:work()
 except BaseException as exc:
  import traceback;write(STUDY/'runner_failure.json',dict(error=str(exc),traceback=traceback.format_exc()));raise
