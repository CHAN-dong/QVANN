import csv,json,math,statistics
from pathlib import Path
import openpyxl
from work9_complete_low_baselines import STUDY,BASE,SOURCE,DATASETS,METHODS,TARGETS,util,read,write

assert (STUDY/'RUN_FINISHED.json').exists(),'experiments still running'
out=STUDY/'work9_low_curves_completed.xlsx'
wb=openpyxl.load_workbook(out);base=openpyxl.load_workbook(BASE)
changed=set();verified={};cells=0;groups=[]
for ds in DATASETS:
 ws=wb[util.DATA[ds]['sheet']]
 for i,method in enumerate(METHODS):
  rec=read(STUDY/'records'/f'{ds}-{method}.json');groups.append((ds,method))
  for r in rec['points']:
   key=r['output'],r['ef'],r['probes']
   if key not in verified:
    folder=Path(r['output']);rows=[x for x in csv.DictReader((folder/'per_query.csv').open()) if (int(x['k']),int(x['ef_search']),int(x['probes']))==(10,r['ef'],r['probes'])]
    assert len(rows)==60 and all(x['verified']=='1' for x in rows)
    assert {(int(x['query']),int(x['repeat'])) for x in rows}=={(q,t) for q in range(100,120) for t in range(3)}
    mean=lambda key:statistics.fmean(float(x[key]) for x in rows)
    expected=dict(recall=mean('recall_at_k'),query_ms=mean('search_ms')+mean('vo_generation_ms'),verify_ms=mean('verification_ms'),vo_kb=mean('vo_bytes')/1000)
    for k,v in expected.items():assert math.isclose(r[k],v,rel_tol=1e-10,abs_tol=1e-8),(ds,method,k)
    attacks=list(csv.DictReader((folder/'tamper_tests.csv').open()));assert attacks and all(x['rejected']=='1' for x in attacks)
    verified[key]=len(rows)
   col=next(c.column for c in ws[16] if isinstance(c.value,(int,float)) and abs(c.value-r['target']*100)<1e-8)
   for start,k in [(20,'vo_kb'),(32,'verify_ms'),(44,'query_ms')]:
    for row,value in [(start+2*i,r['recall']*100),(start+2*i+1,r[k])]:
     assert math.isclose(ws.cell(row,col).value,value,rel_tol=1e-10,abs_tol=1e-8)
     assert ws.cell(row,col).comment
     changed.add((ws.title,row,col));cells+=1
 for row in ws:
  for c in row:
   assert c.value!='待低召回曲线实测',(ws.title,c.coordinate)
preserved=0
for cfg in util.DATA.values():
 a=base[cfg['sheet']];b=wb[cfg['sheet']]
 for row in a:
  for x in row:
   if (a.title,x.row,x.column) in changed:continue
   y=b[x.coordinate]
   assert x.value==y.value and x._style==y._style,(a.title,x.coordinate)
   assert (x.comment.text if x.comment else None)==(y.comment.text if y.comment else None)
   preserved+=1
assert util.sha(SOURCE)==util.sha(out),'original workbook not updated'
result=dict(status='PASS',groups=len(groups),distinct_measured_configurations=len(verified),honest_queries_checked=sum(verified.values()),filled_low_curve_cells=cells,all_other_main_cells_preserved=preserved,pending_low_curve_cells=0,workbook_sha256=util.sha(SOURCE))
write(STUDY/'independent-validation.json',result);print(json.dumps(result,indent=2))
