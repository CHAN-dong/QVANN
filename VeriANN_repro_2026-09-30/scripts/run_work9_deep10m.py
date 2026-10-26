"""Add authenticated DEEP 5M/10M scale measurements to the current workbook."""
import csv
import json
import math
from pathlib import Path
import shutil
import statistics
import sys
import openpyxl
from openpyxl.comments import Comment
from openpyxl.chart import ScatterChart, Series, Reference
import work9_multidata as w
import run_work9_multidata as r

ROOT=w.ROOT
STUDY=ROOT/'experiments/work9-deep10m-20261015'
RESULTS=ROOT/'results/work9-deep10m-20261015'
SOURCE=w.SOURCE
BACKUP=STUDY/'work9_before_deep10m.xlsx'
MERGE_BASELINE=STUDY/'work9_merge_baseline.xlsx'
OUTPUT=STUDY/'work9_completed_deep10m.xlsx'
OLD=ROOT/'experiments/work9-multidata-20261015'
WIDTH=json.loads((OLD/'deep1m-lsh-width.json').read_text())['width']
METHODS=w.METHODS

def initialize():
    STUDY.mkdir(parents=True,exist_ok=True);RESULTS.mkdir(parents=True,exist_ok=True)
    if not BACKUP.exists():shutil.copy2(SOURCE,BACKUP)
    if not OUTPUT.exists():shutil.copy2(BACKUP,OUTPUT)
    # Only selection helpers are reused; their events and raw outputs live in this new study.
    r.STUDY=STUDY;r.RESULTS=RESULTS;w.STUDY=STUDY
    r.DATASETS['deep10m']=('DEEP10M',10000000,96)
    protocol=dict(dataset='official DEEP1B first 10M',sizes=[5000000,10000000],queries=20,
                  query_offset=100,repeats=3,k=10,target_recall=.95,M=16,ef_construction=200,
                  build_threads=16,query_threads=1,LSH_width=WIDTH,
                  width_rule='retain the existing DEEP1M base-only calibrated width across all prefix sizes',
                  query_time='search + actual VO serialization',verification='decode + authentication + replay',
                  vo_unit='decimal KB (1000 bytes)',scope='20-query quick preview',
                  ground_truth='exactly recomputed for each N; same query IDs 100..119')
    (STUDY/'protocol.json').write_text(json.dumps(protocol,indent=2)+'\n')

def scan(method,n,efs,probes,tag):
    out=r.pilot('deep10m',method,n,efs,probes,tag,WIDTH)
    return r.pilot_groups(out)

def calibrate(method,n):
    file=STUDY/f'selection-{method}-{n}.json'
    if file.exists():return json.loads(file.read_text())
    if method=='VLAH':
        buckets=math.ceil(n/20000)
        grid=sorted({1,*[math.ceil(buckets*i/10) for i in range(1,11)]})
        for ef in [100,200,400,800,1600,3200,6400]:
            points=scan(method,n,[ef],grid,f'{n}-ef{ef}-coarse')
            if max(p['recall'] for p in points)>=.95-1e-12:break
        parameter='probes'
    else:
        points=[]
        for phase,grid in enumerate([[10,16,24,32,48,64,96,128,192,256,384,512,768,1024,1536,2048],[4096,8192],[16384,32768]]):
            points+=scan(method,n,grid,[1],f'{n}-phase{phase}')
            if max(p['recall'] for p in points if p['valid'])>=.95-1e-12:break
        parameter='ef'
    for refinement in range(5):
        candidates=[p for p in points if p['valid'] and p['recall']>=.95-1e-12]
        assert candidates,(method,n,'target not reached')
        chosen=min(candidates,key=lambda p:p[parameter])
        lower=max([p[parameter] for p in points if p['valid'] and p[parameter]<chosen[parameter] and p['recall']<.95-1e-12],default=0)
        upper=chosen[parameter]
        # Probe count is refined to an adjacent tested bracket. Other parameters
        # are refined when their coarse recall overshoots 95% by more than 1pp.
        if upper-lower<=1 or (method!='VLAH' and chosen['recall']<=.96+1e-12):break
        step=max(1,math.ceil((upper-lower)/8))
        grid=sorted(set(range(lower+step,upper,step)))
        seen={p[parameter] for p in points}
        grid=[v for v in grid if v not in seen and (parameter!='ef' or v>=10)]
        if not grid:break
        tag=f'{n}-refine{refinement}-{parameter}-{lower}-{upper}'
        points+=scan(method,n,grid if parameter=='ef' else [chosen['ef']],grid if parameter=='probes' else [1],tag)
    chosen=dict(min([p for p in points if p['valid'] and p['recall']>=.95-1e-12],key=lambda p:p[parameter]),target=.95)
    file.write_text(json.dumps(chosen,indent=2)+'\n')
    w.event(action='selected_scale_point',method=method,base=n,**chosen)
    return chosen

def evaluate(method,n,point):
    file=STUDY/f'evaluation-{method}-{n}.json'
    if file.exists():return json.loads(file.read_text())
    out=RESULTS/f'{method}-size-{n}'
    if not (out/'PASS.txt').exists():
        if method=='Ours':
            cmd=[sys.executable,str(ROOT/'scripts/run_ours_benchmark.py'),'--source',str(ROOT/'build/veriann_lut_benchmark.cpp'),
                 '--binary',str(ROOT/'build/veriann_lut_benchmark'),'--E',str(point['ef'])]
        else:
            folder='ANNProof-24FGCS' if method=='ANNProof' else 'VLAH-25TrustCom'
            cmd=[sys.executable,str(ROOT/folder/'run.py'),'--ef-search',str(point['ef']),
                 '--probes',str(point['probes']),'--lsh-width',str(WIDTH)]
        cmd+=['--dataset',str(ROOT/'dataset/deep10m'),'--base',str(n),'--queries','20','--query-offset','100',
              '--repeats','3','--M','16','--ef-construction','200','--build-threads','16','--k','10','--output',str(out),'--skip-build']
        r.launch(cmd,out,STUDY/f'{method}-size-{n}.log')
    rows=w.measurements(out);assert len(rows)==1 and rows[0]['samples']==60 and rows[0]['recall']>=.95-1e-12
    r.compare_pilot(out,[point])
    manifest=json.loads((out/'manifest.json').read_text())
    assert (manifest['base'],manifest['queries'],manifest['query_offset'],manifest['repeats'],manifest['M'],manifest['ef_construction'],manifest['query_threads'])==(n,20,100,3,16,200,1)
    attacks=list(csv.DictReader((out/'tamper_tests.csv').open()))
    assert len(attacks)=={'ANNProof':8,'VLAH':11,'Ours':13}[method]
    assert all(a['rejected']=='1' for a in attacks)
    job=dict(dataset='deep10m',method=method,base=n,output=str(out),selection=point,metrics=rows[0],
             query_count=20,repeats=3,M=16,ef_construction=200,width=WIDTH,tamper_rejections=len(attacks))
    file.write_text(json.dumps(job,indent=2)+'\n')
    checkpoint()
    w.event(action='scale_measurement_passed',method=method,base=n,recall=rows[0]['recall'],query_ms=rows[0]['query_ms'],verify_ms=rows[0]['verify_ms'],vo_kb=rows[0]['vo_kb'])
    return job

def old_point(method,n):
    job=json.loads((OLD/f'evaluation-deep1m-{method}-size-{n}.json').read_text())
    records=w.measurements(Path(job['output']))
    if job.get('filter'):records=[x for x in records if all(x[k]==v for k,v in job['filter'].items())]
    assert len(records)==1
    return records[0]

def preserve_current_workbook():
    """Merge the authorized cells into a stable, backed-up latest workbook."""
    marker=STUDY/'last_written_sha256.txt'
    allowed={w.sha(BACKUP)}
    if marker.exists():allowed.add(marker.read_text().strip())
    current=w.sha(SOURCE)
    if current in allowed:return
    latest=openpyxl.load_workbook(SOURCE)
    assert 'DEEP1M' in latest.sheetnames,'DEEP sheet was renamed; inspect before merging'
    assert [latest['DEEP1M'].cell(53,c).value for c in [9,10]]==['5*10^6','10^7'],'scale columns moved; inspect before merging'
    snapshot=STUDY/f'work9_external_save_{current[:12]}.xlsx'
    if not snapshot.exists():shutil.copy2(SOURCE,snapshot)
    assert w.sha(snapshot)==current and w.sha(SOURCE)==current,'workbook changed during merge preparation'
    shutil.copy2(snapshot,MERGE_BASELINE);shutil.copy2(snapshot,OUTPUT)
    marker.write_text(current+'\n')
    w.event(action='source_rebased_preserving_user_changes',source_sha256=current,snapshot=str(snapshot))

def checkpoint():
    preserve_current_workbook()
    wb=openpyxl.load_workbook(OUTPUT);ws=wb['DEEP1M'];jobs=[json.loads(f.read_text()) for f in sorted(STUDY.glob('evaluation-*.json'))]
    mapping={}
    for row in sum(w.SIZE_ROWS.values(),[]):
        for col in [9,10]:
            ws.cell(row,col,'待测').comment=Comment('DEEP10M已准备；等待对应规模的完整协议测量通过后填入。','实验进度')
    for job in jobs:
        method,n=job['method'],job['base'];metric=w.measurements(Path(job['output']))[0]
        col={5000000:9,10000000:10}[n]
        for row,key in zip(w.SIZE_ROWS[method],['vo_kb','verify_ms','query_ms']):
            c=ws.cell(row,col,metric[key]);c.number_format='0.000'
            c.comment=Comment(f"DEEP1B prefix; N={n}; d=96; {method}; k=10; efSearch/E={metric['ef']}; probes={metric['probes']}; M=16; efConstruction=200.\nPREVIEW: queries100..119 x 3 repeats; samples=60; actual Recall@10={metric['recall']:.4%}.\nQuery=search+VO generation; Verify=decoding+authentication+replay; VO KB=1000 bytes.\nLSH width={WIDTH if method=='VLAH' else 'not used'}.\n{job['output']}/per_query.csv",'实测来源')
            mapping[('DEEP1M',c.coordinate)]=['DEEP1M',c.coordinate,metric[key],key,10,metric['ef'],metric['probes'],job['output']+'/per_query.csv']
    # Only newly authorized 5M/10M cells change; all smaller-scale measurements stay intact.
    for addr in ['D52','D62','D71']:ws[addr]='DEEP1B real prefix subsets: 100K / 500K / 1M / 5M / 10M'
    notes=wb['新增数据实验说明']
    for row in notes.iter_rows(min_row=2):
        if row[0].value=='规模':row[1].value='保留100K/500K/1M/5M/10M；GIST与GloVe的5M/10M不适用。DEEP使用官方DEEP1B前缀扩充到10M，新增5M/10M实测；其余DEEP实验仍为1M。'
    label='DEEP规模扩展'
    text='5M/10M统一20条查询×3次重复，k=10，按约95%Recall选参；M=16、efConstruction=200；LSH宽度沿用原DEEP1M。1M总体/召回/k及其他数据集结果保留。扩展记录：'+str(STUDY)
    rows=[row for row in notes.iter_rows(min_row=2) if row[0].value==label]
    if rows:rows[0][1].value=text
    else:notes.append([label,text])
    source=wb['新增数据单元格来源']
    retained=[list(row) for row in source.iter_rows(min_row=2,values_only=True) if (row[0],row[1]) not in mapping]
    source.delete_rows(2,source.max_row-1)
    for row in retained+list(mapping.values()):source.append(row)
    title='DEEP扩展规模测量'
    if title in wb.sheetnames:del wb[title]
    sheet=wb.create_sheet(title)
    sheet.append(['method','N','k','efSearch/E','probes','actual Recall@10','Query ms','Verify ms','VO KB','samples','raw CSV'])
    groups={}
    for method in METHODS:
        for n in [100000,500000,1000000,5000000,10000000]:
            match=next((j for j in jobs if j['method']==method and j['base']==n),None)
            if n<=1000000:metric=old_point(method,n)
            elif match:metric=w.measurements(Path(match['output']))[0]
            else:continue
            sheet.append([method,n,10,metric['ef'],metric['probes'],metric['recall'],metric['query_ms'],metric['verify_ms'],metric['vo_kb'],metric['samples'],metric['output']+'/per_query.csv'])
            groups.setdefault(method,[]).append(sheet.max_row)
    for row in sheet.iter_rows(min_row=2):row[5].number_format='0.0%'
    for index,(metric,col,unit) in enumerate([('Query time',7,'ms'),('VO size',9,'KB'),('Verify time',8,'ms')]):
        chart=ScatterChart();chart.title='DEEP: Dataset size vs. '+metric
        chart.x_axis.title='Database vectors';chart.x_axis.scaling.min=0;chart.x_axis.scaling.max=10000000
        chart.y_axis.title=metric+' ('+unit+')';chart.y_axis.scaling.min=0;chart.width=22;chart.height=12
        for mi,method in enumerate(METHODS):
            ids=groups[method]
            series=Series(Reference(sheet,min_col=col,min_row=min(ids),max_row=max(ids)),Reference(sheet,min_col=2,min_row=min(ids),max_row=max(ids)),title=method)
            series.marker.symbol=['triangle','circle','square'][mi];series.marker.size=5;chart.series.append(series)
        sheet.add_chart(chart,'M'+str(1+index*24))
    wb.save(OUTPUT)
    marker=STUDY/'last_written_sha256.txt';allowed={w.sha(BACKUP)}
    if marker.exists():allowed.add(marker.read_text().strip())
    delivered=False
    if w.sha(SOURCE) in allowed:
        try:shutil.copy2(OUTPUT,SOURCE);marker.write_text(w.sha(SOURCE)+'\n');delivered=True
        except PermissionError:w.event(action='workbook_locked',completed_copy=str(OUTPUT))
    else:w.event(action='workbook_changed_externally',completed_copy=str(OUTPUT))
    progress=dict(completed_runs=len(jobs),expected_runs=6,measured_added_cells=len(mapping),expected_added_cells=18,original_updated=delivered,output_sha256=w.sha(OUTPUT))
    (STUDY/'progress.json').write_text(json.dumps(progress,indent=2)+'\n')
    w.event(action='deep10m_checkpoint',**progress)

def verify():
    wb=openpyxl.load_workbook(OUTPUT);before=openpyxl.load_workbook(MERGE_BASELINE if MERGE_BASELINE.exists() else BACKUP)
    allowed={'DEEP1M':{f'{col}{row}' for col in 'IJ' for row in sum(w.SIZE_ROWS.values(),[])}|{'D52','D62','D71'}}
    # Source mapping is appended; the notes' capacity statement is intentionally updated.
    for name in before.sheetnames:
        if name in ['新增数据实验说明','新增数据单元格来源','DEEP扩展规模测量']:continue
        assert len(wb[name]._charts)==len(before[name]._charts)
        assert set(map(str,wb[name].merged_cells.ranges))==set(map(str,before[name].merged_cells.ranges))
        for row in before[name]:
            for c in row:
                if c.coordinate not in allowed.get(name,set()):assert wb[name][c.coordinate].value==c.value,(name,c.coordinate)
    jobs=[json.loads(f.read_text()) for f in STUDY.glob('evaluation-*.json')];assert len(jobs)==6
    hashes={};total=0;attacks=0
    for j in jobs:
        out=Path(j['output']);raw=list(csv.DictReader((out/'per_query.csv').open()))
        assert len(raw)==60 and all(x['verified']=='1' for x in raw)
        assert {(int(x['query']),int(x['repeat'])) for x in raw}=={(q,t) for q in range(100,120) for t in range(3)}
        metric=w.measurements(out)[0];assert metric['recall']>=.95-1e-12
        assert (out/'query0.vo').stat().st_size==int(raw[0]['vo_bytes'])
        manifest=json.loads((out/'manifest.json').read_text());key=j['base']
        assert (manifest['base'],manifest['queries'],manifest['query_offset'],manifest['repeats'],manifest['M'],manifest['ef_construction'],manifest['query_threads'])==(key,20,100,3,16,200,1)
        assert manifest['metric']=='l2' and manifest['proof']=='multiproof'
        if j['method']=='VLAH':assert float((out/'lsh_parameters.txt').read_text().splitlines()[0].split()[3])==WIDTH
        if key in hashes:assert hashes[key]==manifest['dataset_sha256']
        hashes[key]=manifest['dataset_sha256']
        env=json.loads((out/'environment.json').read_text())
        for path,expected in env['source_sha256'].items():assert w.sha(ROOT/path)==expected,path
        tests=list(csv.DictReader((out/'tamper_tests.csv').open()));assert all(x['rejected']=='1' for x in tests)
        assert len(tests)=={'ANNProof':8,'VLAH':11,'Ours':13}[j['method']]
        attacks+=len(tests);total+=len(raw)
        for row,name in zip(w.SIZE_ROWS[j['method']],['vo_kb','verify_ms','query_ms']):
            cell=wb['DEEP1M'].cell(row,{5000000:9,10000000:10}[key])
            assert math.isclose(cell.value,metric[name],rel_tol=1e-12) and str(out) in cell.comment.text
        assert json.loads((out/'pilot_equivalence.json').read_text())['status']=='PASS'
    assert total==360 and attacks==64
    assert len(wb['DEEP扩展规模测量']._charts)==3
    # Every previous provenance row must remain available; new cells add exactly 18 rows.
    extension_cells=allowed['DEEP1M']-{'D52','D62','D71'}
    old_sources=[row for row in list(before['新增数据单元格来源'].values)[1:] if not (row[0]=='DEEP1M' and row[1] in extension_cells)]
    new_sources=list(wb['新增数据单元格来源'].values)[1:]
    assert all(row in new_sources for row in old_sources) and len(new_sources)==len(old_sources)+18
    assert w.sha(SOURCE)==w.sha(OUTPUT),'original workbook could not be updated'
    result=dict(status='PASS',new_scale_points=[5000000,10000000],methods=METHODS,query_count=20,repeats=3,
                measured_added_cells=18,formal_measurements=total,tamper_rejections=attacks,
                original_updated=True,workbook_sha256=w.sha(OUTPUT),backup_sha256=w.sha(BACKUP),
                preservation_baseline_sha256=w.sha(MERGE_BASELINE if MERGE_BASELINE.exists() else BACKUP))
    (STUDY/'final_validation.json').write_text(json.dumps(result,indent=2)+'\n')
    lines=['# DEEP 5M / 10M scale preview','', '20 queries (IDs 100..119) x 3 repeats; k=10; target Recall@10 >=95%.',
           'M=16, efConstruction=200; query/VO/verification single-thread; decimal KB.','',
           '| N | Method | Recall % | Query ms | Verify ms | VO KB | efSearch/E | probes |','|---:|---|---:|---:|---:|---:|---:|---:|']
    for j in sorted(jobs,key=lambda x:(x['base'],x['method'])):
        m=w.measurements(Path(j['output']))[0]
        lines.append(f"| {j['base']:,} | {j['method']} | {m['recall']*100:.2f} | {m['query_ms']:.3f} | {m['verify_ms']:.3f} | {m['vo_kb']:.3f} | {m['ef']} | {m['probes']} |")
    lines+=['','The smaller DEEP scale points and all other datasets were retained. The main DEEP overall/recall/k experiments remain at 1M.',
            'Parameters were calibrated on the same 20-query workload; this is a quick preview, not a held-out recall estimate.',
            'PASS: 360 full-protocol measurements, 64 tamper rejections, 18 cells checked against raw CSV.']
    (STUDY/'RESULTS.md').write_text('\n'.join(lines)+'\n')
    print(json.dumps(result),flush=True)

def main():
    initialize()
    assert json.loads((ROOT/'dataset/deep10m/validation.json').read_text())['status']=='PASS'
    for n in [5000000,10000000]:
        for method in METHODS:evaluate(method,n,calibrate(method,n))
    checkpoint();verify()

if __name__=='__main__':main()
