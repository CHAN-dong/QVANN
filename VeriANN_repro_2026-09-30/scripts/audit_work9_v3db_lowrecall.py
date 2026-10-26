"""Check real fixed-query proofs, calibrations and frozen default tables."""
import csv,json,math
from pathlib import Path
import openpyxl
from work9_v3db_lowrecall import STUDY,BASE,SOURCE,DATA,HIGH,LOW,TARGETS,JOBS,points,single

book=STUDY/'work9_lowrecall.xlsx'
w=openpyxl.load_workbook(book);base=openpyxl.load_workbook(BASE)
counts=0;oldcurve=0
for cfg in DATA.values():
 a=base[cfg['sheet']];b=w[cfg['sheet']]
 curve={16,28,40}|set(range(20,28))|set(range(32,40))|set(range(44,52))
 for row in a:
  for c in row:
   if c.row in curve:continue
   d=b[c.coordinate]
   assert c.value==d.value and c._style==d._style,(a.title,c.coordinate)
   assert (c.comment.text if c.comment else None)==(d.comment.text if d.comment else None)
   counts+=1
 for row in list(range(20,26))+list(range(32,38))+list(range(44,50)):
  for j in range(len(HIGH)):
   c=a.cell(row,6+j);d=b.cell(row,6+len(LOW)+j)
   assert c.value==d.value and c._style==d._style,(a.title,row,j)
   assert (c.comment.text if c.comment else None)==(d.comment.text if d.comment else None)
   oldcurve+=1
raw=0
records=points()
for p in records:
 out=Path(p['output']);rows=list(csv.DictReader((out/'per_query.csv').open()))
 row=next(r for r in rows if int(r['query'])==100 and int(r['repeat'])==0)
 assert row['verified']=='1' and p['samples']==1
 for k in ['query_ms','verify_ms']:assert math.isclose(float(row[k]),p[k],rel_tol=1e-12)
 assert math.isclose(float(row['vo_bytes'])/1000,p['vo_kb'],rel_tol=1e-12)
 proof=out/'query0.proof';assert proof.stat().st_size==int(row['vo_bytes'])
 tamper=json.loads((out/'tamper.json').read_text());assert tamper['proof_byte_change_rejected'] and tamper['wrong_query_rejected']
 cal=json.loads(Path(p['recall_source']).read_text())
 point=next(r for r in cal['points'] if r['n_probe']==p['n_probe'] and r['k']==10)
 assert math.isclose(point['recall'],p['recall20'],rel_tol=1e-12)
 raw+=1
metric_cells=0
for ds,cfg in DATA.items():
 ws=w[cfg['sheet']];candidates=[r for r in records if r['dataset']==ds]
 for j,t in enumerate(TARGETS):
  good=[r for r in candidates if r['recall20']>=t-1e-12]
  if not good:continue
  r=min(good,key=lambda r:(r['query_ms'],r['recall20']))
  for row,k in [(26,'vo_kb'),(38,'verify_ms'),(50,'query_ms')]:
   assert math.isclose(ws.cell(row,6+j).value,r['recall20']*100,rel_tol=1e-12)
   assert math.isclose(ws.cell(row+1,6+j).value,r[k],rel_tol=1e-12)
   metric_cells+=2
finished=(STUDY/'RUN_FINISHED.json').exists()
if finished:
 additional=[]
 for ds,jobs in JOBS.items():
  for model,p in jobs:
   r=json.loads((STUDY/'records'/f'{model}-p{p}.json').read_text())
   assert r['status']=='passed' and r['samples']==1,(ds,model,p)
   additional.append(r)
 lines=['# V3DB 低召回补充实验完成','',
  '四个数据集共补充16个配置。Recall为20条搜索校准；每配置固定query100/repeat0的1次真实出证，时间与VO不是均值。',
  f'新运行{sum(not r["reused_existing_sample"] for r in additional)}次真实证明；其余复用相同固定查询的既有有效证明。',
  '', '| 数据集 | M | K | probe | Recall@10 (%) | Query (s), single | Verify (ms), single | VO (KB), single |',
  '|---|---:|---:|---:|---:|---:|---:|---:|']
 for r in additional:
  lines.append(f'| {r["dataset"]} | {r["config"]["M"]} | {r["config"]["K"]} | {r["n_probe"]} | {r["recall20"]*100:.1f} | {r["query_ms"]/1000:.3f} | {r["verify_ms"]:.3f} | {r["vo_kb"]:.3f} |')
 lines+=['',f'全部{raw}个可用曲线点的原始证明记录已核对；非曲线区域{counts}个单元格保持不变，其他方案原高召回曲线{oldcurve}个单元格原样映射保留。',
         '新增目标列为15%–55%；原60%–99%列保留。PQ参数离散，主表显示实际召回率；V3DB每目标采用满足目标的已测最小单次查询时间配置，该选择存在单次计时噪声，不是统计最优结论。',
         'VLAH、ANNProof、Ours默认约95%保持不变，V3DB默认点也保持本轮前配置。主表索引成本仍对应默认配置，不对应所有低召回PQ变体。',
         '工作簿：外层work9.xlsx；完整副本work9_lowrecall.xlsx；备份work9_before_lowrecall.xlsx。']
 (STUDY/'RESULTS.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
result=dict(status='PASS',scope='all planned low-recall experiments' if finished else 'completed points so far',proof_points_checked=raw,v3db_curve_metric_cells_checked=metric_cells,non_curve_cells_unchanged=counts,original_baseline_high_recall_curve_cells_preserved=oldcurve,delivered=single.util.sha(SOURCE)==single.util.sha(book),all_jobs_finished=finished)
single.write(STUDY/'independent-validation.json',result);print(json.dumps(result,indent=2))
