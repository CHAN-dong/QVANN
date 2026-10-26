"""Final evidence and workbook checks, after fill_work9.py --update-original."""
import csv
import json
import math
from pathlib import Path
import openpyxl
from fill_work9 import STUDY, ORIGINAL, OUTPUT, TEMPLATE, sha, load_runs


def main():
    audit=json.loads((STUDY/'workbook_audit.json').read_text())
    assert audit['completed_numeric_cells']==audit['total_target_cells']==206
    assert not audit['pending_cells']
    assert sha(ORIGINAL)==sha(OUTPUT)==audit['output_sha256']
    assert sha(TEMPLATE)==(STUDY/'original_sha256.txt').read_text().strip()
    wb=openpyxl.load_workbook(ORIGINAL,data_only=True)
    template=openpyxl.load_workbook(TEMPLATE,data_only=True)
    assert set(wb['Sheet1'].merged_cells.ranges)==set(template['Sheet1'].merged_cells.ranges)
    for c in audit['cells']:
        value=wb['Sheet1'][c['cell']].value
        assert isinstance(value,(int,float)) and math.isfinite(value) and value>=0
        assert math.isclose(value,c['value'],rel_tol=1e-12)
        assert wb['Sheet1'][c['cell']].comment is not None
        if c['metric']=='recall_percent':assert value<=100
    expected={'实验说明','实验配置与实测召回','单元格来源','资源与尾延迟'}
    assert expected.issubset(wb.sheetnames)
    runs=load_runs()
    folders={Path(r['output']) for r in runs}
    tamper_count=0
    legacy_thread_metadata=[]
    for folder in folders:
        record=next(r for r in runs if Path(r['output'])==folder)
        if 'query_threads' in record['manifest']:
            assert record['manifest']['query_threads']==1
        else:
            # The first two jobs predate the explicit build/query thread fields.
            # Preserve their original manifests rather than inventing metadata.
            assert record['name']=='sift1m-recall-curve'
            environment=json.loads((folder/'environment.json').read_text())
            assert '--build-threads' not in environment['command']
            legacy_thread_metadata.append(str(folder))
        assert record['manifest']['query_offset']==record['query_offset']
        with (folder/'per_query.csv').open() as f:
            rows=list(csv.DictReader(f))
        expected_ids=set(range(record['query_offset'],record['query_offset']+record['queries']))
        assert {int(row['query']) for row in rows}==expected_ids
        assert {int(row['repeat']) for row in rows}==set(range(record['repeats']))
        with (folder/'tamper_tests.csv').open() as f:
            attacks=list(csv.DictReader(f))
        assert attacks and all(row['rejected']=='1' for row in attacks)
        tamper_count+=len(attacks)
    result={'status':'PASS','numeric_cells':206,'merged_ranges':len(wb['Sheet1'].merged_cells.ranges),
            'passed_experiment_jobs':len(folders),'measured_configurations':len(runs),
            'verified_query_measurements':sum(r['samples'] for r in runs),
            'rejected_tamper_tests':tamper_count,'workbook_sha256':sha(ORIGINAL),
            'backup_sha256':sha(TEMPLATE),'sheets':wb.sheetnames,
            'legacy_manifests_without_explicit_thread_fields':legacy_thread_metadata}
    (STUDY/'final_validation.json').write_text(json.dumps(result,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False))


if __name__=='__main__':main()
