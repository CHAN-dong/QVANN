"""Evidence-based table repairs; never smooth observed measurements."""
import csv,hashlib,json,math,shutil,statistics
from copy import copy
from pathlib import Path
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font,PatternFill,Alignment
R=Path(__file__).resolve().parents[1];S=R/'experiments/work9-monotonicity-audit-20261021'
SOURCE=R.parent/'work9.xlsx';BACKUP=S/'work9_before_monotonicity_audit.xlsx'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
S.mkdir(parents=True,exist_ok=True)
if BACKUP.exists():
 assert not (S/'audit.json').exists() and sha(SOURCE)==sha(BACKUP),'Already applied or externally changed; inspect the audit before rerunning.'
else:shutil.copy2(SOURCE,BACKUP)
original_sha=sha(SOURCE)
w=openpyxl.load_workbook(BACKUP);old=openpyxl.load_workbook(BACKUP)
names=['SIFT1M','NYTimes-256','GloVe-100','DEEP1M'];changes=[];annotations=set();anomalies=[]
def change(s,c,value,reason,source_comment=''):
 x=s[c];changes.append(dict(sheet=s.title,cell=c,before=x.value,after=value,reason=reason,old_comment=x.comment.text if x.comment else ''))
 x.value=value;x.comment=Comment(reason+'\n'+source_comment+'\n\n修改前来源：\n'+changes[-1]['old_comment'],'Codex')
 x.fill=PatternFill('solid',fgColor='E2F0D9')
def note(s,c,text):
 x=s[c];x.comment=Comment(text+'\n\n原有来源：\n'+(x.comment.text if x.comment else ''),'Codex');annotations.add((s.title,c))
# Replace stale graph floors with the lowest currently measured 40--80% point.
# This is explicit reuse of a measured operating point, not a new 15--35% result.
for sn in ['SIFT1M','NYTimes-256','DEEP1M']:
 s=w[sn]
 assert all(s.cell(23,c).value>s['K23'].value for c in range(6,11))
 for rr in [23,24,37,38,51,52]:
  src=s.cell(rr,11)
  for cc in range(6,11):
   target=s.cell(17,cc).value
   reason=f'修正旧图配置召回下限与新图低召回点混用造成的倒序。目标列{target}%复用已测40%目标列，实际召回{s["K23"].value}%；未实测达到{target}%，不作为该召回率的新测量。性能沿用该点已标明的ANNProof系数调整值，未重复乘系数。'
   change(s,s.cell(rr,cc).coordinate,src.value,reason,src.comment.text if src.comment else '')
# Re-evaluate the same BIGANN prefix and evaluation IDs, now meeting 95%.
p=json.loads((S/'sift-size-1m-corrected.json').read_text())
rr=list(csv.DictReader(Path(p['csv']).open()))
assert sha(Path(p['csv']))==p['csv_sha256'] and len(rr)==60 and all(r['verified']=='1' for r in rr)
assert {(int(r['query']),int(r['repeat'])) for r in rr}=={(q,t) for q in range(100,120) for t in range(3)}
assert p['recall']>=.95
for cell,key,factor in [('H69','vo_kb',2),('H81','verify_ms',1.5),('H92','query_ms',1.2)]:
 reason=f'修正规模表召回口径：旧点实际87.6%，本次同BIGANN前100万条、query100..119、20条×3次、单线程补测，efSearch={p["ef"]}，实际Recall@10={p["recall"]*100:.4f}%。原始{key}={p[key]:.12g}；表内继续采用用户指定系数×{factor}，属于调整后比较值。'
 change(w['SIFT1M'],cell,p[key]*factor,reason,json.dumps(p,ensure_ascii=False))
panels=[]
for base,metric in [(21,'VO size'),(35,'verify time'),(49,'query time')]:
 for i,method in enumerate(['VLAH','ANNProof','Ours','V3DB']):panels.append(('Recall@10',base+2*i+1,range(6,26),method,metric))
for label,starts,cols in [('Dataset size',[(68,'VO size'),(80,'verify time'),(91,'query time')],range(6,11)),('K',[(104,'VO size'),(116,'verify time'),(127,'query time')],range(6,12))]:
 for start,metric in starts:
  for i,method in enumerate(['VLAH','ANNProof','Ours','V3DB','Chen13-MR']):panels.append((label,start+i,cols,method,metric))
for sn in names:
 s=w[sn]
 for panel,row,cols,method,metric in panels:
  cells=[s.cell(row,c) for c in cols if isinstance(s.cell(row,c).value,(int,float))]
  for a,b in zip(cells,cells[1:]):
   assert math.isfinite(a.value) and a.value>=0
   if b.value>=a.value*(1-1e-6):continue
   src=(a.comment.text if a.comment else '')+'\n'+(b.comment.text if b.comment else '')
   if method=='V3DB':
    reason='V3DB包含不同PQ配置、probe、证明电路/域大小及单次测量；这些条件不保证该指标随横轴单调。'
    if '条件预测' in src:reason+=' 此处还涉及512GiB条件预测，须与实测区分；不依据单调性修改预测中心值。'
   elif panel=='Recall@10' and method=='ANNProof':reason='低召回曲线允许切换M/efConstruction；更稠密图可用更小efSearch达到更高召回。不同图配置的VO及耗时不保证随召回率单调。'
   elif panel=='Dataset size':reason='各规模重新建索引、调搜索参数及校准召回；预算与实际召回不完全相同，VO和耗时不保证仅随N单调。详见两点原始参数批注。'
   else:reason='耗时是独立测量的样本均值，存在缓存、分配和系统波动，不能要求均值严格单调；保留原始观察，不以平滑值替代实测。'
   anomalies.append([sn,panel,method,metric,a.coordinate,a.value,b.coordinate,b.value,(b.value/a.value-1)*100,reason])
   note(s,b.coordinate,'单调性复核：'+reason+f' 前点{a.coordinate}={a.value:.12g}，本点={b.value:.12g}。')
 for r in [21,23,25,27,35,37,39,41,49,51,53,55]:
  vals=[s.cell(r,c).value for c in range(6,26) if isinstance(s.cell(r,c).value,(int,float))]
  assert all(b>=a-1e-8 for a,b in zip(vals,vals[1:])),(sn,r)
def sheet(title,headers,rows):
 s=w.create_sheet(title);s.append(headers)
 for r in rows:s.append(r)
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(bold=True)
 for col in s.columns:s.column_dimensions[col[0].column_letter].width=min(70,max(15,max(len(str(c.value or '')) for c in col)+2))
 for row in s:
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
 return s
sheet('单调性复核说明',['项目','处理'],[
 ['复核范围','四个主表的Recall、K、Dataset size曲线；召回率、单位和原始记录核对。'],
 ['已纠正','SIFT/NYTimes/DEEP ANNProof旧高召回下限混用：15–35%列复用现有最低实测点，实际召回如实显示，不声称新增低召回测量。'],
 ['已补测','SIFT规模100万ANNProof原87.6%不符合约95%口径：20条×3次单线程重测并替换VO/查询/验证。原始补测结果见复核目录JSON与CSV。'],
 ['系数口径','ANNProof仍为query×1.2、VO×2、verify×1.5的用户指定调整值；不是原始实测数值。'],
 ['保留的下降','配置、搜索预算、电路结构和测量噪声可导致非单调；逐项批注与记录。未强制抬高成本或制作平滑实测值。'],
 ['曲线绘图','以实际Recall行作为X轴；目标列不是实际召回。重复使用同一测量点须去重，不当作独立测量。'],
 ['默认配置','其他方法约95%默认值保持原值；V3DB继续低召回默认。'],
 ['备份',str(BACKUP)],['历史附表','既有原始附表保留历史记录；本次修订以单调性修订记录和主表批注为准。']])
sheet('单调性异常复核',['数据集','横轴','方法','指标','前单元格','前值','后单元格','后值','变化%','保留原因'],anomalies)
sheet('单调性修订记录',['数据集','单元格','原值','新值','依据'],[[x['sheet'],x['cell'],x['before'],x['after'],x['reason']] for x in changes])
allowed={(x['sheet'],x['cell']) for x in changes};preserved=0
for os in old:
 ns=w[os.title]
 for row in os:
  for c in row:
   if (os.title,c.coordinate) not in allowed:
    assert c.value==ns[c.coordinate].value,(os.title,c.coordinate)
    preserved+=1
assert sha(SOURCE)==original_sha,'Workbook externally changed; refusing overwrite'
tmp=S/'work9_monotonicity_checked.xlsx';w.save(tmp)
check=openpyxl.load_workbook(tmp)
for x in changes:assert math.isclose(check[x['sheet']][x['cell']].value,x['after'],rel_tol=1e-14,abs_tol=1e-12)
shutil.copy2(tmp,SOURCE)
report=dict(passed=True,numeric_changes=len(changes),anomaly_pairs=len(anomalies),annotated_cells=len(annotations),unchanged_cell_values_verified=preserved,backup_sha256=sha(BACKUP),workbook_sha256=sha(SOURCE),new_measurement=p,changes=changes)
(S/'audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k not in ['changes','new_measurement']},ensure_ascii=False))
