"""Complete the requested four-dataset VO/V table and append measured Raw rows."""
import csv
import json
import math
import shutil
import statistics
import struct
from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, Alignment
import work9_resource_matched as u

R=u.ROOT;Q=R/'experiments/work9-quantization-20261023';S=R/'experiments/work9-quantization-raw-20261023';O=R/'results/work9-quantization-raw-20261023'
SOURCE=R.parent/'work9.xlsx';BACKUP=S/'work9_before_raw_quantization.xlsx'
assert not (S/'APPLIED.json').exists()
assert u.read(S/'FINISHED.json')['passed'] and u.sha(SOURCE)==u.sha(BACKUP)
proto=u.read(Q/'protocol.json');raw=u.read(S/'summary.json');records=u.read(Q/'summary.json')+raw
assert len(records)==20 and len(raw)==4
groups=['Raw']+proto['configurations'];datasets=list(proto['datasets'])

def fmt(value):
    x=Decimal(str(value));one=x.quantize(Decimal('.1'),rounding=ROUND_HALF_UP)
    return str(x.quantize(Decimal('1'),rounding=ROUND_HALF_UP)) if abs(one)>=1000 else format(one,'.1f')

checked=0;proofs=[]
for rec in raw:
    ds=rec['dataset'];path=Path(rec['source']);assert u.sha(path)==rec['csv_sha256'];rows=list(csv.DictReader(path.open()))
    reference={(int(r['query']),int(r['repeat'])):r for r in csv.DictReader((R/'results/work9-quantization-20261023'/ds/'samples.csv').open()) if r['configuration']=='SQ4--SQ6--raw'}
    assert len(rows)==len(reference)==60
    for row in rows:
        old=reference[int(row['query']),int(row['repeat'])]
        assert row['trace_sha256']==old['trace_sha256'] and row['answer']==old['answer'] and row['valid']=='1'
        assert math.isclose(float(row['recall']),float(old['recall']),abs_tol=1e-12)
        assert int(row['record_bytes'])+int(row['proof_bytes'])==int(row['vo_bytes']);checked+=1
    for metric in ['verify_ms','query_ms','vo_bytes','recall']:
        assert math.isclose(rec[metric],statistics.fmean(float(r[metric]) for r in rows),rel_tol=1e-12)
    wire=path.parent/'query0.vo';data=wire.read_bytes();pos=0
    def read():
        global pos
        x=struct.unpack_from('<I',data,pos)[0];pos+=4;return x
    def skip(n):
        global pos
        pos+=n;assert pos<=len(data)
    assert read()==0x56505231;skip(32);assert read()==10 and read()==rec['E'] and read()==10;skip(40)
    n=read();skip(n*(4+4*rec['D']))
    opened=read()
    for _ in range(opened):skip(12);skip(read()*4)
    record_end=pos;count=read();skip(count*32);assert pos==len(data)
    first=next(r for r in rows if int(r['query'])==100 and int(r['repeat'])==0)
    assert record_end==int(first['record_bytes']) and len(data)==int(first['vo_bytes']) and n==int(first['disclosed'])
    proofs.append(dict(dataset=ds,file=str(wire),bytes=len(data),payload_bytes=record_end,proof_nodes=count))
    rec['raw_vectors']=statistics.fmean(int(r['disclosed']) for r in rows);rec['stage1_records']=rec['stage2_records']=0
assert checked==240

title=r'\caption{Impact of quantization configurations across four datasets. All configurations use fixed SQ4 top-1 traversal; Raw transmits unquantized vectors.}'
lines=[r'\begin{table}[t]',r'\centering',title,r'\label{tab:quantization}',r'\vspace{-0.2cm}',r'\footnotesize',r'\setlength{\tabcolsep}{1pt}',r'\renewcommand{\arraystretch}{1.15}',r'\begin{tabular}{@{}l*{4}{cc}@{}}',r'\toprule',r'\multirow{2}{*}{Configuration}',r'& \multicolumn{2}{c}{\textsf{SIFT}} & \multicolumn{2}{c}{\textsf{NYTimes}} & \multicolumn{2}{c}{\textsf{GloVe}} & \multicolumn{2}{c}{\textsf{DEEP}} \\',r'\cmidrule(lr){2-3}\cmidrule(lr){4-5}\cmidrule(lr){6-7}\cmidrule(l){8-9}',r'& VO (KB) & V (ms) & VO (KB) & V (ms) & VO (KB) & V (ms) & VO (KB) & V (ms) \\',r'\midrule']
for mode in groups:
    label=r'\textbf{SQ4--SQ6--raw}' if mode=='SQ4--SQ6--raw' else mode;values=[]
    for ds in datasets:
        rec=next(r for r in records if r['dataset']==ds and r['configuration']==mode)
        values.extend([fmt(rec['vo_kb']),fmt(rec['verify_ms'])])
    lines.append(label+' & '+' & '.join(values)+r' \\')
    if mode=='Raw':lines.append(r'\hline')
lines += [r'\bottomrule',r'\end{tabular}',r'\end{table}','']
latex='\n'.join(lines);assert '-- &' not in latex
target=Q/'quantization_four_datasets.tex';target.write_text(latex);(S/'quantization_four_datasets.tex').write_text(latex)
with (S/'summary-with-raw.csv').open('w',newline='') as f:
    keys=['dataset','configuration','N','D','E','samples','recall','vo_kb','verify_ms','query_ms','source','csv_sha256']
    writer=csv.DictWriter(f,fieldnames=keys);writer.writeheader();writer.writerows({k:r[k] for k in keys} for r in records)

wb=openpyxl.load_workbook(BACKUP);mapping=u.read(Q/'audit.json')['new_tables'];allowed=set();locations={}
for rec in raw:
    ws=wb[rec['sheet']];loc=mapping[rec['dataset']];row=loc['last_data_row']+1;locations[rec['sheet']]=row
    assert all(ws.cell(row,col).value is None for col in range(4,12)),(ws.title,row,'nonempty Raw insertion point')
    vals=['Raw',rec['recall'],rec['query_ms'],rec['vo_kb'],rec['verify_ms'],None,None,60]
    for col,value in enumerate(vals,4):
        cell=ws.cell(row,col,value);cell._style=copy(ws.cell(row-1,col)._style);cell.font=Font(bold=False);cell.alignment=Alignment(wrap_text=True,horizontal='left' if col==4 else 'center')
        if col==5:cell.number_format='0.0%'
        cell.comment=Comment('Raw披露对照组（实测）：固定原SQ4 top-1图遍历、图索引和预算E，仅把VO披露改为所有发现向量的raw float32坐标。不是改用raw距离遍历的HNSW对照。\n'
            '扩展中心认证叶子绑定邻接ID与raw向量；客户端从认证raw向量重建原SQ4码，认证并重放相同遍历、验证候选top-k。\n'
            f"N={rec['N']}, D={rec['D']}, M={rec['M']}, efc={rec['ef_construction']}, E={rec['E']}, k=10；query100..119×3，单线程。\n"
            f"VO={rec['vo_kb']:.9f}KB；V={rec['verify_ms']:.9f}ms；Q={rec['query_ms']:.9f}ms。Q为整体计时，Search/VO gen未单独分拆，故对应两格留空。\n"
            f"逐条与旧默认组的trace和答案一致。\n来源={rec['source']}\nSHA256={rec['csv_sha256']}",'Codex')
        allowed.add((ws.title,cell.coordinate))
    for rr in range(loc['first_data_row'],row+1):
        for col in range(6,11):
            cell=ws.cell(rr,col)
            if isinstance(cell.value,(int,float)):
                cell.number_format='0' if abs(Decimal(fmt(cell.value)))>=1000 else '0.0';allowed.add((ws.title,cell.coordinate))
    ws.row_dimensions[row].height=30
details=wb['量化配置实测记录'];keys=[c.value for c in details[1]]
for rec in raw:details.append([rec.get(k) for k in keys])
notes=wb['量化配置实验说明'];notes.append(['Raw补测','新增Raw披露对照：固定SQ4 top-1 traversal，所有发现向量以raw float32披露；认证仍用expansion-centric策略，直接内容叶子改为绑定raw向量。每数据集20查询×3，单线程，240次完整出证验证，48项负面测试拒绝。与旧四配置逐条trace和答案一致。'])
notes.append(['最新四数据集论文表',str(target)])
review=S/'work9_quantization_with_raw.xlsx';wb.save(review)
old=openpyxl.load_workbook(BACKUP);new=openpyxl.load_workbook(review);preserved=0
for ws in old:
    for row in ws:
        for cell in row:
            other=new[ws.title][cell.coordinate]
            if (ws.title,cell.coordinate) not in allowed:
                assert cell.value==other.value and (cell.comment.text if cell.comment else None)==(other.comment.text if other.comment else None),(ws.title,cell.coordinate)
                assert cell.number_format==other.number_format
            else:
                if cell.value is not None:assert cell.value==other.value,(ws.title,cell.coordinate,'only formatting may change on existing data')
            preserved+=1
result=dict(passed=True,measured_raw_samples=240,matched_old_traces_and_answers=240,rejected_tamper_cases=48,independent_wire_parses=proofs,original_cells_checked=preserved,raw_workbook_rows=locations,latex=str(target),workbook_sha256=u.sha(review),backup_sha256=u.sha(BACKUP))
u.write(S/'summary-with-raw.json',records);u.write(S/'audit.json',result)
assert u.sha(SOURCE)==u.sha(BACKUP)
try:shutil.copy2(review,SOURCE)
except PermissionError:print('AUDITED_REVIEW_READY_SOURCE_LOCKED',flush=True)
else:
    assert u.sha(SOURCE)==result['workbook_sha256'];u.write(S/'APPLIED.json',result);print('APPLIED',flush=True)
print(json.dumps(result,ensure_ascii=False,indent=2))
