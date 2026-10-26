"""Independent Python SHA/Merkle, trace and refinement verification of native VOs."""
import ast,csv,hashlib,heapq,json,math,struct,sys
from pathlib import Path
import numpy as np
import work9_resource_matched as u
R=u.ROOT;S=R/'experiments/authentication-granularity-20261025';O=R/'results/authentication-granularity-20261025'
# Reuse the existing independent pure-Python multiproof and float32 replay code.
helper=R/'scripts/audit_routing_precision_20261024.py'
tree=ast.parse(helper.read_text());exec(compile(ast.Module(body=[x for x in tree.body if isinstance(x,(ast.FunctionDef,ast.ClassDef))],type_ignores=[]),str(helper),'exec'),globals())

def verify(path,a,q,base,lo,span,E,mode):
    r=Reader(path.read_bytes());assert r.u32()==mode and r.u32()==0x56505231 and r.blob(32)==h(q.astype('<f4').tobytes());assert r.u32()==10 and r.u32()==E and r.u32()==10;answer=[r.u32() for _ in range(10)]
    n,d=a['n'],a['dim'];start=r.pos;codes={}
    for _ in range(r.u32()):
        id=r.u32();assert 0<=id<n and id not in codes;codes[id]=r.blob(d//2)
    ids=sorted(codes);c4=np.stack([unpack4(codes[id],d) for id in ids]);ci={id:i for i,id in enumerate(ids)}
    assert np.array_equal(c4[ci[a['entry']]],quant4(np.asarray(base[a['entry']]),lo,span))
    links={};er={};nodes={};edges={};code_leaves={}
    if mode==2:
        for _ in range(r.u32()):
            id,l,pos=r.u32(),r.u32(),r.u32();assert id<n and l<=a['maxlevel'] and pos<a['route_count'];ns=[r.u32() for _ in range(r.u32())];assert ns==sorted(set(ns)) and (id,l) not in links and pos not in er
            links[id,l]=ns;b=b'record/expansion-direct/v1'+bytes.fromhex(a['snapshot'])+U(id)+U(l)+U(len(ns))
            for nb in ns:b+=U(nb)+codes[nb]
            er[pos]=h(b)
    else:
        opened=[(r.u32(),r.u32()) for _ in range(r.u32())];assert opened==sorted(set(opened));needed={id for id,l in opened};records={};bases={}
        for _ in range(r.u32()):
            id=r.u32();assert id in needed and id not in records;edgebase=r.u32() if mode==0 else 0;ls=[];b=(b'granularity/node-meta/v1' if mode==0 else b'granularity/node-adjacency/v1')+bytes.fromhex(a['snapshot'])+U(id)
            if mode==0:b+=U(edgebase)
            nl=r.u32();assert 0<nl<=a['maxlevel']+1;b+=U(nl)
            for l in range(nl):
                degree=r.u32();assert degree<=n;b+=U(degree)
                if mode==0:ls.append(degree)
                else:
                    ns=[r.u32() for _ in range(degree)];assert ns==sorted(set(ns)) and all(x<n for x in ns);ls.append(ns);b+=b''.join(U(x) for x in ns)
            records[id]=ls;bases[id]=edgebase;nodes[id]=h(b)
        assert set(records)==needed
        for id,l in opened:
            ls=records[id];assert id<n and l<len(ls)
            if mode==0:
                ns=[r.u32() for _ in range(ls[l])];assert ns==sorted(set(ns)) and all(nb<n for nb in ns);basepos=bases[id]+sum(ls[:l])
                for j,nb in enumerate(ns):
                    pos=basepos+j;assert pos<a['edge_count'] and pos not in edges;edges[pos]=h(b'granularity/directed-edge/v1'+bytes.fromhex(a['snapshot'])+U(id)+U(l)+U(j)+U(nb))
            else:ns=ls[l]
            links[id,l]=ns
        code_leaves={id:h(b'granularity/sq4/v1'+bytes.fromhex(a['snapshot'])+U(id)+b) for id,b in codes.items()}
    routing_end=r.pos;delta={};dl={};raw={};rl={}
    for _ in range(r.u32()):
        id,pos=r.u32(),r.u32();assert id in codes and id not in delta and pos<n and pos not in dl;b=r.blob(d//4);delta[id]=b;dl[pos]=h(b'record/delta46/v1'+U(id)+b)
    for _ in range(r.u32()):
        id,pos=r.u32(),r.u32();assert id in delta and id not in raw and pos<n and pos not in rl;b=r.blob(d*4);raw[id]=np.frombuffer(b,dtype='<f4').copy();assert np.isfinite(raw[id]).all();rl[pos]=h(b'record/raw/v1'+U(id)+b)
    proof_start=r.pos;siblings=0;targets=0
    def proof(root,count,leaves):
        nonlocal siblings,targets
        siblings+=struct.unpack_from('<I',r.b,r.pos)[0];targets+=len(leaves);r.proof(root,count,leaves)
    if mode==2:proof(a['route'],a['route_count'],er)
    else:
        proof(a['node'],n,nodes)
        if mode==0:proof(a['edges'],a['edge_count'],edges)
        proof(a['codes'],n,code_leaves)
    route_proof=r.pos-proof_start;r.proof(a['delta'],n,dl);r.proof(a['raw'],n,rl);assert r.pos==len(r.b)
    score=scores4(c4,q,lo,span);seen,opened,order,expanded=replay({id:float(x) for id,x in zip(ids,score)},links,a['entry'],a['maxlevel'],E);assert seen==set(codes) and opened==set(links)
    raw_ids=sorted(raw);dist=distances(np.stack([raw[id] for id in raw_ids]),q);best=sorted(zip(map(float,dist),raw_ids));assert [id for _,id in best[:10]]==answer and len(set(answer))==10;tau=np.float32(best[9][0])
    l4=lower(c4,q,lo,span,4);assert set(delta)=={id for id,x in zip(ids,l4) if x<=tau};refined=sorted(delta);c6=[]
    for id in refined:c6.append((c4[ci[id]]<<2)|np.array([(delta[id][j//4]>>((j%4)*2))&3 for j in range(d)],np.uint8))
    l6=lower(np.stack(c6),q,lo,span,6);assert set(raw)=={id for id,x in zip(refined,l6) if x<=tau or id in answer}
    return dict(valid=True,candidates=len(ids),expanded=len(expanded),targets=targets,siblings=siblings,route_proof_bytes=route_proof,route_payload_bytes=routing_end-start,total_bytes=len(r.b),answer=answer)

def main():
    checks=[]
    for ds in ['sift1m','deep1m']:
        d=128 if ds=='sift1m' else 96;base=np.memmap(R/'dataset'/ds/'sift_base.fvecs',dtype='<f4',mode='r').reshape(-1,d+1)[:1000000,1:];queries=np.memmap(R/'dataset'/ds/'sift_query.fvecs',dtype='<f4',mode='r').reshape(-1,d+1)[:,1:];lo=base.min(axis=0);span=base.max(axis=0)-lo
        for M in [8,16,32,64]:
            out=O/f'{ds}-M{M}'
            if not (out/'PASS.txt').exists():continue
            anchors=list(csv.DictReader((out/'anchors.csv').open()));rows=list(csv.DictReader((out/'samples.csv').open()));E=u.read(out/'selection.json')['E']
            for mode,name in enumerate(['Record-wise','Node-wise','ER']):
                a=next(x for x in anchors if x['variant']==name);a={k:int(v) if k in ['n','dim','entry','maxlevel','route_count','edge_count'] else v for k,v in a.items()};wire=out/(name+'.vo');result=verify(wire,a,np.asarray(queries[100]),base,lo,span,E,mode);row=next(x for x in rows if x['variant']==name and x['query']=='100' and x['repeat']=='0')
                for key,col in [('candidates','candidates'),('targets','targets'),('siblings','sibling_hashes'),('route_proof_bytes','route_proof_bytes'),('route_payload_bytes','route_payload_bytes'),('total_bytes','total_bytes')]:assert result[key]==int(row[col]),(ds,M,name,key)
                assert result['answer']==list(map(int,row['answer'].split()));checks.append(dict(dataset=ds,M=M,variant=name,source_sha256=u.sha(wire),**result))
    assert len(checks)==24,'Some planned configurations lack complete native samples'
    import openpyxl
    workbook=R.parent/'authentication_granularity.xlsx';assert u.sha(workbook)==u.sha(S/'authentication_granularity.xlsx');wb=openpyxl.load_workbook(workbook,data_only=True)
    summary=list(csv.DictReader((S/'summary.csv').open()));assert len(summary)==24 and all(r['Status']=='MEASURED' for r in summary)
    plot_cells=0
    for dataset in ['SIFT','DEEP']:
        ws=wb[dataset+'_Plot'];assert len(ws._charts)==3
        for row,M in enumerate([8,16,32,64],2):
            assert ws.cell(row,1).value==M and ws.cell(row,1).data_type=='n'
            for col in range(2,ws.max_column+1):
                metric,variant=ws.cell(1,col).value.split(' | ');record=next(r for r in summary if r['Dataset']==dataset and int(r['M'])==M and r['Variant']==variant);cell=ws.cell(row,col)
                assert cell.data_type=='n' and math.isclose(cell.value,float(record[metric]),rel_tol=1e-12);plot_cells+=1
        for chart in ws._charts:assert [series.tx.v for series in chart.series]==['Record-wise','Node-wise','ER']
    for path,digest in u.read(S/'original-files.json').items():assert u.sha(path)==digest
    u.write(S/'independent-audit.json',dict(passed=True,wire_checks=checks,workbook_sha256=u.sha(workbook),plot_numeric_cells_checked=plot_cells,charts_checked=6,original_files_unchanged=True,reused_python_auditor_sha256=u.sha(helper),auditor_sha256=u.sha(Path(__file__)),scope='All24 sample packets independently rehashed and multiproof-verified; float32 SQ4 replay and progressive final-result certificate checked without any native verifier or server graph access. All plot cells checked against original CSV aggregates; charts and workbook validated.'))
    print('INDEPENDENT_GRANULARITY_AUDIT_PASSED',len(checks),flush=True)
def preflight():
    out=O/'preflight';d=128;base=np.memmap(R/'dataset/sift1m/sift_base.fvecs',dtype='<f4',mode='r').reshape(-1,d+1)[:5000,1:];queries=np.memmap(R/'dataset/sift1m/sift_query.fvecs',dtype='<f4',mode='r').reshape(-1,d+1)[:,1:];lo=base.min(axis=0);span=base.max(axis=0)-lo;E=u.read(out/'selection.json')['E'];anchors=list(csv.DictReader((out/'anchors.csv').open()))
    for mode,name in enumerate(['Record-wise','Node-wise','ER']):
        a=next(x for x in anchors if x['variant']==name);a={k:int(v) if k in ['n','dim','entry','maxlevel','route_count','edge_count'] else v for k,v in a.items()}
        result=verify(out/(name+'.vo'),a,np.asarray(queries[100]),base,lo,span,E,mode);assert result['valid'];print('PREFLIGHT_INDEPENDENT_VERIFIED',name,flush=True)
if __name__=='__main__':
    if '--preflight' in sys.argv:preflight()
    else:main()
