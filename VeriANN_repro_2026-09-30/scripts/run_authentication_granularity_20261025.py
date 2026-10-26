"""Measured-only authentication granularity study. Run with the project's WSL Python."""
import argparse,csv,json,math,platform,shutil,statistics,subprocess,sys,zipfile
from pathlib import Path
import work9_resource_matched as u
R=u.ROOT;S=R/'experiments/authentication-granularity-20261025';O=R/'results/authentication-granularity-20261025'
SOURCE=R/'src/experiments/authentication_granularity.cpp';B=S/'authentication_granularity'
VARIANTS=['Record-wise','Node-wise','ER'];MS=[8,16,32,64]
REFERENCE=u.read(R/'experiments/work9-quantization-20261023/protocol.json')['datasets']
S.mkdir(parents=True,exist_ok=True);O.mkdir(parents=True,exist_ok=True);u.STUDY,u.RESULTS=S,O

def writecsv(path,rows):
    if not rows:return
    with Path(path).open('w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)

def replace_once(text,old,new):
    assert text.count(old)==1,(old[:100],text.count(old));return text.replace(old,new)

def build():
    original=R/'build/veriann_lut_benchmark.cpp';prefix=original.read_text().split('int main(int argc,char**argv)try{')[0]
    generator='Packet generate('+prefix.split('Packet generate(',1)[1].split('// The verifier receives',1)[0]
    generator=replace_once(generator,'Packet generate(','Packet generate_granularity(')
    generator=replace_once(generator,'const float*q,int k,int E){','const float*q,int k,int E,const Store&store,Id mode){')
    generator=replace_once(generator,'p::write_header(b,q,codec.dim,k,E,answer);','p::u32(b,mode);p::write_header(b,q,codec.dim,k,E,answer);out.routing_start=b.size();')
    start=generator.index('    p::u32(b,t.opened.size());');end=generator.index('    std::vector<Id>refined,raw;',start)
    generator=generator[:start]+'    std::set<Id>route_targets,delta_targets,raw_targets;write_routing(out,store,t,mode,route_targets);\n'+generator[end:]
    generator=replace_once(generator,'proof_write(b,*ads.route,route_targets);','write_route_proofs(out,store,mode,route_targets);')
    verifier='bool verify('+prefix.split('bool verify(',1)[1].split('\nvoid attacks(',1)[0]
    verifier=replace_once(verifier,'bool verify(','bool verify_granularity(')
    verifier=replace_once(verifier,'const Anchor&a','const Client&a')
    verifier=replace_once(verifier,'double*auth_ms=nullptr,double*replay_ms=nullptr)','double*auth_ms=nullptr,double*replay_ms=nullptr,Walk*replayed=nullptr)')
    verifier=replace_once(verifier,'p::Reader r{wire};','p::Reader r{wire};p::require(r.u32()==a.mode,"routing mode mismatch");')
    start=verifier.index('    count=r.count(16);');end=verifier.index('    count=r.count(8+a.dim/4);',start)
    verifier=verifier[:start]+'    auto routing=read_routing(r,a,codes,links,route_leaves);\n'+verifier[end:]
    verifier=replace_once(verifier,'check_proof(r,a.route,a.route_count,route_leaves)','check_routing(r,a,routing,route_leaves)')
    verifier=replace_once(verifier,'    p::require(trace.disclosed.size()==codes.size()','    if(replayed)*replayed=trace;\n    p::require(trace.disclosed.size()==codes.size()')
    generated={'granularity_original_include.cpp':prefix,'granularity_generator.cpp':generator,'granularity_verifier.cpp':verifier}
    if B.exists():
        old=u.read(S/'build.json');assert old['source_sha256']==u.sha(SOURCE) and old['original_sha256']==u.sha(original),'Source changed: keep old run artifacts and use a new run directory'
        for name,text in generated.items():assert (S/name).read_text()==text
        return
    for name,text in generated.items():(S/name).write_text(text)
    cmd=u.read(R/'experiments/work9-quantization-20261023/build.json')['command']
    cmd=[str(SOURCE) if str(x).endswith('/scripts/quantization_refinement_probe.cpp') else '-I'+str(S) if str(x).startswith('-I') and str(x).endswith('/experiments/work9-quantization-20261023') else x for x in cmd]
    cmd[cmd.index('-o')+1]=str(B);u.command(cmd,S/'build.log',timeout=300)
    u.write(S/'build.json',dict(command=cmd,source_sha256=u.sha(SOURCE),original_sha256=u.sha(original),binary_sha256=u.sha(B),generated={name:u.sha(S/name) for name in generated}))

def configuration():
    protocol=dict(datasets=['SIFT1M','DEEP1M'],N=1000000,M=MS,efConstruction=200,k=10,query_ids=list(range(100,120)),repeats=3,threads=1,offline_graph_build_threads=8,variants=VARIANTS,
        refinement='Unmodified native SQ4--SQ6--raw; common BFS delta/raw Merkle trees',
        Record_wise='Separate ID-major node metadata tree (all layer degrees and authenticated edge base), directed edge tree (source/layer/slot/neighbor), and SQ4 code tree (node ID/code). Exact edge counts enforce complete adjacency disclosure for opened layers.',
        Node_wise='One complete multi-layer adjacency list per expanded node in an ID-major node tree; all discovered SQ4 codes authenticated separately in the same SQ4 tree used by Record-wise. Unopened layers are authenticated but not scored.',
        ER='Original QVANN Expansion Records per node/layer with neighbor IDs+SQ4 codes jointly committed; deduplicate codes in wire. Remove the common4-byte variant tag to recover byte-identical native QVANN VO.',
        ordering='ID-major routing leaves; edge leaves source/layer/slot order. No learned/BFS routing ordering confound. Refinement ordering remains native BFS.',
        calibration='Search-only coarse E grid4..2048 until >=95%; evaluate every integer between last below and first >=95% sampled budgets; choose closest actual recall to95%, tie prefer >=95%, then lower E. All variants share selected E. Calibration and measurement use the same20 queries; not a held-out accuracy claim.',
        timing='steady_clock; query includes SQ4 LUT setup, traversal, exact candidate ranking, serialization and all multiproofs. Verification includes parsing, authentication, SQ4 traversal replay and final result certification. Single-thread queries; one untimed warmup per query/variant; rotate variant order by query/repeat. Transport and graph construction excluded.',
        units='KB=1000 bytes; ms. Proof size includes serialized4-byte proof counts and32-byte sibling hashes; routing payload includes code/adjacency/node records and necessary metadata, excludes common query/result header and refinement. Total VO includes everything, including common4-byte mode tag.',
        integrity='Every normal VO uses wire-only client verifier (no server graph/data), client/server trace equality, identical refinement bytes and answers/candidates across variants; negative tamper tests; ER/native byte equality and native verifier acceptance.',
        results_policy='Raw actual measurements only, no display multipliers, predictions or smoothing. Existing experiment files/work9.xlsx are read-only. Failed configurations have blank numeric cells and explicit failure status.')
    u.write(S/'protocol.json',protocol)
    env=dict(platform=platform.platform(),python=sys.version,cpu=subprocess.check_output(['lscpu'],text=True),memory=subprocess.check_output(['free','-b'],text=True),compiler=subprocess.check_output(['g++','--version'],text=True),environment=u.environment())
    # Record only performance-related environment fields, not unrelated credentials.
    env['environment']={k:v for k,v in env['environment'].items() if k in ['OMP_NUM_THREADS','OPENBLAS_NUM_THREADS','MKL_NUM_THREADS','RAYON_NUM_THREADS','LD_LIBRARY_PATH']};u.write(S/'environment.json',env)
    guard=S/'original-files.json'
    if not guard.exists():u.write(guard,{str(p):u.sha(p) for p in [R.parent/'work9.xlsx',R/'build/veriann_lut_benchmark.cpp',R/'src/paper_repro_common.hpp',R/'src/veriann_query_lut.hpp']})

def invoke(ds,M,preflight=False):
    cfg=REFERENCE[ds];digest=cfg['manifest']['dataset_sha256'];N=5000 if preflight else 1000000;tag='preflight' if preflight else f'{ds}-M{M}';out=O/tag;out.mkdir(exist_ok=True)
    if (out/'PASS.txt').exists():
        assert u.read(out/'invocation.json')['binary_sha256']==u.sha(B);return
    existing=R/'ANNProof-24FGCS/cache'/f'{digest[:16]}_M{M}_efc200.index'
    cache=existing.parent if existing.exists() and not preflight else O/'graph-cache'/ds
    cache.mkdir(parents=True,exist_ok=True);prior_hash=u.sha(existing) if existing.exists() and not preflight else None
    cmd=[B,'--dataset',R/'dataset'/ds,'--base',N,'--M',M,'--ef-construction',200,'--cache',cache,'--output',out,'--queries',2 if preflight else 20,'--query-offset',100,'--repeats',1 if preflight else 3,'--build-threads',8,'--k',10]
    u.write(out/'invocation.json',dict(command=list(map(str,cmd)),binary_sha256=u.sha(B),historical_graph_sha256_before=prior_hash))
    try:
        u.command(['/usr/bin/time','-v','-o',out/'resources.txt',*cmd],S/(tag+'.log'),bounded=True,timeout=14400)
        assert (out/'PASS.txt').exists();selection=u.read(out/'selection.json');graph=Path(selection['graph']);u.write(out/'graph.json',dict(path=str(graph),sha256=u.sha(graph),historical=bool(prior_hash)))
        if prior_hash:assert u.sha(existing)==prior_hash,'Historical graph changed'
    except Exception as exc:
        u.write(out/'FAILED.json',dict(status='failed',reason=str(exc),command=list(map(str,cmd))));raise

def report():
    import openpyxl
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.chart import LineChart,Reference
    from openpyxl.comments import Comment
    summary=[];samples=[]
    for ds,label in [('sift1m','SIFT'),('deep1m','DEEP')]:
        for M in MS:
            out=O/f'{ds}-M{M}'
            if not (out/'PASS.txt').exists():
                failure=u.read(out/'FAILED.json') if (out/'FAILED.json').exists() else dict(reason='Not completed')
                for variant in VARIANTS:summary.append(dict(Dataset=label,M=M,Variant=variant,E=None,Recall_pct=None,Candidates=None,Routing_Proof_KB=None,Routing_Payload_KB=None,Total_VO_KB=None,Verification_ms=None,Query_ms=None,Routing_Targets=None,Sibling_Hashes=None,Samples=0,Status='FAILED' if (out/'FAILED.json').exists() else 'PENDING',Source='',Reason=failure['reason']))
                continue
            rows=list(csv.DictReader((out/'samples.csv').open()));assert len(rows)==180 and all(r['valid']=='1' for r in rows)
            selection=u.read(out/'selection.json');assert selection['dataset_sha256']==REFERENCE[ds]['manifest']['dataset_sha256']
            anchors=list(csv.DictReader((out/'anchors.csv').open()));assert len(anchors)==3 and all(int(a['n'])==1000000 and int(a['dim'])==(128 if ds=='sift1m' else 96) for a in anchors)
            tests=list(csv.DictReader((out/'tests.csv').open()));assert len(tests)==45 and all(r['rejected']=='1' for r in tests)
            for q in range(100,120):
                for t in range(3):
                    rr=[r for r in rows if int(r['query'])==q and int(r['repeat'])==t];assert len(rr)==3
                    for key in ['E','recall','candidates','expanded','opened','refined','raw','trace_sha256','refinement_sha256','answer']:assert len({r[key] for r in rr})==1,(ds,M,q,t,key)
            for row in rows:
                proofs={'Record-wise':3,'Node-wise':2,'ER':1}[row['variant']];assert int(row['route_proof_bytes'])==int(row['sibling_hashes'])*32+proofs*4
                samples.append(dict(Dataset=label,M=M,**row))
            for variant in VARIANTS:
                rr=[r for r in rows if r['variant']==variant];assert len(rr)==60;mean=lambda col:statistics.fmean(float(r[col]) for r in rr)
                record=dict(Dataset=label,M=M,Variant=variant,E=selection['E'],Recall_pct=mean('recall')*100,Candidates=mean('candidates'),Routing_Proof_KB=mean('route_proof_bytes')/1000,Routing_Payload_KB=mean('route_payload_bytes')/1000,Total_VO_KB=mean('total_bytes')/1000,Verification_ms=mean('verify_ms'),Query_ms=mean('query_ms'),Routing_Targets=mean('targets'),Sibling_Hashes=mean('sibling_hashes'),Samples=60,Status='MEASURED',Source=str(out/'samples.csv'),Reason='')
                assert math.isclose(record['Recall_pct']/100,selection['recall'],abs_tol=1e-12);summary.append(record)
    writecsv(S/'summary.csv',summary);writecsv(S/'samples_all.csv',samples);u.write(S/'summary.json',summary)
    wb=openpyxl.Workbook();raw=wb.active;raw.title='Raw_Data';columns=list(summary[0]);raw.append(columns)
    for r in summary:raw.append([r[k] for k in columns])
    raw.freeze_panes='D2';raw.auto_filter.ref=raw.dimensions
    metrics=['Routing_Proof_KB','Total_VO_KB','Verification_ms','Routing_Payload_KB','Query_ms','Recall_pct','Candidates','Routing_Targets','Sibling_Hashes','E']
    for label in ['SIFT','DEEP']:
        ws=wb.create_sheet(label+'_Plot');headers=['M']+[f'{metric} | {variant}' for metric in metrics for variant in VARIANTS];ws.append(headers)
        for M in MS:
            rr={r['Variant']:r for r in summary if r['Dataset']==label and r['M']==M};ws.append([M]+[rr[v][metric] for metric in metrics for v in VARIANTS])
        ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
        for i,(title,unit) in enumerate([('Routing Merkle proof size','KB'),('Total VO size','KB'),('Verification time','ms')]):
            chart=LineChart();chart.title=label+': '+title;chart.x_axis.title='Graph degree M';chart.y_axis.title=unit;chart.height=11;chart.width=20
            chart.add_data(Reference(ws,min_col=2+i*3,max_col=4+i*3,min_row=1,max_row=5),titles_from_data=True);chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=5))
            for series,variant in zip(chart.series,VARIANTS):series.tx.strRef=None;series.tx.v=variant;series.marker.symbol='circle';series.marker.size=5
            ws.add_chart(chart,'A'+str(8+i*24))
    config=wb.create_sheet('Config');config.append(['Key','Value'])
    for name in ['protocol.json','environment.json','build.json','original-files.json']:
        for k,v in u.read(S/name).items():config.append([name+': '+k,json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v])
    config.append(['Run from inner project','wsl -e /root/.venvs/v3db-work9/bin/python scripts/run_authentication_granularity_20261025.py all'])
    config.append(['Re-export without rerunning','wsl -e /root/.venvs/v3db-work9/bin/python scripts/run_authentication_granularity_20261025.py report'])
    for ds in ['sift1m','deep1m']:
        for M in MS:
            out=O/f'{ds}-M{M}'
            for name in ['invocation.json','selection.json','graph.json','FAILED.json']:
                if (out/name).exists():config.append([f'{ds} M={M} {name}',json.dumps(u.read(out/name),ensure_ascii=False)])
    for ws in wb:
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24476B');cell.alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[1].height=46
        for col in ws.columns:
            letter=col[0].column_letter;ws.column_dimensions[letter].width=22 if ws.title!='Config' else 40 if letter=='A' else 110
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,(int,float)):cell.number_format='0.00' if isinstance(cell.value,float) else '0'
    output=S/'authentication_granularity.xlsx';tmp=output.with_suffix('.tmp.xlsx');wb.save(tmp);tmp.replace(output)
    with zipfile.ZipFile(output) as z:assert z.testzip() is None
    check=openpyxl.load_workbook(output,data_only=True);assert check.sheetnames==['Raw_Data','SIFT_Plot','DEEP_Plot','Config'];assert check['Raw_Data'].max_row==25
    numeric=['M','E','Recall_pct','Candidates','Routing_Proof_KB','Routing_Payload_KB','Total_VO_KB','Verification_ms','Query_ms','Routing_Targets','Sibling_Hashes','Samples']
    for row,record in enumerate(summary,2):
        for key in numeric:
            cell=check['Raw_Data'].cell(row,columns.index(key)+1)
            if record['Status']=='MEASURED':assert cell.data_type=='n' and math.isclose(cell.value,record[key],rel_tol=1e-12)
            else:assert cell.value is None or key in ['M','Samples']
    for path,digest in u.read(S/'original-files.json').items():assert u.sha(path)==digest,('Original file changed',path)
    destination=R.parent/'authentication_granularity.xlsx';shutil.copy2(output,destination)
    audit=dict(passed=True,measured_configurations=sum(r['Status']=='MEASURED' for r in summary),planned_configurations=24,native_samples=len(samples),workbook=str(destination),workbook_sha256=u.sha(destination),numeric_cells_checked=True,variant_trace_and_refinement_identical=True,negative_tests_rejected=True,ER_native_protocol_identical=True,original_files_unchanged=True)
    u.write(S/'audit.json',audit);print('EXPORTED',audit,flush=True)
    lines=['# Impact of Authentication Granularity — 实测结果','','单位：KB=1000 bytes，时间 ms。全部为原始实测；20 条查询×3 次，每个 M 使用相同 E、图、查询、SQ4--SQ6--raw refinement。完整认证定义和复现命令见 README.md。','','| Dataset | M | Variant | E | Recall (%) | Routing proof (KB) | Routing payload (KB) | Total VO (KB) | V (ms) | Q (ms) |','|---|---:|---|---:|---:|---:|---:|---:|---:|---:|']
    for r in summary:
        keys=['E','Recall_pct','Routing_Proof_KB','Routing_Payload_KB','Total_VO_KB','Verification_ms','Query_ms'];values=[f"{r[k]:.2f}" if r[k] is not None else '—' for k in keys]
        lines.append('| '+' | '.join([r['Dataset'],str(r['M']),r['Variant']]+values)+' |')
    lines+=['','## ER 相对其他认证组织的降幅','','负数表示退化，保留原值，不做单调修正。各个 M 重新校准 E，因此不能仅依据 M 推断指标必须单调。','']
    for label in ['SIFT','DEEP']:
        for M in MS:
            group={r['Variant']:r for r in summary if r['Dataset']==label and r['M']==M}
            if any(r['Status']!='MEASURED' for r in group.values()):continue
            er=group['ER']
            for name in ['Record-wise','Node-wise']:
                baseline=group[name];reduction=[100*(1-er[k]/baseline[k]) for k in ['Routing_Proof_KB','Total_VO_KB','Verification_ms','Query_ms']]
                lines.append(f"- {label}, M={M}, ER 相对 {name}：routing proof / Total VO / V / Q 降幅分别为 "+' / '.join(f'{v:.2f}%' for v in reduction)+'。')
    lines+=['','Record-wise 明确采用每条有向边作为独立 record；Node-wise 为整个节点的多层邻接表，而不是每层一个邻接表。结论针对 README 中定义的具体实现与 Merkle 组织，不能外推到所有可能的 record 划分或叶子布局。客户端验证包含认证、SQ4 遍历复现、渐进边界检查和精确 top-k 认证；查询时间还受到各组织序列化实现的影响。','','校准和测量共享 20 条查询，重复计时不增加独立召回率样本量。不同 M 的实际 Recall 可能略有差异；工作簿记录实际数值而不是统一填95%。原有实验与 work9.xlsx 保持只读。','']
    (S/'REPORT.md').write_text('\n'.join(lines),encoding='utf-8')
    u.write(S/'source-fingerprints.json',{str(p.relative_to(R)):u.sha(p) for p in [SOURCE,Path(__file__).resolve(),R/'scripts/audit_authentication_granularity_20261025.py',R/'build/veriann_lut_benchmark.cpp',R/'src/paper_repro_common.hpp',R/'src/veriann_query_lut.hpp']})

def main():
    parser=argparse.ArgumentParser();parser.add_argument('stage',choices=['build','preflight','run','report','all'],default='all',nargs='?');a=parser.parse_args()
    if a.stage in ['all','build','preflight','run']:build();configuration()
    if a.stage in ['all','preflight']:invoke('sift1m',8,True)
    if a.stage in ['all','run']:
        for ds in ['sift1m','deep1m']:
            for M in MS:
                try:invoke(ds,M)
                except Exception as exc:print('CONFIGURATION_FAILED',ds,M,str(exc),flush=True)
        report()
    elif a.stage=='report':report()
if __name__=='__main__':main()
