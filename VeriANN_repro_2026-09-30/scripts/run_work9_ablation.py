"""Real full-wire 1M ablations; resumable runs and precise workbook updates. WSL."""
import csv
from datetime import datetime
import hashlib
import json
import math
import os
from pathlib import Path
import shutil
import statistics
import subprocess
import sys
import openpyxl
from openpyxl.comments import Comment

ROOT=Path(__file__).resolve().parents[1]
STUDY=ROOT/'experiments/work9-ablation-20261015'
RESULTS=ROOT/'results/work9-ablation-20261015'
SOURCE=ROOT.parent/'work9.xlsx'
BACKUP=STUDY/'work9_before_ablation.xlsx'
OUTPUT=STUDY/'work9_completed_ablation.xlsx'
BINARY=ROOT/'build/work9_ablation'
DATASETS={'sift1m':('SIFT1M',128),'gist1m':('GIST1M',960),'glove100':('GloVe-100',100),'deep1m':('DEEP1M',96)}
PAIRS=[(1,3),(2,4),(5,6)]
NAMES={1:'Merkle+IVF',2:'Merkle+HNSW',3:'Merkle+IVF+(i)',4:'Merkle+HNSW+(i)',5:'(i)+(ii)',6:'(i)+(ii)+(iii)'}
ROW_ORDER=[1,2,3,4,5,6]
NOTES='消融实验说明'
RECORDS='消融实测记录'
SOURCES='消融单元格来源'

def sha(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        while b:=f.read(8*1024*1024):h.update(b)
    return h.hexdigest()

def write_json(path,obj):
    path.write_text(json.dumps(obj,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')

def event(**kw):
    kw['time']=datetime.now().astimezone().isoformat()
    with (STUDY/'events.jsonl').open('a',encoding='utf-8') as f:f.write(json.dumps(kw,ensure_ascii=False)+'\n')
    print(json.dumps(kw,ensure_ascii=False),flush=True)

def env():
    e=os.environ.copy();e['OMP_NUM_THREADS']=e['OPENBLAS_NUM_THREADS']='1'
    faiss=ROOT/'third_party/LearnedEarlyTer-20SIGMOD'
    e['LD_LIBRARY_PATH']=':'.join(map(str,[ROOT/'third_party/runtime/lib',faiss,faiss/'lib']))
    return e

def fingerprints():
    files=[BINARY,ROOT/'src/experiments/work9_ablation.cpp',ROOT/'build/multidata_lut_include.cpp',ROOT/'src/veriann_query_lut.hpp',ROOT/'src/paper_repro_common.hpp',ROOT/'src/verifiable_ann_common.hpp']
    # The included protocol has several archived prototype sources.
    files+=sorted((ROOT/'src/experiments').glob('*.cpp'))
    return {str(p.relative_to(ROOT)):sha(p) for p in dict.fromkeys(files)}

def rows(path):return list(csv.DictReader(path.open()))

def table(ws):
    heads=[c.row for row in ws for c in row if c.row>100 and isinstance(c.value,str) and c.value.strip().lower()=='recall@10']
    assert len(heads)==1,(ws.title,heads)
    header=heads[0]
    for offset in range(1,7):
        label=ws.cell(header+offset,4).value or ws.cell(header+offset,5).value
        assert label is not None,(ws.title,header+offset,'missing variant label')
    return header

def initialize():
    STUDY.mkdir(parents=True,exist_ok=True);RESULTS.mkdir(parents=True,exist_ok=True)
    if not BACKUP.exists():shutil.copy2(SOURCE,BACKUP)
    wb=openpyxl.load_workbook(BACKUP)
    mapping={name:table(wb[name]) for name,_ in DATASETS.values()}
    write_json(STUDY/'table_mapping.json',mapping)
    write_json(STUDY/'protocol.json',dict(N=1000000,k=10,query_ids=[100,119],distinct_queries=20,repeats=3,
        target_recall=.95,M=16,efConstruction=200,hash='SHA-256',query_threads=1,build_threads=16,
        IVF=dict(nlist=256,training_rows=65536,kmeans_iterations=12,seed=42),
        pairs=PAIRS,query_time='search + full serialized VO generation',verify_time='decode + authentication + replay',
        units=dict(recall='fraction [0,1]',times='ms',vo='decimal KB = bytes / 1000'),
        scope='20-query preview; parameter calibration and reported recall use the same IDs, not held-out evaluation',
        glove='first 1M of normalized GloVe-100; squared L2 gives cosine ranking',
        interpretation=dict(i='SQ4 -> SQ6 -> raw conservative bounds, coupled to exact traversal',
                            ii='SQ4 top1 budgeted traversal; exact rerank and final-threshold refinement after traversal',
                            iii='expansion-centric authentication, with unchanged traversal and answers')))

def launch(dataset,variants,parameters,tag,pilot=False,n=1000000,queries=20,repeats=3,nlist=256):
    out=RESULTS/dataset/tag;log=STUDY/(dataset+'-'+tag+'.log')
    marker=out/('PILOT_PASS.txt' if pilot else 'PASS.txt')
    cmd=[str(BINARY),'--dataset',str(ROOT/'dataset'/dataset),'--cache',str(ROOT/'ANNProof-24FGCS/cache'),
         '--output',str(out),'--base',str(n),'--queries',str(queries),'--query-offset','100','--repeats',str(repeats),
         '--M','16','--ef-construction','200','--build-threads','16','--k','10',
         '--ef-search',','.join(map(str,parameters)),'--variants',','.join(map(str,variants)),'--nlist',str(nlist),'--pilot',str(int(pilot))]
    current=fingerprints()
    if marker.exists():
        info=json.loads((out/'environment.json').read_text())
        assert info['fingerprints']==current and info['command']==cmd,('stale result',out)
        return out
    if out.exists() and any(out.iterdir()):raise RuntimeError('Inspect incomplete run before resuming: '+str(out))
    out.parent.mkdir(parents=True,exist_ok=True)
    event(action='start',dataset=dataset,variants=variants,pilot=pilot,parameters=parameters,output=str(out))
    with log.open('w') as f:
        completed=subprocess.run(['/usr/bin/time','-v','-o',str(log)+'.resource.txt',*cmd],cwd=ROOT,env=env(),stdout=f,stderr=subprocess.STDOUT)
    if completed.returncode:raise RuntimeError(log.read_text()[-8000:])
    assert marker.exists()
    write_json(out/'environment.json',dict(command=cmd,fingerprints=current,platform=sys.platform,log=str(log),resource=str(log)+'.resource.txt'))
    event(action='completed',dataset=dataset,variants=variants,pilot=pilot,parameters=parameters)
    return out

def pilot_points(out,variants):
    data=rows(out/'pilot.csv');groups={}
    for row in data:groups.setdefault(int(row['parameter']),{}).setdefault(int(row['variant']),[]).append(row)
    points=[]
    for param,byvariant in groups.items():
        assert set(byvariant)==set(variants)
        for v,rs in byvariant.items():assert len(rs)==20 and len({r['query'] for r in rs})==20
        ref=byvariant[variants[0]]
        for v in variants[1:]:
            for a,b in zip(ref,byvariant[v]):
                assert all(a[c]==b[c] for c in ['query','answer','recall_at_k','visited','expanded']),(out,v,a,b)
        points.append(dict(parameter=param,recall=statistics.fmean(float(r['recall_at_k']) for r in ref),pilot=str(out)))
    return points

def calibrate(dataset,pair):
    path=STUDY/f'selection-{dataset}-{pair[0]}-{pair[1]}.json'
    if path.exists():return json.loads(path.read_text())
    if pair[0]==1:phases=[[4,8,16,32,64,128,192,256]]
    else:phases=[[10,16,24,32,48,64,96,128,192,256,384,512,768,1024,1536,2048],[4096,8192],[16384,32768],[65536,131072]]
    points=[]
    for phase,grid in enumerate(phases):
        out=launch(dataset,pair,grid,f'pilot-{pair[0]}-{pair[1]}-phase{phase}',True)
        points+=pilot_points(out,pair)
        if max(p['recall'] for p in points)>=.95-1e-12:break
    for refinement in range(5):
        eligible=[p for p in points if p['recall']>=.95-1e-12]
        assert eligible,(dataset,pair,'95% target not reached')
        chosen=min(eligible,key=lambda p:p['parameter'])
        upper=chosen['parameter'];lower=max([p['parameter'] for p in points if p['parameter']<upper and p['recall']<.95-1e-12],default=0)
        if upper-lower<=1 or chosen['recall']<=.96+1e-12:break
        step=max(1,math.ceil((upper-lower)/8))
        seen={p['parameter'] for p in points};grid=[v for v in range(lower+step,upper,step) if v not in seen and (pair[0]==1 or v>=10)]
        if not grid:break
        out=launch(dataset,pair,grid,f'pilot-{pair[0]}-{pair[1]}-refine{refinement}-{lower}-{upper}',True)
        points+=pilot_points(out,pair)
    write_json(path,dict(chosen,all_points=points));event(action='selected',dataset=dataset,variants=pair,**chosen)
    return dict(chosen,all_points=points)

def audit_pair(out,pair,point,query_count=20,repeats=3):
    comparisons=[];summary=[];attack_total=0
    pilot={}
    if point:
        pilot={(int(r['variant']),int(r['query'])):r for r in rows(Path(point['pilot'])/'pilot.csv') if int(r['parameter'])==point['parameter']}
    reference=None
    for v in pair:
        folder=out/f'variant-{v}';raw=rows(folder/'per_query.csv');answers=rows(folder/'answers.csv')
        assert len(raw)==len(answers)==query_count*repeats
        assert all(r['verified']=='1' for r in raw)
        assert {(int(r['query']),int(r['repeat'])) for r in raw}=={(q,t) for q in range(100,100+query_count) for t in range(repeats)}
        for r,a in zip(raw,answers):
            assert r['query']==a['query'] and r['repeat']==a['repeat'] and r['ef_search']==a['parameter']
            assert len(a['answer'].split())==10
            if point:
                old=pilot[v,int(r['query'])]
                assert a['answer']==old['answer']
                assert all(r[c]==old[c] for c in ['recall_at_k','visited','expanded'])
            for c in ['search_ms','vo_generation_ms','verification_ms','vo_bytes']:assert math.isfinite(float(r[c])) and float(r[c])>0
        if reference is not None:
            for (r,a),(old,oa) in zip(zip(raw,answers),reference):
                assert all(a[c]==oa[c] for c in ['query','repeat','parameter','answer'])
                assert all(r[c]==old[c] for c in ['recall_at_k','visited','expanded'])
                comparisons.append([v,r['query'],r['repeat']])
        else:reference=list(zip(raw,answers))
        attacks=rows(folder/'tamper_tests.csv');assert len(attacks)>=12 and all(r['rejected']=='1' for r in attacks)
        attack_total+=len(attacks)
        mean=lambda key:statistics.fmean(float(r[key]) for r in raw)
        record=dict(variant=v,name=NAMES[v],recall=mean('recall_at_k'),query_ms=mean('server_total_ms'),
                    search_ms=mean('search_ms'),vo_generation_ms=mean('vo_generation_ms'),verify_ms=mean('verification_ms'),
                    authentication_ms=mean('authentication_ms'),replay_ms=mean('replay_ms'),vo_kb=mean('vo_bytes')/1000,
                    parameter=int(raw[0]['ef_search']),visited=mean('visited'),expanded=mean('expanded'),
                    samples=len(raw),tamper_tests=len(attacks),output=str(folder),source_sha256=sha(folder/'per_query.csv'))
        if point:assert record['recall']>=.95-1e-12
        summary.append(record)
    write_json(out/'pair_audit.json',dict(status='PASS',paired_answer_comparisons=len(comparisons),tamper_tests=attack_total,summary=summary))
    return summary

def same_value(a,b):
    if isinstance(a,(int,float)) and isinstance(b,(int,float)):return math.isclose(a,b,rel_tol=1e-12,abs_tol=1e-12)
    return a==b

def checkpoint():
    evaluations=[]
    for file in sorted(STUDY.glob('evaluation-*.json')):evaluations+=json.loads(file.read_text())
    if not evaluations:return
    before_sha=sha(SOURCE);snapshot=STUDY/('work9_snapshot_'+before_sha[:12]+'.xlsx')
    if not snapshot.exists():shutil.copy2(SOURCE,snapshot)
    wb=openpyxl.load_workbook(snapshot);allowed=set();cellrecords=[]
    for record in evaluations:
        name=DATASETS[record['dataset']][0];ws=wb[name];header=table(ws);v=record['variant'];row=header+v
        # Original table labels may occupy D or E, depending on merged layout.
        labelcol=4 if ws.cell(row,4).value is not None else 5
        if v in [2,4]:
            ws.cell(row,labelcol,NAMES[v]);allowed.add((name,row,labelcol))
        for col,label in [(7,'Query Time (ms)'),(8,'VO Size (KB)'),(9,'Verify Time (ms)')]:
            ws.cell(header,col,label);allowed.add((name,header,col))
        for col,key in [(6,'recall'),(7,'query_ms'),(8,'vo_kb'),(9,'verify_ms')]:
            cell=ws.cell(row,col,record[key]);cell.number_format='0.000' if col==6 else '0.0000'
            cell.comment=Comment(f"真实完整VO实验；N=1,000,000, d={DATASETS[record['dataset']][1]}, k=10, M=16, efConstruction=200, SHA-256。\n20条不同查询(ids 100..119) × 3次；单线程。\n{NAMES[v]}; parameter={record['parameter']}; 实际Recall@10={record['recall']:.6f}。\nQuery=search+VO序列化；Verify=解码+认证+重放；ms；VO=bytes/1000 KB。\n原始记录：{record['output']}/per_query.csv\nSHA256={record['source_sha256']}", '实测来源')
            allowed.add((name,row,col));cellrecords.append([name,cell.coordinate,NAMES[v],key,record[key],record['parameter'],record['output'],record['source_sha256']])
    for title in [NOTES,RECORDS,SOURCES]:
        if title in wb:del wb[title]
    notes=wb.create_sheet(NOTES)
    for row in [
        ['项目','说明'],['实验范围','四数据集均使用前1M向量；GloVe原库1,183,514条，本消融固定前1M；其他已完成表格保持原实验规模。'],
        ['样本与召回','每组20条不同查询(ids 100..119) × 3次计时，k=10；单线程。以约95% Recall@10选参数；同一查询集调参及报告，是快速预览，不是独立留出评估。'],
        ['共同配置','HNSW M=16, efConstruction=200；构图16线程，查询/VO/验证1线程；SHA-256；量化版本共同启用LUT、隐式位置多证明及图BFS叶排序。SIFT旧主表的M=8配置不用于本消融。'],
        ['GIST图的限制','沿用已有M16/efConstruction200缓存图：既有入口可达性检查为all_levels 974,208/1,000,000，图并非全入口可达。所有图消融共用该图；真值仍对全部1M精确计算，报告本组20条查询实际召回率。检查见work9-multidata-20261015/gist-M16-efc200-reachability.txt。'],
        ['NSW解释','按用户确认，原表Merkle+NSW沿用当前HNSW，行名已明确改为Merkle+HNSW。'],
        ['(i) 渐进量化','SQ4→SQ6→原始向量的保守距离下界；与IVF扫描/HNSW精确距离遍历耦合，按遍历过程中的动态阈值逐级披露。'],
        ['(ii) 验证解耦遍历','使用SQ4 top1与固定扩展预算E；搜索完成后原始距离重排，并根据最终第k距离生成渐进验证材料。该项改变搜索策略，按约95%召回率另选E。'],
        ['(iii) 按扩展记录认证','仅对实际展开的邻接表与邻居SQ4记录认证；删除(iii)时每个发现节点需要打开其所有层邻接表与SQ4的Merkle叶。两行固定相同E及结果。'],
        ['成对控制','第1/3行共用nprobe且逐条答案相同；第2/4行共用efSearch且逐条答案相同；第5/6行共用E、逐条答案与展开/访问数量相同。'],
        ['IVF','nlist=256；库内均匀取65536条训练，kmeans 12轮、seed=42；完整选中倒排列表成员ID、所有中心和向量/量化数据均真实写入VO，并认证列表完整性及中心承诺。IVF向量/精化树按数据库ID排序，图版本使用BFS排序；各同类成对版本排序相同。'],
        ['查询时间','Query Time=搜索+完整VO生成/序列化；另在消融实测记录拆分search和VO generation。ms；不含索引构建、数据加载、网络。'],
        ['验证时间','VO完整解码+所有Merkle认证+路由/遍历重放+top-k校验；客户端只有可信根及量化元数据，不访问服务端数据库。ms。'],
        ['单位','Recall@10为0到1小数；VO Size为十进制KB=bytes/1000；Query/Verify均为ms。'],
        ['正确性','每条诚实VO必须通过验证；每个版本执行查询绑定/答案/数据/证明/删改等篡改检查；保留序列化VO样本与逐条原始日志。'],
        ['解释限制','这是项目内实现的受控消融，不将IVF/HNSW基础行称为ANNProof或VLAH原论文结果；小样本计时存在波动，各项优化不保证在每个数据集都更快。'],
        ['原始文件',str(RESULTS)],['已完成行数',len(evaluations)],['已填指标单元格',len(cellrecords)]]:notes.append(row)
    rec=wb.create_sheet(RECORDS);keys=['dataset','variant','name','recall','query_ms','search_ms','vo_generation_ms','verify_ms','authentication_ms','replay_ms','vo_kb','parameter','visited','expanded','samples','tamper_tests','output','source_sha256'];rec.append(keys)
    for record in evaluations:rec.append([record.get(k) for k in keys])
    src=wb.create_sheet(SOURCES);src.append(['Sheet','Cell','Variant','Metric','Value','Parameter','Raw folder','SHA256'])
    for row in cellrecords:src.append(row)
    for ws in [notes,rec,src]:ws.freeze_panes='A2';ws.column_dimensions['A'].width=26;ws.column_dimensions['B'].width=100;ws.auto_filter.ref=ws.dimensions
    wb.save(OUTPUT)
    # Verify all existing values outside the exact authorized update set survive.
    before=openpyxl.load_workbook(snapshot);after=openpyxl.load_workbook(OUTPUT)
    assert after.sheetnames[:len(before.sheetnames)-sum(t in before.sheetnames for t in [NOTES,RECORDS,SOURCES])]==[s for s in before.sheetnames if s not in [NOTES,RECORDS,SOURCES]]
    preserved=0
    for ws in before:
        if ws.title in [NOTES,RECORDS,SOURCES]:continue
        for row in ws:
            for cell in row:
                if (ws.title,cell.row,cell.column) in allowed:continue
                assert same_value(cell.value,after[ws.title][cell.coordinate].value),(ws.title,cell.coordinate,'unexpected change')
                preserved+=1
    for sheet,addr,_,_,value,*_ in cellrecords:assert same_value(after[sheet][addr].value,value)
    assert sha(SOURCE)==before_sha,'Workbook externally changed during save; completed copy retained'
    try:shutil.copy2(OUTPUT,SOURCE)
    except PermissionError:
        event(action='workbook_locked',completed_copy=str(OUTPUT));return
    assert sha(SOURCE)==sha(OUTPUT)
    write_json(STUDY/'checkpoint.json',dict(status='PASS',completed_rows=len(evaluations),filled_cells=len(cellrecords),preserved_cells=preserved,workbook_sha256=sha(SOURCE)))
    event(action='workbook_updated',completed_rows=len(evaluations),filled_cells=len(cellrecords))

def smoke():
    for pair in PAIRS:
        out=launch('sift1m',pair,[8 if pair[0]==1 else 64],f'smoke-{pair[0]}-{pair[1]}',False,n=2000,queries=5,repeats=1,nlist=32)
        audit_pair(out,pair,None,5,1)
    event(action='smoke_passed',variants=6)

def main():
    initialize()
    if '--smoke' in sys.argv:smoke();return
    assert all((RESULTS/'sift1m'/f'smoke-{p[0]}-{p[1]}'/'pair_audit.json').exists() for p in PAIRS),'Run --smoke first'
    for dataset in DATASETS:
        for pair in PAIRS:
            file=STUDY/f'evaluation-{dataset}-{pair[0]}-{pair[1]}.json'
            point=calibrate(dataset,pair)
            out=launch(dataset,pair,[point['parameter']],f'formal-{pair[0]}-{pair[1]}')
            summaries=audit_pair(out,pair,point)
            for r in summaries:r['dataset']=dataset
            write_json(file,summaries);checkpoint()
        event(action='dataset_passed',dataset=dataset)
    all_records=[r for file in STUDY.glob('evaluation-*.json') for r in json.loads(file.read_text())]
    assert len(all_records)==24 and sum(r['samples'] for r in all_records)==1440
    checkpoint()
    final=dict(status='PASS',variant_dataset_rows=24,filled_metric_cells=96,formal_measurements=1440,
               tamper_tests=sum(r['tamper_tests'] for r in all_records),paired_answer_comparisons=720,workbook_sha256=sha(SOURCE))
    write_json(STUDY/'final_validation.json',final);event(action='all_complete',**final)

if __name__=='__main__':main()
