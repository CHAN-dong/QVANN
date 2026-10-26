"""Independent wire-length parsers and additive workbook/LaTeX tables."""
import csv
import json
import math
import mmap
import shutil
import statistics
import struct
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
import work9_resource_matched as u

R=u.ROOT;S=R/'experiments/work9-vo-breakdown-20261023';O=R/'results/work9-vo-breakdown-20261023'
SOURCE=R.parent/'work9.xlsx';BACKUP=S/'work9_before_vo_breakdown.xlsx'
assert not (S/'APPLIED.json').exists()
assert u.read(S/'FINISHED.json')['passed'] and u.sha(SOURCE)==u.sha(BACKUP)
protocol=u.read(S/'protocol.json');records=u.read(S/'summary.json');samples=list(csv.DictReader((S/'samples.csv').open()))
assert len(records)==20 and len(samples)==1200

class Reader:
    def __init__(self,data):self.data=data;self.p=0
    def skip(self,n):assert n>=0 and self.p+n<=len(self.data);self.p+=n
    def u32(self):x=struct.unpack_from('<I',self.data,self.p)[0];self.skip(4);return x
    def u8(self):x=self.data[self.p];self.skip(1);return x
    def header(self):
        assert self.u32()==0x56505231;self.skip(32);k=self.u32();ef=self.u32();count=self.u32();assert count==k;self.skip(4*k);return k,ef
    def record(self,dim,edge_bytes):
        self.skip(8);assert self.u32()==dim;self.skip(dim*4)
        for _ in range(self.u32()):self.skip(self.u32()*edge_bytes)
    def proof(self,width):
        before=self.p;self.skip(self.u32()*width);return self.p-before

def parse(path,method,dim):
    with path.open('rb') as handle,mmap.mmap(handle.fileno(),0,access=mmap.ACCESS_READ) as data:
        r=Reader(data);total=len(data)
        if method=='MR-tree':
            assert r.u32()==0x564f3031;n=r.u32();assert r.u32()==dim;k=r.u32();r.skip(128+4*k)
            frontier=opened=points=0
            def node():
                nonlocal frontier,opened,points
                tag=r.u8()
                if tag==0:frontier+=1;r.skip(dim*8+32)
                elif tag==1:
                    opened+=1;count=r.u32();points+=count;r.skip(count*(4+dim*4))
                elif tag==2:
                    opened+=1
                    for _ in range(r.u32()):node()
                else:raise AssertionError('invalid MR node')
            node();auth=128+32*frontier;ef=None;probes=None;local=None;route=None;record=None
        else:
            k,ef=r.header();frontier=opened=points=None;route=0
            if method=='ANNProof':
                begin=r.p;count=r.u32();recbegin=r.p
                for _ in range(count):r.record(dim,4)
                record=r.p-recbegin;r.skip(r.u32()*8);begin=r.p
                probes=1
                for _ in range(r.u32()):r.skip(4);r.proof(40);r.proof(40)
                local=r.p-begin-4;auth=32+r.p-begin
            elif method=='VLAH':
                r.skip(4);route=r.u32();r.skip(route);probes=r.u32();local=record=0
                for _ in range(probes):
                    r.skip(4);count=r.u32();begin=r.p
                    for _ in range(count):r.record(dim,8)
                    record+=r.p-begin;r.skip(r.u32()*8);local+=r.proof(40)
                auth=32+4+route+local
            else:
                variant=r.u32();assert variant==(1 if method=='IVF-ADS' else 2)
                if variant==1:
                    r.skip(256*dim*4);probes=r.u32()
                    for _ in range(probes):r.skip(4);r.skip(r.u32()*4)
                else:probes=1
                count=r.u32()
                for _ in range(count):
                    r.skip(8+dim*4)
                    if variant==2:
                        for _ in range(r.u32()):r.skip(r.u32()*4)
                record=r.p;begin=r.p;r.proof(32)
                if variant==1:r.proof(32)
                local=r.p-begin;auth=32+local
        assert r.p==total,(path,r.p,total)
        return dict(total=total,A=auth,T=total-auth,k=k,ef=ef,probes=probes,local=local,route=route,record=record,frontier=frontier,opened=opened,points=points)

parsed=[]
for rec in records:
    path=Path(rec['source']);assert u.sha(path)==rec['csv_sha256']
    wire=path.parent/('query0.vo')
    assert wire.exists(),wire
    p=parse(wire,rec['method'],rec['D'])
    raw=list(csv.DictReader(path.open()))
    if rec['method']=='MR-tree':
        first=next(x for x in raw if int(x['query'])==100 and int(x['repeat'])==0)
        assert p['A']==int(first['A_bytes']) and p['frontier']==int(first['frontier_nodes'])
        assert p['opened']==int(first['opened_nodes']) and p['points']==int(first['examined'])
    else:
        first=next(x for x in raw if int(x['query'])==100 and int(x['repeat'])==0 and int(x['k'])==p['k'] and int(x['ef_search'])==p['ef'] and int(x['probes'])==p['probes'])
        assert p['local']==int(first['local_proof_bytes']) and p['route']==int(first['route_proof_bytes'])
        assert p['record']==int(first['record_bytes'])
    assert p['total']==int(first['vo_bytes'])
    parsed.append(dict(dataset=rec['dataset'],method=rec['method'],file=str(wire),sha256=u.sha(wire),**p))
for rec in records:
    rows=[x for x in samples if x['dataset']==rec['dataset'] and x['method']==rec['method']]
    assert len(rows)==60 and {(int(x['query']),int(x['repeat'])) for x in rows}=={(q,t) for q in range(100,120) for t in range(3)}
    for x in rows:
        assert int(x['T_raw_bytes'])+int(x['A_raw_bytes'])==int(x['total_raw_bytes'])
        for prefix in ['T','A','total']:assert int(x[prefix+'_display_bytes'])==int(x[prefix+'_raw_bytes'])*int(x['factor'])
    for metric in ['T_raw_bytes','A_raw_bytes','total_raw_bytes','T_display_bytes','A_display_bytes','total_display_bytes','recall']:
        assert math.isclose(rec[metric],statistics.fmean(float(x[metric]) for x in rows),rel_tol=1e-12)

wb=openpyxl.load_workbook(BACKUP);locations={};blue=PatternFill('solid',fgColor='305496')
for ds,(sheet,n,dim) in protocol['datasets'].items():
    ws=wb[sheet];title=ws.max_row+4;header=title+4;locations[ds]=dict(sheet=sheet,title_row=title,header_row=header,last_row=header+6)
    ws.merge_cells(start_row=title,start_column=4,end_row=title,end_column=11);ws.cell(title,4,'VO breakdown（k=10，KB=1000 bytes）').font=Font(size=12,bold=True,color='17365D')
    ws.merge_cells(start_row=title+1,start_column=4,end_row=title+3,end_column=11)
    c=ws.cell(title+1,4,f'N={n:,}, d={dim}；20查询×3重复；T+A=完整VO。T包含向量/ID/邻接表/MBB剪枝证据及一般控制字段；A包含query绑定哈希、独立认证证明及其编码（VLAH包含完整认证路由证明）。'
        'Result-only为10条原始向量及ID，无认证或协议头。ANNProof*的T/A按历史要求各×2，属于派生展示值，原始实测拆分保存在右侧；其他方法不乘系数。'
        'IVF-ADS/HNSW-ADS分别为原Merkle+IVF/Merkle+HNSW消融，使用各自约95%召回配置；MR-tree是100%精确检索。')
    c.alignment=Alignment(wrap_text=True,vertical='center')
    for rr in range(title+1,title+4):ws.row_dimensions[rr].height=25
    labels=['Method','Recall@10','T (KB)','A (KB)','Total (KB)','Raw T (KB)','Raw A (KB)','VO factor']
    for col,label in enumerate(labels,4):
        c=ws.cell(header,col,label);c.font=Font(bold=True,color='FFFFFF');c.fill=blue;c.alignment=Alignment(wrap_text=True,horizontal='center')
    ws.cell(header+1,4,'Result-only');ws.merge_cells(start_row=header+1,start_column=6,end_row=header+1,end_column=7)
    ws.cell(header+1,6,10*(4+4*dim)/1000);ws.cell(header+1,8,10*(4+4*dim)/1000)
    ws.cell(header+1,6).comment=Comment('理论字节计数，不是新性能测量：k×(4-byte ID + d×4-byte float32)。不包含认证、协议头或重复结果。','Codex')
    for index,method in enumerate(protocol['methods'],2):
        rec=next(x for x in records if x['dataset']==ds and x['method']==method);row=header+index
        vals=[method+('*' if rec['factor']==2 else ''),rec['recall'],rec['T_display_bytes']/1000,rec['A_display_bytes']/1000,rec['total_display_bytes']/1000,rec['T_raw_bytes']/1000,rec['A_raw_bytes']/1000,rec['factor']]
        for col,val in enumerate(vals,4):
            c=ws.cell(row,col,val);c.alignment=Alignment(horizontal='left' if col==4 else 'center',wrap_text=True)
            c.number_format='0.0%' if col==5 else '0.000' if 6<=col<=10 else 'General'
            c.comment=Comment(f"VO字节拆分：{method}, N={n}, d={dim}, k=10；查询100..119×3，参数={rec['parameter']}。\n"
                'T为完整VO除去独立认证材料后的其余字节；A含query-binding hash32及认证证明编码。VLAH整段routing proof归A；MR未展开MBB归T，content digest归A，root/signature归A。\n'
                f"原始均值 T={rec['T_raw_bytes']/1000:.9f}KB；A={rec['A_raw_bytes']/1000:.9f}KB；完整={rec['total_raw_bytes']/1000:.9f}KB。\n"
                f"VO展示系数={rec['factor']}；乘系数后为派生值，不能描述成实际序列化VO大小。\n来源={rec['source']}\nSHA256={rec['csv_sha256']}",'Codex')
        ws.row_dimensions[row].height=26

notes=wb.create_sheet('VO拆分实验说明')
for row in [
    ['项目','说明'],['日期','2026-10-23'],['范围','五个方法×四数据集，共20组、1200条实际VO样本的结构计数；全部使用query100..119×3、k=10。'],
    ['规模','SIFT1M；NYTimes289761；GloVe1183514；DEEP1M。GloVe旧消融是1M，本次对IVF-ADS/HNSW-ADS重新校准并在完整规模生成验证。'],
    ['T定义',protocol['definitions']['T']],['A定义',protocol['definitions']['A']],['Result-only',protocol['definitions']['result_only']],
    ['统计','KB=1000bytes。每个数据集/方法60条样本的字节均值。每条保证T+A=实际完整VO。SIFT使用20查询，与旧100查询overall均值可能不同。'],
    ['认证与遍历边界','这是一种明确声明的结构划分，不是唯一划分。A包含独立认证字段及其编码、VLAH认证路由证明；MR的MBB和结构标签计入T。不可只用record_bytes当完整T，它不含所有header/trace。'],
    ['ANNProof*','延续用户先前VO×2要求，展示表T/A各×2且有星号。Raw T/Raw A及原始CSV保留；派生表不代表VO实际扩大发生过。提供raw版LaTeX用于报告真实测量。'],
    ['复用/补测',protocol['source_policy']],['召回','MR精确100%；其余各自沿用或校准>=95%参数，并非每行严格同一实测召回。对应Recall保存在各工作表及附表。'],
    ['字节审计','独立Python mmap解析20个完整query0 VO，交叉核对原始CSV的record/proof/route长度；逐条核对1200个T+A等式及展示系数。'],
    ['备份',str(BACKUP)],['原始明细',str(S/'samples.csv')],['实测LaTeX',str(S/'vo_breakdown_raw.tex')],['当前系数LaTeX',str(S/'vo_breakdown_adjusted.tex')],
]:notes.append(row)
details=wb.create_sheet('VO拆分原始与系数')
keys=['sheet','method','N','D','samples','recall','parameter','T_raw_bytes','A_raw_bytes','total_raw_bytes','factor','T_display_bytes','A_display_bytes','total_display_bytes','source','csv_sha256','status']
details.append(keys)
for rec in records:details.append([rec[x] for x in keys])
for ws in [notes,details]:
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=blue
    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=24
    for row in ws:
        for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
notes.column_dimensions['B'].width=115

for adjusted in [False,True]:
    cap=r'VO breakdown across four datasets ($k=10$, KB). $T$: $\mathcal{VO}_{\mathrm{trav}}$; $A$: $\mathcal{VO}_{\mathrm{auth}}$.'
    if adjusted:cap+=r' ANNProof$^{*}$ values are twice the measured sizes, following the scaling convention used in our comparison tables.'
    lines=[r'\begin{table}[!t]',r'\centering',r'\caption{'+cap+'}',r'\label{tab:vo_breakdown'+('_adjusted' if adjusted else '')+'}',r'\vspace{-0.2cm}',r'\footnotesize',r'\setlength{\tabcolsep}{2pt}',r'\renewcommand{\arraystretch}{1.1}',r'\resizebox{\columnwidth}{!}{',r'\begin{tabular}{@{}l*{4}{rr}@{}}',r'\toprule',r'\multirow{2}{*}{Method} & \multicolumn{2}{c}{\textsf{SIFT}} & \multicolumn{2}{c}{\textsf{NYTimes}} & \multicolumn{2}{c}{\textsf{GloVe}} & \multicolumn{2}{c}{\textsf{DEEP}} \\',r'\cmidrule(lr){2-3}\cmidrule(lr){4-5}\cmidrule(lr){6-7}\cmidrule(l){8-9}',r'& $T$ & $A$ & $T$ & $A$ & $T$ & $A$ & $T$ & $A$ \\',r'\midrule',r'Result-only & \multicolumn{2}{c}{5.16} & \multicolumn{2}{c}{10.28} & \multicolumn{2}{c}{4.04} & \multicolumn{2}{c}{3.88} \\',r'\midrule']
    prefix='display' if adjusted else 'raw'
    for method in protocol['methods']:
        vals=[]
        for ds in protocol['datasets']:
            rec=next(x for x in records if x['dataset']==ds and x['method']==method)
            vals.extend([f"{rec['T_'+prefix+'_bytes']/1000:.2f}",f"{rec['A_'+prefix+'_bytes']/1000:.2f}"])
        label=r'\textsf{'+method+'}'+(r'$^{*}$' if adjusted and method=='ANNProof' else '')
        lines.append(label+' & '+' & '.join(vals)+r' \\')
    lines += [r'\bottomrule',r'\end{tabular}',r'}',r'\end{table}','']
    (S/('vo_breakdown_adjusted.tex' if adjusted else 'vo_breakdown_raw.tex')).write_text('\n'.join(lines))

review=S/'work9_vo_breakdown_updated.xlsx';wb.save(review);old=openpyxl.load_workbook(BACKUP);new=openpyxl.load_workbook(review);preserved=0
for ws in old:
    ns=new[ws.title]
    for row in ws:
        for c in row:
            nc=ns[c.coordinate];assert c.value==nc.value and (c.comment.text if c.comment else None)==(nc.comment.text if nc.comment else None),(ws.title,c.coordinate)
            assert c.number_format==nc.number_format and copy(c.font)==copy(nc.font) and copy(c.fill)==copy(nc.fill) and copy(c.border)==copy(nc.border) and copy(c.alignment)==copy(nc.alignment)
            preserved+=1
for ds,loc in locations.items():
    ws=new[loc['sheet']]
    for i,method in enumerate(protocol['methods'],2):
        row=loc['header_row']+i;assert math.isclose(ws.cell(row,6).value+ws.cell(row,7).value,ws.cell(row,8).value,rel_tol=1e-12)
audit=dict(passed=True,points=20,query_byte_accounts=1200,independently_parsed_VOs=parsed,original_cells_preserved=preserved,new_tables=locations,workbook_sha256=u.sha(review),backup_sha256=u.sha(BACKUP))
u.write(S/'audit.json',audit)
assert u.sha(SOURCE)==u.sha(BACKUP)
try:
    shutil.copy2(review,SOURCE)
except PermissionError:
    print('AUDIT_PASSED; reviewed workbook ready; source locked by spreadsheet app. Apply through saved-workbook COM handoff.',flush=True)
else:
    assert u.sha(SOURCE)==audit['workbook_sha256'];u.write(S/'APPLIED.json',audit);print('APPLIED',flush=True)
print(json.dumps(dict(points=20,query_byte_accounts=1200,old_cells_preserved=preserved,new_tables=locations),ensure_ascii=False,indent=2))
