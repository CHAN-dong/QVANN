"""Fill only offline index columns, preserving all query experiments."""
from pathlib import Path
import json,hashlib,shutil,time,math
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font,PatternFill
from work9_index_study import ROOT,STUDY,OUT,DATA,write,sha
SOURCE=ROOT.parent/'work9.xlsx'
records=[]
for ds,(title,*_) in DATA.items():
    for method in ['VLAH','ANNProof','Ours','V3DB']:
        path=OUT/ds/method/'measurement.json'
        if not path.exists():raise RuntimeError('Missing measurement: '+str(path))
        r=json.loads(path.read_text());assert r['status']=='measured' and r['threads']==1
        assert r['total_bytes']==r['search_bytes']+r['auth_bytes']
        assert all(math.isfinite(r[k]) and r[k]>0 for k in ['total_bytes','construct_seconds'])
        if method=='V3DB':
            for name in ['search','auth']:
                f=path.parent/(name+'.snapshot');assert f.stat().st_size==r[name+'_bytes'] and sha(f)==r[name+'_sha256']
        else:
            f=path.parent/'serving.snapshot';assert f.stat().st_size==r['total_bytes'] and sha(f)==r['snapshot_sha256']
        r['record_path']=str(path);records.append(r)
before_hash=sha(SOURCE);wb=openpyxl.load_workbook(SOURCE)
before={(s.title,c.coordinate):c.value for s in wb for row in s for c in row if c.value is not None}
backup=STUDY/('work9_before_index_'+time.strftime('%Y%m%d_%H%M%S')+'.xlsx');shutil.copy2(SOURCE,backup)
audit=[];allowed=set()
for r in records:
    s=wb[r['title']];row={'VLAH':9,'ANNProof':10,'Ours':11,'V3DB':12}[r['method']]
    assert r['method'] in s.cell(row,4).value
    for addr,value in [('I8','Index Size (MB)'),('J8','Index Construct Time (s)')]:s[addr]=value;allowed.add((s.title,addr))
    for col,key,scale in [('I','total_bytes',1000000),('J','construct_seconds',1)]:
        addr=f'{col}{row}';c=s[addr];c.value=r[key]/scale;c.number_format='0.000';c.font=Font(name='Calibri',size=11,color='006100');c.fill=PatternFill('solid',fgColor='E2F0D9')
        cfg=r['config'] if 'config' in r else dict(M=r['M'],ef_construction=r['ef_construction'],LSH_width=r['width'])
        note=dict(status='实测',dataset=r['dataset'],method=r['method'],metric=key,unit='MB (1,000,000 bytes)' if col=='I' else 'seconds',N=r['N'],D=r['D'],threads=1,build_repeats=1,search_bytes=r['search_bytes'],auth_bytes=r['auth_bytes'],index_seconds=r['index_seconds'],ads_seconds=r['ads_seconds'],source=r['record_path'],scope='Fresh rebuild, no cache load. Index size is serialized serving snapshot, not peak memory; raw vectors included once only if used by serving algorithm. No query proofs/circuits, input I/O, ground truth or hyperparameter tuning. Final audit serialization outside measured construction; original library intermediate cache writes remain included.',config=cfg,rebuild_caveat='Same parameters as query table, freshly rebuilt with 1 thread. Historical graphs trained with more threads may differ in edges/tie assignments. Query results are not overwritten.')
        c.comment=Comment(json.dumps(note,ensure_ascii=False,indent=2),'索引实测来源');allowed.add((s.title,addr));audit.append(dict(sheet=s.title,cell=addr,value=c.value,**note))
for title in ['索引构建实测记录','索引构建说明']:
    if title in wb:del wb[title]
s=wb.create_sheet('索引构建实测记录')
s.append(['Dataset','Method','N','D','Threads','Build repetitions','Search bytes','Authentication bytes','Total MB','Search/index seconds','Authentication seconds','Total construction seconds','Peak RSS GiB','Status','M','efConstruction / K','Source'])
for r in records:s.append([r['title'],r['method'],r['N'],r['D'],1,1,r['search_bytes'],r['auth_bytes'],r['total_bytes']/1e6,r['index_seconds'],r['ads_seconds'],r['construct_seconds'],r['peak_rss_KiB']/1024**2,'实测',r['M'],r.get('ef_construction',r.get('K')),r['record_path']])
s.freeze_panes='C2';s.auto_filter.ref=s.dimensions
notes=[
('指标与单位','I列=Index Size，十进制MB；J列=Index Construct Time，秒。绿色为实测。本次四个数据集×四种方法均要求真实完成后才写实测。'),
('构建口径','已有预处理后的向量文件 → 从空缓存构建搜索索引/分桶/PQ与认证结构。单线程，每组完整构建一次；不是三次平均。排除输入读取、数据集预处理、GT、参数搜索和查询电路/证明。单独最终审计快照写盘不计时；原函数内部生成中间Faiss缓存的写盘保留计入。'),
('大小口径','实际写出的无压缩服务快照字节数，包含搜索表示、服务必需的原始向量和认证结构/元数据。排除重复原数据集、训练中间数组、缓存副本、查询VO和每查询电路。不是峰值内存/RSS，也不是仅Merkle树。'),
('峰值RSS口径','附表RSS是每个独立测量进程的全程峰值，包含输入、训练临时对象及最终快照序列化；不是索引常驻内存。主表构建时间单独计时，不包含最后的审计快照序列化。'),
('ANNProof','直接调用现有Graph和ADS构造器。保存服务实际使用的记录（向量、ID、分层邻接）和两类分片Merkle树及可信锚点；Graph内部用于构建的Faiss副本不重复计入最终服务大小。'),
('VLAH','直接调用现有LSH、分桶Graph、Merkle树和Trie构造器。保存全部桶图记录、LSH参数、Merkle树、桶描述和Trie。'),
('Ours','直接调用当前LUT版Codec、ADS与原HNSW构造器。保存Faiss图（含原始向量与预留邻接槽）、SQ4/增量码、量化参数、三棵Merkle树、位置映射与锚点。查询相关LUT是临时量，不计离线索引。'),
('V3DB','沿用已选M/K/L及25次迭代，从头单线程进行缩放、IVF训练、作者重平衡、全部PQ子空间训练和快照根生成。保存实际数组类型的中心、码本、PQ码、列表ID/偏移、参数与列表/快照根。全局树可由列表根在查询时重建；不假设尚未实现的码压缩。训练用scaled向量不属于服务索引。'),
('配置对应','SIFT图M8/efConstruction100；NYTimes、GloVe、DEEP图M16/efConstruction200，分桶大小20000、ANNProof分片10000；V3DB配置逐条见JSON。本次是相同参数下的新单线程构建；历史多线程图可能存在边或并列分配差异，没有覆盖历史查询结果。'),
('格式限制','各方案保留当前实现的表示：Ours的Faiss预留槽与ANNProof/VLAH实际边记录不同。比较结果包含这种实现/格式差异，不代表各论文理论最小空间。认证和搜索部分已拆列供核查。'),
('证据',str(OUT)),('运行脚本','scripts/work9_index_study.py build / run；scripts/work9_index_v3db.py；scripts/fill_work9_index.py'),('原表备份',str(backup))]
s=wb.create_sheet('索引构建说明')
for item in notes:s.append(item)
for s in [wb['索引构建实测记录'],wb['索引构建说明']]:
    for c in s[1]:c.font=Font(bold=True)
    s.column_dimensions['A'].width=24;s.column_dimensions['B'].width=90
temp=STUDY/'work9_with_index_measurements.xlsx';wb.save(temp)
check=openpyxl.load_workbook(temp)
changed=[(title,addr) for (title,addr),value in before.items() if (title,addr) not in allowed and check[title][addr].value!=value]
assert not changed,changed
assert len(audit)==32 and len(records)==16
assert sha(SOURCE)==before_hash,'Workbook changed during fill; merge needed'
shutil.copy2(temp,SOURCE);assert sha(temp)==sha(SOURCE)
write(STUDY/'cell-audit.json',audit)
write(STUDY/'validation.json',dict(status='PASS',measurements=16,filled_cells=32,predicted_cells=0,unrelated_existing_cells_unchanged=len(before)-len(before.keys()&allowed),source_before_sha256=before_hash,source_after_sha256=sha(SOURCE),backup=str(backup),output=str(SOURCE)))
report=['# 四数据集索引构建实测','', '单线程、每组一次完整空缓存构建。大小为服务快照MB；时间为搜索索引+认证结构秒。全部实测，无预测。细节见工作簿“索引构建说明”。','', '| 数据集 | 方法 | Index MB | 构建 s | 峰值 RSS GiB |','|---|---|---:|---:|---:|']
for r in records:report.append(f"| {r['title']} | {r['method']} | {r['total_bytes']/1e6:.3f} | {r['construct_seconds']:.3f} | {r['peak_rss_KiB']/1024**2:.3f} |")
(STUDY/'RESULTS.md').write_text('\n'.join(report)+'\n')
print(json.dumps(json.loads((STUDY/'validation.json').read_text()),ensure_ascii=False,indent=2))
