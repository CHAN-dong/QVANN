"""Independent read-only audit of the currently completed real measurements."""
import base64,csv,json,math,statistics,zlib
from pathlib import Path
import openpyxl
from work9_resource_matched import ROOT,STUDY,RESULTS,OUTPUT,SOURCE,DATA,TARGETS,KS,SIZES,CURVES,SIZE_ROWS,K_ROWS,METHODS,sha,write

checked={}
def verify_record(r):
 if r.get('status')!='passed' or r['kind']=='index':return
 out=Path(r['output']);key=(str(out),r.get('k',10),r.get('ef'),r.get('probes'),r.get('n_probe'))
 if key in checked:return
 rows=list(csv.DictReader((out/'per_query.csv').open()))
 proof='recall' in rows[0]
 if not proof and r['kind']!='ablation':rows=[x for x in rows if (int(x['k']),int(x['ef_search']),int(x['probes']))==(r['k'],r['ef'],r['probes'])]
 assert len(rows)==60,(out,len(rows))
 assert {(int(x['query']),int(x['repeat'])) for x in rows}=={(q,t) for q in range(100,120) for t in range(3)},out
 assert all(x['verified']=='1' for x in rows),out
 mean=lambda k:statistics.fmean(float(x[k]) for x in rows)
 expected=dict(recall=mean('recall' if proof else 'recall_at_k'),query_ms=mean('query_ms') if proof else mean('search_ms')+mean('vo_generation_ms'),verify_ms=mean('verify_ms' if proof else 'verification_ms'),vo_kb=mean('vo_bytes')/1000)
 for k,v in expected.items():assert math.isclose(r[k],v,rel_tol=1e-10,abs_tol=1e-8),(out,k,r[k],v)
 if proof:
  t=json.loads((out/'tamper.json').read_text());assert t['proof_byte_change_rejected'] and t['wrong_query_rejected']
 else:
  attacks=list(csv.DictReader((out/'tamper_tests.csv').open()));assert attacks and all(x['rejected']=='1' for x in attacks)
 checked[key]=len(rows)

w=openpyxl.load_workbook(OUTPUT)
records=[]
for row in w['资源匹配实测记录'].iter_rows(min_row=2):
 r=json.loads(row[11].value)
 if r.get('encoding')=='zlib+base64':r=json.loads(zlib.decompress(base64.b64decode(r['data'])))
 records.append(r)
cells=0
def match(ws,r,c,value):
 global cells
 got=ws.cell(r,c).value
 assert math.isclose(got,value,rel_tol=1e-10,abs_tol=1e-8),(ws.title,r,c,got,value)
 assert ws.cell(r,c).comment is not None
 cells+=1
for rec in records:
 ds=rec['dataset'];ws=w[DATA[ds]['sheet']];kind=rec['kind'];m=rec.get('method');idx=METHODS.index(m) if m in METHODS else 0
 if kind=='curve':
  for p in rec['points']:
   if p.get('status')!='passed':continue
   r=dict(p,kind=kind);verify_record(r)
   if not any(abs(t-p['target'])<1e-10 for t in TARGETS):continue
   col=6+next(i for i,t in enumerate(TARGETS) if abs(t-p['target'])<1e-10)
   for start,k in CURVES:match(ws,start+2*idx,col,p['recall']*100);match(ws,start+2*idx+1,col,p[k])
  continue
 if rec.get('status')!='passed':continue
 if kind=='index':
  match(ws,9+idx,9,rec['total_bytes']/1e6);match(ws,9+idx,10,rec['construct_seconds']);continue
 verify_record(rec)
 if kind=='overall':
  for col,k in [(5,'recall'),(6,'query_ms'),(7,'vo_kb'),(8,'verify_ms')]:match(ws,9+idx,col,rec[k]*(100 if k=='recall' else 1))
 elif kind=='size':
  for row,k in SIZE_ROWS:match(ws,row+idx,6+SIZES.index(rec['N']),rec[k])
 elif kind=='k':
  for row,k in K_ROWS:match(ws,row+idx,6+KS.index(rec['k']),rec[k])
 elif kind=='ablation':
  ah=next(c.row for row in ws for c in row if c.column==6 and c.row>120 and c.value=='Recall@10')
  for col,k in [(6,'recall'),(7,'query_ms'),(8,'vo_kb'),(9,'verify_ms')]:match(ws,ah+rec['variant'],col,rec[k])
for c in DATA.values():
 ws=w[c['sheet']]
 for row in ws:
  for cell in row:
   if isinstance(cell.value,str):assert '含预测' not in cell.value and '[预测]' not in cell.value,(ws.title,cell.coordinate)
assert sha(SOURCE)==sha(OUTPUT),'Original workbook differs: inspect user edits or workbook lock'
result=dict(status='PASS',scope='completed groups so far; not a completion claim for queued jobs',records=len(records),distinct_raw_groups=len(checked),honest_samples_checked=sum(checked.values()),numeric_cells_checked=cells,workbook_sha256=sha(OUTPUT),all_planned_jobs_finished=(STUDY/'RUN_FINISHED.json').exists())
write(STUDY/'independent-validation.json',result);print(json.dumps(result,indent=2))
