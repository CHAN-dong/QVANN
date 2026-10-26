"""Add four controlled quantization tables, with raw provenance and audits."""
import csv
import json
import math
import shutil
import statistics
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import work9_resource_matched as u

R=u.ROOT; S=R/'experiments/work9-quantization-20261023'; O=R/'results/work9-quantization-20261023'
SOURCE=R.parent/'work9.xlsx'; BACKUP=S/'work9_before_quantization.xlsx'
assert not (S/'APPLIED.json').exists(), 'Already applied'
assert u.read(S/'FINISHED.json')['passed']
assert u.sha(SOURCE)==u.sha(BACKUP), 'Workbook changed during measurement'
protocol=u.read(S/'protocol.json'); summary=u.read(S/'summary.json'); assert len(summary)==16
wb=openpyxl.load_workbook(BACKUP)
changes=[]
blue=PatternFill('solid',fgColor='305496'); light=PatternFill('solid',fgColor='E2EFDA')
line=Side(style='thin',color='B4C6E7')
for ds,cfg in protocol['datasets'].items():
    ws=wb[cfg['sheet']]; original_end=ws.max_row; title=original_end+4; header=title+3
    ws.merge_cells(start_row=title,start_column=4,end_row=title,end_column=11)
    ws.cell(title,4,'Impact of quantization configurations（固定 SQ4 top-1 graph traversal）')
    ws.cell(title,4).font=Font(bold=True,size=12,color='17365D');ws.row_dimensions[title].height=28
    ws.merge_cells(start_row=title+1,start_column=4,end_row=title+2,end_column=11)
    note=ws.cell(title+1,4,
        f"实测；N={cfg['manifest']['base']:,}, k=10, M={cfg['manifest']['M']}, efConstruction={cfg['manifest']['ef_construction']}, E={cfg['E']}。"
        '20条固定查询（100–119）×3次重复，单线程；只改变 refinement。'
        'Q=遍历+候选精确排序+真实VO生成；V=解码+认证+重放+渐进检查；KB=1000 bytes。'
        'Recall用于核对四种配置答案一致；结果与其他表可能因查询样本量/本次运行环境不同而有差异。')
    note.alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[title+1].height=30;ws.row_dimensions[title+2].height=24
    headings=['Configuration','Recall@10','Q (ms)','VO (KB)','V (ms)','Search (ms)','VO gen (ms)','Samples']
    for col,label in enumerate(headings,4):
        c=ws.cell(header,col,label);c.font=Font(bold=True,color='FFFFFF');c.fill=blue;c.alignment=Alignment(wrap_text=True,horizontal='center',vertical='center')
    ws.row_dimensions[header].height=30
    for i,mode in enumerate(protocol['configurations'],1):
        r=next(r for r in summary if r['dataset']==ds and r['configuration']==mode);row=header+i
        values=[mode,r['recall'],r['query_ms'],r['vo_kb'],r['verify_ms'],r['search_ms'],r['vo_generation_ms'],r['samples']]
        for col,value in enumerate(values,4):
            c=ws.cell(row,col,value);c.alignment=Alignment(horizontal='left' if col==4 else 'center',vertical='center',wrap_text=True);c.border=Border(bottom=line)
            c.font=Font(bold=(mode=='SQ4--SQ6--raw'))
            if mode=='SQ4--SQ6--raw':c.fill=light
            if col==5:c.number_format='0.0%'
            elif col in range(6,11):c.number_format='0.000'
            c.comment=Comment('2026-10-23真实测量；未进行预测或性能系数放大。\n'
                f"固定SQ4 top-1 traversal, N={r['N']}, D={r['D']}, M={r['M']}, efConstruction={r['ef_construction']}, E={r['E']}, k=10。\n"
                f"refinement={mode}；查询IDs100..119，每条重复3次，共60样本（20独立查询）。\n"
                '同一查询四种配置的遍历指纹和答案完全一致；SQ4-SQ6-raw的VO字节与原协议完全一致。\n'
                'Q包括服务器检索、精确排序、VO序列化及多重证明生成；V包括客户端VO解码、认证、重放与排名检查；不含网络。\n'
                '固定SQ4、嵌套SQ6/SQ8及增量比特；BFS认证位置；不改变原始维度或图索引。\n'
                f"源文件={r['source']}\nSHA256={r['csv_sha256']}\n"
                f"Q std={r['query_ms_stdev']:.9g}ms；V std={r['verify_ms_stdev']:.9g}ms。",'Codex')
            changes.append(dict(sheet=ws.title,cell=c.coordinate,value=value,dataset=ds,configuration=mode))
        ws.row_dimensions[row].height=32
    cfg['new_table']=dict(title_row=title,header_row=header,first_data_row=header+1,last_data_row=header+4,original_last_row=original_end)

notes=wb.create_sheet('量化配置实验说明')
for row in [
    ['项目','说明'],
    ['日期','2026-10-23'],
    ['实验目的','固定SQ4 top-1 graph traversal，仅比较 SQ4--raw、SQ4--SQ6--raw、SQ4--SQ8--raw、SQ4--SQ6--SQ8--raw 的refinement成本。'],
    ['样本','每个数据集每种配置20条固定查询（100..119）×3次重复；单线程；共16组、960次真实出证与验证。每组仅20条独立查询。'],
    ['数据/参数','沿用四个数据集当前Ours默认图：SIFT 1M/M8/efc100/E128；NYTimes 289761/M16/efc200/E256；GloVe 1183514/M16/efc200/E1536；DEEP 1M/M16/efc200/E64。'],
    ['遍历控制','同一数据集固定图文件、入口、SQ4码、唯一扩展节点预算E、top-1遍历及候选精确排序；逐查询比较trace SHA256与答案，四配置完全一致。'],
    ['Refinement','原有uniform SQ4/SQ6不变，新增嵌套uniform SQ8。4→6传2 bits/dim，4→8传4 bits/dim，6→8传2 bits/dim；通过下界筛选后才传raw float32。'],
    ['认证','沿用direct-content expansion-centric Merkle、BFS refinement ordering、canonical shared multiproof；各增量层有独立域和Merkle树。'],
    ['Q','SQ4遍历+所有发现候选的精确距离排序+完整VO生成与序列化，含查表准备；无网络传输。'],
    ['V','VO解码+Merkle认证+SQ4重放+渐进下界/候选top-k检查；客户端仅有可信根和量化元数据，无服务器图/全量码数组。'],
    ['VO','实际二进制VO长度；包含header、query绑定、答案、各层记录及共享认证路径；KB=1000bytes。'],
    ['计时控制','按query_index+repeat轮换四配置执行次序；每次计时前清除上一查询LUT标记，包含本次查表准备成本。'],
    ['校验','每个数据集60次SQ6默认配置的VO与原实现字节级一致，并由原验证器验证；每配置执行篡改、截断、错误根/查询/预算拒绝检查。'],
    ['Recall','以对应规模精确穷举top10为真值；四配置召回率应一致。本次不为任何配置单独调参。SIFT使用20条查询，因此不要求等于原100查询均值95.4%。'],
    ['原表保护','新表追加到四个数据集原表下方；不覆盖既有实验数据；不应用ANNProof的性能乘数。'],
    ['原文件备份',str(BACKUP)],
    ['复现脚本',str(R/'scripts/run_work9_quantization_20261023.py')],
    ['原始样本',str(O)],
    ['论文表格',str(S/'quantization_sift_deep.tex')],
]:notes.append(row)
details=wb.create_sheet('量化配置实测记录')
keys=['sheet','configuration','N','D','M','ef_construction','E','samples','independent_queries','recall','query_ms','vo_kb','verify_ms','search_ms','vo_generation_ms','raw_vectors','stage1_records','stage2_records','query_ms_stdev','verify_ms_stdev','source','csv_sha256','status']
details.append(keys)
for r in summary:
    details.append([r[k] for k in keys]);details.cell(details.max_row,10).number_format='0.0%'
for ws in [notes,details]:
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=blue
    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=22
    for row in ws:
        for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
notes.column_dimensions['B'].width=110

def latex_pair(datasets):
    a,b=datasets
    titles={'sift1m':r'\textsf{SIFT}','deep1m':r'\textsf{DEEP}','nytimes256':r'\textsf{NYTimes}','glove100':r'\textsf{GloVe}'}
    lines=[r'\begin{table}[!t]',r'\centering',r'\caption{Impact of quantization configurations with fixed SQ4 top-1 graph traversal.}',r'\label{tab:quantization'+('' if a=='sift1m' else '_text')+'}',r'\footnotesize',r'\setlength{\tabcolsep}{2.5pt}',r'\renewcommand{\arraystretch}{1.15}',r'\begin{tabular}{lrrrrrr}',r'\toprule',r'\multirow{2}{*}{Configuration}',f'& \\multicolumn{{3}}{{c}}{{{titles[a]}}} & \\multicolumn{{3}}{{c}}{{{titles[b]}}} '+r'\\',r'\cmidrule(lr){2-4}\cmidrule(l){5-7}',r'& Q (ms) & VO (KB) & V (ms) & Q (ms) & VO (KB) & V (ms) \\',r'\midrule']
    for mode in protocol['configurations']:
        label=r'\textbf{'+mode+'}' if mode=='SQ4--SQ6--raw' else mode
        values=[]
        for ds in [a,b]:
            rr=next(r for r in summary if r['dataset']==ds and r['configuration']==mode)
            values += [f"{rr[k]:.3f}" for k in ['query_ms','vo_kb','verify_ms']]
        lines.append(label+' & '+' & '.join(values)+r' \\')
    return '\n'.join(lines+[r'\bottomrule',r'\end{tabular}',r'\end{table}',''])
(S/'quantization_sift_deep.tex').write_text(latex_pair(('sift1m','deep1m')))
(S/'quantization_nytimes_glove.tex').write_text(latex_pair(('nytimes256','glove100')))
with (S/'summary.csv').open('w',newline='') as f:
    writer=csv.DictWriter(f,fieldnames=keys);writer.writeheader();writer.writerows({k:r[k] for k in keys} for r in summary)
review=S/'work9_quantization_updated.xlsx';wb.save(review)

# Independent readback: preserve all original cells including their comments/formats.
old=openpyxl.load_workbook(BACKUP); new=openpyxl.load_workbook(review);preserved=0
for ws in old:
    ns=new[ws.title]
    assert set(map(str,ws.merged_cells.ranges)).issubset(set(map(str,ns.merged_cells.ranges)))
    for row in ws:
        for c in row:
            nc=ns[c.coordinate]
            assert c.value==nc.value,(ws.title,c.coordinate,'old value modified')
            assert (c.comment.text if c.comment else None)==(nc.comment.text if nc.comment else None),(ws.title,c.coordinate,'old comment modified')
            assert c.number_format==nc.number_format
            assert copy(c.font)==copy(nc.font) and copy(c.fill)==copy(nc.fill) and copy(c.alignment)==copy(nc.alignment) and copy(c.border)==copy(nc.border)
            preserved+=1
total=negative=paired=0
for ds,cfg in protocol['datasets'].items():
    assert u.sha(cfg['graph'])==cfg['graph_sha256']
    rows=list(csv.DictReader((O/ds/'samples.csv').open())); assert len(rows)==240
    gt={int(r['query']):set(map(int,r['answer'].split())) for r in csv.DictReader((O/ds/'ground_truth.csv').open())}
    for mode in protocol['configurations']:
        rr=[r for r in rows if r['configuration']==mode]
        assert {(int(r['query']),int(r['repeat'])) for r in rr}=={(q,t) for q in range(100,120) for t in range(3)}
        for r in rr:
            ans=list(map(int,r['answer'].split()));assert len(set(ans))==len(ans)==10
            assert math.isclose(float(r['recall']),len(set(ans)&gt[int(r['query'])])/10,abs_tol=1e-12) and r['valid']=='1'
            assert math.isclose(float(r['query_ms']),float(r['search_ms'])+float(r['vo_generation_ms']),rel_tol=1e-12)
            total+=1
        idx=protocol['configurations'].index(mode); row=cfg['new_table']['first_data_row']+idx;ws=new[cfg['sheet']]
        for col,metric in [(5,'recall'),(6,'query_ms'),(7,'vo_bytes'),(8,'verify_ms'),(9,'search_ms'),(10,'vo_generation_ms')]:
            expected=statistics.fmean(float(r[metric]) for r in rr)/(1000 if metric=='vo_bytes' else 1)
            assert math.isclose(ws.cell(row,col).value,expected,rel_tol=1e-12,abs_tol=1e-12),(ds,mode,metric)
    for q in range(100,120):
        for rep in range(3):
            group=[r for r in rows if int(r['query'])==q and int(r['repeat'])==rep]
            assert len(group)==4 and len({r['trace_sha256'] for r in group})==len({r['answer'] for r in group})==1
            raw={r['configuration']:int(r['raw_vectors']) for r in group}
            assert raw['SQ4--SQ6--raw']<=raw['SQ4--raw'] and raw['SQ4--SQ8--raw']<=raw['SQ4--SQ6--raw']
            assert raw['SQ4--SQ8--raw']==raw['SQ4--SQ6--SQ8--raw']
            paired+=1
    attack=list(csv.DictReader((O/ds/'tamper.csv').open()));assert all(r['rejected']=='1' for r in attack);negative+=len(attack)
assert total==960 and paired==240
audit=dict(passed=True,measured_operating_points=16,authenticated_queries=total,paired_query_groups=paired,
           rejected_tamper_cases=negative,original_protocol_wire_checks=240,old_cells_preserved=preserved,
           workbook_sha256=u.sha(review),backup_sha256=u.sha(BACKUP),new_tables={x:y['new_table'] for x,y in protocol['datasets'].items()})
u.write(S/'audit.json',audit);u.write(S/'cell-changes.json',changes)
assert u.sha(SOURCE)==u.sha(BACKUP)
shutil.copy2(review,SOURCE);assert u.sha(SOURCE)==audit['workbook_sha256']
u.write(S/'APPLIED.json',audit)
print(json.dumps(audit,ensure_ascii=False,indent=2))
