from pathlib import Path
import openpyxl
source = Path(__file__).resolve().parents[2] / 'work9.xlsx'
workbook = openpyxl.load_workbook(source)
for sheet in workbook:
    print(sheet.title, sheet.max_row, sheet.max_column)
    for row in sheet:
        values = [(c.coordinate,c.value) for c in row if c.value is not None]
        if values:
            print(values)
    print('MERGES', list(map(str,sheet.merged_cells.ranges)))
