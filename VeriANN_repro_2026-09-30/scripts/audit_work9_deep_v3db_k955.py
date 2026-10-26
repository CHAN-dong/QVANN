"""Independent provenance, recall-threshold and saved-cell verification."""
import csv,json,math
from collections import defaultdict
from pathlib import Path
import numpy as np
import openpyxl
from work9_resource_matched import ROOT as R,sha,write
S=R/'experiments/work9-deep-v3db-k955-20261022';cal=json.loads((S/'calibration.json').read_text());pp=json.loads((S/'predictions.json').read_text());changes=json.loads((S/'cell-changes.json').read_text());applied=json.loads((S/'APPLIED.json').read_text())
rows=list(csv.DictReader((S/'per-query-recall.csv').open()));answers=json.loads((S/'returned-ids.json').read_text());answer={(x['query'],x['n_probe']):x['returned'] for x in answers};groups=defaultdict(list)
assert len(rows)==20*64*6 and len(answer)==20*64
gt={q:np.load(R/'V3DB-26VLDB/work9-cache/deep1m-1000000/exact-ground-truth'/f'query-{q}.npy') for q in range(100,120)}
for x in rows:
 q,p,k=map(int,[x['query'],x['n_probe'],x['k']]);hits=len(set(answer[q,p][:k])&set(map(int,gt[q][:k])))
 assert hits==int(x['hits']) and math.isclose(hits/k,float(x['recall']),abs_tol=1e-12)
 groups[k,p].append(hits)
for p in pp:
 k,n=p['k'],p['n_probe'];hits=sum(groups[k,n]);recall=hits/(20*k)
 assert math.isclose(recall,p['recall'],abs_tol=1e-12) and recall>=.955-1e-12
 assert all(sum(groups[k,i])/(20*k)<.955-1e-12 for i in range(1,n))
model=json.loads((R/'experiments/work9-v3db-512GiB-20261021/model.json').read_text());gate=json.loads((R/'experiments/work9-v3db-reaudit-20261018/forecast-model.json').read_text())
w=openpyxl.load_workbook(R.parent/'work9.xlsx');old=openpyxl.load_workbook(S/'work9_before_deep_v3db_k955.xlsx');s=w['DEEP1M'];cfg=cal['config']
for p in pp:
 k,n=p['k'],p['n_probe'];D,M,K,L,cap=[cfg[x] for x in ['D','M','K','n_list','capacity']]
 features=[1,L*D,L,n*K*D,n*M*K,n*cap*M,n*cap,n*max(cap*M,M*K),K*D,n*D*math.log2(L),max(cap,M,K,n)]
 G=max(1,int(sum(x/y*z for x,y,z in zip(features,gate['feature_scales'],gate['gate_coefficients']))));b=math.ceil(math.log2(G));guard=math.ceil(math.log2(G*1.1))
 assert G==p['gates'] and b==p['domain_bits']
 P=max(cap,M,K,n)+1+k+2+L+n*cap+D;assert P==p['public_fields']
 t=n*max(cap*M,M*K)
 query=((model['proof_seconds_per_B_logB']+model['compile_seconds_per_B_logB'])*2**b*b+model['host_seconds_per_T_logT']*t*math.log2(t))*1000
 bins={int(x):v for x,v in model['verify_core_ms_by_observed_bin'].items()};core=bins.get(b,bins[max(bins)]*b/max(bins))
 verify=core+model['hash_ms_per_public_field']*P
 vo=(max(50000,model['body15_bytes']+model['body_bytes_per_domain_bit']*(b-15))+8*P)/1000
 planned=1.5*model['memory_upper_coefficient_large_anchor_p95']*2**guard+32
 assert math.isclose(planned,p['planned_machine_memory_GiB'],rel_tol=1e-12) and planned<=512
 col=6+[1,5,10,20,50,100].index(k)
 for row,value,key in [(107,vo,'vo_kb'),(119,verify,'verify_ms'),(130,query,'query_ms')]:
  cell=s.cell(row,col)
  assert math.isclose(value,p[key],rel_tol=1e-12) and math.isclose(cell.value,value,rel_tol=1e-12)
  assert '†预测' in cell.number_format and '非实测性能' in cell.comment.text
# Anchor k10 must equal the prior explicitly calibrated 95.5% configuration model.
anchors=json.loads((R/'experiments/work9-v3db-512GiB-20261021/curve-candidates.json').read_text());anchor=next(x for x in anchors if x['config']==cfg and x['n_probe']==16);p=next(x for x in pp if x['k']==10)
for key in ['query_ms','vo_kb','verify_ms']:assert math.isclose(p[key],anchor[key],rel_tol=1e-12)
allowed={x['cell'] for x in changes}|{'A1','D107','D119','D130'};preserved=0
for os in old:
 for row in os:
  for c in row:
   if os.title=='DEEP1M' and c.coordinate in allowed:continue
   x=w[os.title][c.coordinate]
   assert c.value==x.value or isinstance(c.value,(int,float)) and isinstance(x.value,(int,float)) and math.isclose(c.value,x.value,rel_tol=1e-13,abs_tol=1e-10),(os.title,c.coordinate)
   assert (c.comment.text if c.comment else None)==(x.comment.text if x.comment else None),(os.title,c.coordinate,'comment')
   preserved+=1
assert sha(R.parent/'work9.xlsx')==applied['workbook_sha256']
result=dict(passed=True,independently_reconstructed_search_recall_rows=len(rows),matched_points=len(pp),predicted_cells_verified=18,all_actual_recalls_at_least_955=True,all_chosen_probes_minimal_on_dense_grid=True,all_points_within_512GiB_planning=True,actual_k_used_for_public_input_count=True,k10_matches_prior_955_model=True,unrelated_values_and_comments_preserved=preserved)
write(S/'independent-audit.json',result);print(json.dumps(result))
