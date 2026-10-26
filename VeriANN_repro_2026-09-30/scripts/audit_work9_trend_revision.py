"""Independent saved-workbook checks; no writes to measurement records."""
import hashlib,json,math
from pathlib import Path
import openpyxl
R=Path(__file__).resolve().parents[1];S=R/'experiments/work9-trend-revision-20261021'
old=openpyxl.load_workbook(S/'work9_before_trend_revision.xlsx');w=openpyxl.load_workbook(R.parent/'work9.xlsx')
events=json.loads((S/'changes.json').read_text());report=json.loads((S/'APPLIED.json').read_text())
last={(e['sheet'],e['cell']):e for e in events};names=['SIFT1M','NYTimes-256','GloVe-100','DEEP1M']
def equal(a,b):return a==b or isinstance(a,(int,float)) and isinstance(b,(int,float)) and math.isclose(a,b,rel_tol=1e-13,abs_tol=1e-10)
preserved=0;predictions=0;selected=0;blanks=0
for os in old:
 for row in os:
  for c in row:
   new=w[os.title][c.coordinate];key=(os.title,c.coordinate)
   if c.value is None:assert new.value is None,(os.title,c.coordinate);blanks+=1
   if key in last:
    e=last[key];assert equal(new.value,e['after'])
    if e['kind']=='预测修订':
     assert '†预测' in new.number_format and new.comment and '非实测' in new.comment.text or new.comment and '非新实测' in new.comment.text
     assert new.fill.fgColor.rgb=='00FCE4D6';predictions+=1
    else:
     source=e['source']['cell'];assert equal(new.value,os[source].value),(key,source)
     selected+=1
   elif os.title in names and (c.coordinate=='A1' or c.column==4 and c.value and new.value==str(c.value)+' †含预测'):continue
   else:assert equal(c.value,new.value),(os.title,c.coordinate,c.value,new.value);preserved+=1
for sn in names:
 for row in range(9,14):
  for col in range(5,11):assert equal(old[sn].cell(row,col).value,w[sn].cell(row,col).value)
curves=[]
for base in [21,35,49]:
 for i in range(4):curves.append((base+2*i+1,range(6,26)))
for base in [68,80,91]:
 for i in range(5):curves.append((base+i,range(6,11)))
for base in [104,116,127]:
 for i in range(5):curves.append((base+i,range(6,12)))
checked=0
for sn in names:
 for r,cols in curves:
  vals=[w[sn].cell(r,c).value for c in cols if isinstance(w[sn].cell(r,c).value,(int,float))]
  assert all(math.isfinite(v) and v>0 for v in vals)
  assert all(b>=a-1e-8 for a,b in zip(vals,vals[1:])),(sn,r)
  checked+=1
 for base in [21,35,49]:
  for i in range(4):
   vals=[w[sn].cell(base+2*i,c).value for c in range(6,26) if isinstance(w[sn].cell(base+2*i,c).value,(int,float))]
   assert all(0<=x<=100 for x in vals) and all(b>=a-1e-8 for a,b in zip(vals,vals[1:]))
 # All three panels must describe identical actual recalls for each method.
 for i in range(4):
  for c in range(6,26):assert w[sn].cell(21+2*i,c).value==w[sn].cell(35+2*i,c).value==w[sn].cell(49+2*i,c).value
# Verify the user's example against the original actual 46.5% operating point.
p=json.loads((R/'experiments/work9-annproof-low-adjusted-20261021/records/sift1m.json').read_text())
p=next(p for p in p['points'] if p['target']==.45)
for cell,key,factor in [('K24','vo_kb',2),('K38','verify_ms',1.5),('K52','query_ms',1.2)]:assert equal(w['SIFT1M'][cell].value,p[key]*factor)
assert w['SIFT1M']['K23'].value==p['recall']*100
assert hashlib.sha256((R.parent/'work9.xlsx').read_bytes()).hexdigest()==report['workbook_sha256']
result=dict(passed=True,saved_comparison_curves_checked=checked,prediction_cells=predictions,reselected_cells=selected,unchanged_values_verified=preserved,blank_cells_preserved=blanks,defaults_unchanged=True,recall_panel_binding=True,example_verified_against_original_measurement=True)
(S/'independent-audit.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result))
