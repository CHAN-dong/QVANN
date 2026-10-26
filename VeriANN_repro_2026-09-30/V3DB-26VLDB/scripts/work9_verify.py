"""Independent workbook/evidence audit; incomplete coverage is never called complete."""
import csv
import json
from pathlib import Path
import statistics
import openpyxl
from work9_calibrate import ROOT, STUDY, DATA
from work9_fill import SOURCE, BACKUP, row_map, sha


def main():
    if (ROOT/'experiments/work9-nytimes-20261017/replacement.json').exists():
        import sys
        sys.path.insert(0,str(ROOT/'scripts'))
        from verify_work9_nytimes import main as replacement_audit
        replacement_audit()
        return
    old=openpyxl.load_workbook(BACKUP)
    wb=openpyxl.load_workbook(SOURCE)
    preserved=0
    for ws in old:
        target=wb[ws.title]
        main_sheet=ws.title in [x[0] for x in DATA.values()]
        for row in ws:
            for c in row:
                if isinstance(c,openpyxl.cell.cell.MergedCell):continue
                new=target.cell(row_map(c.row) if main_sheet else c.row,c.column)
                assert new.value==c.value,(ws.title,c.coordinate,new.coordinate)
                assert new.number_format==c.number_format,(ws.title,c.coordinate,'format')
                assert new._style==c._style,(ws.title,c.coordinate,'style')
                if c.value is not None:preserved+=1
        assert len(target._charts)==len(ws._charts),(ws.title,'old charts changed')
        assert len(target.conditional_formatting)==len(ws.conditional_formatting),(ws.title,'conditional formatting')
        assert len(target.data_validations.dataValidation)==len(ws.data_validations.dataValidation),(ws.title,'data validation')
    audit=json.loads((STUDY/'workbook_audit.json').read_text())
    for item in audit:
        evidence=item.get('source')
        if evidence and evidence.get('status')=='calibration_only':
            calibration=json.loads(Path(evidence['path']).read_text())
            point=next(p for p in calibration['points'] if p['k']==evidence['point']['k'] and p['n_probe']==evidence['point']['n_probe'])
            assert abs(item['value']-point['recall']*evidence['display_scale'])<1e-9
            assert abs(wb[item['sheet']][item['cell']].value-item['value'])<1e-9
    pending=[x for x in audit if x['value'] in ['待测','待校准']]
    limits=[x for x in audit if x['value'] in ['资源受限','超预算未测','内存不足(50GiB上限)']]
    unreachable=[x for x in audit if x['value']=='未达目标']
    passed=[];failed=[]
    for file in (ROOT/'results/work9-v3db-20261016').glob('*/result.json'):
        result=json.loads(file.read_text())
        if result['status']!='passed':failed.append(result);continue
        rows=list(csv.DictReader((file.parent/'per_query.csv').open()))
        assert len(rows)==60 and result['query_count']==20 and result['repeats']==3
        assert {(int(r['query']),int(r['repeat'])) for r in rows}=={(q,r) for q in range(100,120) for r in range(3)}
        assert all(r['verified']=='1' for r in rows)
        for field in ['recall','query_ms','verify_ms']:
            mean=statistics.fmean(float(r[field]) for r in rows)
            assert abs(mean-result[field])<max(1e-10,abs(mean)*1e-12)
        assert abs(statistics.fmean(int(r['vo_bytes'])/1000 for r in rows)-result['vo_kb'])<1e-8
        assert len((file.parent/'query0.proof').read_bytes())==int(rows[0]['vo_bytes'])
        tamper=json.loads((file.parent/'tamper.json').read_text())
        assert all(tamper.values())
        passed.append(result['job'])
    report=dict(original_nonempty_cells_preserved=preserved,workbook_sha256=sha(SOURCE),
                new_cells=len(audit),pending_cells=len(pending),resource_limited_cells=len(limits),
                unreachable_target_cells=len(unreachable),
                passed_20x3_workpoints=passed,resource_limited_workpoints=len(failed),
                all_requested_numeric_measurements_complete=not pending and not limits and not unreachable,
                planned_runner_finished=(STUDY/'RUN_FINISHED.json').exists(),
                note='Resource/target limitations are explicit statuses, not numerical measurements.')
    (STUDY/'final_validation.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report,indent=2))


if __name__=='__main__':main()
