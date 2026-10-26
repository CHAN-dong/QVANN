"""Explicitly labelled conditional forecasts, never synthetic benchmark records."""
import argparse
from copy import copy
import csv
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import statistics as stats
import subprocess
import sys
import numpy as np
from scipy.optimize import nnls
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font,PatternFill
from openpyxl.chart import ScatterChart,Series,Reference

ROOT=Path(__file__).resolve().parents[2]
STUDY=ROOT/'experiments/work9-v3db-estimates-20261017'
CACHE=ROOT/'V3DB-26VLDB/work9-cache'
OLD=ROOT/'experiments/work9-v3db-20261016'
RESULTS=ROOT/'results/work9-v3db-20261016'
SOURCE=ROOT.parent/'work9.xlsx'
OUTPUT=ROOT.parent/'work9_V3DB预测补充.xlsx'
NATIVE='/root/.cache/v3db-target/release/work9-inspect-verifier'
DATA={'sift1m':('SIFT1M',1000000,128,1024),'nytimes256':('NYTimes-256',289761,256,512),
      'glove100':('GloVe-100',1183514,100,1024),'deep1m':('DEEP1M',1000000,96,1024)}
TARGETS=[.60,.65,.70,.75,.80,.85,.90,.95,.97,.98,.99]
KVALS=[1,5,10,20,50,100]
INSERTS=[(11,1),(24,2),(34,2),(44,2),(56,1),(66,1),(75,1),(86,1),(96,1),(105,1)]
FEATURES=['constant','L*D','L','p*K*D','p*M*K','p*n*M','p*n','p*max(n*M,M*K)','K*D','p*D*log2(L)','max(n,M,K,p)']

def rm(r):return r+sum(n for a,n in INSERTS if r>a)
def sha(p):
    with Path(p).open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()
def write(p,obj):
    p.write_text(json.dumps(obj,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
def signature(c,p,k):return tuple(c[x] for x in ['D','M','K','n_list','capacity'])+(p,k)
def features(c,p):
    D,M,K,L,n=[c[x] for x in ['D','M','K','n_list','capacity']]
    return np.array([1,L*D,L,p*K*D,p*M*K,p*n*M,p*n,p*max(n*M,M*K),K*D,p*D*math.log2(L),max(n,M,K,p)],dtype=float)
def pi_count(c,p,k):return max(c['capacity'],c['M'],c['K'],p)+1+k+2+c['n_list']+p*c['capacity']+c['D']

def references():
    refs=[];paths=[]
    for f in sorted(RESULTS.glob('*/per_query.csv')):
        match=re.match(r'(.+)-p(\d+)-k(\d+)-q20-r3-t1$',f.parent.name)
        if not match:continue
        model,p,k=match.groups();meta=CACHE/model/'model.json'
        if not meta.exists():continue
        c=json.loads(meta.read_text())['config']
        if c['dataset'] not in DATA:continue
        rows=list(csv.DictReader(f.open()))
        if not rows:continue
        assert all(r['verified']=='1' for r in rows)
        pairs={(int(r['query']),int(r['repeat'])) for r in rows}
        assert len(pairs)==len(rows) and pairs<={(q,t) for q in range(100,120) for t in range(3)}
        result=json.loads((f.parent/'result.json').read_text()) if (f.parent/'result.json').exists() else {}
        complete=len(rows)==60 and result.get('status')=='passed'
        vk=f.parent/'trusted_verifier.bin';assert vk.exists()
        paths.append(str(vk))
        mean=lambda name:stats.fmean(float(r[name]) for r in rows)
        refs.append(dict(job=f.parent.name,config=c,n_probe=int(p),k=int(k),samples=len(rows),
            full_20x3_passed=complete,source=str(f),source_sha256=sha(f),
            gates=int(rows[0]['gates']),query_ms=mean('query_ms'),verify_ms=mean('verify_ms'),recall_mean_of_available_samples=mean('recall'),
            vo_kb=mean('vo_bytes')/1000,build_seconds=mean('build_seconds'),prove_seconds=mean('prove_seconds'),
            search_ms=mean('search_ms'),peak_rss_GiB=max(int(r['rss_bytes']) for r in rows)/2**30))
    info=[json.loads(x) for x in subprocess.check_output([NATIVE,*paths],text=True).splitlines()]
    bypath={x['path']:x for x in info}
    for r in refs:
        m=bypath[str(Path(r['source']).parent/'trusted_verifier.bin')]
        r['degree_bits']=m['degree_bits'];r['domain']=2**m['degree_bits'];r['public_fields']=m['public_inputs']
        assert r['public_fields']==pi_count(r['config'],r['n_probe'],r['k'])
        r['proof_body_bytes']=r['vo_kb']*1000-8*r['public_fields']
    write(STUDY/'single-thread-reference-samples.json',refs)
    return refs

def gate_data(refs):
    data=json.loads((STUDY/'gate-observations.json').read_text())
    for r in refs:data.append(dict(config=dict(r['config'],n_probe=r['n_probe'],k=r['k']),gates=r['gates'],role='observed formal circuit gate count'))
    for r in json.loads((OLD/'sift-high-precision-gates.json').read_text()):
        data.append(dict(config=dict(D=128,M=r['M'],K=r['K'],n_list=r['n_list'],n_probe=r['n_probe'],capacity=r['capacity'],k=10),gates=r['gates_before_build'],role='retained exact static count'))
    # Attached paper Table9: static constraints, not author hardware timings.
    paper={128:[182225,192846,213475,254433,336214,499721,826743],
           960:[270229,342232,485623,772104,1344931,2490531,4781738]}
    for d,values in paper.items():
        for L,g in zip([128,256,512,1024,2048,4096,8192],values):
            data.append(dict(config=dict(D=d,M=8,K=256,n_list=L,n_probe=L//128,capacity=2**21//L,k=10),gates=g,role='attached paper Table9 exact G; author timing NOT imported'))
    # This local gate count also matches the attached paper Table8.
    data.append(dict(config=dict(D=128,M=8,K=16,n_list=256,n_probe=16,capacity=32,k=64),gates=21198,role='local/paper basic multiset exact G'))
    unique={signature(x['config'],x['config']['n_probe'],x['config']['k']):x for x in data}
    data=list(unique.values());write(STUDY/'gate-model-training.json',data)
    return data

class Model:
    def __init__(self,refs):
        self.refs=refs;self.data=gate_data(refs)
        X=np.stack([features(x['config'],x['config']['n_probe']) for x in self.data]);y=np.array([x['gates'] for x in self.data],dtype=float)
        self.scale=np.maximum(np.max(X,axis=0),1)
        Z=X/self.scale
        self.coef,_=nnls(Z/y[:,None],np.ones(len(y)),maxiter=10000)
        fitted=Z@self.coef;errors=[]
        for i in range(len(y)):
            mask=np.arange(len(y))!=i
            coef,_=nnls(Z[mask]/y[mask,None],np.ones(mask.sum()),maxiter=10000)
            errors.append(abs(Z[i]@coef-y[i])/y[i])
        self.gate_error=max(.10,float(np.quantile(errors,.95))*2)
        self.exact={signature(x['config'],x['config']['n_probe'],10):x['gates'] for x in self.data}
        # Primary time anchors are local single-thread data only, never 256-core paper seconds.
        self.aprove=stats.median(r['prove_seconds']/(r['domain']*r['degree_bits']) for r in refs)
        self.abuild=stats.median(r['build_seconds']/(r['domain']*r['degree_bits']) for r in refs)
        self.ahost=stats.median(max(0,r['query_ms']/1000-r['build_seconds']-r['prove_seconds'])/self.work(r['config'],r['n_probe']) for r in refs)
        self.mem=stats.median(r['peak_rss_GiB']/r['domain'] for r in refs)
        # Byte anchor at b=15 is the local basic proof, same Plonky2 serialization.
        basic=json.loads((ROOT/'V3DB-26VLDB/results/basic-multiset.json').read_text())
        bc=dict(D=128,M=8,K=16,n_list=256,capacity=32)
        self.body15=basic['metrics']['compressed_proof_bytes']-8*pi_count(bc,16,64)
        self.body_slope=stats.median((r['proof_body_bytes']-self.body15)/(r['degree_bits']-15) for r in refs)
        self.hashbench=[json.loads(x) for x in subprocess.check_output([NATIVE,'--hash-bench'],text=True).splitlines()]
        self.hash_ms_per_field=stats.median(x['median_ms']/x['fields'] for x in self.hashbench)
        self.verify_core=stats.median(max(1,r['verify_ms']-1.5*self.hash_ms_per_field*r['public_fields'])/(r['degree_bits']/20) for r in refs)
        self.report=dict(version=1,scope='conditional resource-sufficient single-thread forecast; NOT measured on the 50GiB-limited workstation',
            gate_features=FEATURES,gate_coefficients=self.coef.tolist(),feature_scales=self.scale.tolist(),
            exact_gate_observations=len(y),gate_fit_median_relative_error=float(np.median(abs(fitted/y-1))),
            gate_leave_one_out_median_relative_error=float(np.median(errors)),gate_leave_one_out_p95_relative_error=float(np.quantile(errors,.95)),
            gate_assumed_relative_sensitivity=self.gate_error,
            prove_seconds_per_domain_log2_domain=self.aprove,build_seconds_per_domain_log2_domain=self.abuild,
            host_seconds_per_relation_log2_relation=self.ahost,memory_GiB_per_domain_row=self.mem,
            proof_body_bytes_at_b15=self.body15,proof_body_bytes_per_extra_domain_bit=self.body_slope,
            public_input_hash_microbenchmark=self.hashbench,verify_core_ms_at_b20=self.verify_core,
            paper_time_formula_for_reference_only='Tprove=1.364e-6*GB*log2(GB)+0.261 seconds, attached Table11; do not copy to local single-thread columns',
            limitations=['Only two full 20x3 points and one partial point support local timing. Main observed domains are the same bin; larger domains are theoretical extrapolation, not empirically validated.',
                        'Bands are declared scenario sensitivity ranges, NOT 95% confidence intervals.',
                        'No actual running-time prediction exists on a machine that cannot allocate the circuit; numbers assume sufficient RAM without paging and similar single-core speed.',
                        'Recall comes from existing search calibration only. It is not predicted upward and unreachable targets remain unfilled.',
                        'Size cases without an IVF-PQ index use explicitly assumed parameters, not calibrated equal-recall claims.'])
        write(STUDY/'forecast-model.json',self.report)
        self.cache={}
    @staticmethod
    def work(c,p):
        t=max(p*c['capacity']*c['M'],p*c['M']*c['K'])
        return t*math.log2(max(t,2))
    def body(self,b):return max(50000,self.body15+self.body_slope*(b-15))
    def predict(self,c,p,k,assumed=False):
        key=signature(c,p,k)+(c['N'],c.get('source'),assumed)
        if key in self.cache:return self.cache[key]
        observed=next((r for r in self.refs if signature(r['config'],r['n_probe'],r['k'])==signature(c,p,k) and r['config']['N']==c['N'] and r['config'].get('source')==c.get('source')),None)
        exact=self.exact.get(signature(c,p,10))
        G=exact or max(1,int(features(c,p)/self.scale@self.coef))
        if observed:G=observed['gates']
        b=observed['degree_bits'] if observed else max(1,math.ceil(math.log2(G)))
        lo=b if observed else max(1,math.ceil(math.log2(max(1,G*(1-self.gate_error)))))
        hi=b if observed else math.ceil(math.log2(G*(1+self.gate_error)))
        # Padding overhead can cross a bin boundary; upper scenarios allow one extra bin.
        if not observed:hi=max(hi,b+1)
        B=2**b;P=pi_count(c,p,k);memory=self.mem*B
        query=(self.aprove+self.abuild)*B*b+self.ahost*self.work(c,p)
        verify=self.verify_core*b/20+1.5*self.hash_ms_per_field*P
        vo=(self.body(b)+P*8)/1000
        factor=8 if assumed else (4 if b>22 else 2)
        bands=dict(query_ms=[max(0,query*1000/factor*2**(lo-b)),query*1000*factor*2**(hi-b)],
            verify_ms=[verify/factor,verify*factor],
            vo_kb=[(self.body(lo)*.75+P*8)/1000,(self.body(hi)*1.25+P*8)/1000],
            memory_GiB=[memory/2*2**(lo-b),memory*2*2**(hi-b)])
        kind='预测(参数假设)' if assumed else '预测'
        if observed:
            query=observed['query_ms']/1000;verify=observed['verify_ms'];vo=observed['vo_kb'];memory=observed['peak_rss_GiB']
            if observed['full_20x3_passed']:
                kind='实测20x3'
                bands={m:[v,v] for m,v in [('query_ms',query*1000),('verify_ms',verify),('vo_kb',vo),('memory_GiB',memory)]}
            else:
                kind=f"预测(参考{observed['samples']}/60个已验证样本)"
                rows=list(csv.DictReader(Path(observed['source']).open()))
                expected={(q,t) for q in range(100,120) for t in range(3)};done={(int(x['query']),int(x['repeat'])) for x in rows}
                for metric,column,scale in [('query_ms','query_ms',1),('verify_ms','verify_ms',1),('vo_kb','vo_bytes',1000)]:
                    total=sum(float(x[column])/scale for x in rows);missing=[]
                    for q,t in expected-done:
                        peers=[float(x[column])/scale for x in rows if int(x['query'])==q]
                        missing.append(stats.fmean(peers) if peers else total/len(rows))
                    value=(total+sum(missing))/60
                    bands[metric]=[(total+sum(missing)*.5)/60,(total+sum(missing)*2)/60]
                    if metric=='query_ms':query=value/1000
                    elif metric=='verify_ms':verify=value
                    else:vo=value
        rec=dict(id=f'P{len(self.cache)+1:03}',config=c,n_probe=p,k=k,kind=kind,query_ms=query*1000,
            verify_ms=verify,vo_kb=vo,estimated_peak_memory_GiB=memory,
            gates=G,gate_source='exact static gate count' if exact else 'structural regression prediction',
            padded_domain_bits=b,domain_bits_sensitivity=[lo,hi],scenario_ranges=bands,
            available_verified_samples=observed['samples'] if observed else 0,
            observed_source=observed['source'] if observed else None,
            exceeds_50GiB=memory>50,parameter_assumption=assumed)
        self.cache[key]=rec;return rec

def main():
    p=argparse.ArgumentParser();p.add_argument('--fit-only',action='store_true');a=p.parse_args()
    STUDY.mkdir(parents=True,exist_ok=True)
    refs=references();assert len(refs)>=2
    model=Model(refs)
    print(json.dumps(dict(reference_samples=[(x['job'],x['samples'],x['degree_bits']) for x in refs],model=model.report),ensure_ascii=False),flush=True)
    if a.fit_only:return
    book=SOURCE
    if not book.exists():book=OLD/'work9_with_v3db.xlsx'
    before_sha=sha(book);backup=STUDY/'work9_before_predictions.xlsx'
    if not backup.exists():shutil.copy2(book,backup)
    wb=openpyxl.load_workbook(book);before=openpyxl.load_workbook(book)
    assert all(x[0] in wb for x in DATA.values())
    selected=json.loads((OLD/'selected-models.json').read_text())
    cells=[];allowed=set();curves=[];overalls=[]
    fill=PatternFill('solid',fgColor='FFF2CC')
    sourcepath=str(STUDY/'forecast-points.json')
    def put(sheet,row,col,value,metric,rec=None,note='',cal=False):
        ws=wb[sheet];c=ws.cell(row,col);old=c.value;c.value=value;allowed.add((sheet,c.coordinate))
        estimated=rec and rec['kind']!='实测20x3' and metric!='recall'
        if estimated:
            c.fill=copy(fill);c.font=Font(name=c.font.name,size=c.font.sz,color='9C5700')
            c.number_format=('0.00E+00' if abs(value)>=1e6 else '0.000')+'" [预测]"'
        elif isinstance(value,(int,float)):c.number_format='0.000'
        state='搜索校准' if cal and not (rec and rec['kind']=='实测20x3') else (rec['kind'] if rec else '不适用/目标未达')
        if not isinstance(value,(float,int)):state='不适用/目标未达'
        description=dict(status=state,point_id=rec['id'] if rec else None,
            metric=metric,previous_cell_value=old,note=note,model=str(STUDY/'forecast-model.json'),point_source=sourcepath,
            observed_source=rec.get('observed_source') if rec else None,
            range=rec.get('scenario_ranges',{}).get(metric) if rec else None)
        c.comment=Comment(json.dumps(description,ensure_ascii=False,indent=2),'V3DB预测/实测来源')
        if estimated:c.comment.text+='\n预测，不是本机实测。时间假设RAM充足、无换页且单核算力相当；区间为敏感性情景，不是95%置信区间。'
        cells.append(dict(sheet=sheet,cell=c.coordinate,value=value,estimated=bool(estimated),**description))
    for dataset,(sheet,n,d,L) in DATA.items():
        ws=wb[sheet];sel=selected[dataset];folder=CACHE/sel['model'];cal=json.loads((folder/'calibration.json').read_text());cfg=cal['model']['config']
        points=[p for p in cal['points'] if p['k']==10]
        meets=[p for p in points if p['recall']>=.95-1e-12]
        best=min(meets,key=lambda p:p['n_probe']) if meets else max(points,key=lambda p:p['recall'])
        p95=best['n_probe'];overall=model.predict(cfg,p95,10)
        recall=next((r for r in refs if r['full_20x3_passed'] and r['config']==cfg and r['n_probe']==p95 and r['k']==10),None)
        recall=stats.fmean(float(x['recall']) for x in csv.DictReader(Path(recall['source']).open())) if recall else best['recall']
        overalls.append(dict(dataset=dataset,recall_search_calibration=recall,point_id=overall['id']))
        labels=[12,*[rm(r)+1 for r in [24,34,44,56,66,75,86,96,105]]]
        for row in labels:
            ws.cell(row,4,'V3DB（含预测）');allowed.add((sheet,ws.cell(row,4).coordinate))
        baseline=[r[3].row for r in ws if isinstance(r[3].value,str) and 'V3DB' in r[3].value and 'baseline' in r[3].value]
        assert len(baseline)==1,(sheet,baseline)
        ar=baseline[0];ws.cell(ar,4,'V3DB baseline（含预测）');allowed.add((sheet,ws.cell(ar,4).coordinate))
        for row,col,metric,scale in [(12,5,'recall',100),(12,6,'query_ms',1),(12,7,'vo_kb',1),(12,8,'verify_ms',1),
                                    (ar,6,'recall',1),(ar,7,'query_ms',1),(ar,8,'vo_kb',1),(ar,9,'verify_ms',1)]:
            put(sheet,row,col,recall*scale if metric=='recall' else overall[metric],metric,overall,
                '总体95%目标档；Recall为已有搜索校准，非预测上调。完整实测点优先；否则使用明确标记的条件预测。',metric=='recall')
        for col,target in enumerate(TARGETS,6):
            eligible=[p for p in points if p['recall']>=target-1e-12]
            chosen=min(eligible,key=lambda p:p['n_probe']) if eligible else None
            rec=model.predict(cfg,chosen['n_probe'],10) if chosen else None
            actual=chosen['recall'] if chosen else None
            if rec and rec['kind']=='实测20x3':
                raw=next(r for r in refs if r['source']==rec['observed_source'])
                actual=stats.fmean(float(x['recall']) for x in csv.DictReader(Path(raw['source']).open()))
            for old,metric in [(24,'vo_kb'),(34,'verify_ms'),(44,'query_ms')]:
                rr=rm(old)+1
                put(sheet,rr,col,actual*100 if actual is not None else '未达目标','recall',rec,
                    f'目标{target:.0%}。未完成证明的召回率仅为既有20查询搜索校准；不保证新的证明结果恰好达到该值。',True)
                put(sheet,rr+1,col,rec[metric] if rec else '未达目标',metric,rec,
                    f'目标{target:.0%}；没有达到目标的配置时不预测一个虚假的达标点。')
            if rec:curves.append([sheet,target,actual,rec['query_ms'],rec['verify_ms'],rec['vo_kb'],rec['kind'],rec['id']])
        for col,k in enumerate(KVALS,6):
            rec=model.predict(cfg,p95,k)
            for old,metric in [(86,'vo_kb'),(96,'verify_ms'),(105,'query_ms')]:
                put(sheet,rm(old)+1,col,rec[metric],metric,rec,'固定总体n_probe；作者该路径k主要影响公开输出长度，未假造各k的召回率。')
        for col,size in enumerate([100000,500000,1000000,5000000,10000000],6):
            supported=size<=n or (dataset in ['sift1m','deep1m'] and size<=10000000)
            rec=None;note='规模超过真实可用库容量，不进行复制采样或虚构预测。'
            if supported:
                source='bigann10m' if dataset=='sift1m' else ('deep10m' if dataset=='deep1m' and size>1000000 else dataset)
                small=None
                for f in CACHE.glob(f'{source}-{size}*/calibration.json'):
                    obj=json.loads(f.read_text());c=obj['model']['config']
                    if c['M']==cfg['M'] and c['K']==cfg['K']:
                        eligible=[p for p in obj['points'] if p['k']==10 and p['recall']>=.95-1e-12]
                        if eligible:small=(c,min(eligible,key=lambda p:p['n_probe'])['n_probe']);break
                if small:
                    rec=model.predict(*small,10);note='该规模已有IVF-PQ搜索校准；证明成本仍按实测或预测状态区分。'
                else:
                    c=dict(cfg,N=size,source=source,capacity=1<<(math.ceil(2*size/L)-1).bit_length())
                    rec=model.predict(c,p95,10,assumed=True)
                    note='该规模没有可用的已训练/已校准PQ模型；条件假设沿用全量主表的M/K/n_list/n_probe，容量按作者2N/n_list向上取二次幂。没有声称达95%召回；尤其BIGANN与SIFT1M来源不同。'
            for old,metric in [(56,'vo_kb'),(66,'verify_ms'),(75,'query_ms')]:put(sheet,rm(old)+1,col,rec[metric] if rec else '不适用',metric,rec,note)
        ws['A4']=str(ws['A4'].value or '')+'\nV3DB橙色且带[预测]数字为条件估算，不是实测；区间与参数假设见V3DB预测说明。'
        allowed.add((sheet,'A4'))
    # Keep raw measurement history intact. Refresh the cell-level provenance instead.
    for title in ['V3DB单元格来源','V3DB预测说明','V3DB预测点与区间','V3DB预测校准样本','V3DB预测曲线']:
        if title in wb:del wb[title]
    origin=wb.create_sheet('V3DB单元格来源');origin.append(['sheet','cell','value','state','metric','point ID','reason/source'])
    for r in cells:origin.append([r['sheet'],r['cell'],r['value'],r['status'],r['metric'],r['point_id'],json.dumps(r,ensure_ascii=False)])
    notes=wb.create_sheet('V3DB预测说明')
    statements=[['性质','用户授权的预测补充；橙色[预测]数字不属于实测数据。原始CSV、证明和实测result.json保持不变。'],
        ['时间条件','同等单核算力、RAM足够且无换页。超出本机内存的配置在本机没有已知实际运行时间；预测不能描述为本机已完成测试。'],
        ['四数据集','SIFT1M、NYTimes-256、GloVe-100、DEEP1M；GIST归档不再进入活跃比较。'],
        ['样本','时间模型使用本机单线程已落盘且验证通过的样本；完整20×3的点保留实测；未满60的点仍标预测。'],
        ['门数','32个小形状只构建约束并计数、不生成证明；结合已有大形状门数及原文Table9拟合非负结构模型。模型诊断见forecast-model.json。'],
        ['门数到域','GB=2^ceil(log2(G))；实际已有VerifierCircuitData提供锚点域。预测上界允许额外跨一档域。'],
        ['用时','沿用论文的GB*log2(GB)形态，但系数使用本机单线程build/prove时间；附加relation*log2(relation)的host整理项。没有照搬论文256核服务器的秒数。'],
        ['VO','预测=压缩证明本体的域位数模型+8×公开域元素个数。公开输入数量可由参数结构计算，不能将完整VO当作固定200KB。'],
        ['Verify','本机完整验证时间锚点+log域项+公开输入Poseidon哈希增长；哈希微基准只是模型部件，不是新增完整证明测量。'],
        ['内存','单线程采样峰值RSS按域大小外推，属于容量情景估计；并非论文256核内存的直接换算，亦非严格峰值保证。'],
        ['区间','上下界是主动设定的模型/跨域/硬件敏感性情景，不是统计95%置信区间。大幅外推及未训练规模更不可靠。'],
        ['召回率','只使用已有搜索校准或正式实测，不预测召回率上涨；达不到99%的目标保留未达目标，不能填低召回耗时冒充。'],
        ['规模假设','部分规模没有训练和校准：性能仅在沿用全量M/K/L/p参数假设下成立，未保证等召回95%。超容量NYTimes/GloVe规模保持不适用。'],
        ['公平性','这些预测不能与其他方法实测值混为同一类实验结果，不能据此计算实测加速比、显著性或论文实测结论。'],
        ['原文','附件papers/V3DB-26vldb.pdf，2节binning，6节环境，Table8/9/11。双路AMD EPYC、256硬件核、1TiB RAM；Table11本身是作者预测。'],
        ['可复核文件',str(STUDY)],['历史说明','原V3DB实验说明和V3DB测量记录描述历史实测阶段；本次预测的最新口径以本页和V3DB单元格来源为准。']]
    notes.append(['项目','说明'])
    for r in statements:notes.append(r)
    notes.column_dimensions['A'].width=24;notes.column_dimensions['B'].width=125
    tab=wb.create_sheet('V3DB预测点与区间')
    tab.append(['ID','状态','N','D','M','K','n_list','capacity','n_probe','k','门数','门数来源','域bit',
        'Query ms','Query下界','Query上界','Verify ms','Verify下界','Verify上界','VO KB','VO下界','VO上界',
        'RSS估计GiB','RSS下界','RSS上界','超50GiB','已有样本','原CSV','参数未经该规模校准'])
    for r in model.cache.values():
        c=r['config'];b=r['scenario_ranges'];tab.append([r['id'],r['kind'],c['N'],c['D'],c['M'],c['K'],c['n_list'],c['capacity'],r['n_probe'],r['k'],r['gates'],r['gate_source'],r['padded_domain_bits'],
            r['query_ms'],*b['query_ms'],r['verify_ms'],*b['verify_ms'],r['vo_kb'],*b['vo_kb'],r['estimated_peak_memory_GiB'],*b['memory_GiB'],r['exceeds_50GiB'],r['available_verified_samples'],r['observed_source'],r['parameter_assumption']])
    tab.freeze_panes='N2';tab.auto_filter.ref=tab.dimensions
    refsheet=wb.create_sheet('V3DB预测校准样本');refsheet.append(['job','completed samples','full20x3','domain bit','gates','query ms','verify ms','VO KB','build s','prove s','RSS peak GiB','source','SHA256','available-sample actual Recall@k'])
    for r in refs:refsheet.append([r['job'],r['samples'],r['full_20x3_passed'],r['degree_bits'],r['gates'],r['query_ms'],r['verify_ms'],r['vo_kb'],r['build_seconds'],r['prove_seconds'],r['peak_rss_GiB'],r['source'],r['source_sha256'],r['recall_mean_of_available_samples']])
    chartdata=wb.create_sheet('V3DB预测曲线');chartdata.append(['dataset','target recall','search-calibrated recall','query ms','verify ms','VO KB','status','point ID'])
    for r in curves:chartdata.append(r)
    for i,(dataset,(sheet,*_)) in enumerate(DATA.items()):
        idx=[j+2 for j,r in enumerate(curves) if r[0]==sheet]
        for j,(name,col) in enumerate([('Query ms',4),('Verify ms',5),('VO KB',6)]):
            chart=ScatterChart();chart.title=f'{sheet} V3DB: measured + FORECAST {name}';chart.x_axis.title='Search-calibrated Recall@10';chart.y_axis.title=name
            chart.width=20;chart.height=11;chart.x_axis.numFmt='0%'
            if idx:
                series=Series(Reference(chartdata,min_col=col,min_row=min(idx),max_row=max(idx)),Reference(chartdata,min_col=3,min_row=min(idx),max_row=max(idx)),title='含预测，非纯实测')
                series.graphicalProperties.line.prstDash='dash';series.graphicalProperties.line.solidFill='D98C00';chart.series.append(series)
            chartdata.add_chart(chart,f'J{1+(i*3+j)*23}')
    for tab in [origin,notes,refsheet,chartdata]:tab.freeze_panes='A2'
    wb.save(OUTPUT)
    saved=openpyxl.load_workbook(OUTPUT)
    preserved=0
    for sheet in before:
        if sheet.title.startswith('V3DB'):continue
        now=saved[sheet.title]
        for row in sheet:
            for c in row:
                if isinstance(c,openpyxl.cell.cell.MergedCell) or (sheet.title,c.coordinate) in allowed:continue
                nc=now[c.coordinate]
                assert nc.value==c.value,(sheet.title,c.coordinate,'unrelated value changed')
                assert nc._style==c._style and nc.number_format==c.number_format,(sheet.title,c.coordinate,'unrelated style changed')
                preserved+=c.value is not None
        assert len(now._charts)==len(sheet._charts),(sheet.title,'existing chart changed')
    assert len(cells)==428,len(cells)
    for r in cells:
        c=saved[r['sheet']][r['cell']];assert c.value==r['value'] or isinstance(r['value'],float) and math.isclose(c.value,r['value'],rel_tol=1e-12)
        if r['estimated']:assert '[预测]' in c.number_format and c.comment and '不是本机实测' in c.comment.text
    assert all(sha(r['source'])==r['source_sha256'] for r in refs),'Raw measurement files changed'
    write(STUDY/'forecast-points.json',list(model.cache.values()));write(STUDY/'forecast-cells.json',cells)
    write(STUDY/'overall-forecast-summary.json',[dict(x,**next(r for r in model.cache.values() if r['id']==x['point_id'])) for x in overalls])
    delivered=False
    if sha(SOURCE)==before_sha:
        try:shutil.copy2(OUTPUT,SOURCE);delivered=True
        except PermissionError:pass
    report=dict(status='PASS',output=str(OUTPUT),original_updated=delivered,
        forecast_cells=sum(r['estimated'] for r in cells),full_measurement_cells=sum(r['status']=='实测20x3' and r['metric']!='recall' for r in cells),
        unchanged_unrelated_nonempty_cells=preserved,total_V3DB_cells=len(cells),
        raw_CSVs_unchanged=True,local_reference_sample_counts=[r['samples'] for r in refs],
        source_sha256_before=before_sha,output_sha256=sha(OUTPUT),
        all_values_are_measured=False,note='Conditional forecasts are visibly tagged. Unreachable targets and unavailable sizes remain statuses.')
    write(STUDY/'validation.json',report);print(json.dumps(report,ensure_ascii=False,indent=2))
    lines=['# V3DB 预测补充（不是实测结果）','',
        '橙色且带[预测]的单元格属于条件外推。时间假设内存充足、无换页并具有类似单核算力；在本机内存不足时没有可验证的真实完成时间。',
        '', '| 数据集 | 搜索校准Recall@10 | Query预测/实测 ms | Verify ms | VO KB | RSS容量估计GiB |',
        '|---|---:|---:|---:|---:|---:|']
    for x in overalls:
        r=next(r for r in model.cache.values() if r['id']==x['point_id'])
        lines.append(f"| {x['dataset']} | {x['recall_search_calibration']:.1%} | {r['query_ms']:.3g} | {r['verify_ms']:.3g} | {r['vo_kb']:.3g} | {r['estimated_peak_memory_GiB']:.3g} |")
    lines+=['','区间见工作簿“V3DB预测点与区间”；这是敏感性情景，不是95%置信区间。',
        '',f'参考样本数：{[r["samples"] for r in refs]}。未完成60次的点仍标预测，原始CSV/证明没有改写。',
        '', '原文依据：附件Table9门数与域阶跃模型、Table11作者预测形式。仅使用其结构/门数，不把256核服务器的时间移植到本机。',
        '', '未达到的99%召回目标保留未达目标；超容量规模保留不适用。部分规模没有训练/校准，性能仅在明确列出的参数假设下成立，不能视为等召回测量。',
        '', '只有两个完整点和一个近完整点支持本机耗时，且主要处于同一域档。大幅外推没有实测验证，不能据此主张实测加速比。',
        '',f'输出：{OUTPUT}；原work9.xlsx写回={delivered}。']
    (STUDY/'RESULTS.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')

if __name__=='__main__':main()
