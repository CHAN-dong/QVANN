"""Re-estimate all existing parameter points using new verified domain anchors."""
import os
for k in ['RAYON_NUM_THREADS','OMP_NUM_THREADS','OPENBLAS_NUM_THREADS','MKL_NUM_THREADS']:os.environ[k]='1'
import csv,json,math,statistics as st,subprocess,sys,shutil
from pathlib import Path
import numpy as np
import openpyxl
from openpyxl.comments import Comment
import work9_forecast as old
ROOT=old.ROOT
OLD=ROOT/'experiments/work9-v3db-estimates-20261017'
STUDY=ROOT/'experiments/work9-v3db-reaudit-20261018'
old.STUDY=STUDY
old.OUTPUT=ROOT.parent/'work9_V3DB复核更新.xlsx'
Base=old.Model
class AuditedModel(Base):
    def __init__(self,refs):
        obs=json.loads((STUDY/'gate-observations.json').read_text())
        for f in (ROOT/'results/work9-v3db-reaudit-20261018').glob('*/result.json'):
            r=json.loads(f.read_text())
            if r['status']=='passed':obs.append(dict(config=dict(r['config'],n_probe=r['n_probe'],k=r['k']),gates=r['gates'],role='verified new proof static gate count'))
        old.write(STUDY/'gate-observations.json',obs)
        super().__init__(refs)
        anchors=[]
        for f in (ROOT/'results/work9-v3db-reaudit-20261018').glob('*/result.json'):
            r=json.loads(f.read_text())
            if r['status']!='passed':continue
            rows=list(csv.DictReader((f.parent/'per_query.csv').open()));assert len(rows)==1 and rows[0]['verified']=='1'
            v=json.loads(subprocess.check_output([old.NATIVE,str(f.parent/'trusted_verifier.bin')],text=True))
            row=rows[0];c=r['config'];P=old.pi_count(c,r['n_probe'],r['k']);assert v['public_inputs']==P
            anchors.append(dict(job=r['job'],config=c,n_probe=r['n_probe'],k=r['k'],samples=1,full_20x3_passed=False,role='model calibration only: one query, not a 20x3 benchmark',
                source=str(f.parent/'per_query.csv'),source_sha256=old.sha(f.parent/'per_query.csv'),degree_bits=v['degree_bits'],domain=2**v['degree_bits'],gates=int(row['gates']),
                build_seconds=float(row['build_seconds']),prove_seconds=float(row['prove_seconds']),query_ms=float(row['query_ms']),verify_ms=float(row['verify_ms']),vo_kb=float(row['vo_bytes'])/1000,
                public_fields=P,proof_body_bytes=float(row['vo_bytes'])-8*P,peak_rss_GiB=int(row['rss_bytes'])/2**30))
        assert len(anchors)>=4,'Need all four verified new calibration shapes'
        self.anchors=anchors;allrefs=refs+anchors
        self.old_coefficients=dict(prove=self.aprove,build=self.abuild,host=self.ahost,memory=self.mem,body_slope=self.body_slope)
        # One vote per measured configuration, rather than treating repeated proofs
        # of a single shape as independent evidence for extrapolation.
        self.aprove=st.median(r['prove_seconds']/(r['domain']*r['degree_bits']) for r in allrefs)
        self.abuild=st.median(r['build_seconds']/(r['domain']*r['degree_bits']) for r in allrefs)
        self.ahost=st.median(max(0,r['query_ms']/1000-r['build_seconds']-r['prove_seconds'])/self.work(r['config'],r['n_probe']) for r in allrefs)
        self.mem=st.median(r['peak_rss_GiB']/r['domain'] for r in allrefs)
        self.body_slope=st.median((r['proof_body_bytes']-self.body15)/(r['degree_bits']-15) for r in allrefs)
        # Piecewise observed per-bin verification core; outside measured bins use
        # the original logarithmic-height model anchored at the highest bin.
        self.verify_bins={}
        for b in sorted({r['degree_bits'] for r in allrefs}):
            self.verify_bins[b]=st.median(max(1,r['verify_ms']-1.5*self.hash_ms_per_field*r['public_fields']) for r in allrefs if r['degree_bits']==b)
        checks=[]
        for r in allrefs:
            q=(self.aprove+self.abuild)*r['domain']*r['degree_bits']+self.ahost*self.work(r['config'],r['n_probe'])
            checks.append(dict(job=r['job'],bin=r['degree_bits'],actual_seconds=r['query_ms']/1000,predicted_seconds=q,relative_error=q/(r['query_ms']/1000)-1))
        self.report.update(version=2,new_verified_calibration_shapes=len(anchors),observed_domain_bits=sorted(self.verify_bins),
            prove_seconds_per_domain_log2_domain=self.aprove,build_seconds_per_domain_log2_domain=self.abuild,
            host_seconds_per_relation_log2_relation=self.ahost,memory_GiB_per_domain_row=self.mem,proof_body_bytes_per_extra_domain_bit=self.body_slope,
            verification_core_ms_by_observed_bin=self.verify_bins,previous_coefficients=self.old_coefficients,calibration_reconstruction_checks=checks,
            parameters_policy='Existing main-table M/K/L/probe unchanged. Do not mix retuned query estimates with previously measured index costs.',
            SIFT_parameter_alternative='M128/K256/p16 has existing search calibration at 95%; listed separately, not silently substituted in the main row.',
            note='New calibration samples do not complete any missing 20x3 benchmark. Big-bin times remain unvalidated conditional estimates, not corrected measured values.')
        self.report['limitations']=[
            'Two complete 20x3 points, one 59/60 partial point, and four new single-query proof shapes calibrate the model. New single-query anchors are not complete benchmarks.',
            'Measured domain bits are 18, 19 and 20. Higher bins remain unvalidated extrapolation, even when exact static gate counts are available.',
            'Sufficient RAM, no paging and similar single-core speed are conditions, not properties of the current 50GiB-limited proof process.',
            'Scenario ranges are not statistical confidence intervals; predictions must not support claims of measured speedup.',
            'Existing search-calibrated recall and uncalibrated size assumptions are retained, not newly validated or increased.']
        old.write(STUDY/'new-calibration-anchors.json',anchors);old.write(STUDY/'forecast-model.json',self.report)
        global ACTIVE_MODEL
        ACTIVE_MODEL=self
    def predict(self,c,p,k,assumed=False):
        r=super().predict(c,p,k,assumed)
        if not r['observed_source']:
            bins=sorted(self.verify_bins);b=r['padded_domain_bits']
            if b<=bins[-1]:core=float(np.interp(b,bins,[self.verify_bins[x] for x in bins]))
            else:core=self.verify_bins[bins[-1]]*b/bins[-1]
            v=core+1.5*self.hash_ms_per_field*old.pi_count(c,p,k)
            factor=8 if assumed else (4 if b>22 else 2)
            r['verify_ms']=v;r['scenario_ranges']['verify_ms']=[v/factor,v*factor]
        return r
old.Model=AuditedModel
sys.argv=[sys.argv[0]]
shutil.copy2(old.SOURCE,STUDY/'work9_immediately_before_write.xlsx')
old.main()
assert json.loads((STUDY/'validation.json').read_text())['original_updated'],'Original workbook write blocked; do not claim updated.'
# Add explicit old/new comparison while keeping the original audit files intact.
w=openpyxl.load_workbook(old.SOURCE)
for row in w['V3DB预测说明']:
    if row[0].value=='样本':row[1].value='原有179条有效证明样本（60+60+59），加4个各1次的新校准证明；覆盖域bit18/19/20。新增单次校准不是20×3完整实验。'
    elif row[0].value=='门数':row[1].value='保留32组小形状门数，新增关键配置的静态门数核对，结合原文与既有门数重拟合；不在大形状上运行FFT或证明。'
    elif row[0].value=='用时':row[1].value='保留GB*log2(GB)结构，使用新旧七个配置的单线程实测分项重新标定；每个配置一份权重，保留主机整理项。大于bit20仍属外推。'
    elif row[0].value=='Verify':row[1].value='已测18/19/20档位的验证核心按档位插值；更高档位按最高实测档位的log域趋势外推，再加公开输入哈希项。'
for name in ['V3DB复核差异','V3DB新增校准','V3DB复核说明']:
    if name in w:del w[name]
s=w.create_sheet('V3DB复核差异');s.append(['Dataset/sheet','Cell','Metric','Before','After','Ratio','Status'])
cells=json.loads((STUDY/'forecast-cells.json').read_text())
for c in cells:
    before=c['previous_cell_value'];after=c['value']
    if c['metric'] not in ['query_ms','verify_ms','vo_kb']:continue
    if isinstance(before,(int,float)) and isinstance(after,(int,float)):
        s.append([c['sheet'],c['cell'],c['metric'],before,after,after/before if before else None,c['status']])
s=w.create_sheet('V3DB新增校准');s.append(['Job','Role','Samples','Domain bit','Gates','Query ms','Build s','Prove s','Verify ms','VO KB','Source','SHA256'])
for r in json.loads((STUDY/'new-calibration-anchors.json').read_text()):s.append([r['job'],r['role'],1,r['degree_bits'],r['gates'],r['query_ms'],r['build_seconds'],r['prove_seconds'],r['verify_ms'],r['vo_kb'],r['source'],r['source_sha256']])
s=w.create_sheet('V3DB复核说明')
for row in [
    ['性质','复核固定参数下全部预测。新校准只有各1次真实证明，不是新增20×3主实验。'],
    ['依据','新增较小与中间域档位的真实单线程build/prove/压缩/验证，结合原有179条有效样本重新标定。校准样本按配置平衡；不把同一电路重复60次当作60种外推形状。'],
    ['查询口径','保留原先搜索+主机整理+电路构建+证明生成的完整服务端口径，不通过删除构建或改并行度来降低数值。'],
    ['参数','主表既有M/K/L/probe不变，以保持与已实测索引大小/构建时间配置对应。SIFT较轻PQ备选见报告，不能冒充相同配置的预测公式修正。'],
    ['范围','小档位可检验模型；远超实测档位的高召回、大规模配置依然是条件预测。内存不足的本机没有实测完成时间，不能据此宣称实测加速比。'],
    ['结果核查','未发现秒/毫秒换算或重复叠加证明耗时错误。大小/公开输入沿用实际序列化规则；根据新增证明调整本体斜率，验证核心按已测域档位插值。'],
    ['备份',str(STUDY/'work9_before_reaudit.xlsx')],['原始证据',str(STUDY)]]:s.append(row)
s.column_dimensions['A'].width=24;s.column_dimensions['B'].width=120
if 'V3DB参数备选' in w:del w['V3DB参数备选']
s=w.create_sheet('V3DB参数备选');s.append(['Dataset','Target','Existing search calibration recall','M','K','n_probe','Estimated Query seconds','Model','Scope'])
inventory=json.loads((STUDY/'candidate-configurations.json').read_text())
alternatives=[]
for ds,(_,n,*_) in old.DATA.items():
    for target in [.6,.7,.8,.9,.95,.97,.98,.99]:
        candidates=[r for r in inventory if r['config']['dataset']==ds and r['config']['source']==ds and r['config']['N']==n and r['target']==target]
        if not candidates:continue
        r=min(candidates,key=lambda x:ACTIVE_MODEL.predict(x['config'],x['probe'],10)['query_ms'])
        q=ACTIVE_MODEL.predict(r['config'],r['probe'],10)['query_ms']/1000
        s.append([ds,target,r['recall'],r['config']['M'],r['config']['K'],r['probe'],q,r['model'],'Search-calibrated alternative only. NOT substituted into main table; index costs not remeasured for this alternative.'])
        alternatives.append(dict(dataset=ds,target=target,model=r['model'],recall=r['recall'],probe=r['probe'],query_seconds=q))
old.write(STUDY/'parameter-alternatives.json',alternatives)
temp=STUDY/'work9_reaudited.xlsx';w.save(temp);shutil.copy2(temp,old.SOURCE);shutil.copy2(temp,old.OUTPUT)
check=openpyxl.load_workbook(old.SOURCE);before=openpyxl.load_workbook(STUDY/'work9_immediately_before_write.xlsx')
allowed={(r['sheet'],r['cell']) for r in cells}|{(title,'A4') for title,*_ in old.DATA.values()}
preserved=0
for sheet in before:
    if sheet.title.startswith('V3DB'):continue
    for row in sheet:
        for c in row:
            if (sheet.title,c.coordinate) in allowed:continue
            assert check[sheet.title][c.coordinate].value==c.value,(sheet.title,c.coordinate)
            preserved+=c.value is not None
for r in cells:
    assert check[r['sheet']][r['cell']].value==r['value'] or isinstance(r['value'],float) and math.isclose(check[r['sheet']][r['cell']].value,r['value'],rel_tol=1e-12)
old.write(STUDY/'final-validation.json',dict(status='PASS',updated_original=True,reviewed_cells=len(cells),predicted_cells=sum(r['estimated'] for r in cells),unrelated_nonempty_cells_preserved=preserved,output=str(old.SOURCE),sha256=old.sha(old.SOURCE)))
print(json.dumps(json.loads((STUDY/'final-validation.json').read_text()),ensure_ascii=False))
