import hashlib
import json
from pathlib import Path
import shutil
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
STUDY = ROOT / 'experiments' / 'work9-v3db-20261016'
SOURCE = ROOT.parent / 'work9.xlsx'
STUDY.mkdir(parents=True, exist_ok=True)
backup = STUDY / 'work9_before_v3db.xlsx'
if not backup.exists():
    shutil.copy2(SOURCE, backup)
wb = openpyxl.load_workbook(backup)
inventory = []
for ws in wb:
    info = dict(sheet=ws.title, rows=ws.max_row, cols=ws.max_column,
                charts=len(ws._charts), merged=[str(x) for x in ws.merged_cells.ranges])
    if ws.title in ['SIFT1M', 'GIST1M', 'GloVe-100', 'DEEP1M']:
        info['table_labels'] = [(c.coordinate, c.value) for row in ws for c in row
                                if c.column <= 5 and c.value is not None]
        info['sample_comment'] = ws['E9'].comment.text if ws['E9'].comment else None
    inventory.append(info)
(STUDY / 'workbook_inventory.json').write_text(json.dumps(inventory, indent=2, ensure_ascii=False))
print(json.dumps([{k:v for k,v in x.items() if k in ['sheet','rows','cols','charts']} for x in inventory], ensure_ascii=False))
for x in inventory[:4]:
    print(x['sheet'], x.get('table_labels'))
