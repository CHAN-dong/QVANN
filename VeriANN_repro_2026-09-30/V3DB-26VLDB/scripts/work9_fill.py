"""Rebuild only the four comparison sheets, preserving every old cell via a map."""
from copy import copy
import csv
import hashlib
import json
from pathlib import Path
import re
import shutil
import statistics
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from work9_calibrate import ROOT, REPRO, STUDY, CACHE, DATA

SOURCE=ROOT.parent/'work9.xlsx'
BACKUP=STUDY/'work9_before_v3db.xlsx'
OUTPUT=STUDY/'work9_with_v3db.xlsx'
TARGETS=[.60,.65,.70,.75,.80,.85,.90,.95,.97,.98,.99]
INSERTS=[(11,1),(24,2),(34,2),(44,2),(56,1),(66,1),(75,1),(86,1),(96,1),(105,1)]


def sha(path):
    with path.open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()


def row_map(r):return r+sum(n for after,n in INSERTS if r>after)


def refresh():
    assert BACKUP.exists()
    wb=openpyxl.load_workbook(BACKUP)
    selected=json.loads((STUDY/'selected-models.json').read_text()) if (STUDY/'selected-models.json').exists() else {}
    protocol=json.loads((STUDY/'protocol.json').read_text()) if (STUDY/'protocol.json').exists() else {}
    thread_mode=protocol.get('proof_threads','1')
    audit=[];mappings=[];added=0;numeric=0;statuses={}
    for dataset,(sheet,full_n,d,lists,M) in DATA.items():
        if sheet not in wb:continue
        old=wb[sheet];new=wb.create_sheet(sheet+'__v3db',wb.index(old))
        for key,value in old.column_dimensions.items():new.column_dimensions[key]=copy(value)
        for key,value in old.row_dimensions.items():
            dimension=copy(value);dimension.index=row_map(key)
            new.row_dimensions[row_map(key)]=dimension
        for row in old:
            for cell in row:
                if isinstance(cell,openpyxl.cell.cell.MergedCell):continue
                target=new.cell(row_map(cell.row),cell.column,cell.value)
                target._style=copy(cell._style)
                if cell.comment:target.comment=copy(cell.comment)
                if cell.hyperlink:target._hyperlink=copy(cell.hyperlink)
                if cell.value is not None:
                    mappings.append([sheet,cell.coordinate,target.coordinate,cell.value])
        for area in old.merged_cells.ranges:
            new.merge_cells(start_row=row_map(area.min_row),end_row=row_map(area.max_row),
                            start_column=area.min_col,end_column=area.max_col)
        new.sheet_format=copy(old.sheet_format);new.sheet_properties=copy(old.sheet_properties)
        new.sheet_view.showGridLines=old.sheet_view.showGridLines
        new.freeze_panes='F9'
        wb.remove(old);new.title=sheet
        def newrow(after,count=1):
            first=row_map(after)+1
            for r in range(first,first+count):
                for col in range(4,17):
                    new.cell(r,col)._style=copy(new.cell(row_map(after),col)._style)
            if count==2:new.merge_cells(start_row=first,end_row=first+1,start_column=4,end_column=4)
            elif after!=11:new.merge_cells(start_row=first,end_row=first,start_column=4,end_column=5)
            new.cell(first,4,'V3DB')
            return first
        overall=newrow(11)
        curve=[newrow(r,2) for r in [24,34,44]]
        for rr,label in zip(curve,['VO size','verify time','query time']):
            new.cell(rr,5,'Recall rate');new.cell(rr+1,5,label)
        sizes=[newrow(r) for r in [56,66,75]]
        kvals=[newrow(r) for r in [86,96,105]]
        abl=row_map(old.max_row)+1
        for col in range(4,10):new.cell(abl,col)._style=copy(new.cell(abl-1,col)._style)
        new.merge_cells(start_row=abl,end_row=abl,start_column=4,end_column=5)
        new.cell(abl,4,'V3DB (baseline)')
        model=CACHE/selected.get(dataset,{}).get('model',f'{dataset}-{full_n}')
        cal=json.loads((model/'calibration.json').read_text()) if (model/'calibration.json').exists() else None
        results=[]
        for f in sorted((ROOT/'results/work9-v3db-20261016').glob(f'{dataset}-{full_n}-*/result.json')):
            result=json.loads(f.read_text())
            if result.get('proof_threads','1')!=thread_mode:continue
            if cal and any(result.get('config',{}).get(key)!=cal['model']['config'][key] for key in ['M','K']):continue
            if result.get('status') in ['passed','resource_limit']:results.append(result)
        def find(probe,k):
            eligible=[r for r in results if r['n_probe']==probe and r['k']==k]
            return max(eligible,key=lambda x:x.get('samples',0),default=None)
        def put(row,col,value,record=None,reason=''):
            nonlocal added,numeric
            cell=new.cell(row,col,value);added+=1
            if isinstance(value,(int,float)):numeric+=1
            if isinstance(value,str) and len(value)>8:
                alignment=copy(cell.alignment);alignment.wrap_text=True;cell.alignment=alignment
                new.row_dimensions[row].height=max(new.row_dimensions[row].height or 15,36)
            cell.number_format='0.000'
            details=reason
            if record:
                details+='\n'+json.dumps(record,ensure_ascii=False)
            cell.comment=Comment(details or '等待对应配置真实证明测量；没有复制其他方案数值。','V3DB实验来源')
            audit.append(dict(sheet=sheet,cell=cell.coordinate,value=value,source=record,reason=reason))
        default_probe=max(1,lists//128)
        points=[x for x in (cal or {}).get('points',[]) if x['k']==10]
        meets=[x for x in points if x['recall']>=.95-1e-12]
        overall_probe=min(meets,key=lambda x:x['n_probe'])['n_probe'] if meets else default_probe
        point=find(overall_probe,10)
        cal_point=next((x for x in points if x['n_probe']==overall_probe),None)
        def calibration_evidence(point,display_scale):
            if point is None:return None
            return dict(status='calibration_only',path=str(model/'calibration.json'),
                        point=point,display_scale=display_scale,query_count=20,
                        config=cal['model']['config'])
        def metric_value(record,metric,actual=None,percent=False):
            if record and record['status']=='passed':
                return record[metric]*(100 if percent else 1)
            if metric=='recall' and actual is not None:return actual*(100 if percent else 1)
            if record and record['status']=='resource_limit':
                return '内存不足(50GiB上限)' if record.get('actual_attempt') else '超预算未测'
            return '待测'
        for col,metric in [(5,'recall'),(6,'query_ms'),(7,'vo_kb'),(8,'verify_ms')]:
            value=metric_value(point,metric,cal_point['recall'] if cal_point else None,metric=='recall')
            evidence=calibration_evidence(cal_point,100) if metric=='recall' and not (point and point['status']=='passed') else point
            put(overall,col,value,evidence,'作者 multiset；PQ配置经用户授权调优。总体填实际召回；资源受限时Recall来自20查询校准，未取得20×3证明耗时。')
        for col,metric in [(6,'recall'),(7,'query_ms'),(8,'vo_kb'),(9,'verify_ms')]:
            evidence=calibration_evidence(cal_point,1) if metric=='recall' and not (point and point['status']=='passed') else point
            put(abl,col,metric_value(point,metric,cal_point['recall'] if cal_point else None),evidence,'外部baseline参考行，不属于Ours消融变体；此处recall单位0..1。')
        for col,target in enumerate(TARGETS,6):
            eligible=[x for x in points if x['recall']>=target-1e-12]
            selection=min(eligible,key=lambda x:x['n_probe']) if eligible else None
            measured=find(selection['n_probe'],10) if selection else None
            for rr,metric in zip(curve,['vo_kb','verify_ms','query_ms']):
                actual=measured['recall'] if measured and measured['status']=='passed' else (selection['recall'] if selection else None)
                put(rr,col,actual*100 if actual is not None else ('未达目标' if cal else '待校准'),
                    measured if measured and measured['status']=='passed' else calibration_evidence(selection,100),
                    f'目标{target:.0%}；只扫描已声明n_probe网格，不插值。未完成证明测量时Recall来自20查询校准。'+('网格内未达到，未填入其他召回率的耗时。' if cal and not selection else ''))
                put(rr+1,col,metric_value(measured,metric) if selection else ('未达目标' if cal else '待测'),
                    measured,f'目标{target:.0%}；证据 {model}/calibration.json；实际证明通过后填数值。')
        for col,n in enumerate([100000,500000,1000000,5000000,10000000],6):
            source='bigann10m' if dataset=='sift1m' else ('deep10m' if dataset=='deep1m' and n>1000000 else dataset)
            candidates=[]
            for f in (ROOT/'results/work9-v3db-20261016').glob(f'{source}-{n}-*/result.json'):
                r=json.loads(f.read_text())
                if r.get('proof_threads','1')!=thread_mode:continue
                if r.get('status') in ['passed','resource_limit'] and r['k']==10:
                    if dataset in selected and any(r['config'][key]!=selected[dataset].get(key,256) for key in ['M','K']):continue
                    candidates.append(r)
            r=max(candidates,key=lambda x:x.get('samples',0),default=None)
            for rr,metric in zip(sizes,['vo_kb','verify_ms','query_ms']):
                put(rr,col,metric_value(r,metric) if r else ('不适用' if dataset in ['gist1m','glove100'] and n>full_n else '待测'),r,
                    '使用原表相同数据来源的真实前缀；不从1M数据外推。')
        for col,k in enumerate([1,5,10,20,50,100],6):
            r=find(overall_probe,k)
            for rr,metric in zip(kvals,['vo_kb','verify_ms','query_ms']):
                put(rr,col,metric_value(r,metric),r,'固定总体n_probe；k为返回条数；资源受限配置只做gate预检，未填推算耗时。')
        statuses[sheet]=dict(overall_probe=overall_probe,calibrated=bool(cal),selected_model=model.name,
                             passed_workpoints=sum(r['status']=='passed' for r in results),
                             resource_limited_workpoints=sum(r['status']=='resource_limit' for r in results))
    notes=wb.create_sheet('V3DB实验说明')
    for row in [
        ['项目','说明'],['代码','作者commit a82ac732828591bd49b83d8f31c56fa4bd1cd9b2；work9-author增加压缩证明导出/解码验证、哈希函数绑定，并以索引替代电路外witness整理的嵌套线性查找；256组逐项一致性检查通过。电路和检索算法未修改。'],
        ['原表保留','四个主表每个比较面板新增V3DB，其他原单元格按位置映射保留。旧辅助表的原坐标是备份文件坐标，新坐标见V3DB原表位置映射。'],
        ['单位','Query/Verify为ms，VO为实际压缩证明bytes/1000；Recall为top-k集合重合率。消融表recall为0..1，其余为百分数。'],
        ['Query计时','搜索、候选准备、witness、电路build、prove、压缩和导出；扣除作者内部验证用时。离线索引训练和快照根构建不计入。'],
        ['客户端验证','解码实际压缩证明与预先可信的电路数据，验证证明并检查公开top-k、query和预先固定的快照根；不读取基库。'],
        ['隐私限制','保留作者默认standard_recursion_config，zero_knowledge=false；不宣称隐私保证已复现。多重集挑战绑定的既有源码观察仍适用。'],
        ['固定点','遵循作者代码rescale_database/rescale_query的16-bit全局仿射缩放；分块写盘；原始预处理数据用于精确真值。'],
        ['PQ参数','用户允许调优PQ：依次测试原配置及子向量维度4/2/1的M；选择已测n_probe网格内达到99%的最小M，否则最高实测召回。K固定256。所有试验及最终配置见selected-models.json。'],
        ['统一样本','用户指定四数据集统一20查询(IDs100..119)×3重复；同一工作负载调参，不称为独立留出评估。'],
        ['SIFT可比性','原SIFT部分的100查询测量保持原样；新增V3DB按用户要求使用100..119这20查询。故SIFT横向比较不是严格同查询样本量；其余三数据集沿用原20查询。'],
        ['证明线程',f'当前V3DB证明器线程模式={thread_mode}；native表示作者库默认并行。原表其他方案为单线程；若为native，耗时不是严格等CPU资源比较。搜索/NumPy保持单线程。'],
        ['高召回','只有实际测到目标的n_probe网格点才能填对应数值；未达目标明确标记。总体仍显示作者配置的实测召回，不伪装等召回。'],
        ['采样来源','每个数值的查询数、重复次数、参数、原始CSV均在批注及V3DB单元格来源。重复不会增加独立查询数。'],
        ['内存','重平衡保留作者相同距离和排序，仅将大型临时三维距离张量按行分块；小样本与原函数逐项对比通过。'],
        ['预算','gate-only预检超过2000000门时记超预算未测；在同一50GiB地址空间限制内尝试至2^21域。实际构建/证明分配失败另外标记内存不足，二者不混淆。'],
        ['当前进度',json.dumps(statuses,ensure_ascii=False)],
    ]:notes.append(row)
    notes.column_dimensions['A'].width=22;notes.column_dimensions['B'].width=130
    origin=wb.create_sheet('V3DB单元格来源');origin.append(['sheet','cell','value','reason','measurement_json'])
    for r in audit:origin.append([r['sheet'],r['cell'],r['value'],r['reason'],json.dumps(r['source'],ensure_ascii=False)])
    mapping=wb.create_sheet('V3DB原表位置映射');mapping.append(['sheet','原坐标_备份','新坐标','原值'])
    for r in mappings:mapping.append(r)
    tuning=wb.create_sheet('V3DB调参与资源')
    tuning.append(['dataset','model','N','M','K','max measured grid Recall@10','selected','calibration source'])
    for file in sorted(CACHE.glob('*/calibration.json')):
        record=json.loads(file.read_text());cfg=record['model']['config']
        best=max(p['recall'] for p in record['points'] if p['k']==10)
        tuning.append([cfg['dataset'],file.parent.name,cfg['N'],cfg['M'],cfg['K'],best,
                       selected.get(cfg['dataset'],{}).get('model')==file.parent.name,str(file)])
    measured=wb.create_sheet('V3DB测量记录')
    measured.append(['job','status','N','M','K','n_probe','k','distinct queries','repeats','samples',
                     'Recall@k','Query ms','Verify ms','VO decimal KB','gate count','lookup tuples','reason','source'])
    resources=wb.create_sheet('V3DB资源与尾延迟')
    resources.append(['job','status','proof threads','samples','mean build seconds','mean prove seconds',
                      'query p50 ms','query p95 ms','query p99 ms','verify p50 ms','verify p95 ms',
                      'max prove-sampled RSS GiB','configured VM limit GiB','source'])
    def percentile(values,q):
        values=sorted(values);index=(len(values)-1)*q;lo=int(index);hi=min(lo+1,len(values)-1)
        return values[lo]+(values[hi]-values[lo])*(index-lo)
    for file in sorted((ROOT/'results/work9-v3db-20261016').glob('*/result.json')):
        r=json.loads(file.read_text());cfg=r['config']
        measured.append([r['job'],r['status'],cfg['N'],cfg['M'],cfg['K'],r['n_probe'],r['k'],r.get('query_count'),
                         r.get('repeats'),r.get('samples'),r.get('recall'),r.get('query_ms'),r.get('verify_ms'),r.get('vo_kb'),
                         r.get('gates',r.get('measured_gate_count')),r.get('candidate_lookup_tuples'),r.get('reason'),str(file)])
        if r['status']=='passed':
            rows=list(csv.DictReader((file.parent/'per_query.csv').open()))
            query=[float(x['query_ms']) for x in rows];verify=[float(x['verify_ms']) for x in rows]
            resources.append([r['job'],r['status'],r.get('proof_threads','1'),len(rows),
                statistics.fmean(float(x['build_seconds']) for x in rows),
                statistics.fmean(float(x['prove_seconds']) for x in rows),
                *[percentile(query,q) for q in [.5,.95,.99]],*[percentile(verify,q) for q in [.5,.95]],
                max(int(x['rss_bytes']) for x in rows)/2**30,r.get('address_space_limit_GiB'),str(file.parent/'per_query.csv')])
        else:
            resources.append([r['job'],r['status'],r.get('proof_threads','1'),None,None,None,None,None,None,None,None,None,
                              r.get('address_space_limit_GiB'),str(file)])
    for tab in [origin,mapping,tuning,measured,resources]:
        tab.freeze_panes='D2';tab.auto_filter.ref=tab.dimensions
        for c in tab[1]:c.font=Font(bold=True)
    for sheet,old_addr,new_addr,value in mappings:
        assert wb[sheet][new_addr].value==value,(sheet,old_addr,new_addr)
    replacement=ROOT/'experiments/work9-nytimes-20261017/replacement.json'
    if replacement.exists():
        import sys
        sys.path.insert(0,str(ROOT/'scripts'))
        from work9_nytimes import overlay
        overlay(wb)
    wb.save(OUTPUT)
    if replacement.exists():
        try:shutil.copy2(OUTPUT,replacement.parent/'work9_nytimes.xlsx')
        except PermissionError:pass
    marker=STUDY/'last_written_sha256.txt'
    allowed={sha(BACKUP)}
    reconciliation=ROOT/'experiments/work9-nytimes-20261017/source_reconciliation.json'
    if replacement.exists() and reconciliation.exists():
        check=json.loads(reconciliation.read_text())
        if check['status']=='PASS' and check['source_sha256']==check['archived_source_sha256']:
            allowed.add(check['source_sha256'])
    if marker.exists():allowed.add(marker.read_text().strip())
    delivered=False
    if sha(SOURCE) in allowed:
        try:
            shutil.copy2(OUTPUT,SOURCE);marker.write_text(sha(SOURCE)+'\n');delivered=True
        except PermissionError:pass
    state=dict(new_cells=added,numeric_cells=numeric,original_cells_preserved=len(mappings),
               original_updated=delivered,output=str(OUTPUT),statuses=statuses)
    (STUDY/'workbook_progress.json').write_text(json.dumps(state,indent=2)+'\n')
    (STUDY/'workbook_audit.json').write_text(json.dumps(audit,indent=2,ensure_ascii=False)+'\n')
    print(json.dumps(state,indent=2))
    return state


if __name__=='__main__':refresh()
